#!/usr/bin/env python3
"""Self-test for the slide-deck skill: the example deck must build and pass lint, each known-bad mutation
must fail lint with its expected code. Run after any change to build_deck.py, lint_deck.py or theme.json.

  python3 scripts/test_deck.py [--only list | --only CODE[,CODE]]
Prints `SELF-TEST: <passed>/<total> passed` and exits 1 on any failure. Uses a generated test image, never a real screenshot.
--only list prints every check code the cases exercise (runs nothing); --only CODE[,CODE] runs only the cases tagged
with those codes (each case is tagged at its pick() gate; BUILD and EXAMPLES mark cases that assert no code) and exits 2
on an unknown code.
"""
import argparse, copy, json, os, re, shutil, subprocess, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import brand as BR  # noqa: E402
NB = BR.Brand({}, HERE, "neutral defaults").colors  # every case runs in an empty temp project ($BRAND_PROJECT), so builds use these
EX = json.load(open(os.path.join(ROOT, "assets", "example.deck.json"), encoding="utf-8"))
RX = json.load(open(os.path.join(ROOT, "assets", "rhythm.example.deck.json"), encoding="utf-8"))
VX = json.load(open(os.path.join(ROOT, "assets", "variety.example.deck.json"), encoding="utf-8"))
MX = json.load(open(os.path.join(ROOT, "assets", "minimal.example.deck.json"), encoding="utf-8"))
RDX = json.load(open(os.path.join(ROOT, "assets", "readiness.example.deck.json"), encoding="utf-8"))
GX = json.load(open(os.path.join(ROOT, "assets", "guide.example.deck.json"), encoding="utf-8"))
ICASES = [
    ("ICON_UNKNOWN", lambda d: d["slides"][3]["diagram"]["nodes"][0].update(icon="unicorn-sparkle")),
    ("ICON_PARTIAL", lambda d: d["slides"][3]["diagram"]["nodes"][1].pop("icon")),
    ("ICON_PLACEMENT", lambda d: d["slides"][0].update(icon="home")),
    ("MESSAGE_MISSING", lambda d: d["slides"][2].pop("message")),
    ("MESSAGE_UNKNOWN", lambda d: d["slides"][2].update(message="vibes")),
    ("MESSAGE_FIT", lambda d: d["slides"][4].update(message="numbers")),
    ("FRAME_POSITION", lambda d: d["slides"].insert(3, dict(d["slides"][-1]))),
]
MICASES = [("ICON_CONFLICT", lambda d: d["slides"][6].update(numbered=True))]


def g_second_annotated(d):
    a = dict(d["slides"][5], title="Read the Weekly Variance Chart", side="left")
    d["slides"][5]["side"] = "left"
    d["slides"].insert(6, dict(d["slides"][1], title="Week-Over-Week Changes Explained"))
    d["slides"].insert(7, a)


GCASES = [
    ("OPENER_FIRST", lambda d: d["slides"].pop(0)),
    ("CLOSING_LAST", lambda d: d["slides"].pop()),
    ("META_TYPE", lambda d: d["meta"].update(type="manual")),
    ("META_MISSING", lambda d: [d["meta"].update(type="decision"), d["meta"].pop("decision", None)]),
    ("CONTENTS_MAX", lambda d: d["slides"][0].update(items=["One", "Two"])),
    ("ANATOMY_ZONES", lambda d: d["slides"][2].update(zones=d["slides"][2]["zones"][:1])),
    ("ANATOMY_PIN", lambda d: d["slides"][2]["zones"][0].update(x=1.4)),
    ("ANNOTATIONS_MAX", lambda d: d["slides"][5].update(annotations=d["slides"][5]["annotations"][:1])),
    ("ANNOTATION_WORDS", lambda d: d["slides"][5]["annotations"][0].update(lead="The Solid Rolling Average Trend Line")),
    ("CALLOUT_KIND", lambda d: d["slides"][4]["cards"][0].update(kind="warning")),
    ("CALLOUT_KIND", lambda d: d["slides"][4]["cards"][1].update(kind="hint")),
    ("CALLOUTS_MAX", lambda d: d["slides"][4]["cards"].append(dict(d["slides"][4]["cards"][1], title="Pin Your Filters"))),
    ("ARTICLE_WORDS", lambda d: d["slides"][1]["paragraphs"].append(" ".join(["Each report includes notes on refresh timing and data sources."] * 6))),
    ("CONTACT_MAX", lambda d: d["slides"][6].update(contacts=[])),
    ("DIAGRAM_NODES", lambda d: d["slides"][3]["diagram"].update(nodes=d["slides"][3]["diagram"]["nodes"][:2])),
    ("MIRROR", g_second_annotated),
]
RDCASES = [
    ("CHECKLIST_MAX", lambda d: d["slides"][1]["items"].extend([dict(d["slides"][1]["items"][0], text="Extra item %d" % i) for i in range(3)])),
    ("CHECK_STATUS", lambda d: d["slides"][1]["items"][0].update(status="green")),
    ("ALERT_REASON", lambda d: d["slides"][1]["items"][5].pop("reason")),
    ("CHECK_WORDS", lambda d: d["slides"][1]["items"][0].update(text="Every scheduler in every region finished all of the required training modules on time")),
    ("QA_WORDS", lambda d: d["slides"][9]["pairs"][0].update(q="What if interface testing slips again")),
    ("QA_MAX", lambda d: d["slides"][9]["pairs"].extend([{"q": "Is training done?", "a": "Yes, every scheduler finished."}] * 2)),
    ("DIAGRAM_NODES", lambda d: d["slides"][3]["diagram"]["nodes"].__delitem__(slice(2, None))),
]


def rd_marker(pr):
    from pptx.util import Inches
    g = next(sh for sh in pr.slides[1].shapes if sh.name == "point")
    m = next(k for k in g.shapes if k.name == "point:num"); m.top = m.top - Inches(0.12)


def rd_center(pr):
    from pptx.util import Inches
    for sh in pr.slides[3].shapes:
        if sh.name in ("kpi", "kpi:alert", "divider"):
            sh.left = sh.left + Inches(0.4)


def rd_orphan(pr):
    t = next(sh for sh in pr.slides[11].shapes if sh.name == "title")
    p = t.text_frame.paragraphs[0]
    size, font, col = p.runs[0].font.size, p.runs[0].font.name, p.runs[0].font.color.rgb
    for el in list(p._p):
        if el.tag.endswith("}r") or el.tag.endswith("}br"):
            p._p.remove(el)
    r = p.add_run(); r.text = "Decision Needed: Approve the Go-Live Plan Immediately"
    r.font.size, r.font.name, r.font.bold = size, font, True; r.font.color.rgb = col


RDPCASES = [("MARKER_ALIGN", rd_marker), ("CENTERED", rd_center), ("ORPHAN", rd_orphan)]
MCASES = [
    ("KPI_STYLE", lambda d: d["slides"][1].update(style="glass")),
    ("KPI_STYLE", lambda d: d["slides"][1].update(dark=True)),
    ("SHIFT_LABEL", lambda d: d["slides"][3]["from"].update(label="The Whole Month of June Before")),
    ("SHIFT_CHANGE", lambda d: d["slides"][3].update(change="-34% vs June baseline")),
    ("KPI_VALUE", lambda d: d["slides"][3]["to"].update(value="6.0 percent")),
    ("STEPS_MAX", lambda d: d["slides"][2]["steps"].__delitem__(slice(1, None))),
    ("STEP_WORDS", lambda d: d["slides"][2]["steps"][0].update(label="Verify Every Patient's Coverage Much Earlier")),
    ("ROADMAP_SHAPE", lambda d: d["slides"][5]["diagram"]["lanes"][1]["items"][0].update(start="Apr", end="Feb")),
    ("ROADMAP_OVERLAP", lambda d: d["slides"][5]["diagram"]["lanes"][0]["items"][1].update(start="Mar")),
    ("ROADMAP_SHAPE", lambda d: d["slides"][5]["diagram"].update(periods=["H1", "H2"])),
]


def codes_in(out, level):
    return set(re.findall(r"^%s \[([A-Z_]+)\]" % level, out, re.M))


def d39(tmp):
    """D-39 TERM_BRAND: off-style terms warn with the preferred term, from the shared AP rules plus the project's house
    terms (a terminology file named by the brand profile). Synthetic sentence in the guide example's speaker notes."""
    R = []
    add = lambda ok, label, detail="": R.append((bool(ok), label, detail))
    if not pick("TERM_BRAND"):
        return R
    def lint(deck, name):
        dj = os.path.join(tmp, name + ".json"); json.dump(deck, open(dj, "w"))
        return subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), dj, "--draft"], capture_output=True, text=True).stdout
    kit = os.path.join(tmp, "d39_brand"); os.makedirs(kit, exist_ok=True)
    json.dump({"sets": {"internal": [{"id": "product-case", "kind": "case", "word": "Northwind", "why": "the product name is Northwind"},
                                     {"id": "saved-reports", "kind": "regex", "pattern": "\\bsaved views\\b", "suggest": "saved reports", "why": "the product calls them saved reports"}]}},
              open(os.path.join(kit, "house.json"), "w"))
    json.dump({"name": "Northwind", "terminology": "house.json"}, open(os.path.join(kit, "brand.json"), "w"))
    base = lint(copy.deepcopy(GX), "d39_base")
    d = copy.deepcopy(GX)
    d["meta"]["brand"] = "d39_brand/brand.json"
    d["slides"][1]["notes"] = (d["slides"][1].get("notes") or "") + " The denials view moves on Jan 17th so the team sees the same numbers. Over 40 saved views move with you, and NorthWind billing teams keep their filters."
    out = lint(d, "d39_terms")
    tb = [l for l in out.splitlines() if l.startswith("WARN [TERM_BRAND]")]
    want = ["Jan 17th -> Jan. 17", "Over 40 -> More than 40", "NorthWind -> Northwind", "saved views -> saved reports"]
    add(len(tb) == 4 and all(any(w in l for l in tb) for w in want) and codes_in(out, "FAIL") == codes_in(base, "FAIL"),
        "D-39 TERM_BRAND (synthetic notes, project house terms): 4 warnings with preferred terms, no new FAIL", "\n".join(tb) or out[-300:])
    d1 = copy.deepcopy(d); d1["meta"].pop("brand")
    tb1 = [l for l in lint(d1, "d39_generic").splitlines() if l.startswith("WARN [TERM_BRAND]")]
    add(len(tb1) == 2, "D-39 TERM_BRAND without house terms: only the 2 shared AP rules warn", "\n".join(tb1))
    add(not [l for l in base.splitlines() if l.startswith("WARN [TERM_BRAND]")], "D-39 TERM_BRAND: the guide example has no off-brand terms", base[-300:])
    d2 = copy.deepcopy(d); d2["meta"]["audience"] = "Analytics engineers"
    add("TERM_BRAND" not in codes_in(lint(d2, "d39_dev"), "WARN"), "D-39 TERM_BRAND: developer audiences are skipped, like the jargon checks")
    return R


def v28(tmp):
    """2.8: title widows and one-line titles, empty space, retrieval slides, shorthand, highlights, notes tags,
    screenshot slides and the --draft shot list."""
    R = []
    add = lambda ok, label, detail="": R.append((bool(ok), label, detail))
    if pick("ORPHAN", "TITLE_LINES"):
        # The builder avoids widows in any font, so the widow is written into the built file: a title that the deck's
        # own measurer (font file metrics) wraps to two lines with two short words on the last one.
        from pptx import Presentation
        import lint_deck as LD
        d = copy.deepcopy(EX)
        rc, out = run(d, tmp, "v28_widow")
        pp = os.path.join(tmp, "v28_widow.pptx")
        prs = Presentation(pp)
        sl = prs.slides[len(prs.slides) - 1]
        t = next(sh for sh in sl.shapes if sh.name == "title")
        r0 = t.text_frame.paragraphs[0].runs[0]
        size, tfont, tbold = r0.font.size.pt, r0.font.name, bool(r0.font.bold)
        avail = t.width / 914400 - 0.08
        LD.use_brand(BR.load(source=os.path.join(tmp, "v28_widow.json")))
        fit = LD.Fit()
        pool = "Zenith Explains Performance While Workflow Tools Manage the Work Across Every Team and Region".split()
        wid = None
        for k in range(4, len(pool) + 1):
            for tail in ("of Teams", "by Region", "Every Quarter", "Across Regions", "Throughout Operations", "Simultaneously Today"):
                cand = " ".join(pool[:k]) + " " + tail
                wl = fit.wrap(cand, size, tfont, tbold, avail)
                if len(wl) == 2 and wl[-1] == tail:
                    wid = cand; break
            if wid:
                break
        def widow_title(p):
            sh = next(s for s in p.slides[len(p.slides) - 1].shapes if s.name == "title")
            para = sh.text_frame.paragraphs[0]
            for el in list(para._p):
                if el.tag.endswith("}br"):
                    para._p.remove(el)
            runs = para.runs
            runs[0].text = wid
            for r in runs[1:]:
                r.text = ""
        dst = os.path.join(tmp, "v28_widow_b.pptx")
        if wid:
            pptx_mutation(pp, dst, widow_title)
            outw = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), os.path.join(tmp, "v28_widow.json"), "--pptx", dst], capture_output=True, text=True).stdout
        else:
            outw = "no widow title found for this font"
        add(wid and {"ORPHAN", "TITLE_LINES"} <= codes_in(outw, "FAIL"), "two-word title widow on a table slide (%r): ORPHAN and TITLE_LINES fail" % wid, outw)
    tbl = {"layout": "table", "message": "status", "title": "Three Roles Share the Review Work", "notes": "Each role reviews a different part of the page before sign-off.",
           "table": {"columns": ["Role", "Reviews"], "rows": [["Tester", "Each scenario"], ["Lead", "Open defects"], ["Owner", "Sign-off"]]}}
    if pick("EMPTY_SPACE"):
        d = copy.deepcopy(EX); d["slides"].insert(4, dict(tbl)); rc, out = run(d, tmp, "v28_empty_main")
        add("EMPTY_SPACE" in codes_in(out, "WARN") and "EMPTY_SPACE" not in codes_in(out, "FAIL"), "3-row, 2-column table with no takeaway: EMPTY_SPACE warns", out)
    if pick("EMPTY_SPACE"):
        d = copy.deepcopy(EX); d["appendix"].append(dict(tbl)); rc, out = run(d, tmp, "v28_empty_appx")
        add("EMPTY_SPACE" in codes_in(out, "WARN") and "EMPTY_SPACE" not in codes_in(out, "FAIL"), "empty appendix slide (3-row, 2-column table, about half used): EMPTY_SPACE warns", out)
    if pick("EMPTY_SPACE"):
        d = copy.deepcopy(EX); d["appendix"].append(dict(tbl, table={"columns": ["Role", "Reviews"], "rows": [["Tester", "Each scenario"]]})); rc, out = run(d, tmp, "v28_empty_appx1")
        add("EMPTY_SPACE" in codes_in(out, "FAIL"), "appendix slide under half used (1-row table): EMPTY_SPACE fails", out)
    q = {"layout": "statement", "role": "retrieval", "text": "A colleague wants every open account assigned from a dashboard. Name the workflow need; review the reporting question.",
         "notes": "Answer: a workflow tool assigns accounts; the dashboard explains performance."}
    if pick("RETRIEVAL_FORM"):
        d = copy.deepcopy(GX); d["slides"].insert(3, dict(q)); rc, out = run(d, tmp, "v28_retrieval_bad")
        add("RETRIEVAL_FORM" in codes_in(out, "FAIL"), "quick check that gives its own answer: RETRIEVAL_FORM fails", out)
    if pick("RETRIEVAL_FORM"):
        d = copy.deepcopy(GX); d["slides"].insert(3, dict(q, text="Quick check: a colleague wants every open account assigned from a dashboard. Which tool should do it?"))
        rc, out = run(d, tmp, "v28_retrieval_ok")
        add("LINT:" in out and "RETRIEVAL_FORM" not in codes_in(out, "FAIL"), "'Quick check: …?' with the answer in the notes passes RETRIEVAL_FORM", out)
    st = {"layout": "steps", "message": "sequence", "title": "Run a Guided Test in Three Steps", "notes": "Walk the group through each step on the live page.",
          "steps": [{"label": "Open the Test", "detail": "Open the UX page and select your test."}, {"label": "Rate It", "detail": "Rate ease; Submit the test ID"},
                    {"label": "Finish", "detail": "Select Submit to save your answers."}]}
    swap = lambda dk, s: dk["slides"].__setitem__(next(i for i, x in enumerate(dk["slides"]) if x.get("layout") == "callouts"), s)
    if pick("TELEGRAPHIC"):
        d = copy.deepcopy(GX); swap(d, dict(st)); rc, out = run(d, tmp, "v28_tele")
        add("TELEGRAPHIC" in codes_in(out, "WARN"), "shorthand detail 'Rate ease; Submit the test ID': TELEGRAPHIC warns", out)
    if pick("HIGHLIGHT_REASON"):
        d = copy.deepcopy(GX); swap(d, dict(st, highlight=2)); rc, out = run(d, tmp, "v28_hl")
        add("HIGHLIGHT_REASON" in codes_in(out, "WARN"), "highlighted step that shares no word with the title or takeaway: HIGHLIGHT_REASON warns", out)
    if pick("HIGHLIGHT_REASON"):
        d = copy.deepcopy(GX); swap(d, dict(st, highlight=2, title="Finish a Guided Test in Three Steps")); rc, out = run(d, tmp, "v28_hl_ok")
        add("LINT:" in out and "HIGHLIGHT_REASON" not in codes_in(out, "WARN"), "highlight tied to the title passes HIGHLIGHT_REASON", out)
    if pick("NOTES_TAGS"):
        d = copy.deepcopy(GX); d["slides"][1]["notes"] = "Path: Foundation. Retrieval: ask what the badge means before moving on."
        rc, out = run(d, tmp, "v28_tags")
        add("NOTES_TAGS" in codes_in(out, "WARN"), "design tags in speaker notes ('Path: Foundation', 'Retrieval:'): NOTES_TAGS warns", out)
    if pick("DIAGRAM_DETAIL"):
        d = copy.deepcopy(EX); d["slides"][4]["diagram"]["nodes"][0]["detail"] = "Region filter opens on All for every tester in each of the three regions"
        rc, out = run(d, tmp, "v28_detail14")
        add("LINT:" in out and "DIAGRAM_DETAIL" not in codes_in(out, "FAIL"), "node detail of 14 words is allowed (was 12)", out)
    ui = copy.deepcopy(GX)
    if pick("SCREENSHOT_SLIDES"):
        rc, out = run(ui, tmp, "v28_noshots")
        add("LINT:" in out and "SCREENSHOT_SLIDES" not in codes_in(out, "WARN"), "guide example (fewer than 5 slides walk through a screen): no SCREENSHOT_SLIDES", out)
    for s in ui["slides"]:
        if s.get("layout") == "article":
            s["paragraphs"] = ["Open the report and select the date range you need.", "Click Apply, then check the table before you export it."]
        if s.get("layout") == "callouts":
            for c in s.get("cards", []):
                c["text"] = "Select " + c["text"][0].lower() + c["text"][1:]
    if pick("SCREENSHOT_SLIDES"):
        rc, out = run(ui, tmp, "v28_noshots2")
        add("LINT:" in out and "SCREENSHOT_SLIDES" in codes_in(out, "WARN") and "shot list pending" not in out, "inform deck: 5+ slides walk through a screen, 2 show it: SCREENSHOT_SLIDES warns", out)
    if pick("SCREENSHOT_SLIDES"):
        rc, out = run(ui, tmp, "v28_draft", draft=True)
        sl = os.path.join(tmp, "v28_draft.shotlist.md")
        txt = open(sl, encoding="utf-8").read() if os.path.isfile(sl) else ""
        add("SCREENSHOT_SLIDES" in codes_in(out, "WARN") and "shot list pending" in out and "| ID | Shows | Screen State | Used In | Callouts |" in txt and "1,600 px" in txt and txt.count("| S") >= 3,
            "--draft writes <deck>.shotlist.md (ID, state, slide, callouts, capture rules) and the warning says 'shot list pending'", out + txt[:300])
    def m_gap(pr):
        from pptx.util import Inches
        s = pr.slides[3]; t = next(sh for sh in s.shapes if sh.name == "title"); v = next(sh for sh in s.shapes if sh.name.startswith("visual:"))
        v.top = t.top + Inches(0.6)
    src = os.path.join(tmp, "v28_base.pptx")
    dj = os.path.join(tmp, "v28_base.json"); json.dump(EX, open(dj, "w"))
    if pick("TITLE_GAP"):
        if not os.path.isfile(src):
            subprocess.run([sys.executable, os.path.join(HERE, "build_deck.py"), dj, "--out", src], capture_output=True, text=True)
        dst = os.path.join(tmp, "v28_gap.pptx"); pptx_mutation(src, dst, m_gap)
        l = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), dj, "--pptx", dst], capture_output=True, text=True)
        add("TITLE_GAP" in codes_in(l.stdout, "WARN"), "chart pushed up against the title: TITLE_GAP warns", l.stdout)
    def m_card(pr):
        from pptx.util import Inches
        for s in pr.slides:
            for g in s.shapes:
                if g.shape_type == 6 and g.name == "kpi":
                    for k in g.shapes:
                        if k.name == "card":
                            k.height = k.height * 2
                    return
    if pick("CARD_FILL"):
        if not os.path.isfile(src):
            subprocess.run([sys.executable, os.path.join(HERE, "build_deck.py"), dj, "--out", src], capture_output=True, text=True)
        dst = os.path.join(tmp, "v28_card.pptx"); pptx_mutation(src, dst, m_card)
        l = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), dj, "--pptx", dst], capture_output=True, text=True)
        add("CARD_FILL" in codes_in(l.stdout, "WARN"), "KPI card twice as tall as its text: CARD_FILL warns", l.stdout)
    here = open(os.path.join(HERE, "plain_language.py"), encoding="utf-8").read()
    for sib in ("document-builder", "material-composer", "training-material-designer"):
        cp = os.path.join(os.path.dirname(ROOT), sib, "scripts", "plain_language.py")
        if os.path.isfile(cp) and pick("BUILD"):
            add(here == open(cp, encoding="utf-8").read(), "plain_language.py matches %s's copy" % sib)
    return R


def v29(tmp):
    """2.9: Training Deck Design Standard (tests T1-T15 and R2-1 to R2-10, deck side)."""
    from pptx import Presentation
    from pptx.util import Inches
    from pptx.enum.shapes import MSO_SHAPE
    R = []
    add = lambda ok, label, detail="": R.append((bool(ok), label, detail))
    def build(deck, name, draft=True, need=True):
        dj = os.path.join(tmp, name + ".json"); json.dump(deck, open(dj, "w"))
        out = os.path.join(tmp, name + ".pptx")
        b = subprocess.run([sys.executable, os.path.join(HERE, "build_deck.py"), dj, "--out", out] + (["--draft"] if draft else []), capture_output=True, text=True)
        if not os.path.isfile(out) and not need:
            return dj, None, b.stdout + b.stderr
        if not os.path.isfile(out):
            raise AssertionError("fixture %s did not build:\n%s" % (name, (b.stdout + b.stderr)[-600:]))
        l = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), dj, "--pptx", out] + (["--draft"] if draft else []), capture_output=True, text=True)
        return dj, out, b.stdout + b.stderr + l.stdout + l.stderr
    def relint(dj, pp, draft=True):
        return subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), dj, "--pptx", pp] + (["--draft"] if draft else []), capture_output=True, text=True).stdout
    def comp_idx(dk):
        return next(i for i, x in enumerate(dk["slides"]) if x.get("layout") == "diagram" and x["diagram"].get("kind") == "comparison")
    RX_ = json.load(open(os.path.join(ROOT, "assets", "rhythm.example.deck.json")))
    def ensure(deck, name):
        pp = os.path.join(tmp, name + ".pptx")
        if not os.path.isfile(pp):
            build(deck, name)
        return os.path.join(tmp, name + ".json"), pp
    # T1 accent tabs on callouts, insight boxes and article
    if pick("ACCENT_STYLE", "ACCENT_RADIUS"):
        dj, pp, out = build(GX, "t1")
        pr = Presentation(pp)
        groups = [g for sl in pr.slides for g in sl.shapes if g.shape_type == 6 and g.name in ("callout:card", "callout:alert", "article")]
        tabs = all(any(k.name == "accent:back" for k in g.shapes) for g in groups)
        add(groups and tabs and "ACCENT_STYLE" not in codes_in(out, "FAIL") and "ACCENT_RADIUS" not in codes_in(out, "FAIL"), "T1 callouts and article carry rounded accent tabs; ACCENT_STYLE and ACCENT_RADIUS pass", out[-400:])
    if pick("BUILD"):
        dj2, pp2, out2 = build(RX_, "t1b")
        ins = [g for sl in Presentation(pp2).slides for g in sl.shapes if g.shape_type == 6 and g.name == "takeaway"]
        add(ins and all(any(k.name == "accent:back" for k in g.shapes) for g in ins), "T1 insight boxes carry the left accent tab", out2[-300:])
    # T2 a flat strip on a card edge
    def strip(p):
        for sl in p.slides:
            for g in sl.shapes:
                if g.shape_type == 6 and g.name == "article":
                    c = next(k for k in g.shapes if k.name == "card")
                    sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, c.left, c.top + Inches(0.1), Inches(0.06), c.height - Inches(0.2)); sh.name = "flat:strip"
                    return
    if pick("ACCENT_STYLE"):
        dj, pp = ensure(GX, "t1")
        dst = os.path.join(tmp, "t2.pptx"); pptx_mutation(pp, dst, strip)
        add("ACCENT_STYLE" in codes_in(relint(dj, dst), "FAIL"), "T2 a flat 0.06 in strip on a card's left edge: ACCENT_STYLE fails")
    # T3 comparison heading wraps (the builder keeps headings on one line in any font, so the wrap is put into the built file)
    if pick("MATCHED_HEADINGS"):
        d = copy.deepcopy(RX_); c = d["slides"][comp_idx(d)]["diagram"]
        c["left"]["label"], c["right"]["label"] = "Reporting (Zenith)", "Workflow (Tools Such as Orbit Today)"
        dj, pp, out = build(d, "t3")
        sn = comp_idx(d) + 2
        def wrap_second(p):
            sl = p.slides[sn - 1]
            heads = [k for g in sl.shapes if g.shape_type == 6 for k in g.shapes if k.name == "diagram:label" and k.has_text_frame]
            para = sorted(heads, key=lambda k: k.left)[-1].text_frame.paragraphs[0]
            r0 = para.runs[0]; words = r0.text.split(" ")
            r0.text = words[0]
            para.add_line_break()
            r = para.add_run(); r.text = " ".join(words[1:])
            r.font.size, r.font.bold, r.font.name = r0.font.size, r0.font.bold, r0.font.name
            r.font.color.rgb = r0.font.color.rgb
        dst = os.path.join(tmp, "t3w.pptx"); pptx_mutation(pp, dst, wrap_second)
        out3 = relint(dj, dst)
        add("MATCHED_HEADINGS" not in codes_in(out, "FAIL") and "MATCHED_HEADINGS" in codes_in(out3, "FAIL"),
            "T3 comparison headings 'Reporting (Zenith)' / 'Workflow (Tools Such as Orbit Today)': built on one line; wrapped in the file, MATCHED_HEADINGS fails", out3[-500:])
    # T4 vs marker
    if pick("COMPARISON_CONNECTOR"):
        dj, pp, out = build(RX_, "t4")
        sl = Presentation(pp).slides[comp_idx(RX_) + 1]
        nm = [k.name for g in sl.shapes if g.shape_type == 6 for k in g.shapes]
        add("diagram:vs" in nm and "diagram:arrow" not in nm and "COMPARISON_CONNECTOR" not in codes_in(out, "FAIL"), "T4 comparison uses a dark 'vs' circle, no arrow; COMPARISON_CONNECTOR passes", str(nm))
    # T5 / T6 balance
    d = copy.deepcopy(RX_); c = d["slides"][comp_idx(d)]["diagram"]
    if pick("CARD_FILL_BALANCE", "PAIRED_ALIGNMENT"):
        c["left"]["points"] = ["Two more months of paper forms", "Setup competes with peak season"]
        c["right"]["points"] = ["Same trained team does both setups", "Results visible by January", "One form for every region"]
        dj, pp, out = build(d, "t5")
        add("LINT:" in out and "CARD_FILL_BALANCE" not in codes_in(out, "FAIL") and "PAIRED_ALIGNMENT" not in codes_in(out, "FAIL") and "CARD_FILL_BALANCE" not in codes_in(out, "WARN"), "T5 paired cards with 2 and 3 similar bullets pass (no failure for unequal counts)", out[-400:])
    if pick("CARD_FILL_BALANCE"):
        c["left"]["points"] = ["Wait"]
        c["right"]["points"] = ["The same trained team does both setups during this winter",
                                "Every leader sees the results by January for annual planning",
                                "Schedulers learn one screen instead of three different intake screens"]
        dj, pp, out = build(d, "t6")
        add("CARD_FILL_BALANCE" in codes_in(out, "WARN") and "CARD_FILL_BALANCE" not in codes_in(out, "FAIL"), "T6 one card nearly full, the other nearly empty: CARD_FILL_BALANCE warns (never fails)", out[-400:])
    # T7 divider highlight split and lone word
    d = copy.deepcopy(GX)
    d["slides"].insert(1, {"layout": "section", "title": "Test Ease of Use for Each Partner Team, Then Check the Numbers", "accent": "Check the Numbers"})
    if pick("HIGHLIGHT_SPLIT", "LAST_LINE_ORPHAN"):
        dj, pp, out = build(d, "t7")
        add("LINT:" in out and not ({"HIGHLIGHT_SPLIT", "LAST_LINE_ORPHAN"} & codes_in(out, "FAIL")), "T7 the builder sizes the divider so its highlight stays on one line", out[-300:])
    def big(p):
        from pptx.util import Pt
        for sl in p.slides:
            for sh in sl.shapes:
                if sh.name == "title" and sl.slide_layout.name == "Section" and "Numbers" in sh.text_frame.text:
                    para = sh.text_frame.paragraphs[0]
                    for el in list(para._p):
                        if el.tag.endswith("}br"):
                            para._p.remove(el)
                    for r in para.runs:
                        r.font.size = Pt(54)
                    return
    if pick("HIGHLIGHT_SPLIT", "LAST_LINE_ORPHAN"):
        dj, pp = ensure(d, "t7")
        dst = os.path.join(tmp, "t7b.pptx"); pptx_mutation(pp, dst, big)
        out7 = relint(dj, dst)
        f7 = codes_in(out7, "FAIL")
        add("HIGHLIGHT_SPLIT" in f7 and ({"LAST_LINE_ORPHAN", "ORPHAN"} & f7), "T7 divider 'Test Ease of Use, Then Check the Numbers' set large enough to wrap: HIGHLIGHT_SPLIT and an orphan check fail", out7[-500:])
    # T8 cover subtitle orphan
    d = copy.deepcopy(GX); d["meta"]["subtitle"] = "Partner Training for Wireframe Review and UX Testing \u00b7 (Draft)"
    if pick("LAST_LINE_ORPHAN"):
        dj, pp, out = build(d, "t8")
        add("LINT:" in out and "LAST_LINE_ORPHAN" not in codes_in(out, "FAIL"), "T8 the builder fits the cover subtitle without leaving a word alone", out[-300:])
    def sub32(p):
        from pptx.util import Pt
        for sh in p.slides[0].shapes:
            if sh.name == "subtitle":
                for r in sh.text_frame.paragraphs[0].runs:
                    r.font.size = Pt(32)
    if pick("LAST_LINE_ORPHAN"):
        dj, pp = ensure(d, "t8")
        dst = os.path.join(tmp, "t8b.pptx"); pptx_mutation(pp, dst, sub32)
        out8 = relint(dj, dst)
        add("LAST_LINE_ORPHAN" in codes_in(out8, "FAIL"), "T8 cover subtitle at 32 pt wraps '(Draft)' alone onto its last line: LAST_LINE_ORPHAN fails", out8[-400:])
    # T9 answer slide
    q = {"layout": "statement", "message": "point", "role": "retrieval", "text": "Quick check: a colleague wants every open item assigned from a dashboard. Which tool should do it?", "notes": "Pause for answers, then show the next slide."}
    a_ = {"layout": "statement", "message": "point", "role": "answer", "text": "Answer: a workflow tool assigns the items; the dashboard explains how the work is going.", "notes": "Tie the answer back to the comparison slide."}
    swap = lambda dk, s: dk["slides"].__setitem__(next(i for i, x in enumerate(dk["slides"]) if x.get("layout") == "callouts"), s)
    if pick("RETRIEVAL_ANSWER"):
        d = copy.deepcopy(GX); swap(d, dict(q)); dj, pp, out = build(d, "t9", need=False)
        add("RETRIEVAL_ANSWER" in codes_in(out, "FAIL"), "T9 quick check with no answer slide after it: RETRIEVAL_ANSWER fails", out[-300:])
    if pick("RETRIEVAL_ANSWER"):
        d = copy.deepcopy(GX); i = next(i for i, x in enumerate(d["slides"]) if x.get("layout") == "callouts"); d["slides"][i] = dict(q); d["slides"].insert(i + 1, dict(a_))
        dj, pp, out = build(d, "t9b")
        add("LINT:" in out and "RETRIEVAL_ANSWER" not in codes_in(out, "FAIL"), "T9 quick check followed by an answer slide passes", out[-300:])
    # T11 one step, two names
    if pick("TERM_CONSISTENT"):
        d = copy.deepcopy(GX); dg = next(x for x in d["slides"] if x.get("layout") == "diagram")["diagram"]
        dg["nodes"][0]["label"], dg["nodes"][-1]["label"] = "Approve the Design", "Give Sign-off"
        dj, pp, out = build(d, "t11", need=False)
        add("TERM_CONSISTENT" in codes_in(out, "FAIL"), "T11 'Approve' on one step and 'Sign-off' on another: TERM_CONSISTENT fails", out[-300:])
    # T12 draft mark
    if pick("DRAFT_MARK"):
        dj, pp, out = build(GX, "t12")
        cover = Presentation(pp).slides[0]
        has_mark = any(sh.name == "draft:mark" for sh in cover.shapes)
        def unmark(p):
            for sh in list(p.slides[0].shapes):
                if sh.name == "draft:mark":
                    sh._element.getparent().remove(sh._element)
        dst = os.path.join(tmp, "t12b.pptx"); pptx_mutation(pp, dst, unmark)
        add(has_mark and "DRAFT_MARK" not in codes_in(out, "FAIL") and "DRAFT_MARK" in codes_in(relint(dj, dst), "FAIL"), "T12 --draft puts a draft mark on the cover; without it DRAFT_MARK fails", out[-300:])
    # T13 hero footer
    if pick("FOOTER_CONSISTENT"):
        dj, pp, out = build(RX_, "t13")
        pr = Presentation(pp)
        feet = {next((sh.text_frame.text for sh in sl.slide_layout.shapes if sh.name.startswith("Footer")), "") for sl in pr.slides}
        add(len(feet) == 1 and "internal use only" not in "".join(feet) and "FOOTER_CONSISTENT" not in codes_in(out, "FAIL"), "T13 dark hero slide carries the same footer as every other slide", str(feet))
    # T14 / R2-3 / R2-4 step markers
    d = copy.deepcopy(RX_); proc = next(x for x in d["slides"] if x.get("layout") == "diagram" and x["diagram"].get("kind") == "process")
    proc["diagram"]["nodes"] = proc["diagram"]["nodes"][:4] + [{"label": "Follow Up", "detail": "Visit checked"}]
    proc.pop("takeaway", None)
    k = d["slides"].index(proc) + 2
    if pick("BUILD"):
        dj, pp, out = build(d, "t14")
        nodes = [sh for g in Presentation(pp).slides[k - 1].shapes if g.shape_type == 6 for sh in g.shapes if sh.name == "diagram:node"]
        add(len(nodes) == 5 and not any(getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip() for sh in nodes), "R2-3 process diagram without numbered: hollow dots on the line, no numerals (supersedes T14)")
    if pick("BUILD"):
        proc["diagram"]["numbered"] = True
        dj, pp, out = build(d, "t14b")
        nums = [sh.text_frame.text.strip() for g in Presentation(pp).slides[k - 1].shapes if g.shape_type == 6 for sh in g.shapes if sh.name == "diagram:node" and getattr(sh, "has_text_frame", False)]
        add(nums == ["1", "2", "3", "4", "5"], "R2-4 / T14 process diagram with numbered: true shows circles 1 to 5", str(nums))
    # T15 / R2-7 status tags
    if pick("STATUS_TAG"):
        d = copy.deepcopy(EX); t = d["appendix"][0]["table"]
        t["tag_column"] = 0
        t["rows"][0][0] = "PASS, with the reason in Notes"
        dj, pp, out = build(d, "t15")
        cell = None
        for sl in Presentation(pp).slides:
            for g in sl.shapes:
                if getattr(g, "has_table", False) and g.has_table:
                    cell = g.table.cell(1, 0)
        runs = cell.text_frame.paragraphs[0].runs if cell is not None else []
        ok = len(runs) >= 2 and runs[0].text == "PASS" and runs[0].font.bold and str(runs[0].font.color.rgb).upper() == NB["primary"] and not runs[1].font.bold
        add(ok and "STATUS_TAG" not in codes_in(out, "FAIL"), "T15 / R2-7 'PASS, with the reason in Notes': PASS bold in the primary color, the rest normal; STATUS_TAG passes", str([(r.text, r.font.bold) for r in runs]))
    # R2-1 / R2-2 accent colors
    if pick("ACCENT_CONTRAST"):
        dj, pp, out = build(RX_, "r21")
        pr = Presentation(pp)
        hero = next(sl for sl in pr.slides if sl.slide_layout.name == "Dark")
        hero_acc = [r.font.color.rgb for sh in hero.shapes if sh.name == "title" for p in sh.text_frame.paragraphs for r in p.runs if r.font.color and str(r.font.color.rgb) != "FFFFFF"]
        add(hero_acc and all(str(c).upper() == NB["context"] for c in hero_acc) and "ACCENT_CONTRAST" not in codes_in(out, "FAIL"), "R2-1 hero accent is the context tint on dark, ACCENT_CONTRAST passes", str(hero_acc))
    if pick("ACCENT_CONTRAST"):
        d = copy.deepcopy(RX_); d["slides"].insert(2, {"layout": "section", "title": "The Pilot Results", "accent": "The Pilot Results"})
        dj, pp, out = build(d, "r22")
        sec = next(sl for sl in Presentation(pp).slides if sl.slide_layout.name == "Section")
        cols = {str(r.font.color.rgb).upper() for sh in sec.shapes if sh.name == "title" for p in sh.text_frame.paragraphs for r in p.runs}
        add(cols == {NB["card"]} and "ACCENT_CONTRAST" not in codes_in(out, "FAIL"), "R2-2 divider with its whole title accented uses the card tint on primary and passes contrast", str(cols))
    def low_contrast(p):
        from pptx.dml.color import RGBColor
        for sl in p.slides:
            for sh in sl.shapes:
                if sh.name == "title" and sh.has_text_frame and sl.slide_layout.name == "Dark":
                    sh.text_frame.paragraphs[0].runs[-1].font.color.rgb = RGBColor.from_string(NB["primary"]); return
    if pick("ACCENT_CONTRAST"):
        ensure(RX_, "r21")
        dst = os.path.join(tmp, "r21y.pptx"); pptx_mutation(os.path.join(tmp, "r21.pptx"), dst, low_contrast)
        add("ACCENT_CONTRAST" in codes_in(relint(os.path.join(tmp, "r21.json"), dst), "FAIL"), "ACCENT_CONTRAST fails a low-contrast accent (primary words on the dark hero)")
    # R2-8 file details
    if pick("BUILD"):
        ensure(RX_, "r21")
        cp = Presentation(os.path.join(tmp, "r21.pptx")).core_properties
        add(cp.title == RX_["meta"]["title"] and cp.subject == RX_["meta"].get("subtitle", ""), "R2-8 file title and subject come from the deck's details", "%r %r" % (cp.title, cp.subject))
    def subj(p):
        p.core_properties.subject = "Standardization"
    if pick("CORE_PROPS"):
        ensure(RX_, "r21")
        dst = os.path.join(tmp, "r28.pptx"); pptx_mutation(os.path.join(tmp, "r21.pptx"), dst, subj)
        add("CORE_PROPS" in codes_in(relint(os.path.join(tmp, "r21.json"), dst), "FAIL"), "CORE_PROPS fails the template subject 'Standardization'")
    # R2-9 illustration source line
    d = copy.deepcopy(GX)
    an = next(x for x in d["slides"] if x.get("layout") == "annotated")
    an["image"], an["alt"], an["source"] = "guide_illustration_wireframe.png", "Illustration of the wireframe page (not a live screenshot)", ""
    if pick("ILLUSTRATION_LABEL"):
        dj, pp, out = build(d, "r29", need=False)
        add("ILLUSTRATION_LABEL" in codes_in(out, "FAIL"), "R2-9 stamped guide illustration with no source line: ILLUSTRATION_LABEL fails", out[-300:])
    if pick("ILLUSTRATION_LABEL"):
        an["source"] = "Zenith Tableau Agent User's Guide, Figure 4 (illustration, not a live screenshot)"
        dj, pp, out = build(d, "r29b")
        add("LINT:" in out and "ILLUSTRATION_LABEL" not in codes_in(out, "FAIL"), "R2-9 the same slide with its source line passes")
    # R2-10 status wraps after the separator, never inside a word
    if pick("LABEL_BREAK"):
        d = copy.deepcopy(RX_); proc = next(x for x in d["slides"] if x.get("layout") == "diagram" and x["diagram"].get("kind") == "process")
        proc["diagram"]["nodes"][2]["detail"] = "Status shows FIXED \u00b7 AWAITING RETEST until you retest"
        dj, pp, out = build(d, "r210")
        txt = [sh.text_frame.text for sl in Presentation(pp).slides for g in sl.shapes if g.shape_type == 6 for sh in g.shapes if getattr(sh, "has_text_frame", False) and "RETEST" in sh.text_frame.text]
        add(txt and "FIXED\u00a0\u00b7 AWAITING\u00a0RETEST" in txt[0] and "LABEL_BREAK" not in codes_in(out, "FAIL"), "R2-10 'FIXED \u00b7 AWAITING RETEST' can break only after '\u00b7'; LABEL_BREAK passes", repr(txt[:1]))
    # contractions count as one word (7.4)
    if pick("DIAGRAM_LABEL"):
        d = copy.deepcopy(RX_); proc = next(x for x in d["slides"] if x.get("layout") == "diagram" and x["diagram"].get("kind") == "process")
        proc["diagram"]["nodes"][0]["label"] = "Mark What Doesn't Work"
        dj, pp, out = build(d, "c74")
        add("LINT:" in out and "DIAGRAM_LABEL" not in codes_in(out, "FAIL"), "a 4-word label with a contraction ('Mark What Doesn't Work') passes DIAGRAM_LABEL")
    # shot list: a specific list is kept
    if pick("BUILD"):
        d = copy.deepcopy(GX)
        next(x for x in d["slides"] if x.get("layout") == "annotated")["image"] = "screens/not_supplied_yet.png"
        sl = os.path.join(tmp, "keep.shotlist.md")
        open(sl, "w").write("# Shot List\n\n| ID | Shows | Screen State | Used In | Callouts |\n|---|---|---|---|---|\n| S1 | Toolbar | After Save | slide 7 | 1 Save button |\n")
        dj = os.path.join(tmp, "keep.json"); json.dump(d, open(dj, "w"))
        b = subprocess.run([sys.executable, os.path.join(HERE, "build_deck.py"), dj, "--out", os.path.join(tmp, "keep.pptx"), "--draft"], capture_output=True, text=True)
        add("After Save" in open(sl).read() and "existing list kept" in b.stdout, "an existing, specific shot list is kept, never replaced by a generated one", b.stdout[-300:])
    return R


def pptx_mutation(src, dst, fn):
    from pptx import Presentation as P
    pr = P(src); fn(pr); pr.save(dst)


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
        else:
            yield sh


def m_sibling(pr):
    from pptx.util import Pt
    labs = [sh for sh in walk(pr.slides[2].shapes) if sh.name == "kpi:label"]
    labs[0].text_frame.paragraphs[0].runs[0].font.size = Pt(14)


def m_edge(pr):
    from pptx.util import Inches
    t = next(sh for sh in pr.slides[2].shapes if sh.name == "title"); t.left = Inches(0.3)


def m_near(pr):
    from pptx.util import Inches
    t = next(sh for sh in pr.slides[2].shapes if sh.name == "source"); t.left = t.left + Inches(0.06)


def m_weight(pr):
    from pptx.util import Pt
    d = next(sh for sh in pr.slides[2].shapes if sh.name == "divider"); d.line.width = Pt(4)


PCASES = [("SIBLING_SIZE", m_sibling), ("EDGE_MARGIN", m_edge), ("NEAR_MISS", m_near), ("LINE_WEIGHT", m_weight)]


D25_DECK = json.loads('{"meta": {"type": "inform", "title": "Running a UAT Session", "subtitle": "Facilitating partner data validation on the UAT page", "presenter": "BI / Analytics", "date": "Example portfolio content", "year": 2026, "audience": "BI facilitators", "big_idea": "Facilitators guide partners through Parameters, Testing and Sign-off so every result is checked against the partner\'s own report and every issue is triaged.", "sources": ["Zenith UAT process reference (example portfolio content)"]}, "slides": [{"layout": "summary", "message": "point", "title": "Facilitators Guide Partners Through Data Validation on the UAT Page", "points": ["Walk partners through the Parameters, Testing and Sign-off tabs in order", "Help partners record each result against their own legacy report or extract", "Triage every Fail or Blocked issue with a status and a rationale", "Confirm each dashboard is signed off, or explain what still blocks it"]}, {"layout": "kpi", "message": "numbers", "title": "Every UAT Cycle Has Two Parts, Three Tabs and Three Triage Statuses", "kpis": [{"value": "2", "label": "Assessment parts"}, {"value": "3", "label": "Partner page tabs"}, {"value": "3", "label": "BI triage statuses"}], "source": "Zenith UAT process reference (example portfolio content)", "takeaway": "Your session covers Part 2: partners validate data on the UAT page while BI facilitates.", "style": "minimal"}, {"layout": "diagram", "message": "sequence", "title": "Partners Work Through the Three Tabs in Order", "diagram": {"kind": "process", "nodes": [{"label": "Parameters", "detail": "Partners enter values per scenario in plain language; completed cards turn gray"}, {"label": "Testing", "detail": "Each scenario shows the business question, steps with the partner\'s values and Record result"}, {"label": "Sign-off", "detail": "Partners sign each dashboard off; the tab warns about failed tests or open issues"}]}, "alt": "Three-step process: Parameters, then Testing, then Sign-off."}, {"layout": "statement", "role": "retrieval", "message": "point", "text": "Quick check: a partner says Record result is grayed out on the Testing tab. Why?"}, {"layout": "cards", "role": "answer", "message": "ideas", "title": "Answer: Finish Every Parameter Value, Then Return to Testing", "cards": [{"label": "Open Parameters", "detail": "Find the scenario card that has not turned gray yet"}, {"label": "Enter the missing value", "detail": "Use plain language; Reopen to update keeps the history"}, {"label": "Return to Testing", "detail": "Record result turns on once every value is entered"}]}, {"layout": "split", "message": "number", "title": "Record Each Result Against the Partner\'s Own Report", "points": ["Partners bring the legacy report or extract and never estimate from memory", "Enter whether it matched, both values and the report used", "Passed tests turn gray so partners can see what is left"], "visual": {"kind": "number", "value": "4", "label": "Questions on the result form"}, "source": "Zenith UAT process reference (example portfolio content)"}, {"layout": "statement", "role": "retrieval", "message": "point", "text": "Quick check: a partner marks a test Fail. What must they add before it saves?"}, {"layout": "diagram", "role": "answer", "message": "sequence", "title": "Answer: A Severity and a Note, Then BI Triages the Issue", "diagram": {"kind": "process", "nodes": [{"label": "Fail or Blocked", "detail": "The partner records a result that did not match or could not run"}, {"label": "Severity and note", "detail": "Both are required before the result saves"}, {"label": "Issues found", "detail": "Saving opens an issue on the Issues found list"}, {"label": "BI triage", "detail": "FIX REQUIRED, NOT A DEFECT or DEFERRED TO PHASE 2, with a rationale"}], "highlight": 1}, "alt": "Process: Fail or Blocked, add severity and note, issue opens in Issues found, BI triages it.", "takeaway": "Each Fail or Blocked result becomes a tracked issue that BI triages with a written rationale."}, {"layout": "closing", "title": "Three Habits for Every UAT Session", "asks": [{"action": "Remind partners to bring the legacy report or extract", "owner": "BI facilitator", "date": "Pre-session"}, {"action": "Triage every new issue with a status and rationale", "owner": "BI / Analytics", "date": "After each session"}, {"action": "Check Sign-off for failed tests and open issues", "owner": "BI facilitator", "date": "Each cycle"}]}]}')  # synthetic: portfolio item 09 (example portfolio content), notes replaced below


def d25_deck(answers_as_statements=False, roles=True, closing_title=None):
    """D-25 synthetic fixture: the portfolio item 09 ILT deck with 2 quick check + answer pairs."""
    d = copy.deepcopy(D25_DECK)
    for s in d["slides"]:
        s["notes"] = "Synthetic example note: walk the facilitators through this slide, then pause for their questions."
    if answers_as_statements:  # the deck as first linted: each answer is a statement slide
        d["slides"][4] = {"layout": "statement", "role": "answer", "message": "point", "notes": d["slides"][4]["notes"],
                          "text": "Answer: finish every parameter value on the Parameters tab, then return to Testing."}
        d["slides"][7] = {"layout": "statement", "role": "answer", "message": "point", "notes": d["slides"][7]["notes"],
                          "text": "Answer: add a severity and a note; saving opens an issue that BI triages."}
    if not roles:
        for s in d["slides"]:
            s.pop("role", None)
    if closing_title:
        d["slides"][8]["title"] = closing_title
    return d


def d25(tmp):
    """D-25: ratio rules count a quick check + answer pair as one interactive unit; ORPHAN reads the wrap a viewer sees
    (synthetic inputs from portfolio item 09)."""
    R = []
    add = lambda ok, label, detail="": R.append((bool(ok), label, detail))
    ratio = {"VISUAL_RATIO", "HERO_MAX", "LAYOUT_OVERUSE"}
    if pick("VISUAL_RATIO", "HERO_MAX", "LAYOUT_OVERUSE"):
        rc, out = run(d25_deck(answers_as_statements=True), tmp, "d25_pairs")
        add("LINT:" in out and not ((ratio | {"TEXT_RUN", "LAYOUT_REPEAT", "HERO_ADJACENT"}) & codes_in(out, "FAIL")),
            "D-25 synthetic portfolio 09 deck, 2 statement quick check + answer pairs: VISUAL_RATIO, HERO_MAX, LAYOUT_OVERUSE (and TEXT_RUN, LAYOUT_REPEAT, HERO_ADJACENT) stay quiet", out[-600:])
    if pick("VISUAL_RATIO", "HERO_MAX", "LAYOUT_OVERUSE"):
        rc, out = run(d25_deck(answers_as_statements=True, roles=False), tmp, "d25_pairs_untagged")
        add(ratio <= codes_in(out, "FAIL"),
            "D-25 synthetic same 4 statement slides without retrieval/answer roles: VISUAL_RATIO, HERO_MAX and LAYOUT_OVERUSE fail", out[-600:])
    orig = "Build Three Habits Into Every UAT Session"
    if pick("ORPHAN"):
        rc, out = run(d25_deck(closing_title=orig), tmp, "d25_orphan")
        add("LINT:" in out and "ORPHAN" not in codes_in(out, "FAIL"),
            "D-25 the builder now fits the original slide 10 title so the render leaves no word alone: ORPHAN passes", out[-600:])
    if pick("ORPHAN"):
        need(d25_deck(closing_title=orig), tmp, "d25_orphan")

        def as_built(pr):  # the slide as the portfolio first built it: one line at 43 pt, 0.03 in to spare under Lato Bold x1.06
            t = next(sh for sh in pr.slides[9].shapes if sh.name == "title")
            p = t.text_frame.paragraphs[0]
            font, col = p.runs[0].font.name, p.runs[0].font.color.rgb
            for el in list(p._p):
                if el.tag.endswith("}r") or el.tag.endswith("}br"):
                    p._p.remove(el)
            r = p.add_run(); r.text = orig
            r.font.size, r.font.name, r.font.bold = Pt(43), font, True; r.font.color.rgb = col
        from pptx.util import Pt
        dst = os.path.join(tmp, "d25_orphan_43pt.pptx"); pptx_mutation(os.path.join(tmp, "d25_orphan.pptx"), dst, as_built)
        l = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), os.path.join(tmp, "d25_orphan.json"), "--pptx", dst], capture_output=True, text=True)
        out = l.stdout + l.stderr
        add("ORPHAN" in codes_in(out, "FAIL") and "'Session'" in out,
            "D-25 original slide 10 title 'Build Three Habits Into Every UAT Session' at 43 pt (set past what the builder chose): ORPHAN fails on 'Session'", out[-600:])
    if pick("ORPHAN"):
        rc, out = run(d25_deck(), tmp, "d25_orphan_fixed")
        add("LINT:" in out and "ORPHAN" not in codes_in(out, "FAIL"),
            "D-25 fixed slide 10 title 'Three Habits for Every UAT Session': ORPHAN passes", out[-600:])
    return R


def branded_project(tmp):
    """The brand comes from the project: a Tailwind palette, an SVG logo and a LICENSE give the theme colors, the cover
    and footer logos and the footer line; lint judges colors against that brand. An empty project gets no logo."""
    from pptx import Presentation
    out = []
    proj = tempfile.mkdtemp(prefix="northwind_"); os.makedirs(os.path.join(proj, "public"), exist_ok=True)  # outside tmp, so other cases stay neutral
    open(os.path.join(proj, "package.json"), "w").write('{"name": "northwind-portal"}')
    open(os.path.join(proj, "LICENSE"), "w").write("Copyright (c) 2025 Northwind Traders Ltd.\n")
    open(os.path.join(proj, "tailwind.config.js"), "w").write("module.exports = { theme: { extend: { colors: { primary: { DEFAULT: '#0B7A75', 900: '#06302E' } } } } }")
    open(os.path.join(proj, "public", "logo.svg"), "w").write('<svg xmlns="http://www.w3.org/2000/svg" width="300" height="100"><circle cx="50" cy="50" r="40" fill="#0B7A75"/><rect x="110" y="30" width="170" height="40" fill="#06302E"/></svg>')
    dj = os.path.join(proj, "deck.json"); json.dump(EX, open(dj, "w")); os.makedirs(os.path.join(proj, "screenshots"), exist_ok=True)
    env = dict(os.environ, BRAND_PROJECT=proj)
    pp = os.path.join(proj, "Deck.pptx")
    b = subprocess.run([sys.executable, os.path.join(HERE, "build_deck.py"), dj, "--out", pp, "--draft"], capture_output=True, text=True, env=env)
    if b.returncode:
        return [("branded project builds", False, b.stdout + b.stderr)]
    br = BR.load(source=dj, project=proj)
    pr = Presentation(pp)
    theme = pr.slide_masters[0].part.part_related_by("http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme").blob.decode("utf-8")
    out.append(("branded project: master theme accent1 is the project primary (#%s)" % br.colors["primary"],
                '<a:accent1><a:srgbClr val="%s"/>' % br.colors["primary"] in theme and br.colors["primary"] == "0B7A75", theme[:300]))
    cover = pr.slides[0].slide_layout
    has_svg = BR.raster(os.path.join(proj, "public", "logo.svg")) is not None
    logos = [sh for sh in cover.shapes if sh.name == "Logo"]
    out.append(("branded project: the cover layout carries the project logo" + ("" if has_svg else " (skipped: no SVG rasterizer)"),
                (len(logos) == 1) if has_svg else True, str([sh.name for sh in cover.shapes])))
    feet = {sh.text_frame.text for L in pr.slide_layouts for sh in L.shapes if sh.name == "Footer"}
    out.append(("branded project: every footer reads the LICENSE holder", len(feet) == 1 and "Northwind Traders Ltd." in feet.pop(), str(feet)))
    l = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), dj, "--pptx", pp, "--draft"], capture_output=True, text=True, env=env).stdout
    out.append(("branded project: lint accepts the brand palette (no COLOR, FONT or LOGO FAIL)",
                not ({"COLOR", "FONT", "LOGO_MISSING", "LOGO_ZONE"} & codes_in(l, "FAIL")) and "LOGO_NONE" not in codes_in(l, "WARN") or not has_svg, l[-600:]))
    # the neutral run (this temp folder has no brand) draws no logo and says so
    rc, nl = run(EX, tmp, "neutral_brand", draft=True)
    np_ = Presentation(os.path.join(tmp, "neutral_brand.pptx"))
    no_logo = not any(sh.name == "Logo" for L in np_.slide_layouts for sh in L.shapes)
    out.append(("empty project: no logo drawn, LOGO_NONE warns, no footer line", no_logo and "LOGO_NONE" in codes_in(nl, "WARN")
                and not any(sh.name == "Footer" for L in np_.slide_layouts for sh in L.shapes), nl[-400:]))
    # an off-brand color is still caught
    def recolor(p):
        from pptx.dml.color import RGBColor
        sh = next(sh for sh in p.slides[2].shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip())
        sh.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("FF00FF")
    dst = os.path.join(proj, "Deck_off.pptx"); pptx_mutation(pp, dst, recolor)
    l2 = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), dj, "--pptx", dst, "--draft"], capture_output=True, text=True, env=env).stdout
    out.append(("branded project: an off-brand color (#FF00FF) fails COLOR", "COLOR" in codes_in(l2, "FAIL"), l2[-400:]))
    return out


def font_layout(tmp):
    """Layouts follow the font: the same cards slide built in two fonts gets card heights in proportion to each font's
    line spacing (read from the font file), keeps every card title on one line, and keeps the number markers centered
    on the first line (MARKER_ALIGN). Both fonts ship with the skill, so the case runs on any machine."""
    from pptx import Presentation
    out = []
    if not BR.bundled_dir():
        return [("font-based layout (skipped: no bundled fonts)", True, "")]
    deck = {"meta": dict(EX["meta"]), "slides": [{"layout": "cards", "message": "ideas", "numbered": True, "title": "Three Changes Closed the Biggest Gaps",
             "cards": [{"label": "Call Within One Day", "detail": "Schedulers make the first call by the next business day"},
                       {"label": "One Owner per Referral", "detail": "A named scheduler follows it to the first visit"},
                       {"label": "Text Before Calling", "detail": "Patients know who will call and why"}],
             "notes": "The first change drove most of the gain; the other two kept patients from dropping later."},
            {"layout": "closing", "title": "Adopt the Three Changes This Quarter", "asks": [{"action": "Approve the rollout", "owner": "Ops", "date": "Nov 14"}],
             "notes": "Ask for the rollout approval now so every region can start before the holidays."}]}
    deck["meta"].pop("brand", None)
    heights, metrics = {}, {}
    for label, fonts in (("inter", None), ("mono", {"body": "JetBrains Mono", "display": "JetBrains Mono"})):
        proj = tempfile.mkdtemp(prefix="fontlayout_")
        if fonts:
            json.dump({"fonts": fonts}, open(os.path.join(proj, "brand.json"), "w"))
        dj = os.path.join(proj, "deck.json"); json.dump(deck, open(dj, "w"))
        pp = os.path.join(proj, "d.pptx")
        env = dict(os.environ, BRAND_PROJECT=proj)
        b = subprocess.run([sys.executable, os.path.join(HERE, "build_deck.py"), dj, "--out", pp, "--draft"], capture_output=True, text=True, env=env)
        l = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), dj, "--pptx", pp, "--draft"], capture_output=True, text=True, env=env).stdout
        if b.returncode or not os.path.isfile(pp):
            return [("font-based layout: %s deck builds" % label, False, b.stdout + b.stderr)]
        cards = [sh for g in Presentation(pp).slides[1].shapes if g.name == "cards:card" for sh in g.shapes if sh.name == "card"]
        labels = [sh for g in Presentation(pp).slides[1].shapes if g.name == "cards:card" for sh in g.shapes if sh.name == "diagram:label"]
        heights[label] = cards[0].height / 914400.0 if cards else 0
        metrics[label] = BR.load(project=proj).metrics("body")["line"]
        out.append(("font-based layout (%s): card titles one line, markers aligned, no FAIL" % label,
                    not codes_in(l, "FAIL") and all(len(sh.text_frame.paragraphs) == 1 for sh in labels), l[-500:]))
    ok = metrics["mono"] > metrics["inter"] and heights["mono"] > heights["inter"]
    out.append(("font-based layout: card height follows the font's line spacing (line %.2f vs %.2f, card %.2f vs %.2f in)" %
                (metrics["mono"], metrics["inter"], heights["mono"], heights["inter"]), ok, str(heights)))
    return out


def run(deck, tmp, name, draft=False):
    os.makedirs(os.path.join(tmp, "screenshots"), exist_ok=True)
    dj = os.path.join(tmp, name + ".json")
    json.dump(deck, open(dj, "w"))
    out = os.path.join(tmp, name + ".pptx")
    flags = ["--draft"] if draft else []
    b = subprocess.run([sys.executable, os.path.join(HERE, "build_deck.py"), dj, "--out", out] + flags, capture_output=True, text=True)
    if b.returncode:
        return b.returncode, b.stdout + b.stderr
    l = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), dj, "--pptx", out] + flags, capture_output=True, text=True)
    return l.returncode, l.stdout + l.stderr


def mutate(fn, src=None):
    d = copy.deepcopy(src or EX); fn(d); return d


def only_dv(d):
    s = d["slides"]
    shot = dict(s[5]); shot.pop("series")
    d["slides"] = [s[0], s[1], s[2], s[4], shot, s[10], s[9], s[11]]


def series4(d):
    s = d["slides"]
    extra = [dict(s[6], title="Step Three: The Patient Gets a Confirmation Text"), dict(s[6], title="Step Four: The Visit Shows Up in Every Calendar")]
    s[7:7] = extra


RCASES = [
    ("LAYOUT_REPEAT", lambda d: [s.pop("series") for s in d["slides"][5:7]]),
    ("LAYOUT_REPEAT", lambda d: d["slides"].insert(5, dict(d["slides"][4], title="Pilot Wait Time Stayed Low Through September"))),
    ("ARCHETYPE_RUN", lambda d: d["slides"].__setitem__(7, dict(d["slides"][2], title="Routing Rules Cover Every Referral Type Today"))),
    ("HERO_ADJACENT", lambda d: d["slides"].insert(4, dict(d["slides"][8], title="Schedulers Now Spend Their Time on Patients", accent=None))),
    ("PACING", lambda d: [d["slides"].__setitem__(3, dict(d["slides"][7], title="Schedulers Got Two Hours Back Each Day")), d["slides"].__setitem__(8, dict(d["slides"][1], title="The Remaining Regions Start From the Same Baseline", takeaway=None))]),
    ("LAYOUT_VARIETY", only_dv),
    ("SERIES_MAX", series4),
    ("TAKEAWAY_LAYOUT", lambda d: d["slides"][7].update(takeaway="Both regions can start as soon as leadership approves.")),
    ("TAKEAWAY_ECHO", lambda d: d["slides"][10].update(takeaway="Each region needs two weeks of setup.")),
    ("TAKEAWAY_CONTEXT", lambda d: d["slides"][1].update(context="Goals were set by the operations council in June.")),
    ("HERO_LABEL", lambda d: d["slides"][3].pop("label")),
    ("HERO_TEXT", lambda d: d["slides"][8].update(title="Every Region Across the Whole Company Moves Onto One Single Intake Path by Dec 1", accent=None)),
    ("SOURCE_MISSING", lambda d: d["slides"][3].pop("source")),
]
def tone_run(d):
    s = d["slides"]
    s[7] = {"layout": "kpi", "title": "Adoption Held Above Nine in Ten Calls Each Week", "kpis": [{"value": "93%", "label": "Calls Within One Day"}], "source": "Referral tracker (synthetic example)", "notes": "Adoption stayed high every week of the pilot quarter."}
    s[8] = dict(s[5], title="West Kept Its Lead Through the Whole Quarter", takeaway=None)


def overuse(d):
    s = d["slides"]
    for i, t in ((1, "Calls Now Happen Sooner in Every Clinic"), (5, "Owners Keep Referrals From Slipping Away"), (7, "Texts Raised Answer Rates on First Calls"), (9, "Champions Will Coach Each New Region")):
        s[i] = dict(s[3], title=t, tint=False)


def distinct(d):
    s = d["slides"]
    body = []
    for i in range(4):
        body.append(dict(s[3], title="Change Number %d Closed Another Referral Gap" % (i + 1), tint=False))
        body.append(dict(s[5], title="Region Group %d Improved After the Change" % (i + 1)))
    d["slides"] = [s[0]] + body + [s[-1]]


VCASES = [
    ("TINT_ADJACENT", lambda d: d["slides"].pop(2)),
    ("TONE_RUN", tone_run),
    ("LAYOUT_OVERUSE", overuse),
    ("LAYOUT_DISTINCT", distinct),
    ("MIRROR", lambda d: [d["slides"][1].update(side="left"), d["slides"][7].update(side="left")]),
    ("MIRROR", lambda d: [d["slides"][9].update(side="left"), d["slides"][10].update(side="left")]),
    ("CARDS_MAX", lambda d: d["slides"][3]["cards"].extend([{"label": "Fourth Idea", "detail": "x"}, {"label": "Fifth Idea", "detail": "y"}])),
    ("CARD_WORDS", lambda d: d["slides"][3]["cards"][0].update(label="Call Every Patient Within One Business Day")),
    ("SPLIT_VISUAL", lambda d: d["slides"][1]["visual"].update(kind="chart")),
    ("SPLIT_PANEL", lambda d: d["slides"][1].update(panel="alert")),
    ("FUNNEL_ORDER", lambda d: d["slides"][2]["diagram"]["nodes"].reverse()),
    ("FUNNEL_VALUE", lambda d: d["slides"][2]["diagram"]["nodes"][0].pop("value")),
    ("CYCLE_TAKEAWAY", lambda d: [d["slides"][6]["diagram"]["nodes"].append({"label": "Share", "detail": "Tell other clinics"}), d["slides"][6].update(takeaway="Two-week cycles proved each change before it spread.")]),
    ("TINT_LAYOUT", lambda d: d["slides"][5].update(tint=True)),
    ("DIAGRAM_KIND", lambda d: d["slides"][6]["diagram"].update(kind="pyramid")),
]
ECASES = [
    ("HERO_MAX", lambda d: [d["slides"][1].update(dark=True), d["slides"].__setitem__(4, {"layout": "hero", "title": "Two Fixes Stand Between Us and Release", "text": "Both finish by Oct 28.", "notes": "Breather before the ask so the decision lands clearly."})]),
]


def by_layout(d, lay):
    return next(s for s in d["slides"] if s["layout"] == lay)


CASES = [
    ("BULLETS_MAX", lambda d: by_layout(d, "summary").update(points=["One point here", "Two", "Three", "Four", "Five"])),
    ("TITLE_GENERIC", lambda d: by_layout(d, "kpi").update(title="Overview")),
    ("CHART_TYPE", lambda d: by_layout(d, "chart")["chart"].update(type="pie")),
    ("ALERT_REASON", lambda d: by_layout(d, "chart")["chart"].update(alert={"category": "Claims Edit"})),
    ("ALT_MISSING", lambda d: by_layout(d, "diagram").pop("alt")),
    ("CALLOUTS_MAX", lambda d: by_layout(d, "screenshot")["callouts"].append({"x": 0.5, "y": 0.5, "text": "Fourth callout"})),
    ("CHART_HIGHLIGHT", lambda d: by_layout(d, "chart")["chart"].pop("highlight")),
    ("SOURCE_MISSING", lambda d: by_layout(d, "chart").pop("source")),
    ("TEXT_OVERFLOW", lambda d: by_layout(d, "closing").update(asks=[{"action": "Approve the release with both fixes in flight and confirm that every downstream team has reviewed the change log, the training plan, the communication plan and the rollback plan before the go-live call and then confirm again with every regional director, every clinic lead and every billing supervisor in writing before anyone starts", "owner": "Billing ops lead", "date": "Oct 20"}] * 3)),
    ("VISUAL_RATIO", lambda d: [s.update(layout="bullets", points=["Point one here"]) for s in d["slides"] if s["layout"] in ("kpi", "chart", "diagram")]),
    ("CLOSING_LAST", lambda d: d["slides"].pop()),
    ("SLOP_WORD", lambda d: by_layout(d, "summary")["points"].__setitem__(0, "We leverage UAT to pass 46 scenarios")),
    ("DIAGRAM_NODES", lambda d: by_layout(d, "diagram")["diagram"]["nodes"].extend([{"label": "Extra", "detail": "x"}] * 3)),
    ("LOGO_OPTION", lambda d: d["meta"].update(logo="brand/old_logo.png")),
    ("ACCENT", lambda d: d["meta"].update(title_accent="Not in the title")),
    ("TABLE_SIZE", lambda d: d["slides"].insert(3, dict(d["appendix"][0], title="Seven open items need owners this week", table={"columns": ["A", "B"], "rows": [["1", "2"]] * 7}))),
]


PROPER = [("held revenue review time dropped most for cash posting", "Held Revenue Review Time Dropped Most for Cash Posting"),
          ("UAT passed, and we recommend releasing on Nov 4", "UAT Passed, and We Recommend Releasing on Nov 4"),
          ("quick wins sit in the high-impact, low-effort corner", "Quick Wins Sit in the High-impact, Low-effort Corner"),
          ("decision needed: approve the release", "Decision Needed: Approve the Release"),
          ("what we are working on", "What We Are Working On")]


ONLY = None   # --only CODE[,CODE]: run only cases tagged with one of these codes; None runs every case
SEEN = None   # --only (dry pass): each case's tags are collected here and no case runs


def pick(*codes):
    """Gate for one case, tagged with the check codes it asserts (BUILD or EXAMPLES when it asserts none)."""
    if SEEN is not None:
        SEEN.append(codes)
        return False
    return ONLY is None or not ONLY.isdisjoint(codes)


def need(deck, tmp, name):
    """Build and lint a fixture deck that a selected case reads, when --only skipped the case that makes it."""
    if not os.path.isfile(os.path.join(tmp, name + ".pptx")):
        run(deck, tmp, name)


def vbatch(tmp):
    """VISUAL_BATCH: the contact-sheet review log (visual_batch.py) on synthetic PNG pages; needs no LibreOffice."""
    R = []
    add = lambda ok, label, detail="": R.append((bool(ok), label, detail))
    vb = lambda *a: subprocess.run([sys.executable, os.path.join(HERE, "visual_batch.py")] + [str(x) for x in a], capture_output=True, text=True)
    def pages():
        src = os.path.join(tmp, "vb_pages")
        if not os.path.isdir(src):
            from PIL import Image, ImageDraw
            os.makedirs(src)
            for i in range(1, 7):
                im = Image.new("RGB", (800, 450), "white"); ImageDraw.Draw(im).text((40, 40), "Synthetic page %d" % i, fill="black")
                im.save(os.path.join(src, "slide_%02d.png" % i))
        return src
    def fresh():
        r = vb("render", pages(), "--out", os.path.join(tmp, "vb_out"))
        m = re.search(r"^render: (.+?) \| 6 pages on 2 sheets$", r.stdout, re.M)
        return (m.group(1) if m else None), r.stdout + r.stderr
    if pick("VISUAL_BATCH"):
        d, out = fresh(); ok = False
        if d:
            vb("viewed", d, "--sheet", 1); c = vb("check", d); out += c.stdout
            ok = c.returncode == 1 and "FAIL VISUAL_BATCH: pages on no viewed sheet: [5, 6]" in c.stdout
        add(ok, "VISUAL_BATCH fails when a page is on no viewed sheet (6 pages on 2 sheets, only sheet 1 viewed)", out)
    if pick("VISUAL_BATCH"):
        d, out = fresh(); ok = False
        if d:
            vb("viewed", d, "--sheet", 1, 2); vb("finding", d, "--page", 3, "--text", "label cut off")
            c = vb("check", d); out += c.stdout
            ok = c.returncode == 1 and "unfixed findings: page 3: label cut off" in c.stdout and "no viewed sheet" not in c.stdout
        add(ok, "VISUAL_BATCH fails on an unfixed finding (every sheet viewed)", out)
    if pick("VISUAL_BATCH"):
        d, out = fresh(); ok = False
        if d:
            log = os.path.join(tmp, "vb_deck.visual_qa.md")
            vb("viewed", d, "--sheet", 1, 2); vb("full", d, "--page", 3, "--reason", "small table text")
            vb("finding", d, "--page", 3, "--text", "label cut off"); vb("fixed", d, "--page", 3, "--note", "shorter label")
            c = vb("check", d, "--log", log); out += c.stdout
            txt = open(log, encoding="utf-8").read() if os.path.isfile(log) else ""
            ok = c.returncode == 0 and "visual review: 6 pages on 2 sheets, 1 viewed full size, 1 findings fixed" in c.stdout and "| 3 | 1 | yes | label cut off | yes |" in txt
        add(ok, "VISUAL_BATCH passes once every sheet is viewed and the finding is fixed; check --log writes the table", out)
    if pick("VISUAL_BATCH"):
        d1, o1 = fresh(); d2, o2 = fresh()
        ok = bool(d1 and d2) and d1 != d2 and all(os.path.isfile(os.path.join(x, "sheet_01.png")) for x in (d1, d2))
        add(ok, "each render of the same pages goes to a new folder (%s / %s)" % tuple(os.path.basename(str(x)) for x in (d1, d2)), o1 + o2)
    return R


def cases():
    tmp = tempfile.mkdtemp()
    os.environ["BRAND_PROJECT"] = tmp  # an empty project: every build and lint below uses the neutral brand (NB)
    os.makedirs(os.path.join(tmp, "screenshots"), exist_ok=True)
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (1600, 900), "#F4F4F4")
    dr = ImageDraw.Draw(im); dr.rectangle([40, 40, 1560, 160], fill="#FFFFFF"); dr.rectangle([40, 200, 1000, 860], fill="#FFFFFF")
    dr.text((60, 80), "TEST IMAGE - not a real screenshot", fill="#7D7D7D")
    passed, total, lines = 0, 0, []
    for name, deck, draft, want in (("example_draft", EX, True, 0), ("example_strict_noimg", EX, False, 1)):
        if not pick("EXAMPLES"):
            continue
        total += 1
        rc, out = run(deck, tmp, name, draft)
        ok = (rc == 0) if want == 0 else (rc != 0)
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " %s (exit %d)" % (name, rc) + ("" if ok else "\n" + out[-600:]))
    im.save(os.path.join(tmp, "screenshots", "billing_eo_overview.png"))
    for k in (1, 2):
        im.save(os.path.join(tmp, "screenshots", "intake_form_step%d.png" % k))
    for nm in ("dashboard_overview", "admissions_widget"):
        im.save(os.path.join(tmp, "screenshots", "%s.png" % nm))
    im.save(os.path.join(tmp, "screenshots", "referral_queue.png"))
    if pick("EXAMPLES"):
        total += 1
        rc, out = run(EX, tmp, "example_strict")
        passed += rc == 0; lines.append(("ok  " if rc == 0 else "FAIL") + " example_strict with image (exit %d)" % rc + ("" if rc == 0 else "\n" + out[-800:]))
    sys.path.insert(0, HERE)
    from lint_deck import proper
    for raw, want in PROPER:
        if not pick("BUILD"):
            continue
        total += 1
        ok = proper(raw) == want and proper(want) == want
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " proper case: %r -> %r" % (raw, proper(raw)))
    from pptx import Presentation
    if pick("PROPER_CASE"):
        need(EX, tmp, "example_strict")
        total += 1
        pr = Presentation(os.path.join(tmp, "example_strict.pptx"))
        t = next(sh for sh in pr.slides[2].shapes if sh.name == "title")
        t.text_frame.paragraphs[0].runs[0].text = t.text_frame.text.lower()
        pr.save(os.path.join(tmp, "neg_case.pptx"))
        l = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), os.path.join(tmp, "example_strict.json"), "--pptx", os.path.join(tmp, "neg_case.pptx")], capture_output=True, text=True)
        ok = l.returncode != 0 and "[PROPER_CASE]" in l.stdout
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " negative PROPER_CASE (lower-case title in built deck)")
    if pick("EXAMPLES"):
        total += 1
        rc, out = run(RX, tmp, "rhythm_strict")
        passed += rc == 0; lines.append(("ok  " if rc == 0 else "FAIL") + " rhythm example builds and lints clean (exit %d)" % rc + ("" if rc == 0 else "\n" + out[-800:]))
    if pick("EXAMPLES"):
        need(RX, tmp, "rhythm_strict")
        total += 1
        st = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), os.path.join(tmp, "rhythm_strict.json"), "--story"], capture_output=True, text=True).stdout
        ok = "Rhythm (S summary" in st and "S D V H D V V T H V D C" in st
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " --story prints the rhythm strip" + ("" if ok else "\n" + st[-300:]))
    if pick("EXAMPLES"):
        need(RX, tmp, "rhythm_strict")
        total += 1
        pr = Presentation(os.path.join(tmp, "rhythm_strict.pptx"))
        hero = pr.slides[4]
        names = [sh.name for sh in hero.shapes]
        yel = [sh.name for sh in hero.shapes if NB["context"] in sh._element.xml.upper()]
        take = sum(1 for sl in pr.slides for sh in sl.shapes if sh.name == "takeaway")
        ok = hero.slide_layout.name == "Dark" and "hero:value" in names and yel == ["title"] and take == 4
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " hero on the dark layout, context-tint accent only in its title; 4 insight boxes built" + ("" if ok else " (%s, %s, %s, %d)" % (hero.slide_layout.name, names, yel, take)))
    for code, fn in RCASES:
        if not pick(code):
            continue
        total += 1
        rc, out = run(mutate(fn, RX), tmp, "rneg_" + code.lower())
        ok = rc != 0 and ("[%s]" % code) in out
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " negative %s (rhythm)" % code + ("" if ok else "\n" + out[-600:]))
    for label, fn in (("kpi", lambda d: d["slides"][1].update(takeaway="Testers can move to retest a week early.")),
                      ("dark kpi", lambda d: d["slides"][1].update(dark=True, takeaway="Testers can move to retest a week early."))):
        if not pick("BUILD"):
            continue
        total += 1
        rc, out = run(mutate(fn), tmp, "pos_take_" + label.replace(" ", "_"))
        passed += rc == 0; lines.append(("ok  " if rc == 0 else "FAIL") + " insight box on a %s slide builds and lints clean" % label + ("" if rc == 0 else "\n" + out[-600:]))
    if pick("EXAMPLES"):
        total += 1
        rc, out = run(VX, tmp, "variety_strict")
        passed += rc == 0; lines.append(("ok  " if rc == 0 else "FAIL") + " variety example builds and lints clean (exit %d)" % rc + ("" if rc == 0 else "\n" + out[-800:]))
    if pick("EXAMPLES"):
        need(VX, tmp, "variety_strict")
        total += 1
        try:
            pv = Presentation(os.path.join(tmp, "variety_strict.pptx"))
            sl = list(pv.slides)
            grp = lambda i, nm: next(sh for sh in sl[i].shapes if sh.name == nm)
            pic = lambda i: next(sh for sh in sl[i].shapes if sh.name == "visual:screenshot")
            p3, p9 = grp(2, "split:panel").left / 914400, grp(8, "split:panel").left / 914400
            bands = [any(sh.name == "bg:band" for sh in sl[i].shapes) for i in (4, 9)]
            arrows = sum(1 for sh in grp(7, "visual:diagram").shapes if sh.name == "diagram:arrow" and "tailEnd" in sh._element.xml)
            bars = [sh.width for sh in grp(3, "visual:diagram").shapes if sh.name == "diagram:bar"]
            ok = p3 < 1 < 6 < p9 and pic(11).left > pic(10).left and all(bands) and arrows == 4 and len(bars) == 4 and all(a > b for a, b in zip(bars, bars[1:]))
            detail = "panels %.1f/%.1f, pics %d/%d, bands %s, arrows %d, bars %s" % (p3, p9, pic(10).left, pic(11).left, bands, arrows, bars)
        except Exception as e:
            ok, detail = False, repr(e)
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " split panels and walkthrough mirror, tint bands, 4 cycle arrows, narrowing funnel" + ("" if ok else " (%s)" % detail))
    if pick("EXAMPLES"):
        need(VX, tmp, "variety_strict")
        total += 1
        st = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), os.path.join(tmp, "variety_strict.json"), "--story"], capture_output=True, text=True).stdout
        ok = "Tone (W white" in st and "S L W L N W W N L W W C" in st
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " --story prints the tone strip" + ("" if ok else "\n" + st[-300:]))
    if pick("EXAMPLES"):
        total += 1
        rc, out = run(MX, tmp, "minimal_strict")
        passed += rc == 0; lines.append(("ok  " if rc == 0 else "FAIL") + " minimal example builds and lints clean (exit %d)" % rc + ("" if rc == 0 else "\n" + out[-800:]))
    if pick("BUILD"):
        total += 1
        long = mutate(lambda d: d["slides"][1]["kpis"][1].update(label="Unreimbursed Accounts Receivable Balances"), MX)
        rc, out = run(long, tmp, "harmonize")
        try:
            pm = Presentation(os.path.join(tmp, "harmonize.pptx"))
            szs = {r.font.size.pt for sh in walk(pm.slides[2].shapes) if sh.name == "kpi:label" for p in sh.text_frame.paragraphs for r in p.runs}
            ok = rc == 0 and len(szs) == 1 and max(szs) < 18
        except Exception as e:
            ok, szs = False, repr(e)
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " siblings harmonized to one size after auto-fit" + ("" if ok else " (%s) %s" % (szs, out[-300:])))
    for code, fn in MCASES:
        if not pick(code):
            continue
        total += 1
        rc, out = run(mutate(fn, MX), tmp, "mneg_" + code.lower())
        ok = rc != 0 and ("[%s]" % code) in out
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " negative %s (minimal)" % code + ("" if ok else "\n" + out[-600:]))
    for code, fn in PCASES:
        if not pick(code):
            continue
        need(MX, tmp, "minimal_strict")
        total += 1
        dst = os.path.join(tmp, "pneg_%s.pptx" % code.lower())
        pptx_mutation(os.path.join(tmp, "minimal_strict.pptx"), dst, fn)
        l = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), os.path.join(tmp, "minimal_strict.json"), "--pptx", dst], capture_output=True, text=True)
        ok = l.returncode != 0 and ("[%s]" % code) in l.stdout
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " negative %s (built deck)" % code + ("" if ok else "\n" + l.stdout[-500:]))
    if pick("EXAMPLES"):
        total += 1
        rc, out = run(RDX, tmp, "readiness_strict")
        passed += rc == 0; lines.append(("ok  " if rc == 0 else "FAIL") + " readiness example builds and lints clean (exit %d)" % rc + ("" if rc == 0 else "\n" + out[-800:]))
    if pick("EXAMPLES"):
        need(RDX, tmp, "readiness_strict")
        total += 1
        try:
            pr = Presentation(os.path.join(tmp, "readiness_strict.pptx"))
            E = 914400.0
            pts = [g for g in pr.slides[1].shapes if g.name == "point"]
            offs = []
            for g in pts:
                m = next(k for k in g.shapes if k.name == "point:num"); t = next(k for k in g.shapes if k.name == "body")
                fs = t.text_frame.paragraphs[0].runs[0].font.size.pt
                offs.append(abs((m.top + m.height / 2) / E - (t.top / E + 0.04 + 0.62 * fs / 72)))
            closing = next(sh for sh in pr.slides[11].shapes if sh.name == "title").text_frame.paragraphs[0].text
            cover = next(sh for sh in pr.slides[0].shapes if sh.name == "title").text_frame.paragraphs[0].text
            rows = [k for g in pr.slides[11].shapes if g.name == "ask" for k in g.shapes if k.name == "ask:row"]
            ok = len(offs) == 3 and max(offs) < 0.01 and "Nov\u00a017" in closing and "\v" in cover and all(abs((r.left + r.width) / E - 12.65) < 0.01 for r in rows)
            detail = "offs %s, closing %r, cover %r, rows %s" % ([round(o, 3) for o in offs], closing, cover, [round((r.left + r.width) / E, 2) for r in rows])
        except Exception as e:
            ok, detail = False, repr(e)
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " numbers centered on first lines, Nov 17 kept together, title breaks locked, ask rows full width" + ("" if ok else " (%s)" % detail))
    for code, fn in RDCASES:
        if not pick(code):
            continue
        total += 1
        rc, out = run(mutate(fn, RDX), tmp, "rdneg_" + code.lower())
        ok = rc != 0 and ("[%s]" % code) in out
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " negative %s (readiness)" % code + ("" if ok else "\n" + out[-600:]))
    for code, fn in RDPCASES:
        if not pick(code):
            continue
        need(RDX, tmp, "readiness_strict")
        total += 1
        dst = os.path.join(tmp, "rdp_%s.pptx" % code.lower())
        pptx_mutation(os.path.join(tmp, "readiness_strict.pptx"), dst, fn)
        l = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), os.path.join(tmp, "readiness_strict.json"), "--pptx", dst], capture_output=True, text=True)
        ok = l.returncode != 0 and ("[%s]" % code) in l.stdout
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " negative %s (built deck)" % code + ("" if ok else "\n" + l.stdout[-500:]))
    if pick("LAYOUT_FIT"):
        total += 1
        adv = mutate(lambda d: d["slides"][7].update(layout="bullets", points=["$2.4M avoided in Q3", "94% complete on first try", "2.1 days to first visit"], tint=False), RDX)
        dj = os.path.join(tmp, "advice.json"); json.dump(adv, open(dj, "w"))
        l = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), dj], capture_output=True, text=True)
        ok = "[LAYOUT_FIT]" in l.stdout and "kpi" in l.stdout
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " LAYOUT_FIT advises a kpi slide for a list of numbers" + ("" if ok else "\n" + l.stdout[-400:]))
    if pick("BUILD"):
        total += 1
        texts = []
        for root, _d, fs in os.walk(ROOT):
            for f in fs:
                if "__pycache__" in root:
                    continue
                try:
                    open(os.path.join(root, f), "rb").read().decode("utf-8")
                except UnicodeDecodeError:
                    texts.append(os.path.relpath(os.path.join(root, f), ROOT))
        passed += not texts; lines.append(("ok  " if not texts else "FAIL") + " every packaged file is UTF-8 text (no binary that an upload could corrupt)" + ("" if not texts else " %s" % texts))
    if pick("BUILD"):
        for label, ok, detail in branded_project(tmp) + font_layout(tmp):
            total += 1
            passed += ok; lines.append(("ok  " if ok else "FAIL") + " " + label + ("" if ok else "\n" + detail[-600:]))
    if pick("EXAMPLES"):
        total += 1
        rc, out = run(GX, tmp, "guide_strict")
        passed += rc == 0; lines.append(("ok  " if rc == 0 else "FAIL") + " guide example builds and lints clean (exit %d)" % rc + ("" if rc == 0 else "\n" + out[-800:]))
    if pick("EXAMPLES"):
        need(GX, tmp, "guide_strict")
        total += 1
        try:
            pg = Presentation(os.path.join(tmp, "guide_strict.pptx"))
            E = 914400.0
            cov = next(sh for sh in pg.slides[0].shapes if sh.name == "title")
            cov_c = abs((cov.left + cov.width / 2) / E - 13.333 / 2) < 0.02 and cov.text_frame.paragraphs[0].alignment == 2
            pins = [sh for sh in pg.slides[3].shapes if sh.name == "callout:marker"]
            letters = sorted(sh.text_frame.text for sh in pins)
            zl = sorted(k.text_frame.text for g in pg.slides[3].shapes if g.name == "anatomy:zone" for k in g.shapes if k.name == "anatomy:num")
            notes = [g for g in pg.slides[6].shapes if g.name == "note"]
            offs = []
            for g in notes:
                m = next(k for k in g.shapes if k.name.endswith(":num")); t = next(k for k in g.shapes if k.name == "qa:q")
                fs = t.text_frame.paragraphs[0].runs[0].font.size.pt
                offs.append(abs((m.top + m.height / 2) / E - (t.top / E + 0.04 + 0.62 * fs / 72)))
            orange = NB["alert"]
            org = [sh.name for sh in pg.slides[5].shapes if orange in sh._element.xml.upper()]
            rows = [k for g in pg.slides[7].shapes if g.name == "contact" for k in g.shapes if k.name == "contact:row"]
            rows_c = all(abs((r.left + r.width / 2) / E - 13.333 / 2) < 0.02 for r in rows) and len(rows) == 2
            ok = cov_c and letters == ["A", "B", "C", "D"] and zl == letters and len(offs) == 3 and max(offs) < 0.01 and org == ["callout:alert"] and rows_c
            detail = "cover %s, pins %s, zones %s, note offs %s, orange %s, rows %s" % (cov_c, letters, zl, [round(o, 3) for o in offs], org, rows_c)
        except Exception as e:
            ok, detail = False, repr(e)
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " centered cover, lettered pins match zone cards, annotation markers on first lines, warning is the only orange, contact rows centered" + ("" if ok else " (%s)" % detail))
    if pick("BUILD"):
        total += 1
        import importlib.util as _u
        spec = _u.spec_from_file_location("ld", os.path.join(HERE, "lint_deck.py")); ld = _u.module_from_spec(spec); spec.loader.exec_module(ld)
        ok = ld.proper("Pull a step-by-step report") == "Pull a Step-by-step Report" and ld.proper("UAT and Sign-Off") == "UAT and Sign-off" and ld.proper("Non-SaaS Jul-Sep Users") == "Non-SaaS Jul-Sep Users"
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " one hyphen rule: lowercase after a hyphen unless a proper noun or acronym (Step-by-step, Sign-off, Non-SaaS, Jul-Sep)")
    if pick("BUILD"):
        need(GX, tmp, "guide_strict")
        total += 1
        try:
            sys.path.insert(0, HERE)
            import iconlib
            bad = []
            for nm in iconlib.names():
                try:
                    iconlib.drawingml(nm, NB["primary"], 0.4)
                except Exception as ex:
                    bad.append(nm)
            pg = Presentation(os.path.join(tmp, "guide_strict.pptx"))
            def _walk(shs):
                for sh in shs:
                    if sh.shape_type == 6:
                        yield from _walk(sh.shapes)
                    else:
                        yield sh
            icons = [sh for sl in pg.slides for sh in _walk(sl.shapes) if sh.name.startswith("icon:")]
            pal = set(c.upper() for c in json.load(open(os.path.join(ROOT, "assets", "theme.json")))["colors"].values())
            okx = all("<a:custGeom" in i._element.xml and "<a:noFill/>" in i._element.xml and all(c.upper() in pal for c in re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', i._element.xml)) for i in icons)
            ok = not bad and len(iconlib.names()) >= 200 and len(icons) >= 12 and okx
            detail = "bad %s, set %d, icons %d, native/palette %s" % (bad[:3], len(iconlib.names()), len(icons), okx)
        except Exception as ex:
            ok, detail = False, repr(ex)
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " every bundled Tabler icon converts; built icons are native vector shapes in palette colors" + ("" if ok else " (%s)" % detail))
    if pick("ICON_OVERLAY"):
        total += 1
        ov = os.path.join(tmp, "ovl"); os.makedirs(os.path.join(ov, "screenshots"), exist_ok=True)
        for f in os.listdir(os.path.join(tmp, "screenshots")):
            shutil.copy(os.path.join(tmp, "screenshots", f), os.path.join(ov, "screenshots", f))
        json.dump({"prefix": "tabler", "icons": {"zz-test-shape": '<path fill="none" stroke="currentColor" stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M4 4h16v16H4z"/>'}}, open(os.path.join(ov, "icons.local.json"), "w"))
        gd = mutate(lambda d: d["slides"][3]["diagram"]["nodes"][0].update(icon="zz-test-shape"), GX)
        json.dump(gd, open(os.path.join(ov, "deck.json"), "w"))
        b = subprocess.run([sys.executable, os.path.join(HERE, "build_deck.py"), os.path.join(ov, "deck.json"), "--out", os.path.join(ov, "o.pptx")], capture_output=True, text=True)
        l = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), os.path.join(ov, "deck.json"), "--pptx", os.path.join(ov, "o.pptx")], capture_output=True, text=True)
        json.dump({"prefix": "tabler", "icons": {"zz-filled": '<path fill="currentColor" d="M4 4h16v16H4z"/>'}}, open(os.path.join(ov, "icons.local.json"), "w"))
        l2 = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), os.path.join(ov, "deck.json")], capture_output=True, text=True)
        ok = b.returncode == 0 and "0 failed" in l.stdout and "[ICON_OVERLAY]" in l2.stdout
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " extra Tabler outline icons load from icons.local.json; filled icons are rejected" + ("" if ok else "\n" + (b.stdout + b.stderr + l.stdout + l2.stdout)[-500:]))
    if pick("BUILD"):
        total += 1
        sj = os.path.join(tmp, "story.json")
        json.dump({"meta": GX["meta"], "slides": [{"title": s.get("title", ""), "message": s["message"], "has": ["screenshot", "numbers"]} for s in GX["slides"] if s.get("message")]}, open(sj, "w"))
        pl = subprocess.run([sys.executable, os.path.join(HERE, "plan_deck.py"), sj], capture_output=True, text=True)
        ok = pl.returncode == 0 and pl.stdout.count("->") == len([s for s in GX["slides"] if s.get("message")])
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " plan_deck proposes layouts from a storyline of messages" + ("" if ok else "\n" + (pl.stdout + pl.stderr)[-300:]))
    for code, fn in ICASES:
        if not pick(code):
            continue
        total += 1
        rc, out = run(mutate(fn, GX), tmp, "ineg_" + code.lower())
        ok = rc != 0 and ("[%s]" % code) in out
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " negative %s (guide)" % code + ("" if ok else "\n" + out[-500:]))
    for code, fn in MICASES:
        if not pick(code):
            continue
        total += 1
        rc, out = run(mutate(fn, MX), tmp, "mineg_" + code.lower())
        ok = rc != 0 and ("[%s]" % code) in out
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " negative %s (minimal)" % code + ("" if ok else "\n" + out[-500:]))
    for code, fn in GCASES:
        if not pick(code):
            continue
        total += 1
        rc, out = run(mutate(fn, GX), tmp, "gneg_" + code.lower())
        ok = rc != 0 and ("[%s]" % code) in out
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " negative %s (guide)" % code + ("" if ok else "\n" + out[-600:]))
    for code, fn in VCASES:
        if not pick(code):
            continue
        total += 1
        rc, out = run(mutate(fn, VX), tmp, "vneg_" + code.lower())
        ok = rc != 0 and ("[%s]" % code) in out
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " negative %s (variety)" % code + ("" if ok else "\n" + out[-600:]))
    for code, fn in ECASES:
        if not pick(code):
            continue
        total += 1
        rc, out = run(mutate(fn), tmp, "eneg_" + code.lower())
        ok = rc != 0 and ("[%s]" % code) in out
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " negative %s" % code + ("" if ok else "\n" + out[-600:]))
    for code, fn in CASES:
        if not pick(code):
            continue
        total += 1
        rc, out = run(mutate(fn), tmp, "neg_" + code.lower())
        ok = rc != 0 and ("[%s]" % code) in out
        passed += ok; lines.append(("ok  " if ok else "FAIL") + " negative %s" % code + ("" if ok else "\n" + out[-600:]))
    for ok, label, detail in v28(tmp) + v29(tmp) + d25(tmp) + vbatch(tmp) + d39(tmp):
        total += 1; passed += ok
        lines.append(("ok  " if ok else "FAIL") + " " + label + ("" if ok else "\n" + detail[-600:]))
    return passed, total, lines


def main():
    global ONLY, SEEN
    ap = argparse.ArgumentParser(description="Self-test for the slide-deck skill.")
    ap.add_argument("--only", help="'list' prints the check codes the cases exercise; CODE[,CODE] runs only cases tagged with them")
    a, _rest = ap.parse_known_args()
    everything = 0
    if a.only is not None:
        SEEN = []
        cases()
        tags, SEEN = SEEN, None
        valid = sorted({c for t in tags for c in t})
        if a.only.strip().lower() == "list":
            print("\n".join(valid))
            sys.exit(0)
        ONLY = {c.strip().upper() for c in a.only.split(",") if c.strip()}
        unknown = sorted(ONLY - set(valid))
        if unknown or not ONLY:
            print("ERROR --only: unknown check code(s) %s. Valid codes: %s" % (", ".join(unknown) or "(none given)", ", ".join(valid)), file=sys.stderr)
            sys.exit(2)
        everything = len(tags)
    passed, total, lines = cases()
    print("\n".join(lines))
    print("SELF-TEST: %d/%d passed" % (passed, total) + (" (--only %s: %d of %d cases run)" % (",".join(sorted(ONLY)), total, everything) if ONLY else ""))
    sys.exit(0 if passed == total else 1)


if __name__ == "__main__":
    main()

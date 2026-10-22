#!/usr/bin/env python3
"""doc.json -> Word (.docx) and PDF in the project's brand, for comprehensive or technical documents.

  python3 scripts/build_doc.py doc.json --out <Name>.docx [--pdf] [--draft] [--brand brand.json]
Runs lint_doc.py on doc.json first and refuses to build on any FAIL. With --pdf, LibreOffice exports the PDF
(tagged, bookmarks from headings, fonts embedded) and a second pass writes the true page numbers into the contents.
The brand (logo, colors, fonts, footer) comes from scripts/brand.py: --brand, meta.brand, a brand.json beside doc.json or
at the project root, or what the project's own files define; with none of those the neutral theme applies and no logo
is drawn. assets/doc_theme.json holds the layout and the fallback colors.
"""
import argparse, io, json, os, pathlib, re, shutil, subprocess, sys, tempfile, time
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import brand as BR  # noqa: E402  (project brand: logo, colors, fonts, footer)
import iconlib  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT, WD_SECTION  # noqa: E402,F401
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
T = json.load(open(os.path.join(ROOT, "assets", "doc_theme.json"), encoding="utf-8"))
C, S, F, P = T["colors"], T["sizes"], T.setdefault("fonts", {}), T["page"]  # F stays empty until use_brand(): no family is hard-coded
TW = P["text_width_in"]
LTW = P["height_in"] - 2 * P["margin_in"]  # landscape text width (D-77)
AUTO_ICONS = {"Glossary": "book", "Revision History": "history", "Sources": "link", "Index": "search"}  # meta.appendix_icons
BOXES = {"objectives": {"title": "What You Will Learn", "kind": "note", "icon": "target"},
         "takeaways": {"title": "Key Takeaways", "kind": "tip", "icon": "bulb"}}
SHEET = json.load(open(os.path.join(ROOT, "assets", "sheet_formats.json"), encoding="utf-8"))  # sheet formats, minimums, sizes
S0 = dict(S)  # theme sizes; Builder.__init__ restores them, then a sheet scales its own copy
DOT = "\u25cb"  # step-diagram hollow dot: this glyph in the body font when its file has it (vector text), else a ring picture
FONT_ROLES = ("body", "heading", "display", "mono")
BRAND = []  # the brand of this build (set by use_brand): fonts to measure, metrics and files to embed
CARD_ACCENTS = ("primary", "dark", "secondary", "alert")  # legacy color names (brand.LEGACY) map to these roles
SHEET_CLASSIFICATION = "Proprietary and confidential. For internal use only."  # sheet footer when meta.classification is unset


def sheet_format(doc_type):
    """Rules for a sheet doc_type: default_format overlaid with its entry in assets/sheet_formats.json."""
    return dict(SHEET["default_format"], **SHEET["formats"].get(doc_type or "", {}))


def sheet_cfg(meta):
    """meta.sheet over the defaults (pages_per_section defaults to the format's pages); the minimums never drop below the floors."""
    cfg = dict(SHEET["defaults"], pages_per_section=sheet_format(meta.get("doc_type")).get("pages", SHEET["defaults"]["pages_per_section"])); cfg.update(meta.get("sheet") if isinstance(meta.get("sheet"), dict) else {})
    for k, v in SHEET["floors"].items():
        try:
            cfg[k] = max(float(cfg[k]), v)
        except (TypeError, ValueError):
            cfg[k] = v
    return cfg


def sheet_sizes(cfg):
    """One text scale for the whole document; never below min_body_pt (body) or min_table_pt (tables, captions)."""
    try:
        sc = float(cfg.get("text_scale") or 1.0)
    except (TypeError, ValueError):
        sc = 1.0
    out = {k: round(v * sc, 1) for k, v in SHEET["base_sizes"].items()}
    out["body"] = max(out["body"], cfg["min_body_pt"]); out["h3"] = max(out["h3"], cfg["min_body_pt"])
    for k in ("table", "small", "caption", "code"):
        out[k] = max(out[k], cfg["min_table_pt"])
    out["h2"] = SHEET["fixed_sizes"]["h2"]
    return out


def crop_png(path, crop=None):
    """(image source, height/width ratio). With crop [x0, y0, x1, y1] (fractions of the image) Pillow crops it first."""
    from PIL import Image
    with Image.open(path) as im:
        if not crop:
            return path, im.size[1] / float(im.size[0])
        W, H = im.size
        box = (int(round(crop[0] * W)), int(round(crop[1] * H)), int(round(crop[2] * W)), int(round(crop[3] * H)))
        part = im.crop(box); buf = io.BytesIO(); part.save(buf, "PNG"); buf.seek(0)
        return buf, part.size[1] / float(part.size[0])


def tiny(p, pt=1):
    """Shrink an empty paragraph's mark so a spacer or page-break paragraph takes almost no height."""
    ppr = p._p.get_or_add_pPr(); rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr"); ppr.append(rpr)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(pt * 2))); rpr.append(sz)
    return p


class Box:
    """A table cell that takes blocks like the document body (half-width subsections render into one)."""
    def __init__(self, cell, doc):
        self.cell, self.doc, self.fresh = cell, doc, True

    def add_paragraph(self, text="", style=None):
        if self.fresh:
            self.fresh = False; p = self.cell.paragraphs[0]
            if style:
                p.style = style
            if text:
                p.add_run(text)
            return p
        return self.cell.add_paragraph(text, style)

    def add_table(self, rows, cols):
        first = self.cell.paragraphs[0]._p if self.fresh else None
        self.fresh = False
        t = self.cell.add_table(rows, cols)
        if first is not None:
            self.cell._tc.remove(first)
        return t

    @property
    def paragraphs(self):
        return self.cell.paragraphs

    def __getattr__(self, k):
        return getattr(self.doc, k)


def lead_blocks(s):
    """A section's blocks with the What You Will Learn box after the opener paragraph (first when there is none)."""
    bl = list(s.get("blocks") or [])
    if not s.get("objectives"):
        return bl
    k = 1 if bl and bl[0].get("type") == "paragraph" else 0
    return bl[:k] + [dict(BOXES["objectives"], type="callout", items=list(s["objectives"]))] + bl[k:]


def tail_blocks(s):
    """The Key Takeaways box, drawn after the section's subsections so it is the chapter's last block."""
    return [dict(BOXES["takeaways"], type="callout", items=list(s["takeaways"]))] if s.get("takeaways") else []


def rgb(name):
    return RGBColor.from_string(C.get(name, name))


def use_brand(b):
    """Overlay the project brand on the theme (colors by role and all four font roles, in place, so C and F follow),
    remember it for measuring and embedding, and return it."""
    b.apply(T)
    F.update({k: b.fonts.get(k) or b.fonts["body"] for k in FONT_ROLES})  # apply() fills only the roles a theme already lists
    BRAND[:] = [b]
    _FONTS.clear()
    return b


def brand():
    return BRAND[0] if BRAND else None


def measure_file(role="body", bold=False):
    """A font file for measuring a role's text: the brand's own file, else the bundled fallback (brand.measure_path)."""
    b = brand()
    return b.measure_path(role, bold) if b else BR.find_font(BR.FALLBACK_FONTS["mono" if role == "mono" else "body"], bold)


def metrics(role="body", bold=False):
    """Vertical metrics (fractions of the size) of the font a role is drawn in: line, ascent, descent, cap, x, first."""
    b = brand()
    return b.metrics(role, bold) if b else BR.font_metrics(measure_file(role, bold) or "")


def has_glyph(path, ch):
    """True when the font file's cmap (format 4 or 12) maps the character; False when unreadable."""
    import struct
    try:
        data = open(path, "rb").read()
        n = struct.unpack(">H", data[4:6])[0]
        tabs = {data[12 + 16 * i:16 + 16 * i]: struct.unpack(">II", data[20 + 16 * i:28 + 16 * i]) for i in range(n)}
        off = tabs[b"cmap"][0]
        cp = ord(ch)
        for i in range(struct.unpack(">H", data[off + 2:off + 4])[0]):
            pid, eid, sub = struct.unpack(">HHI", data[off + 4 + 8 * i:off + 12 + 8 * i])
            s = off + sub
            fmt = struct.unpack(">H", data[s:s + 2])[0]
            if fmt == 4 and pid in (0, 3):
                seg = struct.unpack(">H", data[s + 6:s + 8])[0] // 2
                ends = struct.unpack(">%dH" % seg, data[s + 14:s + 14 + 2 * seg])
                starts = struct.unpack(">%dH" % seg, data[s + 16 + 2 * seg:s + 16 + 4 * seg])
                deltas = struct.unpack(">%dh" % seg, data[s + 16 + 4 * seg:s + 16 + 6 * seg])
                ro_at = s + 16 + 6 * seg
                ros = struct.unpack(">%dH" % seg, data[ro_at:ro_at + 2 * seg])
                for k in range(seg):
                    if starts[k] <= cp <= ends[k]:
                        if ros[k] == 0:
                            return (cp + deltas[k]) & 0xFFFF != 0
                        gi = struct.unpack(">H", data[ro_at + 2 * k + ros[k] + 2 * (cp - starts[k]):][:2])[0]
                        return gi != 0
            elif fmt == 12:
                ng = struct.unpack(">I", data[s + 12:s + 16])[0]
                for k in range(ng):
                    a, z, g = struct.unpack(">III", data[s + 16 + 12 * k:s + 28 + 12 * k])
                    if a <= cp <= z:
                        return g + cp - a != 0
        return False
    except Exception:
        return False


def dot_glyph():
    """(font file, advance em, ring center above the baseline em) when the body font's own file draws U+25CB, else None
    (the step diagram then uses a ring picture). Only a font that is really there counts, never a guess by name."""
    b = brand()
    p = b.font_path("body") if b else None
    if not p or not has_glyph(p, DOT):
        return None
    try:
        from PIL import ImageFont
        f = ImageFont.truetype(p, 200)
        x0, y0, x1, y1 = f.getbbox(DOT, anchor="ls")
        return p, f.getlength(DOT) / 200.0, -(y0 + y1) / 2.0 / 200.0
    except Exception:
        return None


def brand_year(meta):
    """Copyright year: the year in meta.date ("2026-10-11", "October 2026"), else this year."""
    yr = re.search(r"\b(?:19|20)\d{2}\b", str(meta.get("date", "")))
    return yr.group(0) if yr else time.strftime("%Y")


def footer_lead(b, meta, classification):
    """(brand footer, classification) for the page footer: the brand's copyright line (empty when the project names no
    owner) and the document's classification, dropped when the brand footer already says it."""
    lead = b.footer(brand_year(meta)) if b else ""
    cls = str(classification or "").strip()
    if cls and cls.lower().rstrip(". ") in lead.lower():
        cls = ""
    return lead, cls


def add_logo(builder, p, box_w, box_h):
    """The brand logo for a white background, fitted inside box_w x box_h inches; returns the height drawn (0 when the
    brand has no logo: nothing is drawn, never a placeholder)."""
    png = builder.brand.logo_png("light")
    if not png:
        return 0.0
    size = builder.brand.logo_size("light") or (box_w, box_h)
    w = min(box_w, box_h * size[0] / float(size[1] or 1))
    r = p.add_run(); r.add_picture(io.BytesIO(png), width=Inches(w)); builder.alt(r, builder.brand.logo_alt())
    return w * size[1] / float(size[0] or 1)


def set_font(run, name=None, size=None, bold=None, italic=None, color=None):
    f = run.font
    name = name or F["body"]
    f.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color: f.color.rgb = rgb(color)
    return run


def para_fmt(p, before=0, after=6, line=1.15, keep=False, align=None, left=None, first=None):
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    if keep: pf.keep_with_next = True
    if align is not None: p.alignment = align
    if left is not None: pf.left_indent = Inches(left)
    if first is not None: pf.first_line_indent = Inches(first)
    return p


def border(el_pr, sides, color, sz=4, tag="w:pBdr"):
    b = el_pr.find(qn(tag))
    if b is None:
        b = OxmlElement(tag); el_pr.append(b)
    for side in sides:
        e = OxmlElement("w:" + side)
        e.set(qn("w:val"), "single" if color else "nil"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "1"); e.set(qn("w:color"), C.get(color, color or "auto"))
        b.append(e)


def shade(cell, color):
    tcPr = cell._element.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), C.get(color, color)); tcPr.append(sh)


def cell_borders(cell, spec):
    tcPr = cell._element.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        color, sz = spec.get(side, (None, 0))
        e = OxmlElement("w:" + side)
        e.set(qn("w:val"), "single" if color else "nil")
        if color:
            e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), C.get(color, color)); e.set(qn("w:space"), "0")
        b.append(e)
    tcPr.append(b)


def cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._element.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("start", left), ("end", right), ("left", left), ("right", right)):
        e = OxmlElement("w:" + k); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)


def table_plain(t, widths):
    tbl = t._element
    tblPr = tbl.tblPr
    for tag in ("w:tblBorders",):
        old = tblPr.find(qn(tag))
        if old is not None: tblPr.remove(old)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    t.autofit = False
    for row in t.rows:
        for i, c in enumerate(row.cells):
            c.width = Inches(widths[i])
    grid = tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i] * 1440)))


def field(p, code, size=None, color=None, bold=False):
    r = p.add_run(); set_font(r, size=size, color=color, bold=bold)
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r._element.append(b)
    r = p.add_run(); set_font(r, size=size, color=color, bold=bold)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " %s " % code; r._element.append(it)
    r = p.add_run(); set_font(r, size=size, color=color, bold=bold)
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); r._element.append(s)
    r = p.add_run("1"); set_font(r, size=size, color=color, bold=bold)
    r = p.add_run(); set_font(r, size=size, color=color, bold=bold)
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r._element.append(e)


def anchor_name(num):
    """Bookmark name for a heading number ("3.2" -> "_h_3_2"; "Appendix B" -> "_h_Appendix_B")."""
    return "_h_" + re.sub(r"[^A-Za-z0-9]", "_", str(num))


def add_bookmark(p, name, bid):
    s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(bid)); s.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(bid))
    ppr = p._p.pPr
    if ppr is not None:
        ppr.addnext(s)
    else:
        p._p.insert(0, s)
    p._p.append(e)


def link_runs(p, anchor, runs):
    """Wrap runs in an internal hyperlink to a heading bookmark (clickable in Word and in the exported PDF)."""
    h = OxmlElement("w:hyperlink"); h.set(qn("w:anchor"), anchor); h.set(qn("w:history"), "1")
    for r in runs:
        p._p.remove(r._element); h.append(r._element)
    p._p.append(h)
    return h


LEGACY = os.environ.get("DOCBUILDER_LEGACY_LAYOUT") == "1"  # test hook only: reproduces the 1.3.0 split callouts and narrow columns
_FONTS = {}
PAD_IN = 240 / 1440 + 0.06  # left + right cell margins (120 twips each) plus slack


def text_w(s, size, bold=False):
    """Width in inches of s in the body font at size pt, measured with Pillow on the brand's font file (or the bundled
    fallback when the brand font isn't on this machine); a conservative estimate only when Pillow is missing."""
    key = (bool(bold), size)
    if key not in _FONTS:
        f = None
        try:
            from PIL import ImageFont
            p = measure_file("body", bold)
            f = ImageFont.truetype(p, int(size * 4)) if p else None
        except Exception:
            f = None
        _FONTS[key] = f
    if _FONTS[key]:
        return _FONTS[key].getlength(s) / 4 / 72
    return sum(0.72 if c.isupper() else 0.58 if c.isdigit() else 0.3 if c in "il.,;:'!|()" else 0.54 for c in s) * size / 72 * (1.06 if bold else 1.0)


def min_widths(cols, rows):
    """Each column's minimum: its longest unbreakable word at the table font size (header bold), plus cell padding."""
    out = []
    for i, c in enumerate(cols):
        best = max([text_w(w, S["table"], True) for w in str(c).split()] + [0.25])
        for r in rows:
            for part in INLINE.split(str(r[i]) if i < len(r) else ""):
                if not part:
                    continue
                bold = part.startswith("**")
                txt = re.sub(r"^\*\*|\*\*$|^`|`$", "", re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", part))
                best = max([best] + [text_w(w, S["table"], bold) for w in txt.split()])
        out.append(best + PAD_IN)
    return out


def fit_widths(widths, need):
    """Widen any column narrower than its longest word, taking the space from columns that have room."""
    w = list(widths)
    if sum(need) >= TW:
        return [TW * n / sum(need) for n in need]
    for _ in range(len(w)):
        short = [i for i in range(len(w)) if w[i] < need[i] - 1e-6]
        if not short:
            break
        for i in short:
            w[i] = need[i]
        rest = [i for i in range(len(w)) if w[i] > need[i] + 1e-6]
        over = sum(w) - TW
        room = sum(w[i] - need[i] for i in rest)
        if not rest or room <= 0:
            break
        for i in rest:
            w[i] -= over * (w[i] - need[i]) / room
    return w


def auto_widths(cols, rows):
    """Share the text width by content length, but never narrower than the longest unbreakable word."""
    def words(s):
        return re.findall(r"\S+", re.sub(r"[*`]", "", str(s)))
    if not LEGACY:
        lens = [min(max(max([len(str(c))] + [len(str(r[i])) for r in rows]), 6), 60) for i, c in enumerate(cols)]
        return fit_widths([TW * x / sum(lens) for x in lens], min_widths(cols, rows))
    need = [max([len(w) for x in [c] + [r[i] for r in rows] for w in words(x)] + [4]) * 0.075 + 0.2 for i, c in enumerate(cols)]
    lens = [min(max(max([len(str(c))] + [len(str(r[i])) for r in rows]), 6), 60) for i, c in enumerate(cols)]
    w = [TW * x / sum(lens) for x in lens]
    for _ in range(len(cols)):
        short = [i for i in range(len(cols)) if w[i] < need[i] - 1e-6]
        if not short:
            break
        for i in short:
            w[i] = need[i]
        rest = [i for i in range(len(cols)) if w[i] > need[i] + 1e-6]
        if not rest:
            break
        spare = TW - sum(w); tot = sum(w[i] for i in rest)
        for i in rest:
            w[i] += spare * w[i] / tot
    s = sum(w)
    return [x * TW / s for x in w]


SRC_LINK = re.compile(r"\[([^\]]+)\]\(((?:https?://|mailto:)[^)\s]+)\)|(https?://[^\s;,)\]]*[^\s;,)\].:])")  # D-27 sheet footer links


def source_links(text):
    """D-27: links in one source line, [(title, url)]: each [title](url) and each bare http(s) URL (shown as itself)."""
    return [(m.group(1) or m.group(3), m.group(2) or m.group(3)) for m in SRC_LINK.finditer(str(text))]


def ext_link(part, p, url, text, size, color="primary"):
    """External hyperlink whose relationship lives on `part` (the body, or a footer for footer links): one underlined run."""
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = p.add_run(text); set_font(r, size=size, color=color); r.font.underline = True
    p._p.remove(r._element); h.append(r._element); p._p.append(h)


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def rich(doc, p, text, size=None, color="text", bold=False, italic=False):
    for part in INLINE.split(text or ""):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_font(p.add_run(part[2:-2]), size=size or S["body"], bold=True, italic=italic, color="dark" if color == "text" else color)
        elif part.startswith("`") and part.endswith("`"):
            set_font(p.add_run(part[1:-1]), name=F["mono"], size=(size or S["body"]) - 1, color=color)
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            rid = doc.part.relate_to(m.group(2), "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
            r = p.add_run(m.group(1)); set_font(r, size=size or S["body"], color="primary"); r.font.underline = True
            p._p.remove(r._element); h.append(r._element); p._p.append(h)
        else:
            set_font(p.add_run(part), size=size or S["body"], bold=bold, italic=italic, color=color)


class Builder:
    def __init__(self, doc_json, base, draft=False, pages=None, brand=None):
        self.d, self.base, self.draft, self.pages = doc_json, base, draft, pages or {}
        self.m = doc_json["meta"]
        self.brand = use_brand(brand or BR.load(project=base))
        self.doc = Document()
        self.fig = self.tab = 0
        self.bid = 0
        self.levels = max(1, min(3, int(self.m.get("toc_levels", 2))))
        self.branded = self.m.get("design") == "branded"
        self.heads = []
        self.icons_png = {}
        self.callouts, self.table_words = [], set()
        self.captions = []
        self.tw = TW  # width the blocks fill: the text width, a half-width cell on a sheet, or LTW in a landscape section
        self.landscape = False
        self.sheet = self.m.get("layout") == "sheet"
        self.icons_hi = {}
        S.clear(); S.update(S0)
        if self.sheet:
            self.cfg, self.fmt = sheet_cfg(self.m), sheet_format(self.m.get("doc_type"))
            S.update(sheet_sizes(self.cfg))

    def styles(self):
        st = self.doc.styles
        n = st["Normal"]; n.font.name = F["body"]; n.font.size = Pt(S["body"]); n.font.color.rgb = rgb("text")
        rpr = n.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), F["body"])
        lang = T.get("a11y", {}).get("lang", "en-US")  # D-64: document language on the defaults and Normal (DOC_A11Y)
        dd = st.element.find(qn("w:docDefaults")); rd = dd.find(qn("w:rPrDefault")) if dd is not None else None
        for rp in [rpr] + ([rd.find(qn("w:rPr"))] if rd is not None and rd.find(qn("w:rPr")) is not None else []):
            lg = rp.find(qn("w:lang"))
            if lg is None:
                lg = OxmlElement("w:lang"); rp.append(lg)
            lg.set(qn("w:val"), lang); lg.set(qn("w:eastAsia"), lang)
            df = rp.find(qn("w:rFonts"))  # document defaults name the brand body font, not the template's theme font
            if df is None:
                df = OxmlElement("w:rFonts"); rp.insert(0, df)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                df.attrib.pop(qn(a), None)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                df.set(qn(a), F["body"])
        for lvl, size, color in ((1, S["h1"], "dark"), (2, S["h2"], "primary"), (3, S["h3"], "dark")):
            h = st["Heading %d" % lvl]
            h.font.name, h.font.size, h.font.bold, h.font.italic = F["heading"], Pt(size), True, False
            h.font.color.rgb = rgb(color)
            hr = h.element.get_or_add_rPr(); hf = hr.find(qn("w:rFonts"))
            if hf is None:
                hf = OxmlElement("w:rFonts"); hr.insert(0, hf)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                hf.set(qn(a), F["heading"])
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if hf.get(qn(a)) is not None:
                    del hf.attrib[qn(a)]
            pf = h.paragraph_format
            pf.space_before, pf.space_after, pf.keep_with_next = Pt({1: 0, 2: 16, 3: 12}[lvl]), Pt({1: 12, 2: 6, 3: 4}[lvl]), True
        for i in (1, 2, 3):
            name = "TOC %d" % i
            s = st[name] if name in [x.name for x in st] else st.add_style(name, 1)
            s.base_style = st["Normal"]
            s.font.size = Pt(S["body"] if i == 1 else S["body"] - 0.5); s.font.bold = i == 1
            s.font.color.rgb = rgb("dark" if i == 1 else "text")
            s.paragraph_format.left_indent = Inches({1: 0, 2: 0.35, 3: 0.7}[i])
            s.paragraph_format.space_before = Pt(8 if i == 1 else 2); s.paragraph_format.space_after = Pt(2)
            s.paragraph_format.tab_stops.add_tab_stop(Inches(TW), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        if self.sheet:
            for lvl, before, after in ((2, 8, 3), (3, 6, 2)):
                pf = st["Heading %d" % lvl].paragraph_format
                pf.space_before, pf.space_after = Pt(before), Pt(after)

    def page(self):
        s = self.doc.sections[0]
        s.page_width, s.page_height = Inches(P["width_in"]), Inches(P["height_in"])
        if self.sheet:
            return self.sheet_page(s)
        for k in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, k, Inches(P["margin_in"]))
        s.header_distance = s.footer_distance = Inches(0.45)
        s.different_first_page_header_footer = True
        para_fmt(s.header.paragraphs[0], after=0)  # no page header: the header stays empty on every page
        fp = s.footer.paragraphs[0]; para_fmt(fp, after=0)
        fp.paragraph_format.tab_stops.add_tab_stop(Inches(TW / 2), WD_TAB_ALIGNMENT.CLEAR)
        fp.paragraph_format.tab_stops.add_tab_stop(Inches(TW), WD_TAB_ALIGNMENT.RIGHT)
        lead, cls = footer_lead(self.brand, self.m, self.m.get("classification", "Internal use only"))
        rest = " · ".join(x for x in (cls, "Version %s" % self.m["version"]) if x)
        if lead and self.branded:  # branded: the brand's footer line in bold dark, the rest muted
            set_font(fp.add_run(lead), size=S["small"] - 1, bold=True, color="dark")
            set_font(fp.add_run(" · " + rest), size=S["small"] - 1, color="muted_text")
        else:
            set_font(fp.add_run(" · ".join(x for x in (lead, rest) if x)), size=S["small"] - 1, color="muted_text")
        set_font(fp.add_run("\tPage "), size=S["small"] - 1, color="muted_text")
        field(fp, "PAGE", S["small"] - 1, "muted_text")
        set_font(fp.add_run(" of "), size=S["small"] - 1, color="muted_text")
        field(fp, "NUMPAGES", S["small"] - 1, "muted_text")

    def sheet_page(self, s):
        """Sheet page: shorter top and bottom margins and the house footer on every page (no cover, no first-page exception)."""
        m, fs = self.m, SHEET["fixed_sizes"]["footer"]
        s.left_margin = s.right_margin = Inches(P["margin_in"])
        s.top_margin, s.bottom_margin = Inches(0.6), Inches(0.9)
        s.header_distance, s.footer_distance = Inches(0.3), Inches(0.4)
        s.different_first_page_header_footer = False
        para_fmt(s.header.paragraphs[0], after=0)
        fp = para_fmt(s.footer.paragraphs[0], after=1, line=1.0)
        lead, cls = footer_lead(self.brand, m, m.get("classification") or SHEET_CLASSIFICATION)
        first = " ".join(x for x in (lead, cls) if x)  # "<brand footer> <classification>", no copyright when the brand has none
        set_font(fp.add_run(" \u00b7 ".join(x for x in (first, "%s v%s" % (m["title"], m["version"]), "Revised %s" % m["date"]) if x)),
                 size=fs, color="muted_text")
        rv, R = para_fmt(s.footer.add_paragraph(), after=1, line=1.0), T.get("review") or {}  # D-65: printed copies line, latest version link, review fields
        set_font(rv.add_run(R.get("printed", "Printed copies may be out of date.")), size=fs, color="muted_text")
        if str(m.get("latest_url") or "").startswith("https://"):
            set_font(rv.add_run(" " + R.get("latest_label", "Latest version: ")), size=fs, color="muted_text")
            ext_link(s.footer.part, rv, m["latest_url"], m["latest_url"], fs)
        for v, lab in self.review_rows():
            set_font(rv.add_run(" \u00b7 %s: %s" % (lab, v)), size=fs, color="muted_text")
        q = para_fmt(s.footer.add_paragraph(), after=0, line=1.0)
        q.paragraph_format.tab_stops.add_tab_stop(Inches(TW), WD_TAB_ALIGNMENT.RIGHT)
        plain = lambda x: re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", str(x)).replace("**", "").replace("`", "")
        linked = [x for x in m.get("sources") or [] if re.search(r"https?://|\[[^\]]+\]\([^)]+\)", str(x))]  # readers only see sources they can open
        src = "Sources: " + "; ".join(plain(x) for x in linked)
        self.sources_in_footer = bool(linked) and text_w(src, fs) <= TW - text_w("Page 99 of 99", fs) - 0.3
        self.sources_on_page = len(linked) if self.sources_in_footer else 0
        self.footer_links = []
        if self.sources_in_footer:  # D-27: each [title](url) or bare URL is a real hyperlink (footer relationship), so the PDF has a link annotation
            set_font(q.add_run("Sources: "), size=fs, color="muted_text")
            for i, x in enumerate(str(x) for x in linked):
                if i:
                    set_font(q.add_run("; "), size=fs, color="muted_text")
                at = 0
                for mt in SRC_LINK.finditer(x):
                    if plain(x[at:mt.start()]):
                        set_font(q.add_run(plain(x[at:mt.start()])), size=fs, color="muted_text")
                    title, url = source_links(mt.group(0))[0]
                    ext_link(s.footer.part, q, url, plain(title), fs); self.footer_links.append(url); at = mt.end()
                if plain(x[at:]):
                    set_font(q.add_run(plain(x[at:])), size=fs, color="muted_text")
        set_font(q.add_run("\tPage "), size=fs, color="muted_text"); field(q, "PAGE", fs, "muted_text")
        set_font(q.add_run(" of "), size=fs, color="muted_text"); field(q, "NUMPAGES", fs, "muted_text")

    def review_rows(self):
        """D-65: [(value, label)] for the meta review fields that are set (doc_theme.json review.fields)."""
        return [(str(self.m[k]), lab) for k, lab in (T.get("review") or {}).get("fields", []) if self.m.get(k)]

    def band(self, title, first, anchor):
        """Sheet title band: the brand logo (when there is one) and format label on white, then a dark band (meta.title,
        the section title as Heading 1, "Version X · Date") closed by a secondary-color rule. Every section after the
        first starts a new page."""
        doc, m, fx = self.doc, self.m, SHEET["fixed_sizes"]
        if not first:
            p = para_fmt(doc.add_paragraph(), after=0, line=1.0); p.paragraph_format.page_break_before = True
            tiny(p)
        t = doc.add_table(rows=2, cols=2); table_plain(t, [2.0, self.tw - 2.0])
        for c in t.rows[0].cells:
            cell_margins(c, 0, 90, 0, 0); cell_borders(c, {})
        p = para_fmt(t.cell(0, 0).paragraphs[0], after=0, line=1.0)
        add_logo(self, p, 1.25, 0.47)
        rc = t.cell(0, 1); rc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        p = para_fmt(rc.paragraphs[0], after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.RIGHT)
        label = m.get("doc_label") or self.fmt.get("label") or str(m["doc_type"]).replace("-", " ").title()
        set_font(p.add_run(label), size=fx["band_doc"], bold=True, color="primary")
        if self.draft:
            set_font(p.add_run(" \u00b7 DRAFT"), size=fx["band_doc"], bold=True, color="alert_text")
        c = t.cell(1, 0).merge(t.cell(1, 1)); shade(c, "dark"); cell_margins(c, 150, 150, 200, 200)
        cell_borders(c, {"bottom": ("secondary", 36)})
        p = para_fmt(c.paragraphs[0], after=2, line=1.0)
        set_font(p.add_run(m["title"]), size=fx["band_doc"], bold=True, color="context")
        h = para_fmt(c.add_paragraph(style="Heading 1"), after=3, line=1.0)
        set_font(h.add_run(title), name=F["heading"], size=fx["band_title"], bold=True, color="white")
        self.bid += 1; add_bookmark(h, anchor, self.bid)
        if m.get("subtitle"):  # D-26: subtitle under the section title in the context tint, smaller than the title; nothing when absent
            p = para_fmt(c.add_paragraph(), after=3, line=1.0)
            sub = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", str(m["subtitle"])).replace("**", "").replace("`", "")
            set_font(p.add_run(sub), size=fx["band_sub"], color="context")
        p = para_fmt(c.add_paragraph(), after=0, line=1.0)
        set_font(p.add_run("Version %s \u00b7 %s" % (m["version"], m["date"])), size=fx["band_meta"], color="white")
        tiny(para_fmt(doc.add_paragraph(), after=4, line=1.0), 4)
        self.heads.append((1, title))

    def sheet_body(self):
        """Each top-level section is a page: band, its own blocks, then level-2 subsections (half-width pairs side by side)."""
        for i, s in enumerate(self.d["sections"], 1):
            self.band(s["title"], i == 1, anchor_name(i))
            self.blocks(lead_blocks(s))
            subs, j = s.get("sections") or [], 0
            while j < len(subs):
                if subs[j].get("width") == "half" and j + 1 < len(subs) and subs[j + 1].get("width") == "half":
                    self.half_pair(subs[j], subs[j + 1]); j += 2
                else:
                    self.sheet_sub(subs[j], 2); j += 1
            self.blocks(tail_blocks(s))

    def sheet_sub(self, s, lvl):
        self.heading(lvl, s["title"], icon=s.get("icon"))
        self.blocks(lead_blocks(s))
        for c in s.get("sections") or []:
            self.sheet_sub(c, min(lvl + 1, 3))
        self.blocks(tail_blocks(s))

    def half_pair(self, a, b):
        """Two half-width subsections side by side: a borderless two-column table, each cell filled like the body."""
        doc, tw, gap = self.doc, self.tw, 0.25
        t = doc.add_table(rows=1, cols=2); table_plain(t, [tw / 2, tw / 2])
        cells = t.rows[0].cells
        cell_margins(cells[0], 0, 0, 108, int(gap / 2 * 1440)); cell_margins(cells[1], 0, 0, int(gap / 2 * 1440), 108)  # 108: text lines up with the body margin
        for c in cells:
            cell_borders(c, {})
        saved = (self.doc, self.tw)
        for c, s in zip(cells, (a, b)):
            self.doc, self.tw = Box(c, saved[0]), tw / 2 - gap / 2 - 0.075
            self.sheet_sub(s, 2)
        self.doc, self.tw = saved
        tiny(para_fmt(doc.add_paragraph(), after=0, line=1.0), 2)

    def icon_hi(self, name, color):
        key = (name, color)
        if key not in self.icons_hi:
            self.icons_hi[key] = iconlib.png(name, C[color], 256)  # high resolution for diagrams
        return self.icons_hi[key]

    def ring_png(self):
        """Hollow ring in the primary color (transparent PNG): the step-diagram dot when the body font has no U+25CB."""
        if "ring" not in self.icons_hi:
            from PIL import Image, ImageDraw
            px, c = 256, RGBColor.from_string(C["primary"])
            im = Image.new("RGBA", (px, px), (0, 0, 0, 0))
            ImageDraw.Draw(im).ellipse([12, 12, px - 12, px - 12], outline=(c[0], c[1], c[2], 255), width=18)
            buf = io.BytesIO(); im.save(buf, "PNG"); self.icons_hi["ring"] = buf.getvalue()
        return self.icons_hi["ring"]

    def step_diagram(self, b):
        """Icons above hollow dots on one horizontal line, labels below, equal spacing. The dots are the body font's own
        U+25CB glyph when its file has one (vector text), else a ring picture; the line is cell borders. The dot rows
        are sized from the font's metrics (two rows = one line of the dot size) and the ring is lifted so its measured
        center sits on the line, whatever the font. Numerals only with emphasize_sequence. Labels shrink in 0.5 pt
        steps (not below the table minimum) until their longest word fits the slot, so no word wraps."""
        doc, steps, tw = self.doc, b["steps"], self.tw
        n, seq, dot = len(steps), bool(b.get("emphasize_sequence")), 18
        fm, glyph = metrics("body"), dot_glyph()
        ring = 0.7 * dot  # picture ring diameter in pt (about the size of a U+25CB glyph)
        adv_pt = glyph[1] * dot if glyph else ring
        center = glyph[2] * dot if glyph else ring / 2  # ring center above the baseline, pt
        lift = (fm["ascent"] * dot - center) - fm["line"] * dot / 2  # baseline sits ascent below the line top; the rule is mid-line
        col = tw / n; dw = adv_pt / 72 + 0.02; sw = (col - dw) / 2  # the dot column is the ring's own advance, so it is never clipped
        self.dot_rows = fm["line"] * dot / 2
        floor = self.cfg["min_table_pt"] if self.sheet else S0["small"]
        words = [w for st in steps for w in str(st["label"]).split()] + (["%d" % n] if seq else [])
        size = S["body"]
        while size - 0.5 >= floor and max(text_w(w, size) for w in words) > col - 0.15:
            size -= 0.5
        t = doc.add_table(rows=4, cols=3 * n); table_plain(t, [sw, dw, sw] * n); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if i in (1, 2):
                row.height, row.height_rule = Pt(self.dot_rows), WD_ROW_HEIGHT_RULE.EXACTLY
        for j, st in enumerate(steps):
            a, mid, z = 3 * j, 3 * j + 1, 3 * j + 2
            for i in (1, 2):
                for k in (a, z):
                    c = t.cell(i, k); cell_margins(c, 0, 0, 0, 0)
                    edge = (j == 0 and k == a) or (j == n - 1 and k == z)
                    cell_borders(c, {"bottom": ("context", 12)} if i == 1 and not edge else {})
            ic, lb, dc = t.cell(0, a).merge(t.cell(0, z)), t.cell(3, a).merge(t.cell(3, z)), t.cell(1, mid).merge(t.cell(2, mid))
            cell_margins(ic, 0, 50, 40, 40); cell_margins(lb, 50, 0, 40, 40); cell_margins(dc, 0, 0, 0, 0)
            for c in (ic, lb, dc):
                cell_borders(c, {})
            p = para_fmt(ic.paragraphs[0], after=0, line=1.0, keep=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            r = p.add_run(); r.add_picture(io.BytesIO(self.icon_hi(st["icon"], "primary")), width=Inches(0.34)); self.alt(r, "%s icon" % st["label"])
            dc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = para_fmt(dc.paragraphs[0], after=0, line=1.0, keep=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            if glyph:
                r = set_font(p.add_run(DOT), name=F["body"], size=dot, color="primary")
            else:
                r = set_font(p.add_run(), size=dot)
                r.add_picture(io.BytesIO(self.ring_png()), width=Pt(ring)); self.alt(r, "Step %d marker" % (j + 1))
            rpr = r._element.get_or_add_rPr(); pos = OxmlElement("w:position"); pos.set(qn("w:val"), str(int(round(2 * lift))))  # half-points: ring center onto the line
            sz = rpr.find(qn("w:sz"))
            sz.addprevious(pos) if sz is not None else rpr.append(pos)
            p = para_fmt(lb.paragraphs[0], after=0, line=1.05, align=WD_ALIGN_PARAGRAPH.CENTER)
            if seq:
                set_font(p.add_run("%d  " % (j + 1)), size=size, bold=True, color="primary")
            set_font(p.add_run(st["label"]), size=size, color="dark")
        tiny(para_fmt(doc.add_paragraph(), after=6, line=1.0), 4)

    def card(self, b):
        """A box with a 4.5 pt left accent bar (brand primary by default) on a light tint; one row, so it never splits."""
        doc, accent = self.doc, BR.LEGACY.get(b.get("accent", "primary"), b.get("accent", "primary"))
        t = doc.add_table(rows=1, cols=1); table_plain(t, [self.tw]); c = t.cell(0, 0)
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        shade(c, "card"); cell_margins(c, 110, 110, 180, 160); cell_borders(c, {"left": (accent, 36)})
        first = [c.paragraphs[0]]
        nxt = lambda: first.pop() if first else c.add_paragraph()
        tone = "alert_text" if accent == "alert" else "dark"  # D-64: title and bullets at 4.5:1 on the card tint; the bar keeps the accent
        if b.get("title"):
            set_font(para_fmt(nxt(), after=2, keep=True).add_run(b["title"]), size=S["body"], bold=True, color=tone)
        if b.get("text"):
            rich(doc, para_fmt(nxt(), after=2), b["text"], size=S["body"])
        for it in b.get("items") or []:
            q = para_fmt(nxt(), after=1, left=0.22, first=-0.18)
            set_font(q.add_run("\u2022  "), size=S["body"], color=tone); rich(doc, q, it, size=S["body"])
        tiny(para_fmt(doc.add_paragraph(), after=6, line=1.0), 4)

    def cover_branded(self):
        """Branded cover: the brand logo on white (when there is one), a dark title band, a secondary-color rule and a
        metadata table with tinted labels. Without a logo the band keeps its place on the page."""
        doc = self.doc
        p = para_fmt(doc.add_paragraph(), after=0)
        h = add_logo(self, p, 1.9, 0.71)
        para_fmt(doc.add_paragraph(), after=0).paragraph_format.space_before = Pt(70 + max(0.0, 0.71 - h) * 72)
        band = doc.add_table(rows=1, cols=1); table_plain(band, [TW]); c = band.cell(0, 0)
        shade(c, "dark"); cell_margins(c, 360, 360, 300, 300)
        p = para_fmt(c.paragraphs[0], after=8)
        set_font(p.add_run(self.m.get("doc_label") or self.m["doc_type"].replace("-", " ").title()), size=S["small"] + 1, bold=True, color="context")
        p = para_fmt(c.add_paragraph(), after=8, line=1.0)
        set_font(p.add_run(self.m["title"]), name=F["heading"], size=S["title"], bold=True, color="white")
        if self.m.get("subtitle"):
            p = para_fmt(c.add_paragraph(), after=0, line=1.15)
            set_font(p.add_run(self.m["subtitle"]), size=S["subtitle"], color="white")
        p = para_fmt(doc.add_paragraph(), after=20)
        set_font(p.add_run(), size=2); border(p._p.get_or_add_pPr(), ["top"], "secondary", 36)
        rows = [("Version", self.m["version"]), ("Date", self.m["date"]), ("Owner", self.m["owner"]), ("Audience", self.m["audience"]),
                ("Classification", self.m.get("classification", "Internal use only"))] + ([("Origination Date", self.m["origination_date"])] if self.m.get("origination_date") else [])
        t = doc.add_table(rows=len(rows), cols=2); table_plain(t, [1.6, TW - 1.6])
        for i, (k, v) in enumerate(rows):
            for j, txt in enumerate((k, v)):
                cc = t.cell(i, j); cell_margins(cc, 80, 80, 120, 120)
                cell_borders(cc, {"bottom": ("white" if j == 0 else "card_line", 6)})
                if j == 0: shade(cc, "card")
                pp = para_fmt(cc.paragraphs[0], after=0)
                set_font(pp.add_run(txt), size=S["body"] - 0.5, bold=j == 0, color="dark" if j == 0 else "text")
        if self.draft:
            p = para_fmt(doc.add_paragraph(), before=18)
            set_font(p.add_run("DRAFT: screenshots or facts still open; not for distribution."), size=S["small"], bold=True, color="alert_text")
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def cover(self):
        if self.branded:
            return self.cover_branded()
        doc = self.doc
        p = para_fmt(doc.add_paragraph(), after=0)
        h = add_logo(self, p, 1.7, 0.64)  # the title sits at the same height with or without a logo
        para_fmt(doc.add_paragraph(), after=0).paragraph_format.space_before = Pt(150 + max(0.0, 0.64 - h) * 72)
        p = para_fmt(doc.add_paragraph(), after=4)
        set_font(p.add_run(self.m.get("doc_label") or self.m["doc_type"].replace("-", " ").title().replace("Sop", "Standard Operating Procedure")), size=S["small"] + 1, bold=True, color="primary")
        p = para_fmt(doc.add_paragraph(), after=6, line=1.0)
        words = self.m["title"].split()
        title = " ".join(words[:-2] + ["\u00a0".join(words[-2:])]) if len(words) >= 3 else self.m["title"]
        set_font(p.add_run(title), name=F["heading"], size=S["title"], bold=True, color="dark")
        if self.m.get("subtitle"):
            p = para_fmt(doc.add_paragraph(), after=0, line=1.1)
            set_font(p.add_run(self.m["subtitle"]), size=S["subtitle"], color="text")
        p = para_fmt(doc.add_paragraph(), before=18, after=24)
        r = p.add_run(); set_font(r, size=2)
        border(p._p.get_or_add_pPr(), ["bottom"], "secondary", 18)
        p.paragraph_format.right_indent = Inches(TW - 1.0)
        rows = [("Version", self.m["version"]), ("Date", self.m["date"]), ("Owner", self.m["owner"]), ("Audience", self.m["audience"]),
                ("Classification", self.m.get("classification", "Internal use only"))] + ([("Origination Date", self.m["origination_date"])] if self.m.get("origination_date") else [])
        t = doc.add_table(rows=len(rows), cols=2); table_plain(t, [1.5, TW - 1.5])
        for i, (k, v) in enumerate(rows):
            for j, txt in enumerate((k, v)):
                c = t.cell(i, j); cell_margins(c, 60, 60, 0, 120)
                cell_borders(c, {"bottom": ("card_line", 4)})
                pp = para_fmt(c.paragraphs[0], after=0)
                set_font(pp.add_run(txt), size=S["body"] - 0.5, bold=j == 0, color="dark" if j == 0 else "text")
        if self.draft:
            p = para_fmt(doc.add_paragraph(), before=18)
            set_font(p.add_run("DRAFT: screenshots or facts still open; not for distribution."), size=S["small"], bold=True, color="alert_text")
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def contents(self):
        doc = self.doc
        p = para_fmt(doc.add_paragraph(), after=12)
        set_font(p.add_run("Contents"), name=F["heading"], size=S["h1"], bold=True, color="dark")
        allentries = self.entries()
        entries = [e for e in allentries if e[0] <= self.levels]
        first = True
        for k, (lvl, num, title, label) in enumerate(entries):
            p = doc.add_paragraph(style="TOC %d" % lvl)
            if first:
                r = p.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r._element.append(fc)
                r = p.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = ' TOC \\o "1-%d" \\h \\z \\u ' % self.levels; r._element.append(it)
                r = p.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r._element.append(fc)
                first = False
            runs = [set_font(p.add_run(label), size=S["body"] if lvl == 1 else S["body"] - 0.5, bold=lvl == 1, color="dark" if lvl == 1 else "text"),
                    set_font(p.add_run("\t%s" % self.pages.get(label, "")), size=S["body"] - 0.5, bold=False, color="text")]
            link_runs(p, anchor_name(num), runs)
            if k == len(entries) - 1:
                r = p.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r._element.append(fc)
        self.toc_labels = [e[3] for e in entries]
        self.all_labels = [e[3] for e in allentries]

    def entries(self):
        """Every heading in reading order: (level, number, title, label). Labels match the built headings."""
        out = []
        def walk(secs, prefix, lvl):
            for i, s in enumerate(secs, 1):
                num = "%s%d" % (prefix, i)
                out.append((lvl, num, s["title"], "%s  %s" % (num, s["title"])))
                walk(s.get("sections", []), num + ".", lvl + 1)
        walk(self.d["sections"], "", 1)
        for letter, title in self.appendices():
            out.append((1, "Appendix " + letter, title, "Appendix %s: %s" % (letter, title)))
        return out

    def resolve(self, ref):
        ref = str(ref).strip()
        for lvl, num, title, label in self.entries():
            if ref == num or ref.lower() == title.lower() or ref == label:
                return num, label
        return None, None

    def index(self):
        """Two columns (continuous section break) so 30-40 entries fit on one page."""
        doc = self.doc
        sec = doc.add_section(WD_SECTION.CONTINUOUS)
        cols = sec._sectPr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols"); sec._sectPr.append(cols)
        cols.set(qn("w:num"), "2"); cols.set(qn("w:space"), "432")
        colw = (TW - 0.3) / 2
        letter = None
        for e in sorted(self.d["index"], key=lambda x: x["term"].lower()):
            first = e["term"][0].upper()
            if first != letter:
                letter = first
                p = para_fmt(doc.add_paragraph(), before=8, after=2, keep=True)
                set_font(p.add_run(first if first.isalpha() else "#"), name=F["heading"], size=S["h3"], bold=True, color="primary")
            num, label = self.resolve(e["section"])
            p = para_fmt(doc.add_paragraph(), after=1, left=0.2)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(colw - 0.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            runs = [set_font(p.add_run(e["term"]), size=S["body"] - 0.5, color="text"),
                    set_font(p.add_run("  (%s)" % num), size=S["small"], color="muted_text"),
                    set_font(p.add_run("\t%s" % self.pages.get(label, "")), size=S["body"] - 0.5, color="text")]
            link_runs(p, anchor_name(num), runs)

    def appendices(self):
        out = [a["title"] for a in self.d.get("appendices") or []]
        if self.d.get("glossary"): out.append("Glossary")
        if self.d.get("revisions"): out.append("Revision History")
        if self.m.get("sources"): out.append("Sources")
        if self.d.get("index"): out.append("Index")
        return [(chr(65 + i), t) for i, t in enumerate(out)]

    def heading(self, lvl, label, new_page=False, anchor=None, icon=None):
        p = self.doc.add_paragraph(style="Heading %d" % lvl)
        if new_page:
            p.paragraph_format.page_break_before = True
        if icon:
            r = p.add_run(); r.add_picture(io.BytesIO(self.icon_png(icon, "dark" if lvl != 2 else "primary")), width=Inches({1: 0.3, 2: 0.22, 3: 0.2}[lvl]))
            self.alt(r, "Section icon"); set_font(p.add_run("  "), name=F["heading"], size=S["h%d" % lvl])
        set_font(p.add_run(label), name=F["heading"], size=S["h%d" % lvl], bold=True, color="dark" if lvl != 2 else "primary")
        if self.branded and lvl == 1:
            border(p._p.get_or_add_pPr(), ["bottom"], "secondary", 12)
            p.paragraph_format.space_after = Pt(14)
        if anchor:
            self.bid += 1; add_bookmark(p, anchor, self.bid)
        self.heads.append((lvl, label))

    def alt(self, run, text):
        for dp in run._element.iter(qn("wp:docPr")):
            dp.set("descr", text); dp.set("title", text[:60])

    def image(self, path, width, alt, para=None, indent=0, crop=None):
        p = para or para_fmt(self.doc.add_paragraph(), before=6, after=4, keep=True, align=WD_ALIGN_PARAGRAPH.CENTER if not indent else None, left=indent or None)
        full = os.path.join(self.base, path)
        if os.path.isfile(full):
            src, ratio = crop_png(full, crop)
            w = min(width, self.tw - indent, 7.6 / ratio * 0.62)
            r = p.add_run(); r.add_picture(src, width=Inches(w)); self.alt(r, alt)
            for pic in r._element.iter(qn("pic:spPr")):
                ln = OxmlElement("a:ln"); ln.set("w", "6350"); sf = OxmlElement("a:solidFill"); clr = OxmlElement("a:srgbClr"); clr.set("val", C["card_line"]); sf.append(clr); ln.append(sf); pic.append(ln)
        else:
            set_font(p.add_run("[Screenshot needed: %s]" % path), size=S["small"], bold=True, color="alert_text")
        return p

    def caption(self, kind, text, indent=0, above=False):
        if kind == "Figure":
            self.fig += 1; n = self.fig
        else:
            self.tab += 1; n = self.tab
        p = para_fmt(self.doc.add_paragraph(), before=2 if not above else 8, after=10 if not above else 4, keep=above,
                     align=WD_ALIGN_PARAGRAPH.CENTER if (kind == "Figure" and not indent) else None, left=indent or None)
        set_font(p.add_run("%s %d. " % (kind, n)), size=S["caption"], bold=True, color="dark")
        self.captions.append([kind, n, text])
        set_font(p.add_run(text), size=S["caption"], color="muted_text")

    def icon_png(self, name, color):
        key = (name, color)
        if key not in self.icons_png:
            self.icons_png[key] = iconlib.png(name, C[color], 96)
        return self.icons_png[key]

    def list_item(self, marker, text, indent=0.0, color="primary", bold_marker=False):
        p = para_fmt(self.doc.add_paragraph(), after=4, left=indent + 0.3, first=-0.3)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(indent + 0.3))
        set_font(p.add_run(marker + "\t"), size=S["body"], bold=bold_marker, color=color)
        rich(self.doc, p, text)
        return p

    def block(self, b):
        k = b["type"]; doc = self.doc
        if k == "paragraph":
            rich(doc, para_fmt(doc.add_paragraph(), after=8), b["text"])
        elif k == "bullets":
            for it in b["items"]:
                self.list_item("•", it)
            para_fmt(doc.paragraphs[-1], after=8)
        elif k == "numbered":
            for i, it in enumerate(b["items"], 1):
                self.list_item("%d." % i, it, bold_marker=True)
            para_fmt(doc.paragraphs[-1], after=8)
        elif k == "procedure":
            if b.get("title"):
                p = para_fmt(doc.add_paragraph(), before=4, after=4, keep=True)
                set_font(p.add_run(b["title"]), size=S["body"], bold=True, color="dark")
            for i, st in enumerate(b["steps"], 1):
                p = self.list_item("%d." % i, st["text"], bold_marker=True)
                if st.get("result") or st.get("image"):
                    p.paragraph_format.keep_with_next = True
                if st.get("result"):
                    q = para_fmt(doc.add_paragraph(), after=4, left=0.3)
                    set_font(q.add_run("Result: "), size=S["body"] - 0.5, bold=True, color="muted_text")
                    rich(doc, q, st["result"], size=S["body"] - 0.5, color="muted_text")
                if st.get("image"):
                    self.image(st["image"], st.get("width_in", 4.8), st.get("alt", st["text"]), indent=0.3)
                    if st.get("caption"):
                        self.caption("Figure", st["caption"], indent=0.3)
            para_fmt(doc.paragraphs[-1], after=10)
        elif k == "table":
            if b.get("caption"):
                self.caption("Table", b["caption"], above=True)
            cols = b["columns"]
            widths = b.get("widths_in")
            if not widths:
                widths = auto_widths(cols, b["rows"])
            elif not LEGACY:
                widths = fit_widths(widths, min_widths(cols, b["rows"]))
            if abs(self.tw - TW) > 1e-6:  # a half-width cell on a sheet, or a landscape section (D-77)
                widths = [w * self.tw / sum(widths) for w in widths]
            self.table_words |= {w.strip(".,;:()!?\"'") for x in cols + [c for r in b["rows"] for c in r]
                                 for w in re.sub(r"[*`]", "", re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", str(x))).split() if len(w.strip(".,;:()!?\"'")) >= 4}
            t = doc.add_table(rows=1 + len(b["rows"]), cols=len(cols)); table_plain(t, widths)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, c in enumerate(cols):
                cell = t.cell(0, j); shade(cell, "dark"); cell_margins(cell); cell_borders(cell, {"bottom": ("dark", 4)})
                set_font(para_fmt(cell.paragraphs[0], after=0).add_run(c), size=S["table"], bold=True, color="white")
            trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); trPr.append(h)
            for i, row in enumerate(b["rows"], 1):
                cs = OxmlElement("w:cantSplit"); t.rows[i]._tr.get_or_add_trPr().append(cs)
                for j, v in enumerate(row):
                    cell = t.cell(i, j); cell_margins(cell); cell_borders(cell, {"bottom": ("card_line", 4)})
                    rich(doc, para_fmt(cell.paragraphs[0], after=0, line=1.1), str(v), size=S["table"])
            para_fmt(doc.add_paragraph(), after=4)
        elif k == "figure":
            self.image(b["image"], b.get("width_in", self.tw), b.get("alt", b["caption"]), crop=b.get("crop"))
            self.caption("Figure", b["caption"])
        elif k == "callout":
            spec = T["callouts"][b["kind"]]
            t = doc.add_table(rows=1, cols=2); table_plain(t, [0.5, self.tw - 0.5])
            if not LEGACY:
                t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))  # one row that can't break: title stays with its body
            ic, tx = t.cell(0, 0), t.cell(0, 1)
            for c in (ic, tx):
                cell_margins(c, 100, 100, 120 if c is ic else 40, 140)
                shade(c, "white")
            cell_borders(ic, {"left": (spec["bar"], 24), "top": ("card_line", 4), "bottom": ("card_line", 4)})
            cell_borders(tx, {"right": ("card_line", 4), "top": ("card_line", 4), "bottom": ("card_line", 4)})
            icolor = "alert" if b["kind"] == "warning" else ("dark" if b["kind"] == "note" else "primary")
            color = "alert_text" if icolor == "alert" else icolor  # D-64: text at 4.5:1; the icon keeps the brand alert color
            r = para_fmt(ic.paragraphs[0], after=0).add_run(); r.add_picture(io.BytesIO(self.icon_png(b.get("icon", spec["icon"]), icolor)), width=Inches(0.24))
            self.alt(r, spec["label"].title() + " icon")
            p = para_fmt(tx.paragraphs[0], after=2)
            set_font(p.add_run(b.get("title") or spec["label"]), size=S["body"] - 0.5, bold=True, color=color)
            if b.get("text"):
                rich(doc, para_fmt(tx.add_paragraph(), after=0), b["text"], size=S["body"] - 0.5)
            for it in b.get("items") or []:
                q = para_fmt(tx.add_paragraph(), after=1, left=0.22, first=-0.18)
                set_font(q.add_run("\u2022  "), size=S["body"] - 0.5, color=color); rich(doc, q, it, size=S["body"] - 0.5)
            plain = lambda s: re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", str(s or "")).replace("**", "").replace("`", "")
            body_txt = " ".join([plain(b.get("text"))] + [plain(i) for i in b.get("items") or []]).split()
            self.callouts.append({"title": b.get("title") or spec["label"], "head": " ".join(body_txt[:4]), "tail": " ".join(body_txt[-4:])})
            para_fmt(doc.add_paragraph(), after=2)
        elif k == "code":
            if b.get("title"):
                p = para_fmt(doc.add_paragraph(), before=4, after=3, keep=True)
                set_font(p.add_run(b["title"]), size=S["body"] - 0.5, bold=True, color="dark")
            t = doc.add_table(rows=1, cols=1); table_plain(t, [self.tw]); c = t.cell(0, 0)
            shade(c, "F4F4F4" if "F4F4F4" not in C else "F4F4F4"); cell_margins(c, 100, 100, 160, 160)
            lines = b["code"].split("\n")
            p = para_fmt(c.paragraphs[0], after=0, line=1.0)
            for i, ln in enumerate(lines):
                r = p.add_run(ln); set_font(r, name=F["mono"], size=S["code"], color="text")
                if i < len(lines) - 1:
                    r.add_break()
            para_fmt(doc.add_paragraph(), after=4)
        elif k == "step_diagram":
            self.step_diagram(b)
        elif k == "card":
            self.card(b)
        elif k == "page_break":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def blocks(self, bl):
        """Blocks in order; a figure keeps its caption with a short paragraph that follows it (40 words or fewer)."""
        for j, b in enumerate(bl):
            self.block(b)
            nxt = bl[j + 1] if j + 1 < len(bl) else None
            if b.get("type") == "figure" and nxt and nxt.get("type") == "paragraph" and len(str(nxt.get("text", "")).split()) <= 40:
                self.doc.paragraphs[-1].paragraph_format.keep_with_next = True

    def orient(self, landscape):
        """D-77: start a new page section in landscape or back in portrait; True when a section break was added."""
        if self.landscape == landscape:
            return False
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        w, h = Inches(P["width_in"]), Inches(P["height_in"])
        sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = (h, w) if landscape else (w, h)
        sec.different_first_page_header_footer = False
        self.landscape, self.tw = landscape, (LTW if landscape else TW)
        self.section_footer(sec)
        return True

    def section_footer(self, sec):
        """D-77: own copy of the house footer with the page number at this section's right margin."""
        import copy as _copy
        src = self.doc.sections[0].footer.paragraphs
        if self.sheet or not src:
            return
        sec.footer.is_linked_to_previous = False
        ft = sec.footer._element
        for p in list(ft.findall(qn("w:p"))):
            ft.remove(p)
        for p in src:
            ft.append(_copy.deepcopy(p._p))
        for tab in list(ft.iter(qn("w:tab"))):
            if tab.get(qn("w:val")) == "right" and abs(self.tw - TW) > 1e-6:
                tab.set(qn("w:pos"), str(int(round(self.tw * 1440))))
                clr = OxmlElement("w:tab"); clr.set(qn("w:val"), "clear"); clr.set(qn("w:pos"), str(int(round(TW * 1440))))
                tab.addprevious(clr)  # the Footer style's own right stop at the portrait width would catch the tab first

    def body(self):
        new_page = self.m["page_break_level"] == 1 if "page_break_level" in self.m else self.m.get("h1_new_page", True)
        def walk(secs, prefix, lvl):
            for i, s in enumerate(secs, 1):
                num = "%s%d" % (prefix, i)
                brk = lvl == 1 and self.orient(s.get("orientation") == "landscape")
                self.heading(lvl, "%s  %s" % (num, s["title"]), new_page=(lvl == 1 and not brk and (new_page or i == 1)), anchor=anchor_name(num), icon=s.get("icon"))
                self.blocks(lead_blocks(s))
                walk(s.get("sections", []), num + ".", lvl + 1)
                self.blocks(tail_blocks(s))
        walk(self.d["sections"], "", 1)
        custom = {a["title"]: a for a in self.d.get("appendices") or []}
        for letter, title in self.appendices():
            self.heading(1, "Appendix %s: %s" % (letter, title), new_page=not self.orient(False), anchor=anchor_name("Appendix " + letter), icon=(custom.get(title) or {}).get("icon") or (AUTO_ICONS.get(title) if self.m.get("appendix_icons") else None))
            if title in custom:
                self.blocks(custom[title].get("blocks", []))
            elif title == "Index":
                self.index()
            elif title == "Glossary":
                self.block({"type": "table", "columns": ["Term", "Definition"], "rows": [[g["term"], g["definition"]] for g in sorted(self.d["glossary"], key=lambda g: g["term"].lower())], "widths_in": [1.8, TW - 1.8]})
            elif title == "Revision History":
                rv = self.review_rows()
                if rv:  # D-65: owner role and review dates above the revisions
                    self.block({"type": "table", "columns": [lab for _, lab in rv], "rows": [[v for v, _ in rv]], "widths_in": [TW / len(rv)] * len(rv)})
                self.block({"type": "table", "columns": ["Version", "Date", "Author", "Change"], "rows": [[r["version"], r["date"], r["author"], r["change"]] for r in self.d["revisions"]], "widths_in": [0.8, 1.1, 1.5, TW - 3.4]})
            else:
                for s in self.m["sources"]:
                    self.list_item("•", s)

    def build(self, out):
        self.styles(); self.page()
        cp = self.doc.core_properties
        cp.title, cp.subject, cp.author = self.m["title"], self.m.get("subtitle", ""), self.m["owner"]
        cp.keywords = self.m["doc_type"]; cp.comments = "Built by document-builder"
        if self.sheet:  # no cover, contents or automatic appendices
            cp.category = self.m.get("classification", "Internal use only")
            if len(self.m.get("sources") or []) > self.sources_on_page:  # unlinked (or unfitted) sources stay in doc properties
                cp.comments += ". Sources: " + "; ".join(str(x) for x in self.m["sources"])
            self.toc_labels, self.all_labels = [], [s["title"] for s in self.d["sections"]]
            self.sheet_body()
        else:
            self.cover(); self.contents(); self.body()
        self.doc.save(out)
        brand_theme(out)
        self.embedded = embed_fonts(out, self.brand)
        return self.heads


EMBED_STYLES = (("embedRegular", False, False), ("embedBold", True, False), ("embedItalic", False, True), ("embedBoldItalic", True, True))
FONT_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font"
FONT_CT = "application/vnd.openxmlformats-officedocument.obfuscatedFont"


def embeddable(path):
    """False when the font's OS/2 fsType forbids embedding (restricted license, bit 1 set alone)."""
    import struct
    try:
        data = open(path, "rb").read()
        n = struct.unpack(">H", data[4:6])[0]
        for i in range(n):
            if data[12 + 16 * i:16 + 16 * i] == b"OS/2":
                off = struct.unpack(">I", data[20 + 16 * i:24 + 16 * i])[0]
                fs = struct.unpack(">H", data[off + 8:off + 10])[0]
                return not (fs & 0x000F) == 0x0002
        return True
    except Exception:
        return False


def obfuscate(data, guid):
    """ECMA-376 Part 1 §17.8.1: XOR the first 32 bytes with the font key (the GUID's 16 bytes in reverse order)."""
    key = bytes.fromhex(guid.strip("{}").replace("-", ""))[::-1]
    head = bytes(b ^ key[i % 16] for i, b in enumerate(data[:32]))
    return head + data[32:]


def embed_fonts(docx_path, b):
    """Embed the brand's font files (every family the document names, each style found) as obfuscated font parts, so
    Word, LibreOffice and recipients draw the right font even when it isn't installed. Returns [(family, style file)]."""
    import uuid, zipfile
    from xml.sax.saxutils import escape
    roles = {}  # family -> a role that uses it (body first)
    for role in ("body", "heading", "mono"):
        if F.get(role):
            roles.setdefault(F[role], role)
    plan, done, n = {}, [], 0
    for fam, role in roles.items():
        for tag, bold, italic in EMBED_STYLES:
            p = b.font_path(role, bold, italic) if b else None
            if not p or not embeddable(p) or str(BR.font_name(p)[0] or "").lower() != fam.lower():
                continue  # never embed a stand-in under another family's name
            n += 1
            guid = "{%s}" % str(uuid.uuid4()).upper()
            plan.setdefault(fam, []).append((tag, "rIdFont%d" % n, guid, "fonts/font%d.odttf" % n, p))
    if not plan:
        return []
    zin = zipfile.ZipFile(docx_path)
    parts = {i.filename: zin.read(i.filename) for i in zin.infolist()}
    zin.close()
    ft = parts["word/fontTable.xml"].decode("utf-8")
    if 'xmlns:r="' not in ft[:600]:
        ft = ft.replace("<w:fonts ", '<w:fonts xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
    rels = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">']
    for fam, items in plan.items():
        emb = "".join('<w:%s r:id="%s" w:fontKey="%s"/>' % (tag, rid, guid) for tag, rid, guid, _, _ in items)
        rx = re.compile(r'<w:font w:name="%s">(.*?)</w:font>' % re.escape(escape(fam, {'"': "&quot;"})), re.S)
        m = rx.search(ft)
        if m:
            inner = re.sub(r"<w:embed\w+ [^>]*/>", "", m.group(1)) + emb
            ft = ft[:m.start()] + '<w:font w:name="%s">%s</w:font>' % (escape(fam, {'"': "&quot;"}), inner) + ft[m.end():]
        else:
            ft = ft.replace("</w:fonts>", '<w:font w:name="%s"><w:charset w:val="00"/><w:family w:val="%s"/><w:pitch w:val="%s"/>%s</w:font></w:fonts>'
                            % (escape(fam, {'"': "&quot;"}), "modern" if fam == F.get("mono") else "swiss", "fixed" if fam == F.get("mono") else "variable", emb))
        for tag, rid, guid, target, p in items:
            parts["word/" + target] = obfuscate(open(p, "rb").read(), guid)
            rels.append('<Relationship Id="%s" Type="%s" Target="%s"/>' % (rid, FONT_REL, target))
            done.append((fam, os.path.basename(p)))
    parts["word/fontTable.xml"] = ft.encode("utf-8")
    old = parts.get("word/_rels/fontTable.xml.rels", b"").decode("utf-8")
    keep = re.findall(r"<Relationship [^>]*/>", old)
    parts["word/_rels/fontTable.xml.rels"] = (rels[0] + "".join(keep) + "".join(rels[1:]) + "</Relationships>").encode("utf-8")
    ct = parts["[Content_Types].xml"].decode("utf-8")
    if 'Extension="odttf"' not in ct:
        ct = ct.replace("</Types>", '<Default Extension="odttf" ContentType="%s"/></Types>' % FONT_CT)
    parts["[Content_Types].xml"] = ct.encode("utf-8")
    st = parts["word/settings.xml"].decode("utf-8")
    if "<w:embedTrueTypeFonts" not in st:  # schema order: right after w:zoom / w:removePersonalInformation when present
        m = re.search(r"<w:zoom[^>]*/>|<w:removePersonalInformation/>|<w:removeDateAndTime/>", st)
        st = st[:m.end()] + "<w:embedTrueTypeFonts/>" + st[m.end():] if m else re.sub(r"(<w:settings[^>]*>)", r"\1<w:embedTrueTypeFonts/>", st, 1)
    parts["word/settings.xml"] = st.encode("utf-8")
    write_parts(docx_path, parts)
    return done


def brand_theme(docx_path):
    """Point the .docx's Office theme at the brand (heading and body fonts as the major and minor theme fonts; dark,
    primary and secondary as dk2, accent1 and accent2), so Word's theme pickers offer the brand, not a template's."""
    import zipfile
    with zipfile.ZipFile(docx_path) as zin:
        parts = {i.filename: zin.read(i.filename) for i in zin.infolist()}
    name = next((n for n in parts if re.match(r"word/theme/theme\d*\.xml$", n)), None)
    if not name or not F.get("body"):
        return
    from xml.sax.saxutils import quoteattr
    th = parts[name].decode("utf-8")
    for tag, fam in (("majorFont", F.get("heading") or F["body"]), ("minorFont", F["body"])):
        th = re.sub(r"(<a:%s>\s*<a:latin typeface=)\"[^\"]*\"" % tag, lambda m: m.group(1) + quoteattr(fam), th, 1)
    for slot, role in (("dk2", "dark"), ("accent1", "primary"), ("accent2", "secondary")):
        th = re.sub(r"(<a:%s>\s*<a:srgbClr val=)\"[0-9A-Fa-f]{6}\"" % slot, lambda m: m.group(1) + '"%s"' % C[role], th, 1)
    parts[name] = th.encode("utf-8")
    write_parts(docx_path, parts)


def write_parts(docx_path, parts):
    import zipfile
    tmp = docx_path + ".part"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for name, data in parts.items():
            zo.writestr(name, data)
    os.replace(tmp, docx_path)


# Windows fallback when LibreOffice is missing: Word exports read-only (the .docx is never saved), tagged, with heading
# bookmarks (CreateBookmarks 1) and document properties, so the PDF checks and the contents page numbers still work.
WORD_PDF_PS = ("$a = New-Object -ComObject Word.Application; $a.Visible = $false; $a.DisplayAlerts = 0; "
               "try { $d = $a.Documents.Open($env:DOC_SRC, $false, $true); "
               "$d.ExportAsFixedFormat($env:DOC_PDF, 17, $false, 0, 0, 1, 1, 0, $true, $true, 1, $true, $true, $false); "
               "$d.Close($false) } finally { $a.Quit() }")


def word_pdf(docx_path, pdf):
    ps = shutil.which("powershell") or shutil.which("pwsh")
    if os.name != "nt" or not ps:
        return False
    env = dict(os.environ, DOC_SRC=os.path.abspath(docx_path), DOC_PDF=os.path.abspath(pdf))
    try:
        subprocess.run([ps, "-NoProfile", "-NonInteractive", "-Command", WORD_PDF_PS], capture_output=True, timeout=300, env=env)
    except (OSError, subprocess.TimeoutExpired):
        return False
    return os.path.isfile(pdf)


WORD_LOG = "exported with Microsoft Word (LibreOffice not found)"


def sheet_outline(pdf, titles):
    """Word bookmarks only headings outside tables, so a sheet's band titles (Heading 1 in the band cell) are missing:
    rewrite the outline as one level-1 bookmark per top-level section, found by text search in order."""
    import fitz
    doc = fitz.open(pdf)
    norm = lambda s: " ".join(str(s).split())
    texts, toc, at = [norm(pg.get_text()) for pg in doc], [], 0
    for t in titles:
        hit = next((i for i in range(at, len(texts)) if norm(t) in texts[i]), None)
        if hit is None:
            break
        toc.append([1, t, hit + 1]); at = hit + 1  # every top-level section starts a new page
    doc.set_toc(toc)
    doc.saveIncr(); doc.close()


def to_pdf(docx_path, outdir):
    """(pdf path or None, log). LibreOffice (on PATH or in its usual install folder) with tagged-PDF and bookmark
    options; on Windows without it, Word through COM. Neither: NOT RUN."""
    from visual_batch import find_soffice
    pdf = os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
    soffice = find_soffice()
    if not soffice:
        if word_pdf(docx_path, pdf):
            return pdf, WORD_LOG
        return None, "NOT RUN: LibreOffice (soffice) is not available, so no PDF; deliver the .docx and say the PDF wasn't made"
    prof = tempfile.mkdtemp(prefix="lo_doc_")
    flt = 'pdf:writer_pdf_Export:{"ExportBookmarks":{"type":"boolean","value":"true"},"UseTaggedPDF":{"type":"boolean","value":"true"},"EmbedStandardFonts":{"type":"boolean","value":"true"}}'
    b = brand()
    env = BR.fontconfig_env(b.font_dirs() if b else ())  # LibreOffice sees the bundled and project fonts; nothing is installed
    r = subprocess.run([soffice, "-env:UserInstallation=" + pathlib.Path(prof).as_uri(), "--headless", "--convert-to", flt, "--outdir", outdir, docx_path], capture_output=True, text=True, timeout=240, env=env)
    shutil.rmtree(prof, ignore_errors=True)
    return (pdf if os.path.isfile(pdf) else None), (r.stdout + r.stderr)[-300:]


def pdf_a11y(pdf, title):
    """D-64 (PDF/UA basics): catalog /Lang, Info /Title = meta.title and DisplayDocTitle, saved in place (DOC_A11Y)."""
    import fitz
    doc = fitz.open(pdf)
    cat = doc.pdf_catalog()
    doc.xref_set_key(cat, "Lang", fitz.get_pdf_str(T.get("a11y", {}).get("lang", "en-US")))
    kind, val = doc.xref_get_key(cat, "ViewerPreferences")
    if kind == "xref":
        doc.xref_set_key(int(val.split()[0]), "DisplayDocTitle", "true")
    else:
        doc.xref_set_key(cat, "ViewerPreferences", "<</DisplayDocTitle true>>")
    keep = ("author", "producer", "creator", "title", "subject", "keywords", "creationDate", "modDate")
    doc.set_metadata(dict({k: v for k, v in (doc.metadata or {}).items() if k in keep and v}, title=title))
    doc.saveIncr(); doc.close()


def heading_pages(pdf, labels):
    """Page of each heading, read from the PDF outline (bookmarks); text search is the fallback."""
    import fitz
    doc = fitz.open(pdf)
    norm = lambda s: " ".join(str(s).split())
    outline = {}
    for lvl, title, page in doc.get_toc():
        outline.setdefault(norm(title), page)
    out, texts, start = {}, None, 2
    for lab in labels:
        key = norm(lab)
        if key in outline and outline[key] > 0:
            out[lab] = outline[key]; continue
        if texts is None:
            texts = [norm(pg.get_text()) for pg in doc]
        for i in range(start, len(texts)):
            if key in texts[i]:
                out[lab] = i + 1; start = i; break
    return out


def sheet_info(pdf, b):
    """Page facts lint needs for a sheet: first page and page count of each top-level section (level-1 bookmarks in
    order; text search is the fallback), plus the usual callout, table-word and caption records."""
    import fitz
    doc = fitz.open(pdf)
    titles = [s["title"] for s in b.d["sections"]]
    starts = [p for lvl, t, p in doc.get_toc() if lvl == 1]
    if len(starts) != len(titles) or any(p < 1 for p in starts):
        norm = lambda s: " ".join(str(s).split())
        texts, starts, at = [norm(pg.get_text()) for pg in doc], [], 0
        for t in titles:
            hit = next((i for i in range(at, len(texts)) if norm(t) in texts[i]), None)
            if hit is None:
                break
            starts.append(hit + 1); at = hit + 1
    n = doc.page_count
    secs = [[t, p, (starts[i + 1] if i + 1 < len(starts) else n + 1) - p] for i, (t, p) in enumerate(zip(titles, starts))]
    return {"sheet": True, "doc_type": b.m.get("doc_type"), "pages": n, "sections": secs, "pages_per_section": b.cfg["pages_per_section"],
            "toc": {}, "all": {t: p for t, p, _ in secs}, "stable": True, "toc_levels": 0, "index": [],
            "first_body": titles[0] if titles else None, "first_appendix": None, "h1": titles, "callouts": b.callouts,
            "table_words": sorted(b.table_words), "captions": b.captions, "page_break_level": 1,
            "sources_in_footer": b.sources_in_footer, "footer_links": b.footer_links}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("doc"); ap.add_argument("--out", required=True); ap.add_argument("--pdf", action="store_true"); ap.add_argument("--draft", action="store_true")
    ap.add_argument("--brand", help="brand profile (default: meta.brand, then brand.json beside doc.json or at the project root, then detected)")
    a = ap.parse_args()
    brand = use_brand(BR.load(source=a.doc, brand=a.brand))
    print(brand.lines()[0] + "; logo: %s" % ("yes" if brand.logo_png("light") else "none"))
    st = brand.font_status()
    print("FONTS: " + ", ".join("%s %s (%s)" % (r, brand.fonts[r], st.get(r, "?")) for r in ("body", "display", "mono")))
    lint = [sys.executable, os.path.join(HERE, "lint_doc.py"), a.doc] + (["--draft"] if a.draft else []) + (["--brand", a.brand] if a.brand else [])
    r = subprocess.run(lint, capture_output=True, text=True)
    if r.returncode != 0:
        codes = set(re.findall(r"^FAIL \[([A-Z_]+)\]", r.stdout, re.M))
        if not codes or codes - set(SHEET["build_nonblocking"]):
            print("NOT BUILT: doc.json failed the design gate; fix these and rebuild.\n" + r.stdout[-3000:]); sys.exit(1)
        print("BUILD CONTINUES: %s is checked again on the built files; lint stays failed until it is fixed.\n%s"
              % (", ".join(sorted(codes)), "\n".join(x for x in r.stdout.splitlines() if x.startswith("FAIL"))[:1500]))
    d = json.load(open(a.doc, encoding="utf-8"))
    base = os.path.dirname(os.path.abspath(a.doc))
    out = os.path.abspath(a.out)
    b = Builder(d, base, a.draft, brand=brand); b.build(out)
    msg = "BUILT %s: %d headings, %d figures, %d tables; fonts embedded: %s" % (out, len(b.heads), b.fig, b.tab, ", ".join(f for _, f in b.embedded) or "none")
    if a.pdf:
        tmp = tempfile.mkdtemp(prefix="docpdf_")
        pdf, log = to_pdf(out, tmp)
        if not pdf:
            print(msg + "\nPDF " + log); sys.exit(0)
        if b.sheet:  # one pass: no contents page numbers to settle
            final = os.path.splitext(out)[0] + ".pdf"
            shutil.copy(pdf, final); shutil.rmtree(tmp, ignore_errors=True)
            if log == WORD_LOG:
                sheet_outline(final, [s["title"] for s in d["sections"]])
            pdf_a11y(final, d["meta"]["title"])
            info = sheet_info(final, b)
            json.dump(info, open(os.path.splitext(out)[0] + ".toc.json", "w"))
            print(msg + "\nBUILT %s: %d pages for %d sections (pages per section: %s)" % (final, info["pages"], len(info["sections"]), ", ".join(str(x[2]) for x in info["sections"])))
            return
        pages = heading_pages(pdf, b.all_labels)
        b2 = Builder(d, base, a.draft, pages, brand=brand); b2.build(out)
        pdf, log = to_pdf(out, tmp)
        pages2 = heading_pages(pdf, b2.all_labels)
        stable = pages2 == pages
        final = os.path.splitext(out)[0] + ".pdf"
        shutil.copy(pdf, final); shutil.rmtree(tmp, ignore_errors=True); pdf_a11y(final, d["meta"]["title"])
        h1 = [lab for lvl, lab in b2.heads if lvl == 1]
        json.dump({"toc": {k: v for k, v in pages2.items() if k in b2.toc_labels}, "all": pages2, "stable": stable, "toc_levels": b2.levels,
                   "index": [[e["term"], b2.resolve(e["section"])[1]] for e in d.get("index") or []],
                   "first_body": h1[0] if h1 else None, "first_appendix": next((x for x in h1 if x.startswith("Appendix ")), None),
                   "callouts": b2.callouts, "table_words": sorted(b2.table_words), "h1": h1, "captions": b2.captions,
                   "page_break_level": 1 if b2.m.get("page_break_level", 1 if b2.m.get("h1_new_page", True) else 0) == 1 else 0}, open(os.path.splitext(out)[0] + ".toc.json", "w"))
        msg += "\nBUILT %s: contents page numbers %s" % (final, "set from the PDF layout" if stable else "SHIFTED on the second pass; rebuild")
    print(msg)


if __name__ == "__main__":
    main()

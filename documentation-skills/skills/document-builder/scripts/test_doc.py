#!/usr/bin/env python3
"""Self-test for document-builder: builds the example to Word and PDF, lints it, and proves the gate fails on bad input.

Every case runs in a temp folder holding an empty brand.json, so it uses the neutral theme and no logo whatever project
the skill sits in; the brand cases (synthetic Northwind project, meta.brand, house terms) build in folders of their own."""
import copy, glob, json, os, re, shutil, subprocess, sys, tempfile
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import brand as BR  # noqa: E402
EX = json.load(open(os.path.join(ROOT, "assets", "user-guide.example.doc.json"), encoding="utf-8"))
TAGGED = []  # (case name, codes) for every case seen, selected or not
ONLY = None  # --only filter: None runs every case


def colors_for(path):
    """The resolved brand colors a build of `path` uses (role -> hex), read the same way build and lint read them."""
    return BR.load(source=path).colors


def font_files():
    """The font file synthetic figures are drawn with: the skill's bundled fallback body font (no family named here)."""
    p = BR.find_font(BR.FALLBACK_FONTS["body"])
    return [p] if p else []


def other_font():
    """(family, regular file, bold file or None) of an installed text font that is not a fallback font, allows
    embedding and has a clearly different line height, to play a project's own font; None when there is none."""
    import build_doc
    fallback_line = BR.font_metrics(BR.find_font(BR.FALLBACK_FONTS["body"]) or "")["line"]
    idx = BR.font_index(BR.system_font_dirs())
    for (fam, style), p in sorted(idx.items(), key=lambda kv: kv[1].lower()):
        if style != "regular" or fam in {f.lower() for f in BR.FALLBACK_FONTS.values()} or "mono" in fam or not p.lower().endswith(".ttf"):
            continue
        if not idx.get((fam, "bold")):  # a text family (regular and bold), not a display face
            continue
        if not build_doc.embeddable(p) or not all(build_doc.has_glyph(p, ch) for ch in "AZaz09"):
            continue
        if abs(BR.font_metrics(p)["line"] - fallback_line) < 0.08:
            continue
        real = BR.font_name(p)[0]
        return real, p, idx.get((fam, "bold"))
    return None


def pdf_engine():
    """True when build_doc.py can make a PDF here: LibreOffice, or Microsoft Word on Windows."""
    from visual_batch import find_soffice
    if find_soffice():
        return True
    if os.name != "nt":
        return False
    try:
        import winreg
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "Word.Application"))
        return True
    except OSError:
        return False


HAS_PDF = pdf_engine()
HAS_LO = bool(__import__("visual_batch").find_soffice())


def only_codes(valid):
    """--only list prints every check code, sorted; --only A,B keeps cases tagged A or B; unknown codes exit 2."""
    global ONLY
    a = sys.argv[1:]
    arg = next((x[7:] for x in a if x.startswith("--only=")), None)
    if "--only" in a:
        i = a.index("--only"); arg = a[i + 1] if i + 1 < len(a) else ""
    if arg is None:
        return
    if arg == "list":
        print("\n".join(sorted(valid))); sys.exit(0)
    req = {c.strip() for c in arg.split(",") if c.strip()}
    if not req or req - set(valid):
        sys.stderr.write("unknown --only code(s): %s. Valid codes: %s\n" % (", ".join(sorted(req - set(valid))) or "(none)", ", ".join(sorted(valid)))); sys.exit(2)
    ONLY = req


def selected(*codes):
    return ONLY is None or bool(ONLY & set(codes))


# Codes per test section: `--only list` prints their union; sections with no selected code are skipped.
LAYOUT_CODES = {"BUILD", "CALLOUT_SPLIT", "WORD_BREAK", "TOC_PAGES", "PAGE_FILL"}
TEXTBOOK_CODES = {"BUILD", "OBJECTIVES", "SECTION_REQUIRED", "FIG_TEXT_SIZE"}
SHEET_CODES = {"LAYOUT", "QS_PAGE_COUNT", "QS_STRUCTURE", "SHEET_PAGE_FIT", "STEP_LABEL", "STEP_COUNT", "FIGURE_CROP", "FIG_TEXT_SIZE", "TOC_FIELD", "BOOKMARKS",
               "DOC_TYPE", "EMPTY_BLOCK", "COLOR", "FONTS", "SOURCES", "SHEET_SUBTITLE", "SOURCE_LINK", "ILLUSTRATION_LABEL"}
ROUND6_CODES = {"DOC_TYPE", "SECTION_REQUIRED", "META_MISSING", "DOC_A11Y", "TERM_BRAND", "BRAND"}


def tags_for(name, tags=None):
    """Explicit tags, else the known codes named in the case title, else BUILD."""
    return set(tags) if tags else ({t for t in re.findall(r"[A-Z][A-Z0-9_]+", name) if t in ALL_CODES} or {"BUILD"})


def shots(d):
    from PIL import Image
    os.makedirs(os.path.join(d, "screenshots"), exist_ok=True)
    for n, (w, h) in {"dashboard_overview": (1600, 1000), "admissions_widget": (1400, 700)}.items():
        Image.new("RGB", (w, h), "#F4F4F4").save(os.path.join(d, "screenshots", n + ".png"))


def lint(p, *extra):
    r = subprocess.run([sys.executable, os.path.join(HERE, "lint_doc.py"), p] + list(extra), capture_output=True, text=True)
    return r.returncode, r.stdout


NEG = [
    ("META_MISSING", lambda d: d["meta"].pop("audience")),
    ("DOC_TYPE", lambda d: d["meta"].update(doc_type="brochure")),
    ("SOURCES", lambda d: d["meta"].pop("sources")),
    ("SOURCES", lambda d: d["meta"].update(sources=[{"title": "A source", "url": "https://example.com"}])),
    ("SECTION_REQUIRED", lambda d: d["sections"].pop()),
    ("REVISIONS", lambda d: d.pop("revisions")),
    ("HEADING_DEPTH", lambda d: d["sections"][0]["sections"][0].update(sections=[{"title": "Deeper", "sections": [{"title": "Too Deep", "blocks": [{"type": "paragraph", "text": "Text here."}]}]}])),
    ("EMPTY_SECTION", lambda d: d["sections"][2]["sections"].append({"title": "Empty One"})),
    ("PROC_STEPS", lambda d: d["sections"][3]["blocks"][0].update(steps=[{"text": "Open it."}])),
    ("PROC_IMPERATIVE", lambda d: d["sections"][3]["blocks"][0]["steps"][0].update(text="You should set the filters first.")),
    ("TABLE_SHAPE", lambda d: d["sections"][4]["blocks"][0]["rows"][0].pop()),
    ("TABLE_CAPTION", lambda d: d["sections"][4]["blocks"][0].pop("caption")),
    ("FIGURE_FILE", lambda d: d["sections"][2]["sections"][1]["blocks"][0].update(image="screenshots/missing.png")),
    ("ALT_TEXT", lambda d: d["sections"][2]["sections"][1]["blocks"][0].update(alt="chart", caption="A chart")),
    ("CALLOUT_KIND", lambda d: d["sections"][0]["blocks"][1].update(kind="danger")),
    ("ICON_UNKNOWN", lambda d: d["sections"][0]["blocks"][1].update(icon="unicorn-sparkle")),
    ("PLACEHOLDER", lambda d: d["sections"][0]["blocks"][0].update(text="TBD")),
    ("SLOP", lambda d: d["sections"][0]["blocks"][0].update(text="Leverage the dashboard to unlock insight.")),
    ("PHI", lambda d: d["sections"][0]["blocks"][0].update(text="Patient MRN: 1234567 was admitted.")),
    ("LINK", lambda d: d["sections"][0]["blocks"][0].update(text="See [the portal](portal/home) for access.")),
    ("INDEX_TARGET", lambda d: d.update(index=[{"term": "Filters", "section": "No Such Section"}])),
    ("TOC_LEVELS", lambda d: d["meta"].update(toc_levels=4)),
    ("APPENDIX", lambda d: d.update(appendices=[{"title": "Index", "blocks": [{"type": "paragraph", "text": "Words here."}]}])),
    ("APPENDIX", lambda d: d.update(appendices=[{"title": "Quick Reference"}])),
    ("CALLOUT_LONG", lambda d: d["sections"][0]["blocks"][1].update(items=["One two three four five six seven."] * 7)),
    ("DESIGN", lambda d: d["meta"].update(design="fancy")),
    ("DUP_HEADING", lambda d: d["sections"][2]["sections"].__setitem__(0, dict(d["sections"][2]["sections"][0], title=d["sections"][2]["title"]))),
    ("PAGE_BREAK_LEVEL", lambda d: d["meta"].update(page_break_level=2)),
]
ALL_CODES = LAYOUT_CODES | TEXTBOOK_CODES | SHEET_CODES | ROUND6_CODES | {"GUIDE_PROFILE", "REVIEW_DATE", "SHEET_FORMAT_KNOWN"} | {"FIGURE_NONE", "TABLE_REDUNDANT"} | {c for c, _ in NEG}


def layout_doc(n_fill=8, n_toc=45):
    """Sections whose callouts land at different heights, a PASS/FAIL/BLOCKED table and a two-page contents."""
    fill = " ".join(["Review each row in the queue and record what you found before you move to the next account."] * 6)
    secs = [{"title": "Introduction", "blocks": [{"type": "paragraph", "text": "This guide shows how layout checks catch problems a reader sees."},
             {"type": "table", "caption": "Test results and what each means", "columns": ["Result", "What It Means"],
              "rows": [["PASS", "The step worked as written and matched the expected result in every check you ran on the page."],
                       ["FAIL", "The step didn't work as written; record what happened, what you expected and the screen where it happened."],
                       ["BLOCKED", "Something outside the test stopped you; record the reason and the person you told about it."]]}]},
            {"title": "Getting Started", "blocks": [{"type": "paragraph", "text": "Open the dashboard and sign in."}, {"type": "page_break"},
             {"type": "paragraph", "text": "A forced break leaves the page above mostly empty in the middle of a chapter."}]}]
    for k in range(1, n_fill + 1):
        secs.append({"title": "Layout Case %d" % k, "blocks": [{"type": "paragraph", "text": fill}] * k + [{"type": "callout", "kind": "tip", "title": "Tip for Case %d" % k,
                     "text": "Keep this whole box on one page so the title stays with its advice for the reader.", "items": ["Check the filters first.", "Write down what you saw.",
                     "Note the time it happened.", "Add the screen name.", "Tell your lead the same day.", "Retest after the fix ships."]}]})
    secs.append({"title": "Reference", "blocks": [{"type": "paragraph", "text": "Short topics follow."}],
                 "sections": [{"title": "Topic %d" % i, "blocks": [{"type": "paragraph", "text": "A short topic used to make the contents run to two pages."}]} for i in range(1, n_toc + 1)]})
    secs += [{"title": "Troubleshooting", "blocks": [{"type": "paragraph", "text": "Rebuild if numbers look wrong."}]}, {"title": "Support", "blocks": [{"type": "paragraph", "text": "Ask your lead."},
             {"type": "procedure", "title": "To ask for help:", "steps": [{"text": "Write down the screen and the time."}, {"text": "Send both to your lead."}]}]}]
    return {"meta": dict(EX["meta"], title="Layout Check Guide", toc_levels=2), "sections": secs, "revisions": EX["revisions"], "glossary": EX["glossary"]}  # D-63: a valid lean guide


def layout_tests(tmp, rec):
    """1.4: CALLOUT_SPLIT and WORD_BREAK fail on the 1.3.0 layout and pass after keep-together and word-width columns."""
    if not HAS_PDF:
        rec(False, "layout checks (LibreOffice needed for the PDF): NOT RUN", tags=LAYOUT_CODES); return
    p = os.path.join(tmp, "layout.json"); json.dump(layout_doc(), open(p, "w"))
    res = {}
    for mode in ("legacy", "fixed"):
        env = dict(os.environ, DOCBUILDER_LEGACY_LAYOUT="1" if mode == "legacy" else "0")
        out = os.path.join(tmp, "layout_%s.docx" % mode)
        b = subprocess.run([sys.executable, os.path.join(HERE, "build_doc.py"), p, "--out", out, "--pdf"], capture_output=True, text=True, env=env, timeout=900)
        rc, txt = lint(p, "--docx", out, "--pdf", out[:-5] + ".pdf")
        res[mode] = (b.returncode, rc, txt)
    lg, fx = res["legacy"][2], res["fixed"][2]
    rec("[CALLOUT_SPLIT]" in lg, "1.3.0 layout: a callout near the bottom of a page splits (CALLOUT_SPLIT fails)", lg[-600:])
    rec("[WORD_BREAK]" in lg and "BLOCKED" in lg, "1.3.0 layout: BLOCKED breaks in a narrow column (WORD_BREAK fails)", lg[-600:])
    rec("[CALLOUT_SPLIT]" not in fx and "[WORD_BREAK]" not in fx and res["fixed"][1] == 0, "1.4.0 layout: callouts stay whole and no table word breaks (lint passes)", fx[-600:], tags=["CALLOUT_SPLIT", "WORD_BREAK"])
    info = json.load(open(os.path.join(tmp, "layout_fixed.toc.json")))
    rec(min(info["toc"].values()) >= 4 and "[TOC_PAGES]" not in fx, "a 50-entry contents runs to two pages and passes TOC_PAGES", json.dumps(info["toc"])[:200])
    pf = [ln for ln in fx.splitlines() if "[PAGE_FILL]" in ln]
    gs = info["all"].get("2  Getting Started")
    rec(len(pf) == 1 and pf[0].startswith("WARN") and ("pages %d end" % gs) in pf[0] and "page_break_level 0" not in pf[0],
        "PAGE_FILL warns only on the mid-chapter page (%s); short chapter-end pages pass at page_break_level 1 and it never suggests 0" % gs, fx[-400:])
    d = layout_doc(); d["meta"]["page_break_level"] = 0
    q = os.path.join(tmp, "runon.json"); json.dump(d, open(q, "w"))
    out = os.path.join(tmp, "runon.docx")
    subprocess.run([sys.executable, os.path.join(HERE, "build_doc.py"), q, "--out", out, "--pdf"], capture_output=True, text=True, timeout=900)
    i2 = json.load(open(os.path.join(tmp, "runon.toc.json")))
    cases = lambda inf: [inf["all"].get("%d  Layout Case %d" % (k + 2, k)) for k in range(1, 9)]
    pb, ro = cases(info), cases(i2)
    rec(len(set(pb)) == 8 and len(set(ro)) < 8 and sum(1 for x in ro if x) == 8, "page_break_level 0: top-level sections run on (shared pages %s vs %s)" % (ro, pb))


def textbook_tests(tmp, rec):
    """1.5: objectives/takeaways boxes, combined Troubleshooting and Support, figure keep-with-next, two-column index,
    FIG_TEXT_SIZE and the textbook example."""
    from PIL import Image, ImageDraw, ImageFont
    from docx import Document
    import glob
    ex = json.load(open(os.path.join(ROOT, "assets", "textbook.example.doc.json"), encoding="utf-8"))
    os.makedirs(os.path.join(tmp, "figures"), exist_ok=True)
    fonts = font_files()
    def fig(name, px):
        im = Image.new("RGB", (1600, 900), "white"); dr = ImageDraw.Draw(im)
        f = ImageFont.truetype(fonts[0], px) if fonts else ImageFont.load_default()
        for i in range(8):
            dr.text((40, 40 + i * px * 2), "Review the release under test and record results", font=f, fill="#414141")
        im.save(os.path.join(tmp, "figures", name))
    fig("overview_illustration.png", 50)
    p = os.path.join(tmp, "textbook.json"); json.dump(ex, open(p, "w"))
    out = os.path.join(tmp, "Textbook.docx")
    b = subprocess.run([sys.executable, os.path.join(HERE, "build_doc.py"), p, "--out", out] + (["--pdf"] if HAS_PDF else []), capture_output=True, text=True, timeout=900)
    pdf = out[:-5] + ".pdf"
    rc, txt = lint(p, "--docx", out, *(["--pdf", pdf] if os.path.isfile(pdf) else []))
    rec(b.returncode == 0 and rc == 0, "textbook.example.doc.json builds and lints 0 failed%s" % (" (Word and PDF)" if os.path.isfile(pdf) else " (Word only: no PDF engine)"), b.stdout + txt)
    doc = Document(out); paras = [x.text.strip() for x in doc.paragraphs]
    body = doc.element.body.xml
    rec(body.count("What You Will Learn") == 3 and body.count("Key Takeaways") == 3, "objectives/takeaways render 3 What You Will Learn and 3 Key Takeaways boxes")
    seg = body[body.rfind(">2  Getting Started<"):body.rfind(">3  The Ground Rules<")] if ">2  Getting Started<" in body else ""  # rfind: skip the contents entries
    rec(bool(seg) and seg.find("What You Will Learn") < seg.find("Open the Dashboard") < seg.rfind("Key Takeaways"),
        "What You Will Learn opens the chapter and Key Takeaways is its last block (after the subsections)", seg[:200])
    cap = next((x for x in doc.paragraphs if x.text.startswith("Figure 1.")), None)
    rec(cap is not None and cap.paragraph_format.keep_with_next is True, "a figure caption keeps with the short paragraph that follows it")
    rec('w:num="2"' in body, "the index is set in two columns")
    if os.path.isfile(pdf):
        info = json.load(open(out[:-5] + ".toc.json"))
        rec(info.get("h1") and info.get("captions") and ["Figure", 1] in [c[:2] for c in info["captions"]], ".toc.json lists chapter headings and caption numbers for cross-reference checks", json.dumps(info.get("captions"))[:200])
    d = copy.deepcopy(ex); d["sections"][1]["objectives"] = "Open the dashboard"; q = os.path.join(tmp, "obj.json"); json.dump(d, open(q, "w"))
    rc, o = lint(q, "--draft"); rec("FAIL [OBJECTIVES]" in o, "objectives must be a list of 2-6 sentences (OBJECTIVES)", o)
    d = copy.deepcopy(ex); q = os.path.join(tmp, "ts.json"); json.dump(d, open(q, "w"))
    rc, o = lint(q, "--draft"); rec("SECTION_REQUIRED" not in o, "one 'Troubleshooting and Support' chapter meets both required sections", o, tags=["SECTION_REQUIRED"])
    d["sections"][4]["title"] = "Troubleshooting"; json.dump(d, open(q, "w"))
    rc, o = lint(q, "--draft"); rec("FAIL [SECTION_REQUIRED]" in o and "Support" in o, "without the combined title, Support is still required", o, tags=["SECTION_REQUIRED"])
    fig("small.png", 17); fig("big.png", 40)
    for name, want in (("small.png", True), ("big.png", False)):
        d = copy.deepcopy(ex); d["sections"][1]["sections"][0]["blocks"][1]["image"] = "figures/" + name
        q = os.path.join(tmp, "fig_%s.json" % name); json.dump(d, open(q, "w"))
        rc, o = lint(q)
        rec(("WARN [FIG_TEXT_SIZE]" in o) == want, "FIG_TEXT_SIZE %s for figure text printing at %s pt" % ("warns" if want else "stays quiet", "5" if want else "11.7"), o)


# Synthetic sheet inputs (D-01 to D-03): neutral report-review text, no PHI, no real names.
QS_POOL = ["Open the test from your queue and read the expected result first.", "Follow each step on the report exactly as written.",
           "Record a result for the step before you move to the next one.", "Add a short note that says what you saw when a step fails.",
           "Attach a screenshot of the page when the result differs.", "Check the filters on the report before you compare totals.",
           "Write the test name and step number in every question you send.", "Save your work after each step so nothing is lost."]
QS_STEPS = [("Open the Test", "clipboard-check"), ("Read Each Step", "eye"), ("Record a Result", "pencil"), ("Attach a Screenshot", "device-desktop"),
            ("Submit the Test", "file-check"), ("Check the Queue", "list-check")]


def qs_filler(n):
    out, i = [], 0
    while len(" ".join(out).split()) < n:
        out.append(QS_POOL[i % len(QS_POOL)]); i += 1
    return " ".join(out)


def qs_role(k, words=220, diagram=4, steps=4, seq=False, get_help=True):
    """One synthetic quick-start role: the six subsections, a step diagram, a card; padded to about `words` words."""
    sys.path.insert(0, HERE); import lint_doc
    subs = [{"title": "Who It's For", "width": "half", "blocks": [{"type": "paragraph", "text": "Synthetic role %d: anyone asked to check a report before it is released." % k}]},
            {"title": "What You Can Do", "width": "half", "blocks": [{"type": "bullets", "items": ["Open an assigned test.", "Record a result for each step."]}]},
            {"title": "Steps", "blocks": [{"type": "step_diagram", "emphasize_sequence": seq, "steps": [{"label": l, "icon": i} for l, i in QS_STEPS[:diagram]]},
                                          {"type": "numbered", "items": [QS_POOL[j % len(QS_POOL)] for j in range(steps)]}]},
            {"title": "Worked Example", "blocks": [{"type": "card", "title": "A Failed Step", "text": "The summary total does not match the detail page. Record **Fail** and note both numbers."}]},
            {"title": "Common Mistakes", "width": "half", "blocks": [{"type": "bullets", "items": ["Submitting before every step has a result.", "Recording a failure without a note."]}]},
            {"title": "Get Help", "width": "half", "blocks": [{"type": "paragraph", "text": "Ask your review lead when a step is unclear."}]}]
    if not get_help:
        subs.pop()
    role = {"title": "Synthetic Role %d" % k, "sections": subs}
    need = words - lint_doc.section_words(role)
    while need > 0:
        chunk = qs_filler(min(need, 100)); words_in = len(chunk.split())
        subs[3]["blocks"].append({"type": "paragraph", "text": chunk}); need -= words_in
    return role


def qs_doc(roles, doc_type="quick-start"):
    return {"meta": {"title": "Synthetic Review Quick Start", "doc_type": doc_type, "layout": "sheet", "audience": "Report testers",
                     "purpose": "Synthetic sheet used by the self-test", "owner": "Report Review Team", "version": "1.0", "date": "2026-10-11",
                     "sources": ["Synthetic self-test content"]}, "sections": roles}


def sheet_tests(tmp, rec):
    """1.6 sheet layout (D-01 to D-03): quick-start pages, step diagram, card, crop, QS_STRUCTURE, QS_PAGE_COUNT, SHEET_PAGE_FIT."""
    import zipfile
    from PIL import Image, ImageDraw, ImageFont
    import glob
    has_pdf = HAS_PDF
    def build(d, name):
        q = os.path.join(tmp, name + ".json"); json.dump(d, open(q, "w"))
        out = os.path.join(tmp, name + ".docx")
        b = subprocess.run([sys.executable, os.path.join(HERE, "build_doc.py"), q, "--out", out] + (["--pdf"] if has_pdf else []), capture_output=True, text=True, timeout=600)
        rc, o = lint(q, "--docx", out, *(["--pdf", out[:-5] + ".pdf"] if has_pdf else []))
        return q, out, b, rc, o
    def src(d, name):
        q = os.path.join(tmp, name + ".json"); json.dump(d, open(q, "w")); return lint(q)
    # D-03 QS_STRUCTURE (source only)
    rc, o = src(qs_doc([qs_role(1, words=200, steps=7, get_help=False)]), "qs_struct_bad")
    rec("FAIL [QS_STRUCTURE] role 1: missing Get Help" in o and "FAIL [QS_STRUCTURE] role 1: 7 steps (maximum 5)" in o,
        "D-03 QS_STRUCTURE negative (synthetic): role missing Get Help with 7 steps prints both FAILs", o, tags=["QS_STRUCTURE"])
    rc, o = src(qs_doc([qs_role(1, words=200), qs_role(2, words=200)]), "qs_struct_ok")
    rec(rc == 0 and "QS_STRUCTURE pass: 2 roles" in o, "D-03 QS_STRUCTURE positive (synthetic): complete roles pass", o, tags=["QS_STRUCTURE"])
    r = qs_role(1, words=200); r["sections"][0], r["sections"][1] = r["sections"][1], r["sections"][0]
    r["sections"][4]["blocks"][0]["items"].append("Skipping the expected result.")
    rc, o = src(qs_doc([r]), "qs_struct_order")
    rec("role 1: sections out of order" in o and "role 1: 3 common mistakes (maximum 2)" in o and "words (maximum 250)" not in o,
        "D-03 QS_STRUCTURE negative (synthetic): order and Common Mistakes limits", o, tags=["QS_STRUCTURE"])
    # D-02 STEP_LABEL, STEP_COUNT, FIGURE_CROP (source only)
    d = qs_doc([qs_role(1, words=200)]); d["sections"][0]["sections"][2]["blocks"][0]["steps"][0]["label"] = "Open the Test Right Away"
    rc, o = src(d, "step_label_bad")
    rec("FAIL [STEP_LABEL]" in o and "5 words (maximum 4)" in o, "D-02 STEP_LABEL negative (synthetic): a 5-word step label fails", o, tags=["STEP_LABEL"])
    rc, o = src(qs_doc([qs_role(1, words=200, diagram=6)]), "step_label_ok")
    rec(rc == 0 and "[STEP_LABEL]" not in o and "[STEP_COUNT]" not in o, "D-02 STEP_LABEL / STEP_COUNT positive (synthetic): six labels of 4 words or fewer pass", o, tags=["STEP_LABEL", "STEP_COUNT"])
    d = qs_doc([qs_role(1, words=200, diagram=2)]); rc, o = src(d, "step_count_bad")
    rec("FAIL [STEP_COUNT]" in o, "D-02 STEP_COUNT negative (synthetic): a 2-step diagram fails", o, tags=["STEP_COUNT"])
    os.makedirs(os.path.join(tmp, "figures"), exist_ok=True)
    fonts = font_files()
    im = Image.new("RGB", (1600, 900), "white"); dr = ImageDraw.Draw(im)
    f = ImageFont.truetype(fonts[0], 30) if fonts else ImageFont.load_default()
    for i in range(20):
        dr.text((30, 20 + i * 42), "Review the release under test and record results", font=f, fill="#414141")
    im.save(os.path.join(tmp, "figures", "wireframe_library_illustration.png"))  # synthetic stand-in for a library illustration
    fig = {"type": "figure", "image": "figures/wireframe_library_illustration.png", "caption": "Illustration of the review page", "alt": "Illustration of the report review page", "width_in": 3}
    d = qs_doc([qs_role(1, words=150)]); d["sections"][0]["sections"][3]["blocks"].append(dict(fig, crop=[0.6, 0, 0.4, 1]))
    rc, o = src(d, "crop_bad")
    rec("FAIL [FIGURE_CROP]" in o, "D-02 FIGURE_CROP negative (synthetic): x0 after x1 fails", o, tags=["FIGURE_CROP"])
    # SHOT-04 ILLUSTRATION_LABEL: an illustration must say so in its caption
    d = qs_doc([qs_role(1, words=150)]); d["sections"][0]["sections"][3]["blocks"].append(dict(fig, caption="The report review page", crop=[0, 0, 0.6, 0.5]))
    rc, o = src(d, "ill_bad")
    rec("FAIL [ILLUSTRATION_LABEL]" in o, "SHOT-04 ILLUSTRATION_LABEL negative (synthetic): an illustration captioned without 'illustration' fails", o, tags=["ILLUSTRATION_LABEL"])
    d = qs_doc([qs_role(1, words=150)]); d["sections"][0]["sections"][3]["blocks"].append(dict(fig, caption="Illustration, not a live screenshot, of the review page", crop=[0, 0, 0.6, 0.5]))
    rc, o = src(d, "ill_ok")
    rec("[ILLUSTRATION_LABEL]" not in o, "SHOT-04 ILLUSTRATION_LABEL positive (synthetic): a stamped illustration caption passes", o, tags=["ILLUSTRATION_LABEL"])
    # D-02 FIG_TEXT_SIZE: FAIL in a sheet, WARN elsewhere, passes once cropped
    d = qs_doc([qs_role(1, words=150)]); d["sections"][0]["sections"][3]["blocks"].append(dict(fig))
    rc, o = src(d, "figtext_sheet")
    rec("FAIL [FIG_TEXT_SIZE] wireframe_library_illustration.png: smallest text about" in o and "(minimum 8 pt); crop or widen it" in o,
        "D-02 FIG_TEXT_SIZE negative (synthetic): illustration at 3 in fails in a sheet", o, tags=["FIG_TEXT_SIZE"])
    d["sections"][0]["sections"][3]["blocks"][-1]["crop"] = [0, 0, 0.35, 0.35]
    rc, o = src(d, "figtext_crop")
    rec("[FIG_TEXT_SIZE]" not in o and "[FIGURE_CROP]" not in o, "D-02 FIG_TEXT_SIZE / FIGURE_CROP positive (synthetic): cropped to the relevant part it passes", o, tags=["FIG_TEXT_SIZE", "FIGURE_CROP"])
    g = copy.deepcopy(EX); g["sections"][2]["sections"][1]["blocks"][0] = dict(fig); shutil.copy(os.path.join(tmp, "figures", "wireframe_library_illustration.png"), os.path.join(tmp, "figures", "x.png"))
    rc, o = src(g, "figtext_guide")
    rec("WARN [FIG_TEXT_SIZE]" in o and "FAIL [FIG_TEXT_SIZE]" not in o, "D-02 FIG_TEXT_SIZE stays WARN outside sheets (synthetic user guide)", o, tags=["FIG_TEXT_SIZE"])
    # D-01 LAYOUT
    d = qs_doc([qs_role(1, words=200)]); d["meta"]["layout"] = "poster"; rc, o = src(d, "layout_bad")
    d2 = qs_doc([qs_role(1, words=200)]); d2["meta"]["sheet"] = {"min_body_pt": 7}; rc2, o2 = src(d2, "layout_floor")
    rec("FAIL [LAYOUT]" in o and "FAIL [LAYOUT] meta.sheet" in o2, "D-01 LAYOUT negative (synthetic): unknown layout and min_body_pt below 9 fail", o + o2, tags=["LAYOUT"])
    d = qs_doc([{"title": "Synthetic Topic", "blocks": [{"type": "paragraph", "text": qs_filler(40)}]}], doc_type="shift-handoff"); rc, o = src(d, "sheet_any_type")
    d2 = qs_doc([qs_role(1, words=200)]); d2["meta"].pop("layout"); rc2, o2 = src(d2, "qs_no_sheet")
    d3 = qs_doc([qs_role(1, words=200)], doc_type="Quick Start"); rc3, o3 = src(d3, "qs_not_kebab")
    rec(rc == 0 and "[DOC_TYPE]" not in o and "FAIL [DOC_TYPE]" in o2 and "FAIL [DOC_TYPE]" in o3,
        "D-01 DOC_TYPE (synthetic): any kebab-case doc_type is valid with layout sheet; quick-start without layout sheet and 'Quick Start' fail", o + o2 + o3, tags=["DOC_TYPE", "LAYOUT"])
    d = qs_doc([qs_role(1, words=200)]); cd = d["sections"][0]["sections"][3]["blocks"][0]; cd.pop("title"); rc, o = src(d, "card_untitled")
    d2 = qs_doc([qs_role(1, words=200)]); d2["sections"][0]["sections"][3]["blocks"][0] = {"type": "card", "title": "Empty", "accent": "pink"}; rc2, o2 = src(d2, "card_bad")
    rec(rc == 0 and "FAIL [EMPTY_BLOCK]" in o2 and "FAIL [COLOR]" in o2, "D-02 card (synthetic): a card without a title passes; an empty card with an unknown accent fails", o + o2, tags=["EMPTY_BLOCK", "COLOR"])
    ex = os.path.join(ROOT, "assets", "quick-start.example.doc.json")
    if not has_pdf:
        rc, o = lint(ex); rec(rc == 0, "D-01 LAYOUT positive: quick-start.example.doc.json lints 0 failed (PDF checks NOT RUN: no LibreOffice)", o, tags=["LAYOUT"]); return
    q, out, b, rc, o = build(json.load(open(ex, encoding="utf-8")), "qs_example")
    rec(b.returncode == 0 and rc == 0 and "QS_PAGE_COUNT pass: 2 roles, 2 pages" in o, "D-01 LAYOUT positive: quick-start.example.doc.json builds to 2 pages and lints 0 failed", b.stdout + o, tags=["LAYOUT"])
    # D-01 QS_PAGE_COUNT: 4 roles -> 4 pages; a 600-word 5th role -> 2 pages
    roles = [qs_role(1, 220, diagram=5), qs_role(2, 220, diagram=4, seq=True), qs_role(3, 220, diagram=3), qs_role(4, 220, diagram=4)]
    q, out, b, rc, o = build(qs_doc(roles), "qs4")
    import fitz
    pdf = out[:-5] + ".pdf"; pd = fitz.open(pdf) if os.path.isfile(pdf) else None
    rec(b.returncode == 0 and pd is not None and pd.page_count == 4 and rc == 0 and " 0 failed" in o and "QS_PAGE_COUNT pass: 4 roles, 4 pages" in o,
        "D-01 QS_PAGE_COUNT positive (synthetic): 4 roles of about 220 words build to exactly 4 pages, 0 failed", b.stdout + o, tags=["QS_PAGE_COUNT"])
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    toc = [t for t in (pd.get_toc() if pd else []) if t[0] == 1]
    rec("TOC" not in xml and [t[2] for t in toc] == [1, 2, 3, 4] and "Contents" not in (pd[0].get_text() if pd else "x"),
        "D-01 TOC_FIELD / BOOKMARKS (synthetic): no TOC field or contents, one level-1 bookmark per role on pages 1-4", json.dumps(toc), tags=["TOC_FIELD", "BOOKMARKS"])
    lead = "Proprietary and confidential. For internal use only. \u00b7 Synthetic Review Quick Start v1.0 \u00b7 Revised 2026-10-11"  # neutral brand: no copyright line
    foot = all(" ".join(pg.get_text().split()).count(lead) == 1 and "\u00a9" not in pg.get_text() and "Page %d of 4" % (pg.number + 1) in " ".join(pg.get_text().split()) for pg in pd) if pd else False
    band = pd is not None and all(t in " ".join(pd[0].get_text().split()) for t in ("Synthetic Review Quick Start", "Synthetic Role 1", "Version 1.0 \u00b7 2026-10-11", "Quick Start"))
    rec(foot and band, "D-01 QS_PAGE_COUNT (synthetic): title band (title, role, Version X \u00b7 Date) and house footer with Page X of Y on every page", tags=["QS_PAGE_COUNT"])
    def diagram(pg):
        dots = [w for w in pg.get_text("words") if w[4] == "\u25cb"]
        if not dots:
            return 0, 0, ["no dots"]
        y0, y1 = min(w[1] for w in dots), max(w[3] for w in dots)
        icons = [i for i in pg.get_image_info() if y0 - 60 < i["bbox"][3] <= y0 + 3]
        labels = [w[4] for w in pg.get_text("words") if y1 < w[1] < y1 + 24]
        return len(dots), len(icons), labels
    n1, i1, l1 = diagram(pd[0]); n2, i2, l2 = diagram(pd[1])
    rec(n1 == 5 and i1 == 5 and l1 and not any(ch.isdigit() for w in l1 for ch in w),
        "D-02 step_diagram (synthetic): 5 hollow-dot glyphs, 5 icons and no numerals on page 1 (%d dots, %d icons)" % (n1, i1), " ".join(l1), tags=["STEP_COUNT"])
    rec(n2 == 4 and {"1", "2", "3", "4"} <= set(l2), "D-02 step_diagram (synthetic): emphasize_sequence adds numerals 1-4", " ".join(l2), tags=["STEP_COUNT"])
    rules = {}  # thin horizontal shapes by height on the page; Word draws the rule as one segment per cell, LibreOffice as one line
    for dw in pd[0].get_drawings():
        if abs(dw["rect"].y0 - dw["rect"].y1) < 2:
            rules[round(dw["rect"].y0)] = rules.get(round(dw["rect"].y0), 0) + dw["rect"].width
    rec(pd[0].get_drawings() != [] and any(w > 50 for w in rules.values()),
        "D-02 step_diagram (synthetic): the connecting line is a vector rule in the PDF", tags=["STEP_COUNT"])
    col = colors_for(q)
    rec(('w:left w:val="single" w:sz="36" w:color="%s"' % col["primary"]) in xml and ('w:fill="%s"' % col["card"]) in xml, "D-02 card (synthetic): 4.5 pt brand primary left bar on the card tint", tags=["QS_PAGE_COUNT"])
    import lint_doc as _L
    body = _L.use_brand(BR.load(source=q)).fonts["body"]
    dots = re.findall(r'<w:r><w:rPr><w:rFonts w:ascii="([^"]+)"[^>]*/>(?:(?!</w:r>).)*<w:t[^>]*>○</w:t>', xml, re.S)
    rec(len(dots) >= 5 and set(dots) == {body} and body in _L.pdf_families(pd),
        "D-02 step_diagram (synthetic): the hollow dots are U+25CB in the brand body font (%s, %d runs), drawn from that font in the PDF" % (body, len(dots)), repr(set(dots)), tags=["STEP_COUNT", "FONTS"])
    q, out, b, rc, o = build(qs_doc(roles + [qs_role(5, 600, diagram=4)]), "qs5")
    rec("BUILD CONTINUES" in b.stdout and "FAIL [QS_PAGE_COUNT] role 5: 2 pages" in o and "role 1:" not in o.split("[QS_PAGE_COUNT]", 1)[-1][:40],
        "D-01 QS_PAGE_COUNT negative (synthetic): a 5th role of about 600 words prints FAIL [QS_PAGE_COUNT] role 5: 2 pages", b.stdout[-400:] + o, tags=["QS_PAGE_COUNT"])
    # D-03 SHEET_PAGE_FIT: a job aid whose second section runs to 2 pages
    ja = qs_doc([{"title": "Before You Start", "blocks": [{"type": "paragraph", "text": qs_filler(60)}]},
                 {"title": "Record Results", "blocks": [{"type": "paragraph", "text": qs_filler(100)} for _ in range(8)]}], doc_type="job-aid")
    ja["meta"]["title"] = "Synthetic Review Job Aid"
    q, out, b, rc, o = build(ja, "jobaid")
    rec("FAIL [SHEET_PAGE_FIT] section 2: 2 pages (maximum 1)" in o and "section 1:" not in o, "D-03 SHEET_PAGE_FIT negative (synthetic): a job-aid section of 2 pages fails", b.stdout[-300:] + o, tags=["SHEET_PAGE_FIT"])
    ja["meta"]["sheet"] = {"pages_per_section": 2}; json.dump(ja, open(q, "w"))
    rc, o = lint(q, "--docx", out, "--pdf", out[:-5] + ".pdf")
    rec(rc == 0 and "SHEET_PAGE_FIT pass: 2 sections, 3 pages" in o, "D-03 SHEET_PAGE_FIT positive (synthetic): pages_per_section 2 passes", o, tags=["SHEET_PAGE_FIT"])
    # D-19 SOURCES: a sheet footer prints only sources the reader can open
    sj = qs_doc([{"title": "Before You Start", "blocks": [{"type": "paragraph", "text": qs_filler(60)}]}], doc_type="job-aid")
    sj["meta"].update(title="Synthetic Sources Job Aid", sources=["Synthetic offline process notes", "[Synthetic Linked Standard](https://example.com/synthetic-standard)"])
    q, out, b, rc, o = build(sj, "sources_sheet")
    ptxt = " ".join(" ".join(pg.get_text().split()) for pg in fitz.open(out[:-5] + ".pdf")) if os.path.isfile(out[:-5] + ".pdf") else ""
    core = zipfile.ZipFile(out).read("docProps/core.xml").decode("utf-8") if os.path.isfile(out) else ""
    rec(rc == 0 and "Sources: Synthetic Linked Standard" in ptxt, "D-19 portfolio clash: SOURCES positive (synthetic): the linked source prints in the sheet footer", b.stdout[-300:] + o, tags=["SOURCES"])
    rec(ptxt and "Synthetic offline process notes" not in ptxt and "Synthetic offline process notes" in core,
        "D-19 portfolio clash: SOURCES negative (synthetic): the unlinked source is not in the sheet PDF text; it stays in the docx core properties", ptxt[-300:], tags=["SOURCES"])
    # D-26 SHEET_SUBTITLE and D-27 SOURCE_LINK: one synthetic 2-role sheet with a subtitle, a markdown-link source and a bare URL
    sl = qs_doc([qs_role(1, 220, diagram=5), qs_role(2, 220)])
    sub, url, bare = "User acceptance testing (UAT) for partners", "https://wiki.example.org/uat-process", "https://example.com/synthetic-standard"
    sl["meta"].update(subtitle=sub, sources=["Synthetic self-test content", "[UAT Process](%s)" % url, "Synthetic standard %s" % bare])
    q, out, b, rc, o = build(sl, "subtitle_links_sheet")
    pdf = out[:-5] + ".pdf"; pd = fitz.open(pdf) if os.path.isfile(pdf) else None
    lines = [(l["bbox"], "".join(s["text"] for s in l["spans"]).strip()) for bl in (pd[0].get_text("dict")["blocks"] if pd else []) for l in bl.get("lines", [])]
    at = lambda t: next((bb for bb, x in lines if x.startswith(t)), None)
    ys = [at("Synthetic Role 1"), at(sub), at("Version 1.0")]
    band_ok = all(ys) and ys[0][3] <= ys[1][1] + 0.5 and ys[1][3] <= ys[2][1] + 0.5
    rec(rc == 0 and "SHEET_SUBTITLE pass: subtitle on page 1" in o and "QS_PAGE_COUNT pass: 2 roles, 2 pages" in o and band_ok,
        "D-26 SHEET_SUBTITLE positive (synthetic): the subtitle prints in the band between the section title and the Version line, no overlap, 1 page per role", b.stdout[-300:] + o + str(ys), tags=["SHEET_SUBTITLE"])
    s2 = copy.deepcopy(sl); s2["meta"]["subtitle"] = "Synthetic subtitle the build never printed"; q2 = os.path.join(tmp, "subtitle_missing.json"); json.dump(s2, open(q2, "w"))
    rc2, o2 = lint(q2, "--docx", out, "--pdf", pdf)
    ns = os.path.join(tmp, "sources_sheet.pdf"); ns_txt = fitz.open(ns)[0].get_text() if os.path.isfile(ns) else "x"
    rec(rc2 == 1 and "FAIL [SHEET_SUBTITLE] pdf: meta.subtitle 'Synthetic subtitle the build never printed' is not on page 1" in o2 and "SHEET_SUBTITLE" not in lint(q.replace("subtitle_links_sheet", "sources_sheet"), "--pdf", ns)[1] and "Version 1.0" in ns_txt,
        "D-26 SHEET_SUBTITLE negative (synthetic): a subtitle missing from page 1 fails; a sheet without meta.subtitle prints nothing and is not checked", o2[-400:], tags=["SHEET_SUBTITLE"])
    zf = zipfile.ZipFile(out); fx = [n for n in zf.namelist() if re.match(r"word/footer\d*\.xml$", n)]
    fxml = " ".join(zf.read(n).decode("utf-8") for n in fx); frels = " ".join(zf.read(n.replace("word/", "word/_rels/") + ".rels").decode("utf-8") for n in fx if n.replace("word/", "word/_rels/") + ".rels" in zf.namelist())
    uris = {l.get("uri") for l in pd[0].get_links()} if pd else set()
    ptxt = " ".join(pd[0].get_text().split()) if pd else ""
    rec(rc == 0 and "SOURCE_LINK pass: 2 linked sources clickable on 2 pages" in o and fxml.count("<w:hyperlink") == 2 and url in frels and bare in frels and 'TargetMode="External"' in frels
        and {url, bare} <= uris and "Sources: UAT Process; Synthetic standard %s" % bare in ptxt and "](" not in ptxt,
        "D-27 SOURCE_LINK positive (synthetic): [UAT Process](url) and a bare URL are footer hyperlinks in the docx and link annotations in the PDF", o[-300:] + str(uris), tags=["SOURCE_LINK"])
    nl = os.path.join(tmp, "nolinks.pdf"); shutil.copy(out[:-5] + ".toc.json", nl[:-4] + ".toc.json")
    if pd:
        pdc = fitz.open(pdf)
        for l in [l for l in pdc[0].get_links() if l.get("uri") == url]:
            pdc[0].delete_link(l)
        pdc.save(nl)
    rc3, o3 = lint(q, "--docx", out, "--pdf", nl)
    rec(rc3 == 1 and "FAIL [SOURCE_LINK] footer: no link annotation to %s on page 1;" % url in o3 and "to %s" % bare not in o3,
        "D-27 SOURCE_LINK negative (synthetic): the linked source's annotation removed from page 1 fails; the bare URL still passes", o3[-400:], tags=["SOURCE_LINK"])


def featured(d):
    """1.1 features: three contents levels, a section icon, a custom appendix and a linked index (textbook profile, D-63)."""
    d["meta"]["toc_levels"] = 3; d["meta"]["profile"] = "textbook"
    d["sections"][0]["blocks"].append({"type": "callout", "kind": "tip", "title": "Key Takeaways", "items": ["Filters apply to every card.", "Export from the toolbar."]})
    d["sections"][0]["blocks"].append({"type": "code", "title": "Sample Request", "code": "Show admissions for the West region."})
    d["sections"][0]["icon"] = "info-circle"
    d["sections"][2]["sections"][0]["sections"] = [{"title": "Reading A Card", "blocks": [{"type": "paragraph", "text": "Each card shows one number and its change from last month."}]}]
    d["appendices"] = [{"title": "Quick Reference", "icon": "list-check", "blocks": [{"type": "table", "caption": "File names", "columns": ["File", "Use"],
                        "rows": [["Census_Dashboard_Export_Full_Detail_Report.xlsx", "Monthly export"], ["Census_Notes.md", "Release notes"]]}]}]
    d["index"] = [{"term": "Filters", "section": "2.2"}, {"term": "KPI ribbon", "section": "KPI Ribbon"}, {"term": "Reading a card", "section": "3.1.1"},
                  {"term": "File names", "section": "Quick Reference"}, {"term": "Support", "section": "Support"}]
    return d


# Round 6 (D-62 policy doc type, D-64 accessibility, D-39 TERM_BRAND, SOP Exceptions for the D-42 drift test)
POLICY_REQ = ["Purpose", "Scope", "Policy Statement", "Roles and Responsibilities", "Definitions", "Exceptions", "Related Documents"]


def d39_doc():
    """Synthetic report whose first section is the kept fixtures/D-39 body (assets/fixtures/D-39.input.md)."""
    md = open(os.path.join(ROOT, "assets", "fixtures", "D-39.input.md"), encoding="utf-8").read()
    body = re.search(r"## Why This Is Changing\n(.+)", md, re.S).group(1).strip()
    sec = lambda t, x: {"title": t, "blocks": [{"type": "paragraph", "text": x}]}
    return {"meta": {"title": "Denials Dashboard Moves to the New Site", "doc_type": "report", "audience": "Revenue operations staff",
                     "purpose": "Tell teams about the move", "owner": "Report owner", "version": "1.0", "date": "October 2026",
                     "owner_role": "Report owner", "last_reviewed": "October 2026", "next_review": "April 2027",  # D-65: no REVIEW_DATE warning in the count
                     "sources": ["Synthetic example content (fixtures/D-39 body)"]},
            "sections": [sec("Why This Is Changing", body), sec("Executive Summary", "The view moves to the new site."),
                         sec("Findings", "Saved views and filters move with you."), sec("Recommendations", "Update your bookmarks after the move.")]}


HOUSE_TERMS = {"format": "terminology", "version": "1.0.0", "source": "Synthetic house style for the self-test",
               "sets": {"internal": [{"id": "product-case", "kind": "case", "word": "LedgerLink", "why": "product name: capital L twice"},
                                     {"id": "associates-lower", "kind": "midsentence", "pattern": "\\bAssociates?\\b", "why": "associates is lowercase mid-sentence"}]}}


def house_project(tmp):
    """A folder whose brand.json names the brand (Northwind) and a house terminology file (house.json)."""
    d = os.path.join(tmp, "house")
    os.makedirs(d, exist_ok=True)
    json.dump(HOUSE_TERMS, open(os.path.join(d, "house.json"), "w", encoding="utf-8"))
    json.dump({"name": "Northwind", "org": "Northwind Traders Ltd.", "terminology": "house.json"}, open(os.path.join(d, "brand.json"), "w", encoding="utf-8"))
    return d


def skill_copy(tmp, name, edit):
    """A throwaway copy of the skill with one file changed (edit(root)), for checks that read the skill's own assets."""
    dst = os.path.join(tmp, name)
    if not os.path.isdir(dst):
        shutil.copytree(ROOT, dst, ignore=shutil.ignore_patterns("__pycache__")); edit(dst)
    return os.path.join(dst, "scripts", "lint_doc.py")


def rezip(src, dst, fn):
    """Copy a .docx with fn(name, xml) applied to every word/*.xml part."""
    import zipfile
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename.startswith("word/") and it.filename.endswith(".xml"):
                data = fn(it.filename, data.decode("utf-8")).encode("utf-8")
            zo.writestr(it, data)


def round6_tests(tmp, rec):
    pol = json.load(open(os.path.join(ROOT, "assets", "policy.example.doc.json"), encoding="utf-8"))
    def src(d, name, *extra):
        q = os.path.join(tmp, name); json.dump(d, open(q, "w")); return lint(q, *extra)
    # D-62 (DOC-01): policy doc type
    if selected("DOC_TYPE", "SECTION_REQUIRED"):
        rc, o = lint(os.path.join(ROOT, "assets", "fixtures", "D-62.input.json"), "--draft")
        miss = [r for r in POLICY_REQ if ("FAIL [SECTION_REQUIRED] -: a policy needs a top-level %r section" % r) in o]
        rec("[DOC_TYPE]" not in o and miss == POLICY_REQ and o.count("FAIL [SECTION_REQUIRED]") == 7,
            "D-62 fixture: doc_type policy passes DOC_TYPE and its 7 required sections are checked", o, tags=["DOC_TYPE", "SECTION_REQUIRED"])
    if selected("SECTION_REQUIRED"):
        d = copy.deepcopy(pol); d["sections"] = [s for s in d["sections"] if s["title"] != "Exceptions"]
        rc, o = src(d, "pol_noexc.json")
        rec(rc != 0 and "FAIL [SECTION_REQUIRED] -: a policy needs a top-level 'Exceptions' section" in o, "D-62 negative: a policy without Exceptions fails SECTION_REQUIRED", o, tags=["SECTION_REQUIRED"])
        sop = copy.deepcopy(pol); sop["meta"].update(doc_type="sop", title="Report Access Review Procedure"); sop["meta"].pop("origination_date")
        keep = {"Purpose": "Purpose", "Scope": "Scope", "Roles and Responsibilities": "Roles and Responsibilities", "Policy Statement": "Procedure",
                "Exceptions": "Exceptions", "Related Documents": "References"}
        sop["sections"] = [dict(s, title=keep[s["title"]]) for s in pol["sections"] if s["title"] in keep]
        rc, o = src(sop, "sop_ok.json")
        rec("SECTION_REQUIRED" not in o, "D-42 parity (synthetic SOP): Purpose, Scope, Roles and Responsibilities, Procedure, Exceptions, References meet SECTION_REQUIRED", o, tags=["SECTION_REQUIRED"])
        sop["sections"] = [s for s in sop["sections"] if s["title"] != "Exceptions"]
        rc, o = src(sop, "sop_noexc.json")
        rec(rc != 0 and "FAIL [SECTION_REQUIRED] -: a sop needs a top-level 'Exceptions' section" in o, "D-42 parity negative (synthetic SOP): an SOP without Exceptions fails SECTION_REQUIRED", o, tags=["SECTION_REQUIRED"])
    if selected("META_MISSING"):
        d = copy.deepcopy(pol); d["meta"].pop("origination_date")
        rc, o = src(d, "pol_noorig.json")
        rec(rc != 0 and "FAIL [META_MISSING] meta: a policy needs meta.origination_date" in o, "D-62 negative: a policy without meta.origination_date fails META_MISSING", o, tags=["META_MISSING"])
    if selected("DOC_TYPE"):
        q = os.path.join(tmp, "policy.json"); json.dump(pol, open(q, "w")); out = os.path.join(tmp, "Policy.docx")
        b = subprocess.run([sys.executable, os.path.join(HERE, "build_doc.py"), q, "--out", out, "--pdf"], capture_output=True, text=True, timeout=600)
        pdf = out[:-5] + ".pdf"
        rc, o = lint(q, "--docx", out, *(["--pdf", pdf] if os.path.isfile(pdf) else []))
        rec(b.returncode == 0 and os.path.isfile(pdf) and rc == 0 and " 0 failed" in o, "D-62 policy.example.doc.json builds to Word and PDF and lints 0 failed", b.stdout + b.stderr + o, tags=["DOC_TYPE"])
    # D-39: TERM_BRAND from the shared terminology.json plus the project's house terms (brand profile
    # `terminology`) and the brand name's capitalization; BRAND only as the fallback
    house = house_project(tmp)
    if selected("TERM_BRAND"):
        rc, o = src(d39_doc(), "d39.json")
        hits = re.findall(r'^WARN \[TERM_BRAND\] 1 Why This Is Changing: "(.+)"$', o, re.M)
        rec(rc == 0 and " 0 failed, 2 warnings" in o and sorted(hits) == sorted(["Over 40 -> More than 40", "Jan 17th -> Jan. 17"]),
            "D-39 TERM_BRAND (synthetic doc from fixtures/D-39 body, no house terms): the shared AP rules give 2 warnings, 0 failed", o, tags=["TERM_BRAND"])
        q = os.path.join(house, "d39.json"); json.dump(d39_doc(), open(q, "w")); rc, o = lint(q)
        want = ["Over 40 -> More than 40", "Jan 17th -> Jan. 17", "Ledgerlink -> LedgerLink", "Associates -> associates", "NorthWind -> Northwind"]
        hits = re.findall(r'^WARN \[TERM_BRAND\] 1 Why This Is Changing: "(.+)"$', o, re.M)
        rec(rc == 0 and " 0 failed, 5 warnings" in o and sorted(hits) == sorted(want),
            "D-39 TERM_BRAND with a brand profile naming house terms (house.json case and midsentence rules) and the brand name: 5 warnings, 0 failed", o, tags=["TERM_BRAND"])
        rc, o = src(pol, "pol_terms.json")
        rec("TERM_BRAND" not in o and rc == 0, "D-39 TERM_BRAND stays quiet on the policy example", o, tags=["TERM_BRAND"])
    if selected("BRAND", "TERM_BRAND"):
        lp = skill_copy(tmp, "noterms", lambda r: os.remove(os.path.join(r, "references", "terminology.json")))
        q = os.path.join(house, "d39b.json"); json.dump(d39_doc(), open(q, "w"))
        r = subprocess.run([sys.executable, lp, q], capture_output=True, text=True); o = r.stdout
        rec(r.returncode != 0 and "FAIL [BRAND] 1 Why This Is Changing: write Northwind (exact capitalization), not NorthWind" in o
            and "TERM_BRAND: not run (terminology.json missing)" in o and "WARN [TERM_BRAND]" not in o,
            "negative BRAND: without terminology.json the brand-name casing rule fails and TERM_BRAND says not run", o, tags=["BRAND", "TERM_BRAND"])
    # D-64 (DOC-03): DOC_A11Y on the built user guide
    if selected("DOC_A11Y"):
        import zipfile
        rc, o = src(d39_doc(), "a11y_theme.json")
        sys.path.insert(0, HERE); import lint_doc as _L
        _L.use_brand(BR.load(source=os.path.join(tmp, "a11y_theme.json")))  # judge the theme the lint run above used
        theme = json.load(open(os.path.join(ROOT, "assets", "doc_theme.json"), encoding="utf-8"))["a11y"]
        worst = min(_L.contrast(f, b) for f, b in theme["text_pairs"])
        rec(("DOC_A11Y pass: %d theme text color pairs at 4.5:1 or more" % len(theme["text_pairs"])) in o and worst >= 4.5,
            "D-64 every theme text color is 4.5:1 or more on its background (lowest %.2f:1)" % worst, o, tags=["DOC_A11Y"])
        lp = skill_copy(tmp, "lowcontrast", lambda r: (lambda f, t: json.dump(dict(t, a11y=dict(t["a11y"], text_pairs=t["a11y"]["text_pairs"] + [["secondary", "white"]])), open(f, "w")))(
            os.path.join(r, "assets", "doc_theme.json"), json.load(open(os.path.join(r, "assets", "doc_theme.json")))))
        r = subprocess.run([sys.executable, lp, os.path.join(tmp, "a11y_theme.json")], capture_output=True, text=True)
        low = "FAIL [DOC_A11Y] theme: text color secondary on white is %.2f:1" % _L.contrast("secondary", "white")
        rec(r.returncode != 0 and _L.contrast("secondary", "white") < 4.5 and low in r.stdout, "D-64 negative (synthetic theme pair): secondary-color text on white fails DOC_A11Y (%s)" % low.split("color ", 1)[1], r.stdout, tags=["DOC_A11Y"])
        g = os.path.join(tmp, "Guide.docx"); p = os.path.join(tmp, "doc.json")
        if not os.path.isfile(g[:-5] + ".pdf"):
            subprocess.run([sys.executable, os.path.join(HERE, "build_doc.py"), p, "--out", g, "--pdf"], capture_output=True, text=True, timeout=600)
        if not os.path.isfile(g[:-5] + ".pdf"):
            rec(False, "D-64 user guide build for DOC_A11Y (Word and PDF): NOT RUN, no PDF (LibreOffice needed)", tags=["DOC_A11Y"]); return
        rc, o = lint(p, "--docx", g, "--pdf", g[:-5] + ".pdf")
        z = zipfile.ZipFile(g); sx = z.read("word/styles.xml").decode(); dx = z.read("word/document.xml").decode()
        heads = len(re.findall(r"<w:tblHeader/>", dx))
        rec('<w:lang w:val="en-US"' in sx and heads >= 4 and re.search(r"DOC_A11Y pass: docx w:lang en-US, %d tables with repeating header rows" % heads, o) and rc == 0,
            "D-64 user guide docx: w:lang en-US, a repeating header row on every table (%d), DOC_A11Y pass" % heads, o, tags=["DOC_A11Y"])
        import fitz
        pd = fitz.open(g[:-5] + ".pdf"); lang = pd.xref_get_key(pd.pdf_catalog(), "Lang"); title = pd.metadata.get("title")
        rec(lang == ("string", "en-US") and title == EX["meta"]["title"] and "DOC_A11Y pass: PDF /Lang en-US, /Title" in o,
            "D-64 user guide PDF: /Lang en-US and /Title from meta.title (%s, %r)" % (lang[1], title), o, tags=["DOC_A11Y"])
        bad = os.path.join(tmp, "NoLang.pdf"); shutil.copy(g[:-5] + ".toc.json", bad[:-4] + ".toc.json")
        pd.xref_set_key(pd.pdf_catalog(), "Lang", "null"); pd.set_metadata({"title": ""}); pd.save(bad); pd.close()
        rc, o = lint(p, "--pdf", bad)
        rec("FAIL [DOC_A11Y] pdf: no /Lang en-US in the PDF catalog" in o and "FAIL [DOC_A11Y] pdf: PDF /Title is ''" in o, "D-64 negative: a PDF without /Lang and /Title fails DOC_A11Y", o, tags=["DOC_A11Y"])
        col = colors_for(p)
        cases = [("NoLang.docx", lambda n, x: re.sub(r"<w:tblHeader/>", "", re.sub(r"<w:lang [^>]*/>", "", x)), ["no document language", "data tables repeat their header row"],
                  "D-64 negative: a docx without w:lang and w:tblHeader fails DOC_A11Y"),
                 ("NoAlt.docx", lambda n, x: re.sub(r'descr="Admissions[^"]*"', 'descr=""', x), ["pictures have no alt text"], "D-64 negative: a picture without alt text fails DOC_A11Y"),
                 ("Gray.docx", lambda n, x: x.replace('<w:color w:val="%s"/>' % col["muted_text"], '<w:color w:val="%s"/>' % col["muted"]), ["text colors %s are not on the 4.5:1 list" % col["muted"]],
                  "D-64 negative: text in the muted (decorative) gray %s fails DOC_A11Y; text uses muted_text" % col["muted"])]
        for name, fn, msgs, title in cases:
            q = os.path.join(tmp, name); rezip(g, q, fn)
            rc, o = lint(p, "--docx", q)
            rec(rc != 0 and all(("FAIL [DOC_A11Y] docx:" in ln and any(m in ln for m in msgs)) for ln in o.splitlines() if "[DOC_A11Y]" in ln)
                and all(any(m in ln for ln in o.splitlines() if ln.startswith("FAIL [DOC_A11Y]")) for m in msgs), title, o, tags=["DOC_A11Y"])


# Round 6 chunk B (D-63 lean guide profile, D-65 review dates, D-66 sheet formats)
ROUND6B_CODES = {"GUIDE_PROFILE", "REVIEW_DATE", "SHEET_FORMAT_KNOWN", "SHEET_PAGE_FIT", "SECTION_REQUIRED"}
ROUND7_CODES = {"FIG_WIDTH", "LAYOUT"}
ALL_CODES |= ROUND7_CODES


def round7_tests(tmp, rec):
    """D-77 (GAP-26): landscape top-level sections and FIG_WIDTH. All inputs synthetic."""
    from PIL import Image, ImageDraw
    from docx import Document as _Doc
    shots(tmp)
    os.makedirs(os.path.join(tmp, "figures"), exist_ok=True)
    im = Image.new("RGB", (2700, 900), "white"); dr = ImageDraw.Draw(im)
    for i in range(6):
        dr.rectangle([40 + i * 440, 60, 420 + i * 440, 840], outline="#004B87", width=6)
    im.save(os.path.join(tmp, "figures", "wide_matrix.png"))
    fig = {"type": "figure", "image": "figures/wide_matrix.png", "caption": "Wide approval matrix across six stages", "alt": "Six columns showing each approval stage side by side", "width_in": 9.0}
    def doc(orient="landscape", sub=False):
        d = copy.deepcopy(EX)
        s = {"title": "Approval Matrix", "blocks": [{"type": "paragraph", "text": "The matrix below shows every approval stage side by side."}, dict(fig)]}
        if sub:
            d["sections"][0].setdefault("sections", []).append(dict(s, orientation=orient))
        else:
            if orient:
                s["orientation"] = orient
            d["sections"].insert(1, s)
        return d
    q, out, b = build(tmp, doc(), "Landscape")
    rc, o = lint(q, "--docx", out)
    secs = [(x.orientation, round(x.page_width.inches, 2)) for x in _Doc(out).sections] if b.returncode == 0 else []
    land = [i for i, (o_, w) in enumerate(secs) if o_ == 1 and w == 11.0]
    widths = [round(s.width.inches, 2) for s in _Doc(out).inline_shapes] if b.returncode == 0 else []
    from docx.oxml.ns import qn as _qn
    tabs = [(t.get(_qn("w:val")), t.get(_qn("w:pos"))) for t in _Doc(out).sections[land[0]].footer._element.iter(_qn("w:tab")) if t.get(_qn("w:val"))] if land else []
    rec(("right", "12960") in tabs and ("clear", "9360") in tabs, "D-77 FIG_WIDTH positive (synthetic): the landscape footer puts the page number at the 9.0 in right margin (tabs %s)" % tabs, "", tags=["FIG_WIDTH"])
    rec(b.returncode == 0 and "[FIG_WIDTH]" not in o and "[LAYOUT]" not in o and len(land) == 1 and land[0] + 1 < len(secs) and secs[land[0] + 1] == (0, 8.5) and 9.0 in widths,
        "D-77 FIG_WIDTH / LAYOUT positive (synthetic): a landscape section prints a 9.0 in figure and the next section is portrait (sections %s, figure widths %s)" % (secs, widths), b.stderr[-400:] + o, tags=["FIG_WIDTH", "LAYOUT"])
    q = os.path.join(tmp, "land_none.json"); json.dump(doc(orient=None), open(q, "w")); rc, o = lint(q)
    rec("FAIL [FIG_WIDTH]" in o and "6.5 in" in o, "D-77 FIG_WIDTH negative (synthetic): a 9.0 in figure in a portrait section fails", o, tags=["FIG_WIDTH"])
    q = os.path.join(tmp, "land_sub.json"); json.dump(doc(sub=True), open(q, "w")); rc, o = lint(q)
    rec("FAIL [LAYOUT]" in o and "orientation" in o, "D-77 LAYOUT negative (synthetic): orientation on a subsection fails", o, tags=["LAYOUT"])


def docx_text(path):
    """(heading-ish paragraph texts, whole body xml) of a built .docx."""
    from docx import Document
    doc = Document(path)
    return [p.text.strip() for p in doc.paragraphs if p.text.strip()], doc.element.body.xml


def build(tmp, d, name, pdf=False):
    q = os.path.join(tmp, name + ".json"); json.dump(d, open(q, "w"))
    out = os.path.join(tmp, name + ".docx")
    b = subprocess.run([sys.executable, os.path.join(HERE, "build_doc.py"), q, "--out", out] + (["--pdf"] if pdf else []), capture_output=True, text=True, timeout=900)
    return q, out, b


def round6b_tests(tmp, rec):
    def src(d, name, *extra):
        q = os.path.join(tmp, name); json.dump(d, open(q, "w")); return lint(q, *extra)
    G = ["GUIDE_PROFILE"]
    if selected("GUIDE_PROFILE"):
        # D-63 (DOC-02): lean profile; the user-guide example is the lean example
        q, out, b = build(tmp, copy.deepcopy(EX), "Lean")
        rc, o = lint(q, "--docx", out)
        paras, body = docx_text(out) if b.returncode == 0 else ([], "")
        heads = [p for p in paras if re.match(r"^(\d+(\.\d+)*  |Appendix [A-Z]: )", p)]
        boxes = [p for p in paras if p in ("What You Will Learn", "Key Takeaways", "Index") or p.endswith(": Index")]
        ok = (b.returncode == 0 and rc == 0 and "GUIDE_PROFILE pass: lean (meta.profile): 2 numbered task topic(s)" in o and "2.1  Open the Dashboard" in heads and "4  Exporting Data" in heads
              and "6  Reference" in heads and "Appendix A: Glossary" in heads and 'TOC \\o' in body and "Contents" in paras and not boxes and "What This Guide Covers" in " ".join(heads))
        rec(ok, "D-63 lean guide (meta.profile lean) builds with numbered task topics ['2.1  Open the Dashboard', '4  Exporting Data'], a Reference chapter, contents and glossary; no objectives or takeaways boxes, no index",
            b.stdout + b.stderr + o + repr(heads[:30]) + repr(boxes), tags=G)
        tb = json.load(open(os.path.join(ROOT, "assets", "textbook.example.doc.json"), encoding="utf-8"))
        if not os.path.isfile(os.path.join(tmp, "figures", "overview_illustration.png")):
            from PIL import Image
            os.makedirs(os.path.join(tmp, "figures"), exist_ok=True); Image.new("RGB", (1600, 900), "white").save(os.path.join(tmp, "figures", "overview_illustration.png"))
        q, out, b = build(tmp, tb, "Textbook63")
        rc, o = lint(q, "--docx", out)
        paras, body = docx_text(out) if b.returncode == 0 else ([], "")
        ok = b.returncode == 0 and rc == 0 and "GUIDE_PROFILE pass: textbook (meta.profile)" in o and body.count("What You Will Learn") == 3 and body.count("Key Takeaways") == 3 and any(p.endswith("Index") and p.startswith("Appendix") for p in paras)
        rec(ok, "D-63 meta.profile textbook still gets objectives and takeaways boxes and an index", b.stdout + b.stderr + o, tags=G)
        d = copy.deepcopy(EX); d["meta"].pop("profile")
        rc, o = src(d, "d63_default.json", "--draft")
        rec(rc == 0 and "GUIDE_PROFILE pass: lean (default for a user-guide for business users)" in o, "D-63 default: a user guide for business users with no meta.profile is checked as lean", o, tags=G)
        d = copy.deepcopy(EX); d["sections"] = [s for s in d["sections"] if s["title"] != "Reference"]
        rc, o = src(d, "d63_noref.json", "--draft")
        rec(rc != 0 and "FAIL [GUIDE_PROFILE] -: a lean guide needs a top-level 'Reference' chapter" in o, "D-63 negative: a lean guide without a Reference chapter fails GUIDE_PROFILE", o, tags=G)
        d = copy.deepcopy(EX); d["sections"][1]["objectives"] = ["Open the dashboard from the portal.", "Set the filters you need."]
        rc, o = src(d, "d63_obj.json", "--draft")
        rec(rc != 0 and "FAIL [GUIDE_PROFILE] 2 Getting Started: a lean guide has no 'What You Will Learn' box" in o, "D-63 negative: a lean guide with an objectives box fails GUIDE_PROFILE", o, tags=G)
        d = copy.deepcopy(EX); d["sections"][0]["sections"][1]["title"] = "What You Will Learn"; d["index"] = [{"term": "Filters", "section": "2.2"}]
        rc, o = src(d, "d63_index.json", "--draft")
        rec(rc != 0 and "FAIL [GUIDE_PROFILE] index: a lean guide has no index" in o and "FAIL [GUIDE_PROFILE] 1.2 What You Will Learn: 'What You Will Learn' is a textbook box" in o,
            "D-63 negative: a lean guide with an index and a What You Will Learn heading fails GUIDE_PROFILE", o, tags=G)
        d = copy.deepcopy(EX); d["sections"][3]["blocks"].append({"type": "callout", "kind": "tip", "title": "Key Takeaways", "items": ["Export from the toolbar.", "Keep files inside the network."]})
        rc, o = src(d, "d63_kt.json", "--draft")
        rec(rc != 0 and "FAIL [GUIDE_PROFILE] 4 Exporting Data: 'Key Takeaways' is a textbook box" in o, "D-63 negative (synthetic): a lean guide with a Key Takeaways callout fails GUIDE_PROFILE", o, tags=G)
        d = copy.deepcopy(EX); d["meta"].pop("profile"); d["meta"]["audience"] = "Analytics engineers who maintain the dashboard"; d["sections"] = [s for s in d["sections"] if s["title"] != "Reference"]
        rc, o = src(d, "d63_dev.json", "--draft")
        rec(rc == 0 and "[GUIDE_PROFILE]" not in o and "GUIDE_PROFILE: not checked (developer audience" in o, "D-63 developer audience (synthetic): no default profile, so the same guide is not checked as lean", o, tags=G)
        d = copy.deepcopy(EX); d["meta"]["profile"] = "compact"
        rc, o = src(d, "d63_bad.json", "--draft")
        rec(rc != 0 and "FAIL [GUIDE_PROFILE] meta: profile must be one of lean, textbook" in o, "D-63 negative (synthetic): an unknown meta.profile fails GUIDE_PROFILE", o, tags=G)

    R = ["REVIEW_DATE"]
    if selected("REVIEW_DATE"):
        # D-65 (DOC-04): owner role, review dates, printed copies line on sheets
        d = copy.deepcopy(EX); d["meta"].pop("next_review")
        rc, o = src(d, "d65_none.json", "--draft")
        rec(rc == 0 and "WARN [REVIEW_DATE] meta: next review date missing: set meta.next_review (with meta.owner_role and meta.last_reviewed)" in o,
            "D-65 REVIEW_DATE: a doc.json without meta.next_review warns (not a failure)", o, tags=R)
        rc, o = src(copy.deepcopy(EX), "d65_set.json", "--draft")
        rec(rc == 0 and "[REVIEW_DATE]" not in o and "REVIEW_DATE pass: owner role Dashboard owner, Clinic Operations Analytics, last reviewed October 2026, next review October 2027" in o,
            "D-65 REVIEW_DATE stays quiet when meta.next_review is set", o, tags=R)
        out = os.path.join(tmp, "Lean.docx")
        if not os.path.isfile(out):
            _, out, _ = build(tmp, copy.deepcopy(EX), "Lean")
        from docx import Document
        tabs = [[[c.text for c in r.cells] for r in t.rows] for t in Document(out).tables] if os.path.isfile(out) else []
        ri = next((i for i, t in enumerate(tabs) if t[0] == ["Owner Role", "Last Reviewed", "Next Review"]), -1)
        vi = next((i for i, t in enumerate(tabs) if t[0][:2] == ["Version", "Date"]), -1)
        rec(ri >= 0 and vi == ri + 1 and tabs[ri][1] == ["Dashboard owner, Clinic Operations Analytics", "October 2026", "October 2027"],
            "D-65 user guide docx: the Revision History shows Owner Role, Last Reviewed and Next Review above the revisions", repr(tabs[ri:vi + 1] if ri >= 0 else tabs[-2:]), tags=R)
        qs = json.load(open(os.path.join(ROOT, "assets", "quick-start.example.doc.json"), encoding="utf-8"))
        url = "https://example.com/review-quick-start"
        want = "Printed copies may be out of date. Latest version: %s \u00b7 Owner Role: Review lead \u00b7 Last Reviewed: October 2026 \u00b7 Next Review: April 2027" % url
        foot = lambda p: re.sub(r"\s+", " ", " ".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", "".join(s.footer._element.xml for s in Document(p).sections))))
        pdf_ok = bool(shutil.which("soffice") or shutil.which("libreoffice"))
        d = copy.deepcopy(qs); d["meta"]["latest_url"] = url
        q, out, b = build(tmp, d, "QSReview", pdf=pdf_ok)
        ft = foot(out) if b.returncode == 0 else ""
        detail = b.stdout + b.stderr + ft
        ok = b.returncode == 0 and want in ft
        if ok and pdf_ok:
            import fitz
            pd = fitz.open(out[:-5] + ".pdf"); rc, o = lint(q, "--docx", out, "--pdf", out[:-5] + ".pdf")
            ok = rc == 0 and all(want in " ".join(pg.get_text().split()) and url in [l.get("uri") for l in pg.get_links()] for pg in pd)
            detail += o
        rec(ok, "D-65 sheet footer (synthetic latest_url): printed copies line, latest version link, owner role, last reviewed and next review%s" % (" on every PDF page, lint 0 failed" if pdf_ok else " (docx only: LibreOffice missing)"), detail, tags=R)
        q, out, b = build(tmp, copy.deepcopy(qs), "QSNoUrl")
        ft = foot(out) if b.returncode == 0 else ""
        rec(b.returncode == 0 and "Printed copies may be out of date. \u00b7 Owner Role: Review lead" in ft and "Latest version" not in ft,
            "D-65 sheet footer without meta.latest_url: printed copies line, no Latest version label", b.stderr + ft, tags=R)
        d = copy.deepcopy(qs); d["meta"]["latest_url"] = "www.example.com/quick-start"
        rc, o = src(d, "d65_badurl.json", "--draft")
        rec(rc != 0 and "FAIL [LINK] meta.latest_url: latest_url must be a full https:// address" in o, "D-65 negative (synthetic): a meta.latest_url without https:// fails", o, tags=R)

    K = ["SHEET_FORMAT_KNOWN"]
    if selected("SHEET_FORMAT_KNOWN", "SHEET_PAGE_FIT", "SECTION_REQUIRED"):
        # D-66 (DOC-05): sheet rules for the catalog sheet types
        rc, o = src(dg_doc(pages=2), "d66_dg2.json")
        rec(rc == 0 and " 0 failed" in o and "SHEET_FORMAT_KNOWN pass: dashboard-guide rules (2 pages per section, 8 required sections)" in o,
            "D-66 row scenario (synthetic dashboard-guide sheet, pages_per_section 2): SHEET_FORMAT_KNOWN pass, 0 failed", o, tags=K)
        rc, o = src(dg_doc(pages=3), "d66_dg3.json")
        rec(rc != 0 and "FAIL [SHEET_PAGE_FIT] meta.sheet: pages_per_section 3 is over the dashboard-guide limit of 2 pages" in o,
            "D-66 negative (synthetic): pages_per_section 3 on a dashboard-guide fails SHEET_PAGE_FIT", o, tags=K + ["SHEET_PAGE_FIT"])
        rc, o = src(dg_doc(pages=2, drop="Known Limitations"), "d66_dgmiss.json")
        rec(rc != 0 and "FAIL [SECTION_REQUIRED] section 1: a dashboard-guide section needs a level-2 'Known Limitations' section" in o,
            "D-66 negative (synthetic): a dashboard-guide without Known Limitations fails SECTION_REQUIRED", o, tags=K + ["SECTION_REQUIRED"])
        fq = qs_doc([{"title": "Report Review Questions", "sections": [{"title": t, "blocks": [{"type": "paragraph", "text": "Synthetic answer text for the self-test."}]} for t in ("Questions", "Get Help")]}], doc_type="faq-sheet")
        fq["meta"].update(title="Synthetic Review FAQ", sheet={"pages_per_section": 2})
        rc, o = src(fq, "d66_faq.json")
        rc3, o3 = src(dict(fq, meta=dict(fq["meta"], sheet={"pages_per_section": 3})), "d66_faq3.json")
        rec(rc == 0 and "SHEET_FORMAT_KNOWN pass: faq-sheet rules (2 pages per section)" in o and "FAIL [SHEET_PAGE_FIT] meta.sheet: pages_per_section 3 is over the faq-sheet limit of 2 pages" in o3,
            "D-66 faq-sheet (synthetic) allows two pages per section (3 fails)", o + o3, tags=K + ["SHEET_PAGE_FIT"])
        def drop_dg(root):
            p = os.path.join(root, "assets", "sheet_formats.json"); sf = json.load(open(p, encoding="utf-8")); sf["formats"].pop("dashboard-guide"); json.dump(sf, open(p, "w"), indent=1)
        lp = skill_copy(tmp, "skill_nodg", drop_dg)
        q = os.path.join(tmp, "d66_unknown.json"); json.dump(dg_doc(pages=1), open(q, "w"))
        r = subprocess.run([sys.executable, lp, q], capture_output=True, text=True)
        rec(r.returncode != 0 and "FAIL [SHEET_FORMAT_KNOWN] meta: 'dashboard-guide' is a catalog sheet type with no entry in assets/sheet_formats.json" in r.stdout,
            "D-66 negative: a catalog sheet type with no sheet_formats.json entry fails SHEET_FORMAT_KNOWN", r.stdout + r.stderr, tags=K)
        comp = os.path.join(os.path.dirname(ROOT), "material-composer", "scripts")
        if not os.path.isfile(os.path.join(comp, "catlib.py")):
            rec(False, "D-66 sheet_formats.json matches the composer catalog: NOT RUN (material-composer not installed alongside)", tags=K)
        else:
            sys.path.insert(0, comp); import catlib
            sys.path.insert(0, HERE); from build_doc import SHEET as SH, sheet_cfg as cfg_of, sheet_format as fmt_of
            cat = catlib.load(); seen, bad = [], []
            for n, e in cat["formats"].items():
                if catlib.renderer_of(cat, e)[0] != "doc-sheet":
                    continue
                dt = e.get("doc_type") or n; seen.append(dt)
                cp = e.get("pages") or cat["families"][e["family"]].get("pages", 1)
                cs = [a["title"] for a in e.get("anatomy", []) if a.get("required")]
                f = SH["formats"].get(dt)
                bs = None if f is None else (f.get("sections") or (fmt_of(dt).get("structure") or {}).get("subsections"))
                if f is None or int(cfg_of({"doc_type": dt})["pages_per_section"]) != cp or (bs is not None and bs != cs):
                    bad.append(dt)
            ok = not bad and len(seen) == 18 and sorted(SH["catalog_sheet_types"]) == sorted(seen)
            rec(ok, "D-66 sheet_formats.json matches the composer catalog: %d sheet types, pages and required sections (mismatches: %s)" % (len(seen), ", ".join(bad) or "none"), repr(bad), tags=K)

def dg_doc(pages=2, drop=None):
    """Synthetic dashboard-guide sheet (D-66): one top-level section with the catalog's 8 level-2 sections."""
    want = ["Purpose", "Access", "How It Is Organized", "Definitions", "How to Use It", "Data Sources and Refresh", "Known Limitations", "Get Help"]
    sec = {"title": "Synthetic Review Dashboard", "sections": [{"title": t, "blocks": [{"type": "paragraph", "text": "Synthetic text about %s for the self-test." % t.lower()}]} for t in want if t != drop]}
    d = qs_doc([sec], doc_type="dashboard-guide"); d["meta"].update(title="Synthetic Review Dashboard Guide", sheet={"pages_per_section": pages})
    return d


NW_SVG = ('<svg xmlns="http://www.w3.org/2000/svg" width="360" height="120" viewBox="0 0 360 120">'
          '<rect x="6" y="6" width="108" height="108" rx="18" fill="#0B7A75"/><rect x="134" y="38" width="220" height="44" rx="8" fill="#14313A"/></svg>')


def pictures(docx_path):
    """[(alt text, part name)] for every picture in the body, headers and footers of a .docx."""
    import zipfile
    z = zipfile.ZipFile(docx_path)
    return [(a, n) for n in z.namelist() if re.match(r"word/(document|header\d*|footer\d*)\.xml$", n)
            for a in re.findall(r'<wp:docPr [^>]*descr="([^"]*)"', z.read(n).decode("utf-8"))]


def part_text(docx_path, pattern):
    import zipfile
    z = zipfile.ZipFile(docx_path)
    return " ".join("".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", z.read(n).decode("utf-8"))) for n in z.namelist() if re.match(pattern, n))


def brand_tests(tmp, rec):
    """The brand comes from the project: none (neutral, no logo), a detected one (Tailwind color, SVG logo, LICENSE
    owner) and an explicit meta.brand profile. All synthetic."""
    import zipfile
    import build_doc
    B = ["BUILD"]
    # 1. no brand: the neutral build in tmp has no logo and lint warns LOGO_NONE
    g, p = os.path.join(tmp, "Guide.docx"), os.path.join(tmp, "doc.json")
    rc, o = lint(p, "--docx", g)
    logos = [a for a, _ in pictures(g) if a == "Logo" or a.endswith(" logo")] if os.path.isfile(g) else ["missing"]
    rec(rc == 0 and not logos and "WARN [LOGO_NONE] docx:" in o and "©" not in part_text(g, r"word/footer\d*\.xml$"),
        "a project with no brand builds with no logo or copyright line (LOGO_NONE warns, 0 failed)", o + repr(logos), tags=B)
    # 2. detected brand: Tailwind primary, public/logo.svg, LICENSE copyright holder
    nw = tempfile.mkdtemp(prefix="doctest_nw_")
    try:
        os.makedirs(os.path.join(nw, ".git")); os.makedirs(os.path.join(nw, "public")); shots(nw)
        open(os.path.join(nw, "tailwind.config.js"), "w").write("module.exports = { theme: { extend: { colors: { primary: '#0B7A75' } } } }\n")
        open(os.path.join(nw, "public", "logo.svg"), "w").write(NW_SVG)
        open(os.path.join(nw, "LICENSE"), "w").write("MIT License\n\nCopyright (c) 2025 Northwind Traders Ltd.\n")
        q = os.path.join(nw, "doc.json"); json.dump(EX, open(q, "w")); out = os.path.join(nw, "Guide.docx")
        b = subprocess.run([sys.executable, os.path.join(HERE, "build_doc.py"), q, "--out", out], capture_output=True, text=True, timeout=600)
        br = BR.load(source=q)
        rc, o = lint(q, "--docx", out)
        ok = b.returncode == 0 and os.path.isfile(out)
        xml = (zipfile.ZipFile(out).read("word/document.xml").decode() + zipfile.ZipFile(out).read("word/styles.xml").decode()) if ok else ""
        cover = [a for a, _ in pictures(out) if a == br.logo_alt()] if ok else []
        foot = part_text(out, r"word/footer\d*\.xml$") if ok else ""
        rec(ok and br.colors["primary"] == "0B7A75" and ('w:val="0B7A75"' in xml or 'w:color="0B7A75"' in xml) and "logo: yes" in b.stdout,
            "Northwind project (synthetic): the build picks up the Tailwind primary 0B7A75 and uses it in the .docx", b.stdout + b.stderr, tags=B)
        rec(ok and len(cover) == 1 and br.logo_alt() == "Northwind Traders Ltd. logo", "Northwind project: one cover picture, the project logo (alt %r)" % br.logo_alt(), repr(pictures(out) if ok else []), tags=B)
        rec(ok and "Northwind Traders Ltd." in foot and ("©%s Northwind Traders Ltd." % build_doc.brand_year(EX["meta"])) in foot,
            "Northwind project: the footer carries the LICENSE owner (%s)" % foot[:80], foot, tags=B)
        rec(rc == 0 and "LOGO_NONE" not in o and "[LOGO]" not in o and "BRAND: Northwind Traders Ltd. (detected in" in o,
            "Northwind project: lint judges the same brand (palette, logo on the cover), 0 failed", o, tags=B)
        sq = os.path.join(nw, "sheet.json"); json.dump(qs_doc([qs_role(1, 200), qs_role(2, 200)]), open(sq, "w")); so = os.path.join(nw, "Sheet.docx")
        b = subprocess.run([sys.executable, os.path.join(HERE, "build_doc.py"), sq, "--out", so], capture_output=True, text=True, timeout=600)
        rc, o = lint(sq, "--docx", so)
        n = len([a for a, _ in pictures(so) if a == br.logo_alt()]) if b.returncode == 0 else 0
        rec(b.returncode == 0 and rc == 0 and n == 2 and "Northwind Traders Ltd." in part_text(so, r"word/footer\d*\.xml$"),
            "Northwind sheet: the logo in each of the 2 title bands and the owner in the footer, 0 failed (%d logos)" % n, b.stdout + o, tags=B)
        nb = os.path.join(nw, "Override.docx"); json.dump(dict(EX, meta=dict(EX["meta"], design="branded")), open(q, "w"))
        r = subprocess.run([sys.executable, lp_of("build_doc.py"), q, "--out", nb, "--brand", os.path.join(tmp, "brand.json")], capture_output=True, text=True, timeout=600)
        rec(r.returncode == 0 and os.path.isfile(nb) and not [a for a, _ in pictures(nb) if a.endswith(" logo")] and "logo: none" in r.stdout,
            "--brand overrides the detected brand (an empty profile builds the branded design with no logo)", r.stdout + r.stderr, tags=B)
    finally:
        shutil.rmtree(nw, ignore_errors=True)
    # 3. explicit profile via meta.brand (legacy color names accepted), custom footer and a PNG logo
    from PIL import Image, ImageDraw
    kit = os.path.join(tmp, "kit"); os.makedirs(kit, exist_ok=True)
    im = Image.new("RGBA", (400, 100), (0, 0, 0, 0)); ImageDraw.Draw(im).rectangle([10, 10, 390, 90], fill=(122, 31, 92, 255)); im.save(os.path.join(kit, "mark.png"))
    json.dump({"name": "Contoso", "org": "Contoso Group", "footer": "©{year} {org}. Internal use only.", "logo": "mark.png",
               "colors": {"blue": "#7A1F5C"}}, open(os.path.join(kit, "brand.json"), "w", encoding="utf-8"))
    d = copy.deepcopy(EX); d["meta"]["brand"] = "kit/brand.json"; q = os.path.join(tmp, "metabrand.json"); json.dump(d, open(q, "w")); out = os.path.join(tmp, "MetaBrand.docx")
    b = subprocess.run([sys.executable, os.path.join(HERE, "build_doc.py"), q, "--out", out], capture_output=True, text=True, timeout=600)
    rc, o = lint(q, "--docx", out)
    br = BR.load(source=q); ok = b.returncode == 0 and os.path.isfile(out)
    foot = part_text(out, r"word/footer\d*\.xml$") if ok else ""
    want = br.footer(build_doc.brand_year(d["meta"]))
    rec(ok and rc == 0 and br.colors["primary"] == "7A1F5C" and "7A1F5C" in zipfile.ZipFile(out).read("word/styles.xml").decode()
        and [a for a, _ in pictures(out)].count("Contoso logo") == 1 and want in foot and foot.lower().count("internal use only") == 1,
        "meta.brand profile (synthetic Contoso kit): legacy color name maps to primary, its logo on the cover, its footer once (%r)" % foot[:90], b.stdout + o + foot, tags=B)


def embedded_parts(docx_path):
    """{family: [(style tag, de-obfuscated font bytes)]} read back from a .docx's fontTable and font parts."""
    import zipfile
    z = zipfile.ZipFile(docx_path)
    ft = z.read("word/fontTable.xml").decode("utf-8")
    rels = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', z.read("word/_rels/fontTable.xml.rels").decode("utf-8"))) if "word/_rels/fontTable.xml.rels" in z.namelist() else {}
    out = {}
    for fam, inner in re.findall(r'<w:font w:name="([^"]+)">(.*?)</w:font>', ft, re.S):
        for tag, rid, key in re.findall(r'<w:(embed\w+) r:id="([^"]+)" w:fontKey="([^"]+)"/>', inner):
            raw = z.read("word/" + rels[rid])
            k = bytes.fromhex(key.strip("{}").replace("-", ""))[::-1]
            out.setdefault(fam, []).append((tag, bytes(b ^ k[i % 16] for i, b in enumerate(raw[:32])) + raw[32:]))
    return out


def dot_row_twips(docx_path):
    """Exact row heights (twips) of the step-diagram dot rows in a .docx."""
    import zipfile
    return sorted({int(v) for v in re.findall(r'<w:trHeight w:val="(\d+)" w:hRule="exact"/>', zipfile.ZipFile(docx_path).read("word/document.xml").decode("utf-8"))})


def font_tests(tmp, rec):
    """Fonts come from the project: the bundled fallback body font is embedded and used by the PDF converter even when it
    isn't installed; a project's own font file is named, embedded and drives the font-based layout (dot rows)."""
    import zipfile
    import build_doc
    B = ["BUILD", "FONTS"]
    g, p = os.path.join(tmp, "Guide.docx"), os.path.join(tmp, "doc.json")
    nb = BR.load(source=p)
    emb = embedded_parts(g) if os.path.isfile(g) else {}
    reg = next((raw for tag, raw in emb.get(nb.fonts["body"], []) if tag == "embedRegular"), b"")
    src = open(nb.font_path("body"), "rb").read() if nb.font_path("body") else b"x"
    z = zipfile.ZipFile(g) if os.path.isfile(g) else None
    ok = z is not None and reg == src and "<w:embedTrueTypeFonts/>" in z.read("word/settings.xml").decode() and 'Extension="odttf"' in z.read("[Content_Types].xml").decode()
    rec(ok and nb.fonts["body"] == BR.FALLBACK_FONTS["body"] and nb.font_status()["body"] == "bundled" and BR.FALLBACK_FONTS["mono"] in emb,
        "no project fonts: the bundled %s and %s are embedded (obfuscated parts read back byte for byte)" % (BR.FALLBACK_FONTS["body"], BR.FALLBACK_FONTS["mono"]), repr(list(emb)), tags=B)
    pdf = g[:-5] + ".pdf"
    if os.path.isfile(pdf):
        import fitz
        import lint_doc as _L
        _L.use_brand(nb)
        installed = (BR.FALLBACK_FONTS["body"].lower(), "regular") in BR.font_index(BR.system_font_dirs())
        rec(BR.FALLBACK_FONTS["body"] in _L.pdf_families(fitz.open(pdf)),
            "the PDF is drawn in the embedded %s (%s on this machine)" % (BR.FALLBACK_FONTS["body"], "also installed" if installed else "not installed, so the converter used the .docx's embedded copy"), tags=B)
    of = other_font()
    if not of:
        rec(False, "project font file case: NOT RUN (no installed text font to play the project's font)", tags=B); return
    fam, regf, boldf = of
    pj = tempfile.mkdtemp(prefix="doctest_font_")
    try:
        os.makedirs(os.path.join(pj, ".git")); os.makedirs(os.path.join(pj, "fonts"))
        for f in (regf, boldf):
            if f:
                shutil.copy(f, os.path.join(pj, "fonts", os.path.basename(f)))
        q = os.path.join(pj, "sheet.json"); json.dump(qs_doc([qs_role(1, 200), qs_role(2, 200)]), open(q, "w")); out = os.path.join(pj, "Sheet.docx")
        b = subprocess.run([sys.executable, os.path.join(HERE, "build_doc.py"), q, "--out", out], capture_output=True, text=True, timeout=600)
        rc, o = lint(q, "--docx", out)
        pb = BR.load(source=q)
        ok = b.returncode == 0 and os.path.isfile(out)
        doc = zipfile.ZipFile(out).read("word/document.xml").decode() + zipfile.ZipFile(out).read("word/styles.xml").decode() if ok else ""
        emb = embedded_parts(out) if ok else {}
        mine = [raw for _, raw in emb.get(fam, [])]
        rec(ok and pb.fonts["body"] == fam and pb.font_status()["body"] == "project file" and ('w:ascii="%s"' % fam) in doc
            and open(regf, "rb").read() in mine and ("body %s (project file)" % fam) in b.stdout,
            "project font (synthetic project shipping fonts/%s): the .docx names %s and embeds the project's own file" % (os.path.basename(regf), fam), b.stdout + b.stderr + repr(list(emb)), tags=B)
        rec(rc == 0 and "[FONTS]" not in o and "[FONT_EMBED]" not in o, "project font: lint accepts the project font and its embedding, 0 failed", o, tags=B)
        nq = os.path.join(tmp, "sheet_neutral.json"); json.dump(qs_doc([qs_role(1, 200), qs_role(2, 200)]), open(nq, "w")); no = os.path.join(tmp, "SheetNeutral.docx")
        subprocess.run([sys.executable, os.path.join(HERE, "build_doc.py"), nq, "--out", no], capture_output=True, text=True, timeout=600)
        want = lambda br: int(round(br.metrics("body")["line"] * 18 / 2 * 20))
        mine_rows, neutral_rows = (dot_row_twips(out) if ok else []), (dot_row_twips(no) if os.path.isfile(no) else [])
        near = lambda w, rows: any(abs(w - v) <= 1 for v in rows)
        rec(mine_rows and neutral_rows and near(want(pb), mine_rows) and near(want(nb), neutral_rows) and abs(want(pb) - want(nb)) > 2,
            "font-based layout: the step-diagram dot rows follow each font's line height (%s %s twips, %s %s twips)" % (fam, mine_rows, nb.fonts["body"], neutral_rows), "", tags=B)
    finally:
        shutil.rmtree(pj, ignore_errors=True)


def lp_of(name):
    return os.path.join(HERE, name)


def main():
    only_codes(ALL_CODES)
    tmp = tempfile.mkdtemp(prefix="doctest_"); shots(tmp)
    open(os.path.join(tmp, "brand.json"), "w").write("{}\n")  # neutral brand for every case here; also stops the project-root walk
    passed = total = 0; lines = []
    def rec(ok, name, detail="", tags=None):
        nonlocal passed, total
        tags = tags_for(name, tags); TAGGED.append((name, sorted(tags)))
        if not selected(*tags):
            return
        total += 1; passed += ok; lines.append(("ok  " if ok else "FAIL") + " " + name + ("" if ok else "\n" + detail[-600:]))
    p = os.path.join(tmp, "doc.json"); json.dump(EX, open(p, "w"))
    if selected("BUILD"):
        b = subprocess.run([sys.executable, os.path.join(HERE, "build_doc.py"), p, "--out", os.path.join(tmp, "Guide.docx"), "--pdf"], capture_output=True, text=True, timeout=600)
        rec(b.returncode == 0 and os.path.isfile(os.path.join(tmp, "Guide.docx")), "example builds to .docx", b.stdout + b.stderr)
        pdf_ok = os.path.isfile(os.path.join(tmp, "Guide.pdf"))
        if HAS_PDF:
            rec(pdf_ok, "example exports to PDF with %s" % ("LibreOffice" if HAS_LO else "Microsoft Word"), b.stdout + b.stderr)
        else:
            lines.append("NOT RUN PDF export (no LibreOffice or Word)")
        rc, out = lint(p, "--docx", os.path.join(tmp, "Guide.docx"), *(["--pdf", os.path.join(tmp, "Guide.pdf")] if pdf_ok else []))
        rec(rc == 0 and " 0 failed" in out, "built Word%s file lints clean (TOC field, alt text, fonts, palette%s)" % (" and PDF" if pdf_ok else "", ", bookmarks, embedded fonts, contents page numbers" if pdf_ok else ""), out)
        if pdf_ok:
            import fitz
            doc = fitz.open(os.path.join(tmp, "Guide.pdf"))
            foot = [pg for pg in doc if pg.number > 0]
            ok = all(any("Page %d of %d" % (pg.number + 1, doc.page_count) in " ".join(pg.get_text().split()) for _ in [0]) for pg in foot)
            rec(ok, "every page after the cover shows 'Page X of Y' in the footer")
            top = [" ".join(pg.get_textbox(fitz.Rect(0, 0, pg.rect.width, 72 * 0.9)).split()) for pg in foot]
            rec(not any(top), "no page header: nothing printed above the top margin on any page", repr([t for t in top if t][:2]))
            info = json.load(open(os.path.join(tmp, "Guide.toc.json")))
            txt = [" ".join(pg.get_text().split()) for pg in doc]
            rec(info["stable"] and bool(info["toc"]) and all(" ".join(lab.split()) in txt[n - 1] for lab, n in info["toc"].items()), "contents page numbers match where each heading lands", json.dumps(info)[:400])
            cov = doc[0].get_text()
            rec("User\u00a0Guide" in cov or "User Guide" in cov.replace("\n", " "), "cover title keeps its last two words together")
        q = os.path.join(tmp, "featured.json"); json.dump(featured(copy.deepcopy(EX)), open(q, "w"))
        fx = os.path.join(tmp, "Featured.docx")
        b = subprocess.run([sys.executable, os.path.join(HERE, "build_doc.py"), q, "--out", fx] + (["--pdf"] if pdf_ok else []), capture_output=True, text=True, timeout=600)
        rc, out = lint(q, "--docx", fx, *(["--pdf", fx[:-5] + ".pdf"] if pdf_ok else []))
        rec(b.returncode == 0 and rc == 0 and " 0 failed" in out, "1.1 features build and lint clean (3-level contents, section icon, appendix, index)", b.stdout + b.stderr + out)
        if b.returncode == 0:
            import re as _re
            from docx import Document as _D
            body = _D(fx).element.body.xml
            anchors = set(_re.findall(r'w:name="(_h_[^"]+)"', body)); links = _re.findall(r'w:anchor="([^"]+)"', body)
            rec(len(links) >= 20 and set(links) <= anchors and "_h_3_1_1" in links, "contents and index entries link to heading bookmarks in Word (%d links)" % len(links))
            rec('TOC \\o "1-3"' in body, "contents field lists three heading levels")
            if pdf_ok:
                import fitz
                pdoc = fitz.open(fx[:-5] + ".pdf"); info = json.load(open(fx[:-5] + ".toc.json"))
                first = min(info["toc"].values())
                gl = [l for i in range(1, first - 1) for l in pdoc[i].get_links() if l["kind"] == fitz.LINK_GOTO]
                ok = len(gl) >= len(info["toc"]) and all(0 <= l["page"] < pdoc.page_count for l in gl)
                rec(ok, "PDF contents entries are working links (%d links for %d entries)" % (len(gl), len(info["toc"])))
                last = [" ".join(pg.get_text().split()) for pg in pdoc][-1]
                golast = [l for l in pdoc[-1].get_links() if l["kind"] == fitz.LINK_GOTO]
                pages_ok = all(str(info["all"][lab]) in last for _, lab in info["index"])
                rec("Appendix" in last and len(golast) >= 5 and pages_ok, "PDF index entries link to their sections and show the right pages")
                tabs = [t for t in _D(fx).tables if t.rows[0].cells[0].text == "File"]
                rec(bool(tabs) and tabs[0].columns[0].width.inches >= 48 * 0.075, "auto column width fits the longest unbreakable word")
                alltext = " ".join(" ".join(pg.get_text().split()) for pg in pdoc)
                rec("Export from the toolbar." in alltext and "Sample Request" in alltext, "callout bullet items and code block titles render")
        q2 = os.path.join(tmp, "branded.json"); bd = copy.deepcopy(EX); bd["meta"]["design"] = "branded"; json.dump(bd, open(q2, "w"))
        bx = os.path.join(tmp, "Branded.docx")
        b2 = subprocess.run([sys.executable, os.path.join(HERE, "build_doc.py"), q2, "--out", bx] + (["--pdf"] if pdf_ok else []), capture_output=True, text=True, timeout=600)
        rc2, out2 = lint(q2, "--docx", bx, *(["--pdf", bx[:-5] + ".pdf"] if pdf_ok else []))
        rec(b2.returncode == 0 and rc2 == 0 and " 0 failed" in out2, "branded design builds and lints clean", b2.stdout + b2.stderr + out2)
        if b2.returncode == 0:
            from docx import Document as _D2
            dd = _D2(bx)
            hdr = dd.sections[0].header._element.xml
            rec(('w:fill="%s"' % colors_for(q2)["dark"]) in dd.tables[0]._tbl.xml and "<w:t" not in hdr and "pic:pic" not in hdr, "branded: dark cover band and no page header")
    for code, fn in (n for n in NEG if selected(n[0])):
        d = copy.deepcopy(EX); fn(d); q = os.path.join(tmp, "neg_%s.json" % code.lower()); json.dump(d, open(q, "w"))
        rc, out = lint(q)
        rec(rc != 0 and ("[%s]" % code) in out, "negative %s" % code, out)
    if selected("FIGURE_NONE"):
        proc = lambda i: {"type": "procedure", "title": "To run task %d:" % i, "steps": [{"text": "Open the page and select **Tests**."}, {"text": "Select **Submit**."}]}
        d = copy.deepcopy(EX); d["sections"][3]["blocks"] = [proc(i) for i in range(8)]
        def strip_figs(o):
            if isinstance(o, dict):
                if "blocks" in o: o["blocks"] = [b for b in o["blocks"] if b.get("type") != "figure"]
                for st in o.get("steps", []) or []: st.pop("image", None)
                for v in o.values(): strip_figs(v)
            elif isinstance(o, list):
                for v in o: strip_figs(v)
        strip_figs(d["sections"]); q = os.path.join(tmp, "nofig.json"); json.dump(d, open(q, "w"))
        rc, out = lint(q)
        rec(rc != 0 and "FAIL [FIGURE_NONE]" in out, "negative FIGURE_NONE (user guide, 8 UI procedures, 0 figures, no shot list)", out)
        open(os.path.join(tmp, "nofig.shotlist.md"), "w").write("# Shot list\n"); d["meta"]["shotlist"] = "nofig.shotlist.md"; json.dump(d, open(q, "w"))
        rc, out = lint(q)
        rec("FAIL [FIGURE_NONE]" not in out and "WARN [FIGURE_NONE]" in out, "FIGURE_NONE with meta.shotlist: WARN 'screenshots pending', not a silent pass", out)
    if selected("TABLE_REDUNDANT"):
        d = copy.deepcopy(EX); t = {"type": "table", "caption": "Roles and what each does", "columns": ["Role", "Does"], "rows": [["Tester", "Runs tests"], ["Lead", "Signs off"]]}
        d["sections"][0]["blocks"] += [t, dict(t, caption="Roles and when each joins", columns=["Role", "Joins"])]; q = os.path.join(tmp, "redund.json"); json.dump(d, open(q, "w"))
        rc, out = lint(q)
        rec("WARN [TABLE_REDUNDANT]" in out, "warn TABLE_REDUNDANT (two role tables in one section)", out)
    if selected(*LAYOUT_CODES):
        layout_tests(tmp, rec)
    if selected(*TEXTBOOK_CODES):
        textbook_tests(tmp, rec)
    if selected(*SHEET_CODES):
        sheet_tests(tmp, rec)
    if selected(*ROUND6_CODES):
        round6_tests(tmp, rec)
    if selected(*ROUND6B_CODES):
        round6b_tests(tmp, rec)
    if selected(*ROUND7_CODES):
        round7_tests(tmp, rec)
    if selected("BUILD"):
        brand_tests(tmp, rec)
        font_tests(tmp, rec)
        bad = []
        for dp, _, fs in os.walk(ROOT):
            if "__pycache__" in dp:
                continue
            for f in fs:
                try:
                    open(os.path.join(dp, f), encoding="utf-8").read()
                except UnicodeDecodeError:
                    bad.append(f)
        rec(not bad, "every packaged file is plain UTF-8 text (survives upload)", str(bad))
        sys.path.insert(0, HERE); import iconlib
        rec(all(len(iconlib.png(n, BR.DEFAULT_COLORS["primary"], 32)) > 60 for n in iconlib.names()), "every bundled Tabler icon renders to PNG for Word")
    print("\n".join(lines)); print("SELF-TEST: %d/%d passed" % (passed, total))
    shutil.rmtree(tmp, ignore_errors=True)
    sys.exit(0 if passed == total else 1)


if __name__ == "__main__":
    main()

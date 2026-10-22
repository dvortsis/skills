#!/usr/bin/env python3
"""Design gate for document-builder.

  python3 scripts/lint_doc.py doc.json [--docx Name.docx] [--pdf Name.pdf] [--draft] [--brand brand.json]
Source checks always; built-file checks when --docx / --pdf are given. Exit 1 on any FAIL.
Colors and fonts are checked against the project brand the build used (scripts/brand.py; same lookup order).
"""
import argparse, json, os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import brand as BR  # noqa: E402  (project brand: logo, colors, fonts, footer)
import iconlib  # noqa: E402
import plain_language  # noqa: E402  (shared copy: TERM_BRAND reads references/terminology.json)
from build_doc import lead_blocks, tail_blocks, BOXES  # noqa: E402  (one definition of the objectives/takeaways boxes)
from build_doc import SHEET, sheet_cfg, sheet_format, crop_png, CARD_ACCENTS, source_links  # noqa: E402  (sheet rules shared with the build)
from build_doc import T, use_brand as _theme_brand  # noqa: E402  (one theme dict for build helpers and lint; the brand overlays it)

HERE = os.path.dirname(os.path.abspath(__file__))
BLOCKS = {"paragraph", "bullets", "numbered", "procedure", "table", "figure", "callout", "code", "page_break", "step_diagram", "card"}
KEBAB = re.compile(r"^[a-z][a-z0-9]*(-[a-z0-9]+)*$")
SLOP = re.compile(r"\b(leverage[sd]?|unlock(s|ed)?|seamless(ly)?|robust|synerg(y|ies)|empower(s|ed)?|cutting-edge|best-in-class|game-chang\w+)\b", re.I)
PLACE = re.compile(r"\b(TBD|TODO|lorem ipsum|First Last|xx+)\b|\[insert", re.I)
PHI = re.compile(r"\b(MRN|medical record number)\s*[:#]?\s*\d{5,}|\b\d{3}-\d{2}-\d{4}\b|\bDOB\s*[:#]?\s*\d", re.I)
META = ("title", "doc_type", "audience", "purpose", "owner", "version", "date")
A11Y = T.get("a11y") or {}


def use_brand(b):
    """Make build and lint judge one brand: its palette and fonts on the shared theme, and its house terms."""
    _theme_brand(b)
    if b.terminology:
        plain_language.use_project_terms(b.terminology)
    return b


def pdf_families(doc):
    """Brand font families a PDF really draws with. Font names count directly; Word renames a font it took from the
    .docx's embedded fonts to ___WRD_EMBED_SUB_n, so those are identified by the family name inside the font's own
    name table (ASCII or UTF-16)."""
    want = {T["fonts"][k] for k in ("body", "heading", "mono") if T["fonts"].get(k)}
    found, seen = set(), set()
    for pg in doc:
        for f in pg.get_fonts():
            if f[0] in seen:
                continue
            seen.add(f[0])
            name = f[3].split("+", 1)[-1].replace(" ", "").lower()
            found |= {w for w in want if w.replace(" ", "").lower() in name}
            if "wrd_embed_sub" in name:
                try:
                    buf = doc.extract_font(f[0])[3] or b""
                except Exception:
                    buf = b""
                found |= {w for w in want if w.encode("ascii", "ignore") in buf or w.encode("utf-16-be") in buf}
    return found


def brand_names(b):
    """Names whose exact capitalization BRAND / TERM_BRAND check: the brand's name and org from an explicit profile
    (detected names such as a README title are too loose to enforce)."""
    if not b or not str(b.origin).startswith("profile"):
        return []
    return [n for n in dict.fromkeys((b.name, b.org)) if n and re.search(r"[A-Za-z]", n)]


def name_case_hits(texts, names):
    """(found, where, name) for each spelling of a brand name with other capitalization."""
    out = []
    for n in names:
        rx = re.compile(r"(?<![\w@./-])%s(?![\w@-]|\.\w)" % re.escape(n), re.I)  # not inside a domain such as name.com
        for where, s in texts:
            out += [(m.group(0), where, n) for m in rx.finditer(plain(s)) if m.group(0) != n]
    return out


def luminance(hexv):
    v = [int(hexv[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    v = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in v]
    return 0.2126 * v[0] + 0.7152 * v[1] + 0.0722 * v[2]


def contrast(a, b):
    """WCAG 2.1 contrast ratio of two theme color names (or 6-digit hex values)."""
    la, lb = sorted((luminance(T["colors"].get(a, a)), luminance(T["colors"].get(b, b))), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def num(v):
    return ("%.1f" % v).rstrip("0").rstrip(".")


def plain(s):
    return re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", str(s or "")).replace("**", "").replace("`", "")


def block_words(b):
    parts = [b.get(k) for k in ("text", "title", "caption", "code")] + list(b.get("items") or [])
    parts += [x for st in b.get("steps") or [] for x in (st.get("text"), st.get("result"), st.get("label"))]
    parts += list(b.get("columns") or []) + [c for r in b.get("rows") or [] for c in r]
    return sum(len(plain(x).split()) for x in parts if x)


def section_words(s, top=True):
    """Words a reader sees in a section: subsection titles and every block (the band's section title excluded)."""
    n = 0 if top else len(plain(s.get("title")).split())
    n += sum(block_words(b) for b in lead_blocks(s) + tail_blocks(s))
    return n + sum(section_words(c, False) for c in s.get("sections") or [])


class Lint:
    def __init__(self, draft=False):
        self.n = 0; self.fails = []; self.warns = []; self.draft = draft
        self.passes, self.sheet, self.fig_tw, self.meta = [], False, None, {}
        self.sec_tw = None  # D-77: landscape text width inside a landscape top-level section
        self.texts, self.notes = [], []  # (where, text) the reader sees, for TERM_BRAND (BRAND when terminology.json is missing)

    def check(self, ok, code, where, msg, warn=False):
        self.n += 1
        if not ok:
            (self.warns if warn else self.fails).append("%s [%s] %s: %s" % ("WARN" if warn else "FAIL", code, where, msg))
        return ok

    def text(self, s, where):
        if not isinstance(s, str):
            return
        self.check(not PLACE.search(s), "PLACEHOLDER", where, "placeholder text %r" % (PLACE.search(s).group(0) if PLACE.search(s) else ""))  # check: PLACEHOLDER FAIL
        self.check(not SLOP.search(s), "SLOP", where, "filler word %r; say what it does" % (SLOP.search(s).group(0) if SLOP.search(s) else ""))  # check: SLOP FAIL
        self.check(not PHI.search(s), "PHI", where, "possible patient identifier; remove it")  # check: PHI FAIL
        if where != "sources" and (where, s) not in self.texts:
            self.texts.append((where, s))
        for m in re.finditer(r"\[[^\]]+\]\(([^)]+)\)", s):
            self.check(m.group(1).startswith(("https://", "http://", "mailto:")), "LINK", where, "link target %r must be a full https:// or mailto: address" % m.group(1))  # check: LINK FAIL
        if len(s.split()) > 120:
            self.check(False, "PARA_LONG", where, "%d words in one paragraph; split it or use a list" % len(s.split()), warn=True)  # check: PARA_LONG WARN

    def source(self, d, base):
        m = d.get("meta", {})
        self.meta = m
        for k in META:
            self.check(bool(m.get(k)), "META_MISSING", "meta", "meta.%s is required" % k)  # check: META_MISSING FAIL
        dt = m.get("doc_type")
        self.sheet = m.get("layout") == "sheet"
        self.check(m.get("layout") in (None, "sheet"), "LAYOUT", "meta", 'layout is "sheet" (one page per top-level section) or left out')  # check: LAYOUT FAIL
        if self.sheet:
            self.sheet_meta(d)
        known = dt in T["doc_types"] or (self.sheet and bool(KEBAB.match(str(dt or ""))))
        if not self.check(known, "DOC_TYPE", "meta", "doc_type must be one of %s, or any kebab-case id (quick-start, job-aid, ...) with layout sheet" % ", ".join(T["doc_types"])):  # check: DOC_TYPE FAIL
            return
        for k in ("title", "subtitle", "purpose", "owner_role", "last_reviewed", "next_review"):
            self.text(m.get(k), "meta.%s" % k)
        self.review(m)
        self.check(bool(m.get("sources")), "SOURCES", "meta", "meta.sources lists where the facts came from (request, attachments, documents found with connected search tools or named by the user)")  # check: SOURCES FAIL
        self.check(all(isinstance(s, str) and s.strip() for s in m.get("sources") or []), "SOURCES", "meta", "each source is one line of text; write links as [title](https://...)")  # check: SOURCES FAIL
        for s in m.get("sources") or []:
            if isinstance(s, str):
                self.text(s, "sources")
        secs = d.get("sections") or []
        self.check(len(secs) >= (1 if self.sheet else 2), "SECTIONS", "-", "a sheet needs at least one top-level section" if self.sheet else "a document needs at least two top-level sections")  # check: SECTIONS FAIL
        tops = [s.get("title", "") for s in secs]
        both = any(t.strip().lower() == "troubleshooting and support" for t in tops)  # one combined chapter meets both
        for req in (sheet_format(dt).get("required", []) if self.sheet else T["doc_types"][dt]["required"]):
            self.check(any(req.lower() == t.lower() for t in tops) or (both and req.lower() in ("troubleshooting", "support")), "SECTION_REQUIRED", "-", "a %s needs a top-level %r section" % (dt, req))  # check: SECTION_REQUIRED FAIL
        if not self.sheet and T["doc_types"][dt]["revisions"]:
            self.check(bool(d.get("revisions")), "REVISIONS", "-", "a %s needs a revisions list (version, date, author, change)" % dt)  # check: REVISIONS FAIL
        for k in ([] if self.sheet else T["doc_types"][dt].get("meta_required", [])):  # D-62: policy needs meta.origination_date
            self.check(bool(m.get(k)), "META_MISSING", "meta", "a %s needs meta.%s" % (dt, k))  # check: META_MISSING FAIL
        for r in d.get("revisions") or []:
            self.check(all(r.get(k) for k in ("version", "date", "author", "change")), "REVISIONS", "revisions", "each revision needs version, date, author and change")  # check: REVISIONS FAIL
        for g in d.get("glossary") or []:
            self.check(bool(g.get("term")) and bool(g.get("definition")), "GLOSSARY", "glossary", "each entry needs term and definition")  # check: GLOSSARY FAIL
            self.texts += [("glossary", str(g[k])) for k in ("term", "definition") if g.get(k)]
        self.fig_files = []
        self.counts = {"h": 0, "fig": 0, "tab": 0, "callout": 0, "toc": 0}
        self.check(m.get("design", "standard") in ("standard", "branded"), "DESIGN", "meta", "design is standard (default) or branded")  # check: DESIGN FAIL
        lv = m.get("toc_levels", 2)
        self.check(lv in (1, 2, 3), "TOC_LEVELS", "meta", "toc_levels must be 1, 2 or 3 (heading levels listed in the contents)")  # check: TOC_LEVELS FAIL
        levels = lv if lv in (1, 2, 3) else 2
        self.check(m.get("page_break_level", 1) in (0, 1), "PAGE_BREAK_LEVEL", "meta", "page_break_level is 1 (each top-level section starts a page, default) or 0 (sections run on)")  # check: PAGE_BREAK_LEVEL FAIL
        def walk(ss, prefix, lvl, parent=None):
            seen = set()
            for i, s in enumerate(ss, 1):
                where = "%s%d %s" % (prefix, i, s.get("title", "?"))
                if lvl == 1:
                    self.sec_tw = T["page"]["height_in"] - 2 * T["page"]["margin_in"] if s.get("orientation") == "landscape" else None
                if "orientation" in s:
                    self.check(lvl == 1 and not self.sheet and s["orientation"] in ("landscape", "portrait"), "LAYOUT", where, 'orientation is "landscape" or "portrait", on a top-level section of a document (not a sheet)')  # check: LAYOUT FAIL
                self.counts["h"] += 1
                if lvl <= levels:
                    self.counts["toc"] += 1
                if s.get("icon"):
                    self.check(iconlib.exists(s["icon"]), "ICON_UNKNOWN", where, "icon %r is not in the bundled Tabler set" % s["icon"])  # check: ICON_UNKNOWN FAIL
                self.check(bool(s.get("title")), "HEADING", where, "every section needs a title")  # check: HEADING FAIL
                self.check(lvl <= 3, "HEADING_DEPTH", where, "headings go three levels deep at most (1, 1.1, 1.1.1)")  # check: HEADING_DEPTH FAIL
                t = s.get("title", "")
                if t:
                    self.texts.append((where, t))
                self.check(t[:1].isupper() or t[:1].isdigit(), "HEADING_CASE", where, "start headings with a capital letter")  # check: HEADING_CASE FAIL
                self.check(len(t.split()) <= 10, "HEADING_LONG", where, "keep headings to 10 words or fewer")  # check: HEADING_LONG FAIL
                self.check(t.lower() not in seen, "DUP_HEADING", where, "two sibling sections share the title %r" % t)  # check: DUP_HEADING FAIL
                self.check(parent is None or t.strip().lower() != parent.strip().lower(), "DUP_HEADING", where, "heading repeats its parent %r one level down; give the child its own name or remove the level" % parent)  # check: DUP_HEADING FAIL
                seen.add(t.lower())
                self.check(bool(s.get("blocks")) or bool(s.get("sections")), "EMPTY_SECTION", where, "section has no content")  # check: EMPTY_SECTION FAIL
                for key, lo, hi in (("objectives", 2, 6), ("takeaways", 2, 6)):
                    if key in s:
                        v = s[key]
                        self.check(isinstance(v, list) and lo <= len(v) <= hi and all(isinstance(x, str) and x.strip() for x in v), key.upper(), where,
                                   "%s is a list of %d-%d short sentences (rendered as the %r box)" % (key, lo, hi, BOXES[key]["title"]))
                if "width" in s:
                    self.check(s["width"] in ("half", "full"), "LAYOUT", where, 'width is "half" (side by side with the next half-width subsection on a sheet) or "full"')  # check: LAYOUT FAIL
                self.fig_tw = (T["page"]["text_width_in"] - 0.25) / 2 - 0.075 if self.sheet and s.get("width") == "half" else None
                for b in lead_blocks(s) + tail_blocks(s):
                    self.block(b, where, base)
                firsts = [tuple(sorted({str(r[0]).strip().lower() for r in b.get("rows", []) if r})) for b in s.get("blocks", []) if b.get("type") == "table"]
                dup = [f for f in set(firsts) if firsts.count(f) > 1 and len(f) >= 2]
                self.check(not dup, "TABLE_REDUNDANT", where, "two tables in this section list the same first-column values (%s); merge them" % ", ".join(dup[0][:4]) if dup else "", warn=True)  # check: TABLE_REDUNDANT WARN
                walk(s.get("sections", []), "%s%d." % (prefix, i), lvl + 1, t)
        walk(secs, "", 1)
        procs = figs = 0
        def scan(ss):
            nonlocal procs, figs
            for s in ss:
                for b in s.get("blocks", []):
                    if b.get("type") == "procedure":
                        procs += any(re.search(r"\*\*[^*]+\*\*", str(st.get("text", ""))) for st in b.get("steps", []))
                        figs += sum(1 for st in b.get("steps", []) if st.get("image"))
                    figs += b.get("type") == "figure"
                scan(s.get("sections", []))
        scan(secs)
        shot = m.get("shotlist")
        shot_ok = bool(shot) and os.path.isfile(shot if os.path.isabs(shot) else os.path.join(base, shot))
        self.check(procs < 3 or figs or shot_ok, "FIGURE_NONE", "-", "%d procedures quote UI labels but the document has no figures; add the user's screenshots or set meta.shotlist to the shot list" % procs,
                   warn=dt != "user-guide")  # check: FIGURE_NONE FAIL # check: FIGURE_NONE WARN
        if procs >= 3 and not figs and shot_ok:
            self.check(False, "FIGURE_NONE", "-", "screenshots pending (meta.shotlist %s); the document isn't complete until they are added" % shot, warn=True)  # check: FIGURE_NONE WARN
        reserved = {"glossary", "revision history", "sources", "index"}
        apps = []
        for j, ap in enumerate(d.get("appendices") or [], 1):
            w = "appendix %d %s" % (j, ap.get("title", "?"))
            self.check(bool(ap.get("title")) and bool(ap.get("blocks")), "APPENDIX", w, "each appendix needs a title and blocks")  # check: APPENDIX FAIL
            self.check(str(ap.get("title", "")).lower() not in reserved, "APPENDIX", w, "%r is built automatically; use the glossary, revisions, sources or index keys" % ap.get("title"))  # check: APPENDIX FAIL
            self.check(len(str(ap.get("title", "")).split()) <= 10, "HEADING_LONG", w, "keep headings to 10 words or fewer")  # check: HEADING_LONG FAIL
            if ap.get("icon"):
                self.check(iconlib.exists(ap["icon"]), "ICON_UNKNOWN", w, "icon %r is not in the bundled Tabler set" % ap["icon"])  # check: ICON_UNKNOWN FAIL
            self.text(ap.get("title"), w)
            for bl in ap.get("blocks") or []:
                self.block(bl, w, base)
            apps.append(str(ap.get("title", "")).lower())
        apps += [t for k, t in (("glossary", "glossary"), ("revisions", "revision history")) if d.get(k)] + (["sources"] if m.get("sources") else [])
        nums, titles = set(), {}
        def collect(ss, prefix):
            for i, s in enumerate(ss, 1):
                n = "%s%d" % (prefix, i); nums.add(n); titles.setdefault(str(s.get("title", "")).lower(), []).append(n)
                collect(s.get("sections", []), n + ".")
        collect(secs, "")
        for t in apps:
            titles.setdefault(t, []).append("appendix")
        for e in d.get("index") or []:
            ref = str(e.get("section", "")).strip()
            hit = bool(e.get("term")) and (ref in nums or len(titles.get(ref.lower(), [])) == 1)
            self.check(hit, "INDEX_TARGET", "index", "index entry %r must point to one existing section (number or unique title), not %r" % (e.get("term"), ref))  # check: INDEX_TARGET FAIL
            self.text(e.get("term"), "index")
        terms = [str(e.get("term", "")).lower() for e in d.get("index") or []]
        self.check(len(terms) == len(set(terms)), "INDEX_DUP", "index", "duplicate index terms", warn=True)  # check: INDEX_DUP WARN
        self.counts["apps"] = len(apps) + (1 if d.get("index") else 0)
        if self.sheet:  # no automatic appendices on a sheet
            self.counts["apps"] = 0
            self.sheet_known(d, dt)
            self.structure(d, dt)
        else:
            self.guide_profile(d, dt)

    def review(self, m):
        """REVIEW_DATE (D-65, warn): meta.next_review (with owner_role and last_reviewed) so readers see when the doc is
        checked again; meta.latest_url (sheets print it after 'Printed copies may be out of date.') is a full link."""
        if self.check(bool(m.get("next_review")), "REVIEW_DATE", "meta", "next review date missing: set meta.next_review (with meta.owner_role and meta.last_reviewed)", warn=True):  # check: REVIEW_DATE WARN
            miss = [k for k in ("owner_role", "last_reviewed") if not m.get(k)]
            self.check(not miss, "REVIEW_DATE", "meta", "next review is set; add %s" % " and ".join("meta." + k for k in miss), warn=True)  # check: REVIEW_DATE WARN
            if not miss:
                self.passes.append("REVIEW_DATE pass: owner role %s, last reviewed %s, next review %s" % (m["owner_role"], m["last_reviewed"], m["next_review"]))
        if m.get("latest_url") is not None:
            self.check(str(m["latest_url"]).startswith("https://"), "LINK", "meta.latest_url", "latest_url must be a full https:// address")  # check: LINK FAIL

    def guide_profile(self, d, dt):
        """GUIDE_PROFILE (D-63): lean guides (numbered task topics, a Reference chapter, contents, glossary; no objectives
        or takeaways boxes, no index) or textbook. Default from doc_theme.json profiles; lint only."""
        P, m = T.get("profiles") or {}, d.get("meta") or {}
        secs, prof = d.get("sections") or [], m.get("profile")
        flat = []
        def walk(ss, prefix):
            for i, s in enumerate(ss, 1):
                flat.append(("%s%d" % (prefix, i), s)); walk(s.get("sections") or [], "%s%d." % (prefix, i))
        walk(secs, "")
        used = [k for k in P.get("textbook_signals", []) if (d.get(k) if k == "index" else any(k in s for _, s in flat))]
        if prof is not None:
            if not self.check(prof in P.get("values", []), "GUIDE_PROFILE", "meta", "profile must be one of %s (or left out)" % ", ".join(P.get("values", []))):  # check: GUIDE_PROFILE FAIL
                return
            how = "meta.profile"
        elif used:
            prof, how = "textbook", "inferred from %s" % ", ".join(used)
        elif dt in P.get("default_types", []) and not plain_language.is_developer(m.get("audience")):
            prof, how = "lean", "default for a %s for business users" % dt
        else:
            if dt in P.get("default_types", []):
                self.notes.append("GUIDE_PROFILE: not checked (developer audience; set meta.profile lean or textbook to check)")
            return
        if prof == "textbook":
            self.passes.append("GUIDE_PROFILE pass: textbook (%s)" % how); return
        lean, before = P.get("lean") or {}, len(self.fails)
        forbid = {t.lower() for t in lean.get("forbid_titles", [])}
        tasks = ["%s  %s" % (n, s.get("title", "")) for n, s in flat if any(b.get("type") == "procedure" for b in s.get("blocks") or [])]
        self.check(len(tasks) >= lean.get("min_task_topics", 1), "GUIDE_PROFILE", "-", "a lean guide needs numbered task topics: a section with a procedure block for each task")  # check: GUIDE_PROFILE FAIL
        refs = {t.lower() for t in lean.get("reference_titles", [])}
        self.check(any(str(s.get("title", "")).strip().lower() in refs for s in secs), "GUIDE_PROFILE", "-", "a lean guide needs a top-level %s chapter (lookup tables, settings, definitions)" % " or ".join(repr(t) for t in lean.get("reference_titles", [])))  # check: GUIDE_PROFILE FAIL
        self.check(bool(d.get("glossary")), "GUIDE_PROFILE", "-", "a lean guide needs a glossary")  # check: GUIDE_PROFILE FAIL
        for n, s in flat:
            where = "%s %s" % (n, s.get("title", "?"))
            for k in ("objectives", "takeaways"):
                self.check(k not in s, "GUIDE_PROFILE", where, "a lean guide has no %r box; remove %s or set meta.profile textbook" % (BOXES[k]["title"], k))  # check: GUIDE_PROFILE FAIL
            named = [s.get("title")] + [b.get("title") for b in s.get("blocks") or [] if b.get("type") == "callout"]
            for t in named:
                self.check(str(t or "").strip().lower() not in forbid, "GUIDE_PROFILE", where, "%r is a textbook box or index title; a lean guide has none (say 'What This Guide Covers', or set meta.profile textbook)" % t)  # check: GUIDE_PROFILE FAIL
        for ap in d.get("appendices") or []:
            self.check(str(ap.get("title", "")).strip().lower() not in forbid, "GUIDE_PROFILE", "appendix %s" % ap.get("title"), "a lean guide has no index or takeaways appendix")  # check: GUIDE_PROFILE FAIL
        self.check(not d.get("index"), "GUIDE_PROFILE", "index", "a lean guide has no index; remove index or set meta.profile textbook")  # check: GUIDE_PROFILE FAIL
        if len(self.fails) == before:
            self.passes.append("GUIDE_PROFILE pass: lean (%s): %d numbered task topic(s), Reference chapter, contents, glossary; no objectives, takeaways or index" % (how, len(tasks)))

    def term_brand(self):
        """TERM_BRAND (D-39, warn): off-style terms from the shared references/terminology.json plus the project's house
        terms (brand profile `terminology`) and the brand name's capitalization, in headings, block text, meta and
        glossary. Without the shared list it is not run and the BRAND casing rule (FAIL) on the brand name is the fallback."""
        names = brand_names(getattr(self, "brand", None))
        hits = plain_language.term_hits(self.texts, sets=("internal",))
        if hits is None:
            self.notes.append("TERM_BRAND: not run (terminology.json missing)")
            if not names:
                self.notes.append("BRAND: not run (the brand profile names no brand or org)")
            bad = name_case_hits(self.texts, names)
            for found, where, n in bad:
                self.check(False, "BRAND", where, "write %s (exact capitalization), not %s" % (n, found))  # check: BRAND FAIL
            self.n += len(self.texts) * bool(names) - len(bad)
            return
        hits = hits + [(f, w, n, "brand name") for f, w, n in name_case_hits(self.texts, names)]
        self.n += len(self.texts)
        for found, where, preferred, why in hits:
            self.warns.append('WARN [TERM_BRAND] %s: "%s -> %s"' % (where, found, preferred))  # check: TERM_BRAND WARN

    def a11y_theme(self):
        """DOC_A11Y (D-64): every theme text color pair the build uses (doc_theme.json a11y.text_pairs) at 4.5:1 or more."""
        pairs, mn = A11Y.get("text_pairs") or [], A11Y.get("min_contrast", 4.5)
        self.check(bool(pairs), "DOC_A11Y", "theme", "doc_theme.json a11y.text_pairs lists no text colors")  # check: DOC_A11Y FAIL
        bad = [(f, b, contrast(f, b)) for f, b in pairs if contrast(f, b) < mn]
        for f, b, r in bad:
            self.check(False, "DOC_A11Y", "theme", "text color %s on %s is %.2f:1; WCAG 2.1 AA needs %g:1" % (f, b, r, mn))  # check: DOC_A11Y FAIL
        if pairs and not bad:
            self.passes.append("DOC_A11Y pass: %d theme text color pairs at %g:1 or more on their backgrounds" % (len(pairs), mn))

    def sheet_meta(self, d):
        """meta.sheet keys and floors; a sheet has no appendices, glossary or index."""
        sh = d["meta"].get("sheet", {})
        keys, fl, (lo, hi) = set(SHEET["defaults"]), SHEET["floors"], SHEET["text_scale_range"]
        isnum = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
        ok = isinstance(sh, dict) and set(sh) <= keys and all(isnum(v) for v in sh.values())
        if ok:
            pps = sh.get("pages_per_section", 1)
            ok = int(pps) == pps and pps >= 1 and lo <= sh.get("text_scale", 1.0) <= hi and all(sh.get(k, v) >= v for k, v in fl.items())
        self.check(ok, "LAYOUT", "meta.sheet", "sheet takes %s; pages_per_section a whole number of 1 or more, text_scale %g-%g, min_body_pt %g+, min_table_pt %g+, min_figure_pt %g+"
                   % (", ".join(sorted(keys)), lo, hi, fl["min_body_pt"], fl["min_table_pt"], fl["min_figure_pt"]))  # check: LAYOUT FAIL
        self.check(not (d.get("appendices") or d.get("glossary") or d.get("index")), "LAYOUT", "-", "a sheet has no appendices, glossary or index; put that content in a section or a separate document")  # check: LAYOUT FAIL

    def sheet_known(self, d, dt):
        """D-66 SHEET_FORMAT_KNOWN: a catalog sheet type (sheet_formats.json catalog_sheet_types) has its own entry;
        its pages limit meta.sheet.pages_per_section and its sections are level-2 titles in every top-level section."""
        fmts = SHEET.get("formats") or {}
        if dt not in fmts:
            known = dt in (SHEET.get("catalog_sheet_types") or [])
            self.check(not known, "SHEET_FORMAT_KNOWN", "meta", "%r is a catalog sheet type with no entry in assets/sheet_formats.json; add its pages and sections from the format catalog" % dt)  # check: SHEET_FORMAT_KNOWN FAIL
            if not known:
                self.notes.append("SHEET_FORMAT_KNOWN: %s is not a catalog sheet type; default sheet rules" % dt)
            return
        fmt, before = sheet_format(dt), len(self.fails)
        sh = d["meta"].get("sheet") if isinstance(d["meta"].get("sheet"), dict) else {}
        pps, mx = sh.get("pages_per_section"), fmt.get("pages")
        if mx and isinstance(pps, (int, float)) and not isinstance(pps, bool):
            self.check(pps <= mx, fmt["page_code"], "meta.sheet", "pages_per_section %s is over the %s limit of %d page%s" % (num(pps), dt, mx, "" if mx == 1 else "s"))  # check: SHEET_PAGE_FIT FAIL # check: QS_PAGE_COUNT FAIL
        want = fmt.get("sections") or []
        key = lambda t: " ".join(str(t or "").replace("\u2019", "'").lower().split())
        for n, s in enumerate(d.get("sections") or [], 1):
            have = {key(x.get("title")) for x in s.get("sections") or []}
            for t in want:
                self.check(key(t) in have, "SECTION_REQUIRED", "%s %d" % (fmt["section_noun"], n), "a %s %s needs a level-2 %r section" % (dt, fmt["section_noun"], t))  # check: SECTION_REQUIRED FAIL
        if len(self.fails) == before:
            p = mx or SHEET["defaults"]["pages_per_section"]
            self.passes.append("SHEET_FORMAT_KNOWN pass: %s rules (%d page%s per section%s)" % (dt, p, "" if p == 1 else "s", ", %d required sections" % len(want) if want else ""))

    def structure(self, d, dt):
        """Sheet anatomy from assets/sheet_formats.json (quick-start: QS_STRUCTURE): level-2 titles in order, step and
        item limits, words per section (band and footer excluded)."""
        fmt = sheet_format(dt); st = fmt.get("structure")
        if not st:
            return
        code, noun, want = st["code"], fmt["section_noun"], st["subsections"]
        canon = {w.lower(): w for w in want}
        for k, vs in (st.get("aliases") or {}).items():
            canon.update({v.lower().replace("\u2019", "'"): k for v in vs})
        key = lambda t: str(t or "").strip().lower().replace("\u2019", "'")
        before, secs = len(self.fails), d.get("sections") or []
        for n, s in enumerate(secs, 1):
            where = "%s %d" % (noun, n)
            subs = s.get("sections") or []
            names = [canon.get(key(x.get("title")), x.get("title")) for x in subs]
            for title in want:
                self.check(title in names, code, where, "missing %s" % title)  # check: QS_STRUCTURE FAIL
            seen = [x for x in names if x in want]
            self.check(seen == [w for w in want if w in seen], code, where, "sections out of order")  # check: QS_STRUCTURE FAIL
            by = dict(zip(names, subs))
            if st.get("steps_title") in by:
                k = next((len(b.get("items") or b.get("steps") or []) for b in by[st["steps_title"]].get("blocks") or [] if b.get("type") in ("numbered", "procedure")), 0)
                self.check(k <= st["steps_max"], code, where, "%d steps (maximum %d)" % (k, st["steps_max"]))  # check: QS_STRUCTURE FAIL
                self.check(k >= st["steps_min"], code, where, "%d steps (minimum %d)" % (k, st["steps_min"]))  # check: QS_STRUCTURE FAIL
            for title, mx in (st.get("max_items") or {}).items():
                if title in by:
                    k = sum(len(b.get("items") or []) or (1 if b.get("text") else 0) for b in by[title].get("blocks") or [])
                    self.check(k <= mx, code, where, "%d %s (maximum %d)" % (k, title.lower(), mx))  # check: QS_STRUCTURE FAIL
            w = section_words(s)
            self.check(w <= st["max_words"], code, where, "%d words (maximum %d)" % (w, st["max_words"]))  # check: QS_STRUCTURE FAIL
        if secs and not any(("[%s]" % code) in f for f in self.fails[before:]):
            self.passes.append("%s pass: %d %ss" % (code, len(secs), noun))

    def block(self, b, where, base):
        k = b.get("type")
        if not self.check(k in BLOCKS, "BLOCK_TYPE", where, "unknown block type %r" % k):  # check: BLOCK_TYPE FAIL
            return
        if k == "paragraph":
            self.check(bool(b.get("text")), "EMPTY_BLOCK", where, "empty paragraph"); self.text(b.get("text"), where)  # check: EMPTY_BLOCK FAIL
        elif k in ("bullets", "numbered"):
            its = b.get("items") or []
            self.check(len(its) >= 2, "LIST_SHORT", where, "a list needs at least two items")  # check: LIST_SHORT FAIL
            self.check(len(its) <= 9, "LIST_LONG", where, "%d items; group them or use a table" % len(its), warn=True)  # check: LIST_LONG WARN
            for it in its: self.text(it, where)
        elif k == "procedure":
            st = b.get("steps") or []
            self.check(2 <= len(st) <= 12, "PROC_STEPS", where, "a procedure has 2-12 steps (%d); split long tasks" % len(st))  # check: PROC_STEPS FAIL
            for j, s in enumerate(st, 1):
                tx = s.get("text", "")
                self.check(bool(tx), "PROC_STEP", where, "step %d is empty" % j)  # check: PROC_STEP FAIL
                self.check(len(tx.split()) <= 40, "PROC_STEP_LONG", where, "step %d has %d words; one action per step (40 max)" % (j, len(tx.split())))  # check: PROC_STEP_LONG FAIL
                self.check(not re.match(r"(?i)(you (should|can|will|need)|the user|users? should|it is|there is)\b", tx), "PROC_IMPERATIVE", where, "step %d: start with the action verb (Select, Enter, Open)" % j)  # check: PROC_IMPERATIVE FAIL
                self.text(tx, where); self.text(s.get("result"), where)
                if s.get("image"):
                    self.figure({"image": s["image"], "caption": s.get("caption", "x x x"), "alt": s.get("alt", tx)}, where, base)
        elif k == "table":
            cols = b.get("columns") or []
            rows = b.get("rows") or []
            self.counts["tab"] += 1
            self.check(1 <= len(cols) <= 6, "TABLE_COLS", where, "tables take 1-6 columns (%d); split it" % len(cols))  # check: TABLE_COLS FAIL
            self.check(len(rows) >= 1, "TABLE_EMPTY", where, "table has no rows")  # check: TABLE_EMPTY FAIL
            self.check(all(len(r) == len(cols) for r in rows), "TABLE_SHAPE", where, "every row needs exactly %d cells" % len(cols))  # check: TABLE_SHAPE FAIL
            self.check(bool(b.get("caption")), "TABLE_CAPTION", where, "tables need a caption that says what they show")  # check: TABLE_CAPTION FAIL
            if b.get("widths_in"):
                self.check(len(b["widths_in"]) == len(cols) and abs(sum(b["widths_in"]) - (self.sec_tw or T["page"]["text_width_in"])) < 0.05, "TABLE_WIDTH", where, "widths_in must have one width per column and add up to %.1f in" % (self.sec_tw or T["page"]["text_width_in"]))  # check: TABLE_WIDTH FAIL
            for r in rows:
                for c in r: self.text(str(c), where)
        elif k == "figure":
            self.figure(b, where, base)
        elif k == "callout":
            kind = b.get("kind")
            self.counts["callout"] += 1
            self.check(kind in T["callouts"], "CALLOUT_KIND", where, "callout kind must be one of %s" % ", ".join(T["callouts"]))  # check: CALLOUT_KIND FAIL
            items = b.get("items") or []
            self.check(bool(b.get("text")) or bool(items), "EMPTY_BLOCK", where, "empty callout (needs text or items)")  # check: EMPTY_BLOCK FAIL
            words = len((b.get("text") or "").split()) + sum(len(str(i).split()) for i in items)
            self.check(words <= 60 and len(items) <= 6, "CALLOUT_LONG", where, "callouts stay under 60 words and 6 items")  # check: CALLOUT_LONG FAIL
            for i in items:
                self.text(i, where)
            if b.get("icon"):
                self.check(iconlib.exists(b["icon"]), "ICON_UNKNOWN", where, "icon %r is not in the bundled Tabler set" % b["icon"])  # check: ICON_UNKNOWN FAIL
            self.text(b.get("text"), where); self.text(b.get("title"), where)
        elif k == "code":
            self.check(bool(b.get("code")), "EMPTY_BLOCK", where, "empty code block")  # check: EMPTY_BLOCK FAIL
            self.text(b.get("title"), where)
            self.check(max((len(x) for x in (b.get("code") or "").split("\n")), default=0) <= 90, "CODE_WIDE", where, "code lines over 90 characters wrap badly; break them")  # check: CODE_WIDE FAIL
        elif k == "step_diagram":
            steps = b.get("steps") or []
            self.check(3 <= len(steps) <= 6, "STEP_COUNT", where, "a step diagram has 3-6 steps (%d); put the detail in the numbered steps" % len(steps))  # check: STEP_COUNT FAIL
            for j, st in enumerate(steps, 1):
                lab = str(st.get("label") or "").strip()
                self.check(bool(lab), "EMPTY_BLOCK", where, "step %d has no label" % j)  # check: EMPTY_BLOCK FAIL
                self.check(len(lab.split()) <= 4, "STEP_LABEL", where, "step %d label %r has %d words (maximum 4)" % (j, lab, len(lab.split())))  # check: STEP_LABEL FAIL
                self.check(iconlib.exists(st.get("icon") or ""), "ICON_UNKNOWN", where, "step %d icon %r is not in the bundled Tabler set" % (j, st.get("icon")))  # check: ICON_UNKNOWN FAIL
                self.text(lab, where)
        elif k == "card":
            items = b.get("items") or []
            self.check(bool(b.get("text")) or bool(items), "EMPTY_BLOCK", where, "empty card (needs text or items)")  # check: EMPTY_BLOCK FAIL
            self.check(BR.LEGACY.get(b.get("accent", "primary"), b.get("accent", "primary")) in CARD_ACCENTS, "COLOR", where, "card accent must be one of %s" % ", ".join(CARD_ACCENTS))  # check: COLOR FAIL
            for x in [b.get("title"), b.get("text")] + list(items):
                self.text(x, where)

    def figure(self, b, where, base):
        self.counts["fig"] += 1
        img = b.get("image")
        exists = bool(img) and os.path.isfile(os.path.join(base, img))
        self.check(exists or self.draft, "FIGURE_FILE", where, "screenshot %r is missing; ask the user for it or build with --draft" % img)  # check: FIGURE_FILE FAIL
        self.check(len((b.get("caption") or "").split()) >= 3, "FIGURE_CAPTION", where, "figures need a caption of 3+ words")  # check: FIGURE_CAPTION FAIL
        self.check(len((b.get("alt") or b.get("caption") or "").split()) >= 4, "ALT_TEXT", where, "figures need alt text (4+ words) for screen readers")  # check: ALT_TEXT FAIL
        if b.get("width_in") is not None:
            tw = self.fig_tw or self.sec_tw or T["page"]["text_width_in"]
            wi = b["width_in"]
            self.check(isinstance(wi, (int, float)) and not isinstance(wi, bool) and 0 < wi <= tw + 0.01, "FIG_WIDTH", where,
                       "width_in %s is wider than this section's text width (%.1f in); put wide figures in a top-level section with orientation: landscape" % (wi, tw))  # check: FIG_WIDTH FAIL
        if img:
            self.check(not re.search(r"(?i)stock|clipart|generated", img), "FIGURE_SOURCE", where, "use only screenshots the user supplied")  # check: FIGURE_SOURCE FAIL
            if b.get("illustration") or re.search(r"(?i)illustration", os.path.basename(img)):
                self.check("illustration" in (b.get("caption") or "").lower(), "ILLUSTRATION_LABEL", where, 'a screen illustration needs "illustration" in its caption (for example "Illustration, not a live screenshot, of ...")')  # check: ILLUSTRATION_LABEL FAIL
        crop = b.get("crop")
        if crop is not None:
            ok = isinstance(crop, list) and len(crop) == 4 and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in crop)
            ok = ok and 0 <= crop[0] < crop[2] <= 1 and 0 <= crop[1] < crop[3] <= 1
            self.check(ok, "FIGURE_CROP", where, "crop is [x0, y0, x1, y1] as fractions of the image: 0 <= x0 < x1 <= 1 and 0 <= y0 < y1 <= 1")  # check: FIGURE_CROP FAIL
            crop = crop if ok else None
        if exists:
            self.fig_text(os.path.join(base, img), dict(b, crop=crop), where)

    def fig_text(self, path, b, where):
        """FIG_TEXT_SIZE (warn): smallest text in the figure at its printed width, estimated from glyph heights."""
        try:
            import figtext
            src, ratio = crop_png(path, b.get("crop"))
            tw = self.fig_tw or self.sec_tw or T["page"]["text_width_in"]
            width = min(b.get("width_in", tw), tw, 7.6 / ratio * 0.62)  # the builder's printed width (after any crop)
            pt = figtext.min_text_pt(src, width)
        except Exception:
            pt = None
        if pt is not None and self.sheet:  # sheets: FAIL under min_figure_pt (8 pt)
            mn = sheet_cfg(self.meta)["min_figure_pt"]
            self.check(pt >= mn, "FIG_TEXT_SIZE", os.path.basename(path), "smallest text about %s pt (minimum %s pt); crop or widen it" % (num(pt), num(mn)))  # check: FIG_TEXT_SIZE FAIL
        elif pt is not None:
            self.check(pt >= 7, "FIG_TEXT_SIZE", where, "text in %s prints at about %.1f pt (under 7 pt); crop to the area that matters, split the figure or widen it" % (os.path.basename(path), pt), warn=True)  # check: FIG_TEXT_SIZE WARN

    def docx(self, path, d):
        try:
            from docx import Document
            doc = Document(path)
        except Exception as ex:
            self.check(False, "DOCX_OPEN", path, "can't open: %r" % ex); return  # check: DOCX_OPEN FAIL
        body = doc.element.body.xml
        if self.sheet:
            self.check("TOC \\o" not in body, "TOC_FIELD", "docx", "a sheet has no contents: remove the TOC field")  # check: TOC_FIELD FAIL
        else:
            self.check("TOC \\o" in body, "TOC_FIELD", "docx", "the contents must be a Word TOC field (updates with F9)")  # check: TOC_FIELD FAIL
        heads = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
        if self.sheet:  # band titles and half-width subsections sit in table cells
            heads = re.findall(r'<w:pStyle w:val="Heading[1-3]"/>', body)
        apps = self.counts.get("apps", 0)
        self.check(len(heads) == self.counts["h"] + apps, "HEADINGS", "docx", "%d heading paragraphs, expected %d" % (len(heads), self.counts["h"] + apps))  # check: HEADINGS FAIL
        anchors = set(re.findall(r'<w:bookmarkStart [^>]*w:name="([^"]+)"', body))
        links = re.findall(r'<w:hyperlink [^>]*w:anchor="([^"]+)"', body)
        want = self.counts["toc"] + apps + len(d.get("index") or [])
        self.check(self.sheet or (len(links) >= want and all(l in anchors for l in links)), "TOC_LINKS", "docx",
                   "%d internal links for %d contents and index entries; every one must point to a heading bookmark" % (len(links), want))  # check: TOC_LINKS FAIL
        pics = re.findall(r"<wp:docPr [^>]*>", body)
        self.check(all('descr="' in p and 'descr=""' not in p for p in pics), "ALT_TEXT", "docx", "every picture needs alt text")  # check: ALT_TEXT FAIL
        fonts = set(re.findall(r'w:ascii="([^"]+)"', body))
        allowed = {T["fonts"][k] for k in ("body", "heading", "mono") if T["fonts"].get(k)}
        self.check(fonts <= allowed, "FONTS", "docx", "only the brand fonts (%s): found %s" % (", ".join(sorted(allowed)), ", ".join(sorted(fonts))))  # check: FONTS FAIL
        self.embedded_fonts(path)
        cols = set(c.upper() for c in re.findall(r'w:(?:color|fill)="([0-9A-Fa-f]{6})"', body)) | set(c.upper() for c in re.findall(r'<w:color w:val="([0-9A-Fa-f]{6})"', body))
        pal = set(v.upper() for v in T["colors"].values()) | {"F4F4F4"}
        self.check(cols <= pal, "COLOR", "docx", "off-palette colors %s" % ", ".join(sorted(cols - pal)))  # check: COLOR FAIL
        self.check(doc.core_properties.title == d["meta"]["title"], "DOC_TITLE", "docx", "document title property must match meta.title")  # check: DOC_TITLE FAIL
        hx = " ".join(h._element.xml for sec in doc.sections for h in (sec.header, sec.first_page_header, sec.even_page_header))
        self.check(not re.search(r"<w:t[ >]|pic:pic|<w:pBdr>", hx), "NO_PAGE_HEADER", "docx", "page headers must be empty (no text, logo or rule)")  # check: NO_PAGE_HEADER FAIL
        self.logo(pics, d)
        self.a11y_docx(doc, d, body, pics)

    def embedded_fonts(self, path):
        """FONT_EMBED: when the brand's body font file is on this machine (project file, bundled fallback or installed)
        and allows embedding, the .docx carries it (fontTable embedRegular + an obfuscated font part); FONT_MISSING
        warns when the file isn't here, so readers need the font installed."""
        import zipfile
        from build_doc import embeddable
        b = getattr(self, "brand", None)
        if b is None:
            return
        fam, p = T["fonts"]["body"], b.font_path("body")
        if not p:
            self.check(False, "FONT_MISSING", "docx", "the body font %s isn't on this machine, so it can't be embedded; readers without it see a substitute (ship the font files in the project or name them in brand.json)" % fam, warn=True)  # check: FONT_MISSING WARN
            return
        if not embeddable(p):
            self.notes.append("FONT_EMBED: not run (%s does not allow embedding)" % os.path.basename(p))
            return
        z = zipfile.ZipFile(path)
        ft = z.read("word/fontTable.xml").decode("utf-8") if "word/fontTable.xml" in z.namelist() else ""
        m = re.search(r'<w:font w:name="%s">(.*?)</w:font>' % re.escape(fam), ft, re.S)
        ok = bool(m and re.search(r"<w:embedRegular [^>]*r:id=", m.group(1))) and any(n.startswith("word/fonts/") for n in z.namelist())
        self.check(ok, "FONT_EMBED", "docx", "the body font %s is not embedded; rebuild with build_doc.py" % fam)  # check: FONT_EMBED FAIL

    def logo(self, pics, d):
        """The brand logo on the cover (sheets: in each title band) when the brand has one; LOGO_NONE warns when the
        project has no logo, so the cover has none."""
        b = getattr(self, "brand", None)
        if b is None:
            return
        alt = 'descr="%s"' % b.logo_alt().replace("&", "&amp;").replace('"', "&quot;")
        n = sum(1 for p in pics if alt in p)
        if b.logo_png("light"):
            want = len(d.get("sections") or []) if self.sheet else 1
            self.check(n >= max(1, want), "LOGO", "docx", "the brand logo is missing: %d of %d expected (%s); rebuild with build_doc.py" % (n, max(1, want), "one per title band" if self.sheet else "on the cover"))  # check: LOGO FAIL
        else:
            where = "title band" if self.sheet else "cover"
            why = ("the project logo %s could not be read (an SVG needs PyMuPDF), so the %s has none" % (os.path.basename(b.logo), where) if b.logo
                   else "no project logo found, so the %s has none; add a logo to the brand profile (brand.json \"logo\") if the project has one" % where)
            self.check(False, "LOGO_NONE", "docx", why, warn=True)  # check: LOGO_NONE WARN

    def a11y_docx(self, doc, d, body, pics):
        """DOC_A11Y (D-64): w:lang, a repeating header row on every data table, alt text on every picture, and text
        colors only from the 4.5:1 list."""
        from docx.oxml.ns import qn
        before, lang = len(self.fails), A11Y.get("lang", "en-US")
        sx = doc.styles.element.xml
        self.check(('<w:lang w:val="%s"' % lang) in sx, "DOC_A11Y", "docx", "no document language: styles need w:lang %s" % lang)  # check: DOC_A11Y FAIL
        fill = lambda tc: (tc.find(qn("w:tcPr")) is not None and tc.find(qn("w:tcPr")).find(qn("w:shd")) is not None
                           and str(tc.find(qn("w:tcPr")).find(qn("w:shd")).get(qn("w:fill"))).upper() == T["colors"]["dark"])
        data = rep = 0
        for tbl in doc.element.body.iter(qn("w:tbl")):
            trs = tbl.findall(qn("w:tr"))
            if len(trs) < 2 or not all(fill(tc) for tc in trs[0].findall(qn("w:tc"))) or not trs[0].findall(qn("w:tc")):
                continue  # layout tables (cover, callouts, cards, bands) have no header row
            data += 1
            pr = trs[0].find(qn("w:trPr"))
            rep += pr is not None and pr.find(qn("w:tblHeader")) is not None
        want = self.counts.get("tab", 0) + (0 if self.sheet else bool(d.get("glossary")) + bool(d.get("revisions")))
        self.check(rep == data and data >= want, "DOC_A11Y", "docx", "%d of %d data tables repeat their header row (w:tblHeader); expected %d or more" % (rep, max(data, want), want))  # check: DOC_A11Y FAIL
        noalt = sum(1 for p in pics if 'descr="' not in p or 'descr=""' in p)
        self.check(not noalt, "DOC_A11Y", "docx", "%d of %d pictures have no alt text" % (noalt, len(pics)))  # check: DOC_A11Y FAIL
        ok_fg = {T["colors"].get(f, f).upper() for f, _ in A11Y.get("text_pairs") or []}
        fx = body + " ".join(f._element.xml for sec in doc.sections for f in (sec.footer, sec.first_page_footer, sec.even_page_footer))
        odd = sorted({c.upper() for c in re.findall(r'<w:color w:val="([0-9A-Fa-f]{6})"', fx)} - ok_fg)
        self.check(not odd, "DOC_A11Y", "docx", "text colors %s are not on the 4.5:1 list (doc_theme.json a11y.text_pairs)" % ", ".join(odd))  # check: DOC_A11Y FAIL
        if len(self.fails) == before:
            self.passes.append("DOC_A11Y pass: docx w:lang %s, %d tables with repeating header rows, alt text on %d pictures, text colors from the 4.5:1 list" % (lang, rep, len(pics)))

    def pdf(self, path, d):
        try:
            import fitz
            doc = fitz.open(path)
        except Exception as ex:
            self.check(False, "PDF_OPEN", path, "can't open: %r" % ex); return  # check: PDF_OPEN FAIL
        self.check(doc.page_count >= (1 if self.sheet else 3), "PDF_PAGES", "pdf", "only %d pages" % doc.page_count)  # check: PDF_PAGES FAIL
        not_emb = sorted({f[3] for pg in doc for f in pg.get_fonts() if f[1] in ("n/a", "")})
        self.check(not not_emb, "FONTS_EMBEDDED", "pdf", "fonts not embedded: %s" % ", ".join(not_emb))  # check: FONTS_EMBEDDED FAIL
        allf = " ".join(sorted({f[3] for pg in doc for f in pg.get_fonts()}))
        body = T["fonts"]["body"]
        self.check(body in pdf_families(doc), "PDF_FONT", "pdf", "the PDF doesn't use %s (font substitution): %s" % (body, allf[:120]))  # check: PDF_FONT FAIL
        toc = doc.get_toc()
        apps = self.counts.get("apps", 0)
        self.check(sum(1 for t in toc if t[0] == 1) >= len(d["sections"]) + apps, "BOOKMARKS", "pdf", "PDF bookmarks missing (%d level-1)" % sum(1 for t in toc if t[0] == 1))  # check: BOOKMARKS FAIL
        self.check((doc.metadata or {}).get("title") == d["meta"]["title"], "PDF_TITLE", "pdf", "PDF title metadata must match meta.title")  # check: PDF_TITLE FAIL
        kind, val = doc.xref_get_key(doc.pdf_catalog(), "Lang"); lang = A11Y.get("lang", "en-US"); title = (doc.metadata or {}).get("title")
        ok1 = self.check(kind == "string" and val == lang, "DOC_A11Y", "pdf", "no /Lang %s in the PDF catalog (found %s)" % (lang, val if kind != "null" else "none"))  # check: DOC_A11Y FAIL
        ok2 = self.check(bool(title) and title == d["meta"]["title"], "DOC_A11Y", "pdf", "PDF /Title is %r; set it to meta.title" % (title or ""))  # check: DOC_A11Y FAIL
        if ok1 and ok2:
            self.passes.append("DOC_A11Y pass: PDF /Lang %s, /Title %r" % (val, title))
        text = " ".join(" ".join(pg.get_text().split()) for pg in doc)
        self.check("Error!" not in text and "[Screenshot needed" not in text or self.draft, "PDF_TEXT", "pdf", "field error or missing-screenshot marker in the PDF")  # check: PDF_TEXT FAIL
        tj = os.path.splitext(path)[0] + ".toc.json"
        if os.path.isfile(tj):
            info = json.load(open(tj))
            if self.sheet:
                self.layout(doc, info); self.sheet_pages(doc, info, d); self.sheet_band_links(doc, info, d); return
            self.check(info.get("stable") and bool(info.get("toc")), "TOC_PAGES", "pdf", "contents page numbers didn't settle; rebuild")  # check: TOC_PAGES FAIL
            texts = [" ".join(pg.get_text().split()) for pg in doc]
            first = min((info.get("toc") or {"x": 3}).values())
            contents = " ".join(texts[1:max(2, first - 1)])
            golinks = sum(1 for pg in doc for l in pg.get_links() if l.get("kind") in (fitz.LINK_GOTO, fitz.LINK_NAMED))
            want = len(info.get("toc") or {}) + len(info.get("index") or [])
            self.check(golinks >= want, "TOC_LINKS", "pdf", "%d working internal links in the PDF for %d contents and index entries" % (golinks, want))  # check: TOC_LINKS FAIL
            for term, lab in info.get("index") or []:
                self.check(bool(lab) and lab in (info.get("all") or {}), "INDEX_PAGES", "pdf", "index entry %r has no page" % term)  # check: INDEX_PAGES FAIL
            for lab, pgn in (info.get("toc") or {}).items():
                ok = pgn <= doc.page_count and " ".join(lab.split()) in texts[pgn - 1]
                self.check(ok, "TOC_PAGES", "pdf", "contents says %r is on page %s, but it isn't" % (lab, pgn))  # check: TOC_PAGES FAIL
                self.check(" ".join(lab.split()) in contents, "TOC_ENTRY", "pdf", "%r missing from the contents page" % lab)  # check: TOC_ENTRY FAIL
            self.layout(doc, info)
        else:
            self.check(False, "TOC_PAGES", "pdf", "no .toc.json next to the PDF; build with build_doc.py --pdf")  # check: TOC_PAGES FAIL

    def sheet_pages(self, doc, info, d):
        """QS_PAGE_COUNT (quick-start) or SHEET_PAGE_FIT (other sheets): each top-level section within pages_per_section."""
        fmt = sheet_format(d["meta"].get("doc_type")); mx = int(sheet_cfg(d["meta"])["pages_per_section"])
        code, noun, secs, want = fmt["page_code"], fmt["section_noun"], info.get("sections") or [], len(d.get("sections") or [])
        before = len(self.fails)
        self.check(len(secs) == want, code, "pdf", "found %d of %d %ss in the PDF; rebuild with build_doc.py --pdf" % (len(secs), want, noun))  # check: QS_PAGE_COUNT FAIL # check: SHEET_PAGE_FIT FAIL
        for n, (title, start, k) in enumerate(secs, 1):
            self.check(k <= mx, code, "%s %d" % (noun, n), fmt["page_fail"].format(k=k, m=mx))  # check: QS_PAGE_COUNT FAIL # check: SHEET_PAGE_FIT FAIL
        if len(self.fails) == before:
            self.passes.append("%s pass: %s" % (code, fmt["page_pass"].format(n=len(secs), p=doc.page_count)))

    def sheet_band_links(self, doc, info, d):
        """D-26 SHEET_SUBTITLE: meta.subtitle is printed on page 1. D-27 SOURCE_LINK: every linked source's URL is a link
        annotation on every page (footer), unless the sources did not fit the footer and stay in the doc properties."""
        import fitz
        norm, m = (lambda s: " ".join(str(s).split())), d["meta"]
        sub = norm(plain(m.get("subtitle")))
        if sub:
            ok = doc.page_count > 0 and sub in norm(doc[0].get_text())
            if self.check(ok, "SHEET_SUBTITLE", "pdf", "meta.subtitle %r is not on page 1; rebuild with build_doc.py --pdf" % sub[:60]):  # check: SHEET_SUBTITLE FAIL
                self.passes.append("SHEET_SUBTITLE pass: subtitle on page 1")
        urls = [u for s in m.get("sources") or [] for _, u in source_links(s)]
        if not urls:
            return
        if info.get("sources_in_footer") is False:
            self.passes.append("SOURCE_LINK pass: sources did not fit the footer; kept in doc properties"); return
        key = lambda u: str(u or "").rstrip("/")
        pages = [{key(l.get("uri")) for l in pg.get_links() if l.get("kind") == fitz.LINK_URI} for pg in doc]
        before = len(self.fails)
        for u in urls:
            miss = [str(i) for i, got in enumerate(pages, 1) if key(u) not in got]
            self.check(not miss, "SOURCE_LINK", "footer", "no link annotation to %s on page %s; write the source as [title](https://...) and rebuild" % (u, ", ".join(miss[:5])))  # check: SOURCE_LINK FAIL
        if len(self.fails) == before:
            self.passes.append("SOURCE_LINK pass: %d linked sources clickable on %d pages" % (len(urls), len(pages)))

    def layout(self, doc, info):
        """What a reader sees: callouts kept whole, no words broken inside table cells, no mostly empty body pages."""
        norm = lambda s: " ".join(str(s).split())
        texts = [norm(pg.get_text()) for pg in doc]
        start = max(0, (info.get("all") or {}).get(info.get("first_body"), 1) - 1)
        full, offs = "", []
        for t in texts:
            offs.append(len(full)); full += t + " "
        page_of = lambda pos: max(i for i, o in enumerate(offs) if o <= pos)
        pos = offs[start] if start < len(offs) else 0
        for c in info.get("callouts") or []:
            h = full.find(norm(c["head"]), pos)
            if h < 0:
                continue
            ti = full.rfind(norm(c["title"]), max(pos, h - 400), h + len(norm(c["head"])))
            tl = full.find(norm(c["tail"]), h)
            if ti < 0 or tl < 0:
                continue
            pt, pb = page_of(ti), page_of(tl)
            self.check(pt == pb, "CALLOUT_SPLIT", "pdf", "callout %r starts on page %d but its body ends on page %d; keep a callout on one page" % (c["title"], pt + 1, pb + 1))  # check: CALLOUT_SPLIT FAIL
            pos = tl + len(norm(c["tail"]))
        tw = set(info.get("table_words") or [])
        broken = set()
        for i, pg in enumerate(doc):
            ws = pg.get_text("words")
            for a, b in zip(ws, ws[1:]):
                if a[5] == b[5] and a[6] != b[6] and not a[4].endswith("-"):
                    joined = (a[4] + b[4]).strip(".,;:()!?\"'")
                    if joined in tw and a[4].strip(".,;:()!?\"'") not in tw:
                        broken.add((joined, i + 1, a[4] + " / " + b[4]))
        for joined, pnum, how in sorted(broken):
            self.check(False, "WORD_BREAK", "pdf", "%r is broken across lines in a table on page %d (%s); widen the column or shorten the text" % (joined, pnum, how))  # check: WORD_BREAK FAIL
        if info.get("sheet"):
            return  # each sheet section is a page by design; QS_PAGE_COUNT / SHEET_PAGE_FIT judge the fit
        self.n += 1
        apx = (info.get("all") or {}).get(info.get("first_appendix")) or doc.page_count + 1
        starts = {(info.get("all") or {}).get(h) for h in info.get("h1") or []}
        chapter_end = info.get("page_break_level", 1) == 1  # a chapter's last page may be short when each chapter starts a page
        low = []
        for i in range(start, min(apx - 1, doc.page_count)):
            if chapter_end and (i + 2) in starts:
                continue
            pg = doc[i]; H = pg.rect.height
            ys = [b[3] for b in pg.get_text("blocks") if b[4].strip() and b[3] < H - 0.95 * 72]
            ys += [im["bbox"][3] for im in pg.get_image_info() if im["bbox"][3] < H - 0.95 * 72]
            if ys and (max(ys) - 72) / (H - 2 * 72) < 0.45:
                low.append(i + 1)
        self.check(not low, "PAGE_FILL", "pdf", "mid-chapter pages %s end above 45%% of the page; check for a figure or table pushed to the next page, or merge short sections" % ", ".join(map(str, low)), warn=True)  # check: PAGE_FILL WARN


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("doc"); ap.add_argument("--docx"); ap.add_argument("--pdf"); ap.add_argument("--draft", action="store_true")
    ap.add_argument("--brand", help="brand profile (default: the one the build uses)")
    a = ap.parse_args()
    try:
        d = json.load(open(a.doc, encoding="utf-8"))
    except (OSError, ValueError) as ex:
        print("LINT: 1 checks, 1 failed, 0 warnings\nFAIL [JSON] %s: %s" % (a.doc, ex)); sys.exit(1)
    L = Lint(a.draft)
    L.brand = use_brand(BR.load(source=a.doc, brand=a.brand))
    L.notes.append(L.brand.lines()[0])
    L.a11y_theme()
    L.source(d, os.path.dirname(os.path.abspath(a.doc)))
    L.term_brand()
    if a.docx:
        if L.check(os.path.isfile(a.docx), "DOCX_MISSING", a.docx, "file not found"):  # check: DOCX_MISSING FAIL
            L.docx(a.docx, d)
    if a.pdf:
        if L.check(os.path.isfile(a.pdf), "PDF_MISSING", a.pdf, "file not found"):  # check: PDF_MISSING FAIL
            L.pdf(a.pdf, d)
    print("LINT: %d checks, %d failed, %d warnings" % (L.n, len(L.fails), len(L.warns)))
    for x in L.fails + L.warns:
        print(x)
    for x in L.passes + L.notes:
        print(x)
    sys.exit(1 if L.fails else 0)


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Self-tests for build_form.py (rows D-13, D-46, D-48..D-51), build_module.py (rows D-14, D-28 SCORM 1.2, D-52 resume)
and build_sheet.py (row D-53 Excel tables).

  python3 test_builders.py                 run every case
  python3 test_builders.py --only CODE[,CODE]   run cases tagged with these codes (--only list prints the codes)
Each gate code has a positive and a negative case named after its request row. Cases write to a temp dir.
Prints "<n> passed, <k> failed"; exits 1 on any failure. Node cases print UNVERIFIED (never a pass) when node is missing.
"""
import copy, json, os, re, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_form as BF
import brand as BR
import build_module as BM

TMP = tempfile.mkdtemp(prefix="composer_builders_test_")
# Hermetic brand: every case (and every CLI it runs) resolves the brand of an empty project, so results never depend on
# the folder the tests run in. The branded-project cases point BRAND_PROJECT at their own synthetic project.
os.environ["BRAND_PROJECT"] = tempfile.mkdtemp(prefix="composer_brand_empty_")
CASES = []


def case(name, code):
    def deco(fn):
        CASES.append((name, code, fn)); return fn
    return deco


def quiet():
    lines = []
    return lines, lines.append


def has(g, level, code, text=None):
    return any(lv == level and c == code and (text is None or text in m) for lv, c, m in g.msgs)


# ---------------------------------------------------------------- D-13 fixtures (synthetic)
def survey():
    return {"meta": {"title": "Synthetic Training Pulse Survey", "kind": "survey", "audience": "Synthetic course attendees",
                     "purpose": "Rate the synthetic course so the next run improves.", "owner": "Synthetic Training Team", "version": "1.0",
                     "date": "2026-10-11", "sources": ["Synthetic example: course feedback request"], "anonymous": True, "phi_approved": False, "est_minutes": 3},
            "sections": [{"title": "Your Rating", "items": [
                {"id": "Q1", "type": "single", "label": "Which session did you attend?", "required": True, "options": ["Morning", "Afternoon"]},
                {"id": "Q2", "type": "scale", "label": "The course met my needs.", "required": True, "scale": {"min": 1, "max": 5, "low": "Strongly disagree", "high": "Strongly agree"}},
                {"id": "Q3", "type": "scale", "label": "How satisfied are you with the training?", "required": True, "scale": {"min": 1, "max": 5, "low": "Very dissatisfied", "high": "Very satisfied"}},
                {"id": "Q4", "type": "scale", "label": "How satisfied are you with the trainer?", "required": True, "scale": {"min": 1, "max": 5, "low": "Very dissatisfied", "high": "Very satisfied"}},
                {"id": "Q5", "type": "multi", "label": "Which topics should get more time?", "options": ["Lookups", "Appeals", "Portals"]},
                {"id": "Q6", "type": "dropdown", "label": "How did you hear about the course?", "options": ["Manager", "Newsletter", "Coworker"]},
                {"id": "Q7", "type": "longtext", "label": "What should we change?"},
                {"id": "Q8", "type": "checkbox", "label": "I would take a follow-up course."}]}]}


def row_scenario_survey():
    s = survey()
    items = s["sections"][0]["items"]
    items[2]["label"] = "How satisfied are you with the training and the trainer?"
    del items[3]
    items[1]["scale"] = {"min": 1, "max": 7, "low": "Strongly disagree", "high": "Strongly agree"}
    return s


def form():
    return {"meta": {"title": "Synthetic Equipment Request Form", "kind": "form", "audience": "Synthetic team members", "purpose": "Request a synthetic laptop dock.",
                     "owner": "Synthetic Support Desk", "version": "1.0", "date": "2026-10-11", "sources": ["Synthetic example: equipment checklist"], "anonymous": False, "phi_approved": False},
            "sections": [{"title": "Request", "items": [
                {"id": "F1", "type": "text", "label": "Department", "required": True},
                {"id": "F2", "type": "date", "label": "Date needed", "required": True},
                {"id": "F3", "type": "single", "label": "Request type", "options": ["New", "Replacement"]}]},
                {"title": "Approval", "items": [
                    {"id": "F4", "type": "signature", "label": "Manager signature", "required": True},
                    {"id": "F5", "type": "date", "label": "Date signed", "required": True}]}]}


@case("D-13 SURVEY_ITEM row scenario: Q3 \"the training and the trainer\" plus 1-5 and 1-7 scales fails twice", "SURVEY_ITEM")
def _():
    g = BF.gate(row_scenario_survey())
    return has(g, "FAIL", "SURVEY_ITEM", "Q3 asks two things (\"the training\" and \"the trainer\")") and \
        any(m == "mixed scales (1-5, 1-7)" for lv, c, m in g.msgs if c == "SURVEY_ITEM")


@case("D-13 SURVEY_ITEM row scenario: build refuses on FAIL without --draft and writes nothing", "SURVEY_ITEM")
def _():
    out = os.path.join(TMP, "refused_survey")
    lines, say = quiet()
    g, files = BF.run(row_scenario_survey(), out, say=say)
    return files is None and not os.path.exists(out + ".pdf") and lines[-1].startswith("refused") and re.match(r"FORM GATE: \d+ checks, 2 failed", lines[-2])


@case("D-13 SURVEY_ITEM negative: leading wording, one-option single and scale without end labels fail", "SURVEY_ITEM")
def _():
    s = survey(); it = s["sections"][0]["items"]
    it[1]["label"] = "Don't you agree the course met your needs?"
    it[0]["options"] = ["Morning"]
    it[3]["scale"] = {"min": 1, "max": 5, "low": "Very dissatisfied"}
    g = BF.gate(s)
    return has(g, "FAIL", "SURVEY_ITEM", "leading wording \"Don't you agree\"") and has(g, "FAIL", "SURVEY_ITEM", "Q1: single needs at least 2 options (has 1)") \
        and has(g, "FAIL", "SURVEY_ITEM", "Q4: scale needs both end labels")


@case("D-13 SURVEY_ITEM positive: fixed survey (split items, one 1-5 scale) passes", "SURVEY_ITEM")
def _():
    g = BF.gate(survey())
    return "SURVEY_ITEM" not in g.codes() and g.failed == 0


@case("D-13 SURVEY_ITEM positive: compound_ok allows a deliberate compound item", "SURVEY_ITEM")
def _():
    s = row_scenario_survey(); it = s["sections"][0]["items"]
    it[2]["compound_ok"] = True; it[1]["scale"]["max"] = 5
    return "SURVEY_ITEM" not in BF.gate(s).codes()


@case("D-13 FORM_FIELDS_MATCH row scenario: fixed survey builds a PDF whose field count equals the question count", "FORM_FIELDS_MATCH")
def _():
    out = os.path.join(TMP, "fixed_survey")
    lines, say = quiet()
    g, files = BF.run(survey(), out, say=say)
    import fitz
    names = {w.field_name.split(".")[0] for p in fitz.open(out + ".pdf") for w in p.widgets()}
    n_items = sum(1 for _ in BF.all_items(survey()))
    return files and g.failed == 0 and len(names) == n_items == 8 and not has(g, "FAIL", "FORM_FIELDS_MATCH")


@case("D-13 FORM_FIELDS_MATCH negative: a PDF missing one item's widget fails", "FORM_FIELDS_MATCH")
def _():
    import fitz
    src = os.path.join(TMP, "fm_src.pdf"); BF.build_pdf(survey(), src)
    doc = fitz.open(src)
    page = next(p for p in doc if any(w.field_name == "Q7" for w in p.widgets()))
    page.delete_widget(next(w for w in page.widgets() if w.field_name == "Q7")); bad = os.path.join(TMP, "fm_bad.pdf"); doc.save(bad)
    g = BF.fields_match(bad, survey())
    return has(g, "FAIL", "FORM_FIELDS_MATCH", "PDF has 7 field names for 8 items; missing: Q7")


@case("D-13 FORM_FIELD negative: PHI field (date of birth) fails unless phi_approved", "FORM_FIELD")
def _():
    f = form(); f["sections"][0]["items"].append({"id": "F9", "type": "date", "label": "Date of birth"})
    g1 = BF.gate(f)
    f["meta"]["phi_approved"] = True
    g2 = BF.gate(f)
    return has(g1, "FAIL", "FORM_FIELD", "F9 asks for PHI (\"Date of birth\")") and not has(g2, "FAIL", "FORM_FIELD")


@case("D-13 FORM_FIELD negative: empty label, duplicate id and MRN fail; signature without date warns", "FORM_FIELD")
def _():
    f = form(); a = f["sections"][0]["items"]
    a.append({"id": "F1", "type": "text", "label": ""})
    a.append({"id": "F7", "type": "text", "label": "MRN"})
    del f["sections"][1]["items"][1]
    g = BF.gate(f)
    return has(g, "FAIL", "FORM_FIELD", "F1: empty label") and has(g, "FAIL", "FORM_FIELD", "duplicate id F1") \
        and has(g, "FAIL", "FORM_FIELD", "F7 asks for PHI (\"MRN\")") and has(g, "WARN", "FORM_FIELD", "F4: signature without a date item")


@case("D-13 FORM_FIELD positive: clean form (signature paired with a date) has no FORM_FIELD findings", "FORM_FIELD")
def _():
    g = BF.gate(form())
    return not [m for m in g.msgs if m[1] == "FORM_FIELD"]


@case("D-13 PLACEHOLDER negative: TBD and First Last fail", "PLACEHOLDER")
def _():
    s = survey(); s["meta"]["owner"] = "TBD"; s["sections"][0]["items"][6]["help"] = "Sign as First Last"
    g = BF.gate(s)
    return has(g, "FAIL", "PLACEHOLDER", "owner: placeholder text 'TBD'") and has(g, "FAIL", "PLACEHOLDER", "'First Last'")


@case("D-13 PLACEHOLDER positive: words containing todo/tbd letters do not trip the rule", "PLACEHOLDER")
def _():
    s = survey(); s["meta"]["purpose"] = "Rate the course; your feedback on the outbound toddler-free agenda helps."
    return "PLACEHOLDER" not in BF.gate(s).codes()


@case("D-13 SOURCES_MISSING negative/positive: empty sources fail, listed sources pass", "SOURCES_MISSING")
def _():
    s = survey(); s["meta"]["sources"] = []
    return has(BF.gate(s), "FAIL", "SOURCES_MISSING") and "SOURCES_MISSING" not in BF.gate(survey()).codes()


@case("D-13 SOURCES_MISSING negative: blank strings do not count as sources", "SOURCES_MISSING")
def _():
    s = survey(); s["meta"]["sources"] = ["  "]
    return has(BF.gate(s), "FAIL", "SOURCES_MISSING")


@case("D-13 TITLE_CASE negative: lowercase title fails with the corrected title", "TITLE_CASE")
def _():
    s = survey(); s["meta"]["title"] = "synthetic training pulse survey for sign-off"
    return has(BF.gate(s), "FAIL", "TITLE_CASE", "use \"Synthetic Training Pulse Survey for Sign-off\"")


@case("D-13 TITLE_CASE positive: Proper Case title passes", "TITLE_CASE")
def _():
    return "TITLE_CASE" not in BF.gate(survey()).codes()


@case("D-13 OUTPUT: PDF is Letter, brand body font embedded, logo only when the brand has one, footer Page X of Y, tooltips = labels; xlsx sheets, fill, frozen header", "OUTPUT")
def _():
    import fitz
    from openpyxl import load_workbook
    out = os.path.join(TMP, "out_form")
    lines, say = quiet()
    g, files = BF.run(form(), out, say=say)
    doc = fitz.open(out + ".pdf")
    p0 = doc[0]
    text = "".join(p.get_text() for p in doc)
    fonts = {f[3] for p in doc for f in p.get_fonts()}
    tips = {w.field_name: w.field_label for p in doc for w in p.widgets()}
    T = BF.cur(); fam = BR.font_name(T["brand"].measure_path("body"))[0]  # the brand body font, else the bundled fallback
    ok_pdf = abs(p0.rect.width - 612) < 1 and abs(p0.rect.height - 792) < 1 and any(fam in f or fam.replace(" ", "") in f for f in fonts) \
        and "Page 1 of %d" % len(doc) in text and "Synthetic Equipment Request Form v1.0" in text and "MM/DD/YYYY" in text \
        and tips.get("F1") == "Department" and bool(p0.get_images()) == bool(T["brand"].logo_png("light")) and any(x.startswith("PDF font: %s (embedded" % fam) for x in lines)
    wb = load_workbook(out + ".xlsx")
    q = wb["Questions"]
    ok_x = wb.sheetnames == ["Questions", "About"] and [c.value for c in q[1]] == ["ID", "Section", "Question", "Type", "Required", "Options/Scale"] \
        and q.freeze_panes == "A2" and wb["About"].freeze_panes == "A2" and q["A1"].fill.fgColor.rgb.endswith(T["colors"]["dark"]) and q.max_row == 6
    return bool(files) and ok_pdf and ok_x


@case("D-13 OUTPUT negative: --draft builds despite FAIL and marks DRAFT; no item is split across pages", "OUTPUT")
def _():
    import fitz
    s = row_scenario_survey()
    for k in range(9, 40):
        s["sections"][0]["items"].append({"id": "X%d" % k, "type": "longtext", "label": "Synthetic comment box %d" % k, "rows": 3})
    out = os.path.join(TMP, "draft_long")
    lines, say = quiet()
    g, files = BF.run(s, out, draft=True, say=say)
    doc = fitz.open(out + ".pdf")
    split = False
    for p in doc:   # every label on a page must have its widget on the same page
        txt = p.get_text()
        for w in p.widgets():
            nm = w.field_name.split(".")[0]
            lab = next(it["label"] for _, it in BF.all_items(s) if it["id"] == nm)
            if lab[:20] not in txt.replace("\n", " "):
                split = True
    return bool(files) and g.failed == 2 and "DRAFT" in doc[0].get_text() and len(doc) > 2 and not split


# ---------------------------------------------------------------- D-14 fixtures (synthetic)
def q(stem, n=3, fb=None, correct=0, multi=False):
    fb = fb if fb is not None else list(range(n))
    return {"stem": stem, "multi": multi, "options": [{"text": "Synthetic option %d" % (i + 1), "correct": i == correct,
                                                       "feedback": ("Synthetic feedback %d." % (i + 1)) if i in fb else ""} for i in range(n)]}


def module(screens=None):
    return {"meta": {"title": "Synthetic Claim Status Basics", "audience": "Synthetic new team members",
                     "objectives": ["Identify the synthetic claim status codes", "Choose the next step for a pending claim"],
                     "duration_min": 10, "version": "1.0", "date": "2026-10-11", "sources": ["Synthetic example: claim status outline"]},
            "screens": screens or [
                {"id": "S1", "type": "intro", "title": "Welcome", "body": "Synthetic intro."},
                {"id": "S2", "type": "content", "title": "Status Codes", "body": "Synthetic content.", "bullets": ["Pending", "Paid"]},
                {"id": "S3", "type": "question", "title": "Quick Check", "body": "Choose one.", "question": q("Which code means paid?")},
                {"id": "S4", "type": "content", "title": "Next Steps", "body": "Synthetic content."},
                {"id": "S5", "type": "scenario", "title": "A Pending Claim", "body": "Synthetic scenario.",
                 "choices": [{"text": "Wait", "correct": False, "feedback": "Synthetic: waiting delays payment."},
                             {"text": "Check status", "correct": True, "feedback": "Synthetic: right."}]},
                {"id": "S6", "type": "summary", "title": "Summary", "body": "Synthetic summary."}]}


def row_scenario_module():
    sc = [{"id": "S1", "type": "intro", "title": "Welcome", "body": "Synthetic intro."},
          {"id": "S2", "type": "question", "title": "Warm Up", "body": "Choose one.", "question": q("Synthetic warm-up?")}]
    sc += [{"id": "S%d" % k, "type": "content", "title": "Topic %d" % k, "body": "Synthetic text-only screen."} for k in range(3, 9)]
    sc += [{"id": "S9", "type": "question", "title": "Check", "body": "Choose one.", "question": q("Synthetic check?", 4, fb=[0, 1])},
           {"id": "S10", "type": "content", "title": "Wrap", "body": "Synthetic text."},
           {"id": "S11", "type": "practice", "title": "Practice", "body": "Choose one.", "question": q("Synthetic practice?")},
           {"id": "S12", "type": "summary", "title": "Summary", "body": "Synthetic summary."}]
    return module(sc)


@case("D-14 ELEARN_INTERACTION row scenario: 12 screens, S3-S8 text-only and feedback on 2 of 4 options fails twice", "ELEARN_INTERACTION")
def _():
    m = row_scenario_module()
    g = BM.gate(m)
    msgs = [x for lv, c, x in g.msgs if c == "ELEARN_INTERACTION" and lv == "FAIL"]
    return len(m["screens"]) == 12 and "screens S3-S8 have no interaction (maximum 4 in a row)" in msgs and "S9: options C, D have no feedback" in msgs and len(msgs) == 2


@case("D-14 ELEARN_INTERACTION negative: a question with no correct option and a question screen with no question fail", "ELEARN_INTERACTION")
def _():
    m = module(); m["screens"][2]["question"]["options"][0]["correct"] = False
    m["screens"][3]["type"] = "question"
    g = BM.gate(m)
    return has(g, "FAIL", "ELEARN_INTERACTION", "S3: question has no correct option") and has(g, "FAIL", "ELEARN_INTERACTION", "S4: question screen has no question")


@case("D-14 ELEARN_INTERACTION positive: fixed module (4 text screens in a row max, feedback on every option) passes", "ELEARN_INTERACTION")
def _():
    m = row_scenario_module()
    m["screens"][5] = {"id": "S6", "type": "practice", "title": "Mid Check", "body": "Choose one.", "question": q("Synthetic mid check?")}
    m["screens"][8]["question"] = q("Synthetic check?", 4)
    g = BM.gate(m)
    return "ELEARN_INTERACTION" not in g.codes() and g.failed == 0


@case("D-14 ELEARN_OBJECTIVE negative: understand / know objectives fail", "ELEARN_OBJECTIVE")
def _():
    m = module(); m["meta"]["objectives"] = ["Understand claim status codes", "Learners will be able to know the next step", "Be aware of payer rules"]
    g = BM.gate(m)
    return has(g, "FAIL", "ELEARN_OBJECTIVE", "objective 1 starts with \"understand\"") and has(g, "FAIL", "ELEARN_OBJECTIVE", "objective 2 starts with \"know\"") \
        and has(g, "FAIL", "ELEARN_OBJECTIVE", "objective 3 starts with \"be aware")


@case("D-14 ELEARN_OBJECTIVE positive: observable verbs pass", "ELEARN_OBJECTIVE")
def _():
    m = module(); m["meta"]["objectives"] = ["Identify the codes", "You will be able to complete the checklist"]
    return "ELEARN_OBJECTIVE" not in BM.gate(m).codes() and not [x for x in BM.gate(m).msgs if x[1] == "ELEARN_OBJECTIVE"]


@case("D-14 ELEARN_OFFLINE negative: an https link or remote image fails and no file is written", "ELEARN_OFFLINE")
def _():
    m = module(); m["screens"][1]["body"] = "See the [synthetic policy](https://example.invalid/policy)."
    m["screens"][3].update({"image": "https://example.invalid/pic.png", "alt": "Synthetic picture"})
    out = os.path.join(TMP, "offline_bad")
    lines, say = quiet()
    g, files = BM.run(m, out, say=say)
    return has(g, "FAIL", "ELEARN_OFFLINE", "built HTML loads external URLs") and files is None and not os.path.exists(out + ".html")


@case("D-14 ELEARN_OFFLINE positive: built HTML has no http(s) in src/href/url() and inlines images", "ELEARN_OFFLINE")
def _():
    from PIL import Image
    Image.new("RGB", (40, 20), "#1F5FAD").save(os.path.join(TMP, "synthetic_box.png"))
    m = module(); m["screens"][1].update({"image": "synthetic_box.png", "alt": "Synthetic blue box"})
    out = os.path.join(TMP, "offline_ok")
    lines, say = quiet()
    g, files = BM.run(m, out, base_dir=TMP, say=say)
    doc = open(out + ".html", encoding="utf-8").read()
    return bool(files) and "ELEARN_OFFLINE" not in g.codes() and not re.search(r"https?://", doc) and 'src="data:image/png;base64,' in doc


@case("D-14 ELEARN_STRUCTURE negative: duplicate screen id and image without alt fail", "ELEARN_STRUCTURE")
def _():
    m = module(); m["screens"][1]["id"] = "S1"; m["screens"][3]["image"] = "missing_synthetic.png"
    g = BM.gate(m, TMP)
    return has(g, "FAIL", "ELEARN_STRUCTURE", "duplicate screen id S1") and has(g, "FAIL", "ELEARN_STRUCTURE", "S4: image needs alt text") \
        and has(g, "FAIL", "ELEARN_STRUCTURE", "image file not found")


@case("D-14 ELEARN_STRUCTURE positive: clean module passes; a missing summary screen is added with a WARN", "ELEARN_STRUCTURE")
def _():
    m = module(); g1 = BM.gate(m)
    m2 = module(); del m2["screens"][-1]
    g2 = BM.gate(m2)
    return "ELEARN_STRUCTURE" not in g1.codes() and has(g2, "WARN", "ELEARN_STRUCTURE", "Summary screen") and BM.screens_of(m2)[-1]["type"] == "summary"


@case("D-14 PLACEHOLDER negative: TODO and lorem fail", "PLACEHOLDER")
def _():
    m = module(); m["screens"][1]["body"] = "TODO write this"; m["screens"][2]["question"]["options"][1]["feedback"] = "Lorem ipsum"
    g = BM.gate(m)
    return has(g, "FAIL", "PLACEHOLDER", "'TODO'") and has(g, "FAIL", "PLACEHOLDER", "S3 feedback")


@case("D-14 PLACEHOLDER positive: clean module has none", "PLACEHOLDER")
def _():
    return "PLACEHOLDER" not in BM.gate(module()).codes()


@case("D-14 SOURCES_MISSING negative/positive: empty sources fail, listed sources pass", "SOURCES_MISSING")
def _():
    m = module(); m["meta"]["sources"] = []
    return has(BM.gate(m), "FAIL", "SOURCES_MISSING") and "SOURCES_MISSING" not in BM.gate(module()).codes()


@case("D-14 TITLE_CASE negative/positive: lowercase module title fails with correction; Proper Case passes", "TITLE_CASE")
def _():
    m = module(); m["meta"]["title"] = "synthetic claim status basics"
    return has(BM.gate(m), "FAIL", "TITLE_CASE", "use \"Synthetic Claim Status Basics\"") and "TITLE_CASE" not in BM.gate(module()).codes()


@case("D-14 OUTPUT: HTML has N screens, Back/Next buttons, progress, aria-live, labels; storyboard front matter and table", "OUTPUT")
def _():
    from html.parser import HTMLParser
    out = os.path.join(TMP, "out_module")
    lines, say = quiet()
    g, files = BM.run(module(), out, say=say)
    doc = open(out + ".html", encoding="utf-8").read()
    tags = []
    class P(HTMLParser):
        def handle_starttag(self, t, a): tags.append((t, dict(a)))
    P().feed(doc)
    n = sum(1 for t, a in tags if t == "section" and "data-screen" in a)
    ids = {a.get("id"): t for t, a in tags}
    labels = {a.get("for") for t, a in tags if t == "label"}
    sb = open(out + ".storyboard.md", encoding="utf-8").read()
    return n == 6 and ids.get("back") == "button" and ids.get("next") == "button" and "Screen 1 of 6" in doc \
        and any(a.get("aria-live") == "polite" and "feedback" in a.get("class", "") for t, a in tags) \
        and all(a.get("id") in labels for t, a in tags if t == "input") and ":focus-visible" in doc and "Module complete" in doc \
        and not any(t in ("link",) or (t == "script" and "src" in a) for t, a in tags) \
        and sb.startswith("---\ntitle: ") and "duration_min: 10" in sb and "| Screen | Type | On Screen | Interaction | Feedback |" in sb \
        and re.search(r"^MODULE GATE: \d+ checks, 0 failed$", "\n".join(lines), re.M)


@case("D-14 OUTPUT negative: build refuses on FAIL without --draft; --draft writes the files", "OUTPUT")
def _():
    out = os.path.join(TMP, "refused_module")
    l1, s1 = quiet(); g1, f1 = BM.run(row_scenario_module(), out, say=s1)
    l2, s2 = quiet(); g2, f2 = BM.run(row_scenario_module(), out, draft=True, say=s2)
    return f1 is None and l1[-1].startswith("refused") and f2 and os.path.exists(out + ".html") and "DRAFT" in open(out + ".html").read()


# ---------------------------------------------------------------- D-28 SCORM 1.2 package (synthetic)
def scorm_module(pass_percent=None, scored=True):
    m = module([
        {"id": "S1", "type": "intro", "title": "Welcome", "body": "Synthetic intro."},
        {"id": "S2", "type": "question" if scored else "practice", "title": "Quick Check", "body": "Choose one.", "question": q("Synthetic single question?")},
        {"id": "S3", "type": "practice", "title": "Practice", "body": "Choose one.", "question": q("Synthetic practice?")},
        {"id": "S4", "type": "question" if scored else "practice", "title": "Second Check", "body": "Choose all.",
         "question": {"stem": "Synthetic multi question?", "multi": True, "options": [
             {"text": "Synthetic A", "correct": True, "feedback": "Synthetic: yes."},
             {"text": "Synthetic B", "correct": True, "feedback": "Synthetic: yes."},
             {"text": "Synthetic C", "correct": False, "feedback": "Synthetic: no."}]}},
        {"id": "S5", "type": "scenario", "title": "A Pending Claim", "body": "Synthetic scenario.",
         "choices": [{"text": "Wait", "correct": False, "feedback": "Synthetic: waiting delays payment."},
                     {"text": "Check status", "correct": True, "feedback": "Synthetic: right."}]},
        {"id": "S6", "type": "summary", "title": "Summary", "body": "Synthetic summary."}])
    m["meta"]["title"] = "Synthetic SCORM Check Module"
    if pass_percent is not None:
        m["meta"]["pass_percent"] = pass_percent
    return m


def scorm_build(name, mod=None):
    out = os.path.join(TMP, name)
    l, s = quiet()
    g, files = BM.run(mod or scorm_module(), out, say=s, scorm=True)
    return g, files, out, l


def rezip(src, dst, edit):
    import zipfile
    with zipfile.ZipFile(src) as z:
        items = {n: z.read(n) for n in z.namelist()}
    items = edit(items) or items
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in items.items():
            z.writestr(n, b)
    return dst


@case("D-28 SCORM_PACKAGE positive: --scorm writes Name_SCORM12.zip with imsmanifest.xml and index.html at the root plus the brand's font files under fonts/; gate passes", "SCORM_PACKAGE")
def _():
    import zipfile
    g, files, out, l = scorm_build("scorm_ok")
    z = zipfile.ZipFile(out + "_SCORM12.zip")
    man = z.read("imsmanifest.xml").decode("utf-8")
    sg = BM.scorm_check(out + "_SCORM12.zip")
    return (files and files[-1].endswith("_SCORM12.zip") and sorted(z.namelist()) == sorted(["imsmanifest.xml", "index.html"] + [BM.font_url(p) for _f, _w, p in BM.font_assets()])
            and all('href="%s"' % BM.font_url(p) in man for _f, _w, p in BM.font_assets())
            and "SCORM_PACKAGE" not in g.codes() and sg.failed == 0 and "<schema>ADL SCORM</schema>" in man
            and "<schemaversion>1.2</schemaversion>" in man and man.count('adlcp:scormtype="sco"') == 1 and man.count("<organization ") == 1
            and man.count("<item ") == 1 and 'href="index.html"' in man and "<adlcp:masteryscore>80</adlcp:masteryscore>" in man
            and any(x.startswith("SCORM_PACKAGE GATE:") and "0 failed" in x for x in l))


@case("D-28 SCORM_PACKAGE positive: index.html in the zip is the same self-contained module as Name.html and passes ELEARN_OFFLINE", "SCORM_PACKAGE")
def _():
    import zipfile
    g, files, out, l = scorm_build("scorm_same")
    page = zipfile.ZipFile(out + "_SCORM12.zip").read("index.html").decode("utf-8")
    off = BM.offline_check(page, BM.Gate("T"))
    return page == open(out + ".html", encoding="utf-8").read() and off.failed == 0 and all(k in page for k in BM.SCORM_NEEDS)


@case("D-28 SCORM_PACKAGE positive: without --scorm no zip is written (offline build unchanged); pass_percent reaches the page and manifest", "SCORM_PACKAGE")
def _():
    import zipfile
    out = os.path.join(TMP, "scorm_off")
    l, s = quiet(); g, files = BM.run(scorm_module(), out, say=s)
    g2, f2, out2, _ = scorm_build("scorm_pp", scorm_module(pass_percent=70))
    man = zipfile.ZipFile(out2 + "_SCORM12.zip").read("imsmanifest.xml").decode("utf-8")
    return (len(files) == 2 and not os.path.exists(out + "_SCORM12.zip") and 'data-pass-percent="80"' in open(out + ".html").read()
            and 'data-pass-percent="70"' in open(out2 + ".html").read() and "<adlcp:masteryscore>70</adlcp:masteryscore>" in man)


@case("D-28 SCORM_PACKAGE negative: a manifest whose SCO points at a missing href fails", "SCORM_PACKAGE")
def _():
    g, files, out, l = scorm_build("scorm_href")
    def edit(items):
        items["imsmanifest.xml"] = items["imsmanifest.xml"].replace(b'adlcp:scormtype="sco" href="index.html"', b'adlcp:scormtype="sco" href="start.html"')
    bad = rezip(out + "_SCORM12.zip", out + "_badhref.zip", edit)
    sg = BM.scorm_check(bad)
    return has(sg, "FAIL", "SCORM_PACKAGE", 'SCO href "start.html" is not in the zip')


@case("D-28 SCORM_PACKAGE negative: a module without the adapter (stripped) fails on LMSInitialize, lesson_status and score.raw", "SCORM_PACKAGE")
def _():
    g, files, out, l = scorm_build("scorm_strip")
    def edit(items):
        page = items["index.html"].decode("utf-8")
        items["index.html"] = re.sub(r"/\* scorm12-adapter:start.*?scorm12-adapter:end \*/", "", page, flags=re.S).encode("utf-8")
    sg = BM.scorm_check(rezip(out + "_SCORM12.zip", out + "_stripped.zip", edit))
    return has(sg, "FAIL", "SCORM_PACKAGE", "no SCORM 1.2 adapter (missing LMSInitialize, cmi.core.lesson_status, cmi.core.score.raw)")


@case("D-28 SCORM_PACKAGE negative: broken XML, schemaversion 1.3, two SCOs, an unlisted file and a nested zip each fail", "SCORM_PACKAGE")
def _():
    g, files, out, l = scorm_build("scorm_misc")
    src = out + "_SCORM12.zip"
    def broken(items): items["imsmanifest.xml"] = b"<manifest><metadata>"
    def v13(items): items["imsmanifest.xml"] = items["imsmanifest.xml"].replace(b"<schemaversion>1.2<", b"<schemaversion>1.3<")
    def two(items): items["imsmanifest.xml"] = items["imsmanifest.xml"].replace(b"</resources>", b'<resource identifier="RES-2" type="webcontent" adlcp:scormtype="sco" href="index.html"/></resources>')
    def extra(items): items["notes.txt"] = b"synthetic extra file"
    def nested(items): items["old_package.zip"] = b"PK synthetic"
    r = {k: BM.scorm_check(rezip(src, out + "_%s.zip" % k, f)) for k, f in
         [("broken", broken), ("v13", v13), ("two", two), ("extra", extra), ("nested", nested)]}
    return (has(r["broken"], "FAIL", "SCORM_PACKAGE", "does not parse as XML") and has(r["v13"], "FAIL", "SCORM_PACKAGE", "schemaversion 1.2")
            and has(r["two"], "FAIL", "SCORM_PACKAGE", "exactly one SCO resource (has 2)")
            and has(r["extra"], "FAIL", "SCORM_PACKAGE", "not listed in the manifest: notes.txt")
            and has(r["nested"], "FAIL", "SCORM_PACKAGE", "nested archive(s)"))


@case("D-28 SCORM_PACKAGE negative: when the package gate fails the build refuses and removes the zip", "SCORM_PACKAGE")
def _():
    keep = BM.SCORM_NEEDS
    try:
        BM.SCORM_NEEDS = keep + ("synthetic-token-not-in-page",)
        g, files, out, l = scorm_build("scorm_refuse")
    finally:
        BM.SCORM_NEEDS = keep
    return files is None and not os.path.exists(out + "_SCORM12.zip") and any("refused" in x for x in l) and "SCORM_PACKAGE" in g.codes()


# Node proof of the SCORM 1.2 adapter: the module's own script runs against a minimal DOM shim and a fake window.API
# that records every call (same approach as the D-14 JS smoke check, results/render_forms_20261011_182200/checks/run.js).
NODE_SHIM = r"""
const fs=require('fs'); const tree=JSON.parse(fs.readFileSync(process.argv[2])); const js=fs.readFileSync(process.argv[3],'utf8'); const sc=JSON.parse(process.argv[4]);
class CL{constructor(e){this.e=e} get l(){return (this.e.attrs['class']||'').split(/\s+/).filter(Boolean)} add(c){const l=this.l; if(!l.includes(c)){l.push(c); this.e.attrs['class']=l.join(' ')}} remove(c){this.e.attrs['class']=this.l.filter(x=>x!==c).join(' ')} contains(c){return this.l.includes(c)} toggle(c,on){if(on===undefined)on=!this.contains(c); on?this.add(c):this.remove(c); return on}}
class El{constructor(n,parent){this.tag=n.tag;this.attrs=Object.assign({},n.attrs||{});this.parent=parent;this.children=[];this._text=n.text;this.ls={};this.style={};this.classList=new CL(this);
  if(this.tag==='input'){this.checked=false} (n.children||[]).forEach(c=>this.children.push(new El(c,this)))}
 get hidden(){return 'hidden' in this.attrs} set hidden(v){if(v)this.attrs.hidden='';else delete this.attrs.hidden}
 get disabled(){return 'disabled' in this.attrs} set disabled(v){if(v)this.attrs.disabled='';else delete this.attrs.disabled}
 set className(v){this.attrs['class']=v} get className(){return this.attrs['class']||''}
 getAttribute(k){return k in this.attrs?this.attrs[k]:null} setAttribute(k,v){this.attrs[k]=String(v)}
 get textContent(){return this._text!==undefined?this._text:this.children.map(c=>c.textContent).join('')}
 set textContent(v){this.children=[];this._text=undefined; if(v!=='') this.children.push(new El({text:v},this))}
 appendChild(c){c.parent=this;this.children.push(c);return c}
 all(){let o=[];for(const c of this.children){if(c.tag){o.push(c);o=o.concat(c.all())}}return o}
 matchesOne(sel){const m=sel.trim().match(/^([a-z0-9]*)((?:\.[\w-]+)*)((?:\[[^\]]+\])*)$/i); if(!m) throw new Error('sel '+sel);
   if(m[1]&&this.tag!==m[1])return false; for(const c of (m[2].match(/\.[\w-]+/g)||[]))if(!this.classList.contains(c.slice(1)))return false;
   for(const a of (m[3].match(/\[[^\]]+\]/g)||[])){const k=a.slice(1,-1).split('=')[0]; if(!(k in this.attrs))return false} return true}
 matches(sel){return sel.split(',').some(s=>{const parts=s.trim().split(/\s+/); if(!this.matchesOne(parts[parts.length-1]))return false; let p=this.parent,i=parts.length-2; while(i>=0&&p){if(p.tag&&p.matchesOne(parts[i]))i--; p=p.parent} return i<0})}
 querySelectorAll(sel){return this.all().filter(e=>e.matches(sel))} querySelector(sel){return this.querySelectorAll(sel)[0]||null}
 addEventListener(t,f){(this.ls[t]=this.ls[t]||[]).push(f)} dispatch(t){(this.ls[t]||[]).forEach(f=>f({target:this,preventDefault(){}}))}
 click(){if(this.disabled)return; this.dispatch('click')} focus(){doc.activeElement=this}
}
const root=new El(tree,null); const html=root.querySelector('html');
const doc={documentElement:html,activeElement:null,ls:{},querySelectorAll:s=>root.querySelectorAll(s),querySelector:s=>root.querySelector(s),
 getElementById:id=>root.all().find(e=>e.attrs.id===id)||null,createElement:t=>new El({tag:t},null),createTextNode:t=>new El({text:t},null),
 addEventListener(t,f){(this.ls[t]=this.ls[t]||[]).push(f)}};
const calls=[]; const store=Object.assign({'cmi.core.lesson_status':sc.prior||'not attempted'},sc.store||{});
const API={LMSInitialize(a){calls.push(['LMSInitialize',a]);return 'true'}, LMSFinish(a){calls.push(['LMSFinish',a]);return 'true'},
 LMSGetValue(k){calls.push(['LMSGetValue',k]);return store[k]||''}, LMSSetValue(k,v){calls.push(['LMSSetValue',k,v]);store[k]=v;return 'true'},
 LMSCommit(a){calls.push(['LMSCommit',a]);return 'true'}, LMSGetLastError(){return '0'}, LMSGetErrorString(){return ''}, LMSGetDiagnostic(){return ''}};
const wls={}; const win={scrollTo(){}, addEventListener(t,f){(wls[t]=wls[t]||[]).push(f)}};
function chain(n){let w=win; for(let k=0;k<n;k++){const p={}; w.parent=p; w=p;} w.parent=w; return w;}
if(sc.api==='self'){win.API=API; win.parent=win;} else if(sc.api==='opener'){win.parent=win; win.opener={API:API};}
else if(sc.api&&sc.api.startsWith('parent')){chain(parseInt(sc.api.slice(6))).API=API;} else {win.parent=win;}
global.document=doc; global.window=win;
let error=null;
try{
  new Function(js)();
  const $=id=>doc.getElementById(id); let k=0; const cur=doc.querySelector('.screen.current');
  var boot={counter:$('counter').textContent, screen:cur?cur.getAttribute('data-id'):null, checked:doc.querySelectorAll('input').filter(x=>x.checked).map(x=>x.attrs.id),
    feedback:doc.querySelectorAll('.feedback').map(f=>f.textContent).filter(Boolean), calls:calls.length};
  if(sc.steps){for(const st of sc.steps){if(st[0]==='next')$('next').click(); else if(st[0]==='pick'){const x=$(st[1]); if(x.attrs.type==='radio')doc.querySelectorAll('input').forEach(y=>{if(y.attrs.name===x.attrs.name)y.checked=false;}); x.checked=true; x.dispatch('change');} else if(st[0]==='unload'){['beforeunload','pagehide','unload'].forEach(t=>(wls[t]||[]).forEach(f=>f({})));}}
    throw 'steps-done';}
  for(const f of doc.querySelectorAll('fieldset.q')){const ins=f.querySelectorAll('input'); const scored=f.getAttribute('data-scored')==='true'; const wrong=scored&&(sc.wrongFirst||[]).includes(k); if(scored)k++;
    if(f.getAttribute('data-multi')==='true'){ if(wrong){ins.forEach(i=>i.checked=i.getAttribute('data-correct')!=='true'); f.querySelector('button.check').click();}
      ins.forEach(i=>i.checked=i.getAttribute('data-correct')==='true'); f.querySelector('button.check').click(); }
    else { if(wrong){const w=ins.find(i=>i.getAttribute('data-correct')!=='true'); w.checked=true; w.dispatch('change');}
      const r=ins.find(i=>i.getAttribute('data-correct')==='true'); r.checked=true; r.dispatch('change'); }}
  for(const g of doc.querySelectorAll('.choices')){g.querySelector('button.choice').click();}
  if(sc.unloadFirst){(wls.beforeunload||[]).forEach(f=>f({}));}
  const N=doc.querySelectorAll('.screen[data-screen]').length; for(let s=0;s<N;s++) $('next').click();
  var atFinish=calls.length, counter=$('counter').textContent;
  $('restart').click(); for(let s=0;s<N;s++) $('next').click(); ['beforeunload','pagehide','unload'].forEach(t=>(wls[t]||[]).forEach(f=>f({})));
}catch(e){if(e!=='steps-done')error=String(e&&e.stack||e);}
console.log(JSON.stringify({calls:calls, atFinish:atFinish, counter:counter, error:error, listeners:Object.keys(wls).length, boot:boot, store:store}));
"""


def node_run(mod, scenario, name):
    import shutil, subprocess
    from html.parser import HTMLParser
    node = shutil.which("node")
    if not node:
        return None
    out = os.path.join(TMP, name)
    l, s = quiet(); BM.run(mod, out, say=s)
    page = open(out + ".html", encoding="utf-8").read()

    class P(HTMLParser):
        def __init__(self):
            super().__init__(); self.root = {"tag": "#root", "attrs": {}, "children": []}; self.st = [self.root]
        def handle_starttag(self, t, a):
            n = {"tag": t, "attrs": {k: (v if v is not None else "") for k, v in a}, "children": []}
            self.st[-1]["children"].append(n)
            if t not in ("meta", "img", "input", "br", "link", "hr"):
                self.st.append(n)
        def handle_endtag(self, t):
            for i in range(len(self.st) - 1, 0, -1):
                if self.st[i]["tag"] == t:
                    del self.st[i:]; break
        def handle_data(self, d):
            if self.st[-1]["tag"] not in ("script", "style") and d.strip():
                self.st[-1]["children"].append({"text": d})
    p = P(); p.feed(page)
    json.dump(p.root, open(out + ".tree.json", "w"))
    open(out + ".js", "w", encoding="utf-8").write(re.search(r"<script>(.*?)</script>", page, re.S).group(1))
    open(os.path.join(TMP, "scorm_shim.js"), "w", encoding="utf-8").write(NODE_SHIM)
    r = subprocess.run([node, os.path.join(TMP, "scorm_shim.js"), out + ".tree.json", out + ".js", json.dumps(scenario)],
                       capture_output=True, text=True, timeout=60)
    return json.loads(r.stdout.strip().splitlines()[-1])


BOOKMARK = ("cmi.core.lesson_location", "cmi.suspend_data")


def lms_seq(res):
    """LMS calls without reads and without the D-52 bookmark saves (bookmark sets plus the commit that follows them)."""
    out, after = [], False
    for c in res["calls"]:
        if c[0] == "LMSGetValue" or (c[0] == "LMSSetValue" and c[1] == "cmi.core.exit"):
            continue
        if c[0] == "LMSSetValue" and c[1] in BOOKMARK:
            after = True; continue
        if c[0] == "LMSCommit" and after:
            after = False; continue
        after = False; out.append(c)
    return out


def bookmarks(res):
    return [c for c in res["calls"] if c[0] == "LMSSetValue" and c[1] in BOOKMARK]


def ends_with(seq, raw, status):
    tail = ([["LMSSetValue", "cmi.core.score.raw", raw]] if raw is not None else []) + [
        ["LMSSetValue", "cmi.core.lesson_status", status], ["LMSCommit", ""], ["LMSFinish", ""]]
    return seq[-len(tail):] == tail


@case("D-28 SCORM_PACKAGE Node: fake window.API records LMSInitialize, lesson_status incomplete, ..., score.raw, passed/failed, LMSCommit, LMSFinish", "SCORM_PACKAGE")
def _():
    ok_run = node_run(scorm_module(), {"api": "self"}, "node_pass")
    if ok_run is None:
        return "UNVERIFIED"
    miss = node_run(scorm_module(), {"api": "self", "wrongFirst": [0]}, "node_fail")
    low = node_run(scorm_module(pass_percent=50), {"api": "self", "wrongFirst": [1]}, "node_low")
    a, b, c = lms_seq(ok_run), lms_seq(miss), lms_seq(low)
    start = [["LMSInitialize", ""], ["LMSSetValue", "cmi.core.lesson_status", "incomplete"]]
    return (not ok_run["error"] and a[:2] == start and ends_with(a, "100", "passed") and ["LMSSetValue", "cmi.core.score.min", "0"] in a
            and ["LMSSetValue", "cmi.core.score.max", "100"] in a and len(ok_run["calls"]) == ok_run["atFinish"]
            and ok_run["counter"] == "Module complete" and b[:2] == start and ends_with(b, "50", "failed") and ends_with(c, "50", "passed"))


@case("D-28 SCORM_PACKAGE Node: no API present runs offline with no LMS calls and reaches Module complete", "SCORM_PACKAGE")
def _():
    res = node_run(scorm_module(), {"api": "none"}, "node_none")
    if res is None:
        return "UNVERIFIED"
    return not res["error"] and res["calls"] == [] and res["counter"] == "Module complete" and res["listeners"] == 0


@case("D-28 SCORM_PACKAGE Node: findAPI walks parents (found at 3, not at 9 levels) and the opener; unload commits and finishes once", "SCORM_PACKAGE")
def _():
    p3 = node_run(scorm_module(), {"api": "parent3"}, "node_p3")
    if p3 is None:
        return "UNVERIFIED"
    p9 = node_run(scorm_module(), {"api": "parent9"}, "node_p9")
    op = node_run(scorm_module(), {"api": "opener"}, "node_op")
    un = node_run(scorm_module(), {"api": "self", "unloadFirst": True}, "node_unload")
    u = lms_seq(un)
    return (ends_with(lms_seq(p3), "100", "passed") and p9["calls"] == [] and not p9["error"] and ends_with(lms_seq(op), "100", "passed")
            and u == [["LMSInitialize", ""], ["LMSSetValue", "cmi.core.lesson_status", "incomplete"], ["LMSCommit", ""], ["LMSFinish", ""]])


@case("D-28 SCORM_PACKAGE Node: no scored questions reports completed without a score; a prior pass is never downgraded", "SCORM_PACKAGE")
def _():
    nos = node_run(scorm_module(scored=False), {"api": "self"}, "node_noscore")
    if nos is None:
        return "UNVERIFIED"
    kept = node_run(scorm_module(), {"api": "self", "prior": "passed", "wrongFirst": [0, 1]}, "node_kept")
    n, k = lms_seq(nos), lms_seq(kept)
    return (ends_with(n, None, "completed") and not any(c[1].startswith("cmi.core.score") for c in n if c[0] == "LMSSetValue")
            and ["LMSSetValue", "cmi.core.lesson_status", "incomplete"] not in k and ["LMSSetValue", "cmi.core.lesson_status", "failed"] not in k
            and k[-3:] == [["LMSSetValue", "cmi.core.score.raw", "0"], ["LMSCommit", ""], ["LMSFinish", ""]])


# ---------------------------------------------------------------- D-52 SCORM resume (synthetic)
def resume_store(loc, answers):
    return {"cmi.core.lesson_location": loc, "cmi.suspend_data": json.dumps({"q": answers}, separators=(",", ":"))}


@case("D-52 SCORM_PACKAGE row scenario (synthetic): fake LMS returns lesson_location '4' and suspend_data with an answer; module reopens on screen 4 with the answer restored", "SCORM_PACKAGE")
def _():
    res = node_run(scorm_module(), {"api": "self", "store": resume_store("4", {"0": [1, "0"]}), "steps": []}, "node_resume")
    if res is None:
        return "UNVERIFIED"
    b = res["boot"]
    gets = [c[1] for c in res["calls"] if c[0] == "LMSGetValue"]
    return (not res["error"] and b["counter"] == "Screen 4 of 6" and b["screen"] == "S4" and b["checked"] == ["S2-A"]
            and b["feedback"] == ["Correct. Synthetic feedback 1."] and "cmi.core.lesson_location" in gets and "cmi.suspend_data" in gets
            and bookmarks(res) == [])


@case("D-52 SCORM_PACKAGE positive (synthetic): each answer and screen change sets lesson_location and compact suspend_data, then LMSCommit; a restored first-try miss still counts", "SCORM_PACKAGE")
def _():
    run = node_run(scorm_module(), {"api": "self", "steps": [["pick", "S2-B"], ["next"], ["pick", "S2-A"]]}, "node_save")
    if run is None:
        return "UNVERIFIED"
    c = run["calls"]
    seq = [x for x in c if x[0] != "LMSGetValue"][2:]
    sd = json.loads(run["store"]["cmi.suspend_data"])
    want = [["LMSSetValue", "cmi.core.lesson_location", "1"], ["LMSSetValue", "cmi.suspend_data", '{"q":{"0":[0,"1"]}}'], ["LMSCommit", ""],
            ["LMSSetValue", "cmi.core.lesson_location", "2"], ["LMSSetValue", "cmi.suspend_data", '{"q":{"0":[0,"1"]}}'], ["LMSCommit", ""],
            ["LMSSetValue", "cmi.core.lesson_location", "2"], ["LMSSetValue", "cmi.suspend_data", '{"q":{"0":[0,"0"]}}'], ["LMSCommit", ""]]
    again = node_run(scorm_module(), {"api": "self", "store": resume_store("2", sd["q"])}, "node_again")
    return (not run["error"] and seq == want and run["store"]["cmi.core.lesson_location"] == "2" and sd == {"q": {"0": [0, "0"]}}
            and not again["error"] and again["boot"]["screen"] == "S2" and again["boot"]["checked"] == ["S2-A"]
            and ends_with(lms_seq(again), "50", "failed"))


@case("D-52 SCORM_PACKAGE positive (synthetic): the built package passes the resume checks and its longest suspend_data stays within 4096 characters", "SCORM_PACKAGE")
def _():
    g, files, out, l = scorm_build("scorm_resume_ok")
    sg = BM.scorm_check(out + "_SCORM12.zip")
    page = open(out + ".html", encoding="utf-8").read()
    return (sg.failed == 0 and BM.suspend_worst(page) == len('{"q":{"0":[0,"012"],"1":[0,"012"],"2":[0,"012"]}}')
            and 'data-suspend-limit="4096"' in page and all(k in page for k in BM.RULES["scorm_resume_needs"]))


@case("D-52 SCORM_PACKAGE negative (synthetic): an adapter that never saves the place fails; 300 four-option questions overflow the 4096-character suspend_data", "SCORM_PACKAGE")
def _():
    g, files, out, l = scorm_build("scorm_resume_bad")
    def nosave(items):
        page = items["index.html"].decode("utf-8")
        items["index.html"] = page.replace("cmi.core.lesson_location", "cmi.core.x").replace("cmi.suspend_data", "cmi.x").encode("utf-8")
    def big(items):
        page = items["index.html"].decode("utf-8")
        fs = '<fieldset class="q"><input type="radio"><input type="radio"><input type="radio"><input type="radio"></fieldset>'
        items["index.html"] = page.replace("</main>", fs * 300 + "</main>").encode("utf-8")
    a = BM.scorm_check(rezip(out + "_SCORM12.zip", out + "_nosave.zip", nosave))
    b = BM.scorm_check(rezip(out + "_SCORM12.zip", out + "_big.zip", big))
    return (has(a, "FAIL", "SCORM_PACKAGE", "does not save the learner's place (missing cmi.core.lesson_location, cmi.suspend_data)")
            and has(b, "FAIL", "SCORM_PACKAGE", "SCORM 1.2 allows 4096") and not has(b, "FAIL", "SCORM_PACKAGE", "does not save"))


@case("D-52 SCORM_PACKAGE negative (synthetic): bad bookmarks (lesson_location '99' or 'abc', broken or foreign suspend_data) start on screen 1 with nothing checked", "SCORM_PACKAGE")
def _():
    runs = [node_run(scorm_module(), {"api": "self", "store": st, "steps": []}, "node_badbm%d" % n) for n, st in enumerate([
        {"cmi.core.lesson_location": "99", "cmi.suspend_data": "{not json"},
        {"cmi.core.lesson_location": "abc", "cmi.suspend_data": '{"q":{"7":[1,"0"],"x":[1,"0"],"1":"0"}}'},
        {"cmi.core.lesson_location": "0", "cmi.suspend_data": "[]"}])]
    if runs[0] is None:
        return "UNVERIFIED"
    return all(not r["error"] and r["boot"]["counter"] == "Screen 1 of 6" and r["boot"]["screen"] == "S1" and r["boot"]["checked"] == []
               and r["boot"]["feedback"] == [] for r in runs)


@case("D-52 SCORM_PACKAGE positive (synthetic): closing before Finish sets cmi.core.exit 'suspend' just before LMSCommit and LMSFinish; finishing sets exit '' and never 'suspend'", "SCORM_PACKAGE")
def _():
    early = node_run(scorm_module(), {"api": "self", "steps": [["pick", "S2-A"], ["next"], ["unload"]]}, "node_exit_early")
    if early is None:
        return "UNVERIFIED"
    done = node_run(scorm_module(), {"api": "self"}, "node_exit_done")
    e = [x for x in early["calls"] if x[0] != "LMSGetValue"]
    d = [x for x in done["calls"] if x[0] == "LMSSetValue" and x[1] == "cmi.core.exit"]
    return (not early["error"] and e[-3:] == [["LMSSetValue", "cmi.core.exit", "suspend"], ["LMSCommit", ""], ["LMSFinish", ""]]
            and early["store"]["cmi.core.lesson_location"] == "2" and early["store"]["cmi.core.exit"] == "suspend"
            and not done["error"] and d == [["LMSSetValue", "cmi.core.exit", ""]] and done["store"]["cmi.core.exit"] == "")


@case("D-52 SCORM_PACKAGE positive (synthetic): suspend_data over the limit drops whole answers from the end and stays valid JSON", "SCORM_PACKAGE")
def _():
    keep = BM.RULES["scorm_suspend_limit"]
    try:
        BM.RULES["scorm_suspend_limit"] = 25
        res = node_run(scorm_module(), {"api": "self", "steps": [["pick", "S2-A"], ["next"], ["next"], ["pick", "S3-A"]]}, "node_trunc")
    finally:
        BM.RULES["scorm_suspend_limit"] = keep
    if res is None:
        return "UNVERIFIED"
    sds = [x[2] for x in res["calls"] if x[0] == "LMSSetValue" and x[1] == "cmi.suspend_data"]
    return (not res["error"] and sds[-1] == '{"q":{"0":[1,"0"]}}' and all(len(s) <= 25 and json.loads(s) for s in sds)
            and len('{"q":{"0":[1,"0"],"1":[1,"0"]}}') > 25)


# ---------------------------------------------------------------- D-46..D-51 form rows (synthetic)
import subprocess
EX = os.path.join(BF.SKILL, "assets", "examples")


def cli(*args):
    r = subprocess.run([sys.executable, os.path.join(HERE, "build_form.py")] + list(args), capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr


def dump(obj, name):
    p = os.path.join(TMP, name); json.dump(obj, open(p, "w")); return p


def example(name):
    return json.load(open(os.path.join(EX, name + ".form.json")))


def satisfaction_followup(days=90):
    s = survey(); s["meta"]["format"] = "follow-up-evaluation"; s["meta"]["send_after_days"] = days
    return s


@case("D-46 FOLLOWUP_ITEMS row scenario (synthetic): --check on a follow-up-evaluation form with only satisfaction items warns on the three topics and send_after_days 90", "FOLLOWUP_ITEMS")
def _():
    rc, out = cli(dump(satisfaction_followup(), "fu_satisfaction.form.json"), "--check")
    return rc == 0 and "WARN [FOLLOWUP_ITEMS] missing: frequency of use, barriers, result" in out \
        and "WARN [FOLLOWUP_ITEMS] meta.send_after_days is 90; send a follow-up evaluation 30-60 days after the training" in out


@case("D-46 FOLLOWUP_ITEMS positive: shipped synthetic_followup_evaluation.form.json has no findings, 0 failed, and builds", "FOLLOWUP_ITEMS")
def _():
    f = example("synthetic_followup_evaluation")
    g = BF.gate(f)
    lines, say = quiet()
    g2, files = BF.run(f, os.path.join(TMP, "fu_example"), say=say)
    return f["meta"]["format"] == BF.RULES["followup_format"] and "FOLLOWUP_ITEMS" not in g.codes("WARN") and not g.msgs and bool(files) and g2.failed == 0


@case("D-46 FOLLOWUP_ITEMS boundaries: 30 and 60 days pass, 29 and unset warn, a partial form names only the missing topic, other formats are not checked", "FOLLOWUP_ITEMS")
def _():
    f = example("synthetic_followup_evaluation")
    days = {}
    for d in (30, 60, 29, None):
        f["meta"]["send_after_days"] = d
        days[d] = has(BF.gate(f), "WARN", "FOLLOWUP_ITEMS", "send_after_days")
    del f["sections"][1]
    g = BF.gate(f)
    return days == {30: False, 60: False, 29: True, None: True} and has(g, "WARN", "FOLLOWUP_ITEMS", "missing: barriers (") \
        and "FOLLOWUP_ITEMS" not in BF.gate(survey()).codes("WARN")


def widgets_of(doc, name):
    for p in doc:
        ws = [w for w in p.widgets() if w.field_name.split(".")[0] == name]
        if ws:
            return p, ws
    return None, []


def label_under(page, w, text):
    import fitz
    r = w.rect
    return bool(page.search_for(text, clip=fitz.Rect(r.x0 - 40, r.y1, r.x1 + 40, r.y1 + 34)))


@case("D-48 SCALE_LABELS row scenario: --check assets/examples/synthetic_training_survey.form.json, 0 failed, no SCALE_LABELS warning, every point labelled", "SCALE_LABELS")
def _():
    rc, out = cli(os.path.join(EX, "synthetic_training_survey.form.json"), "--check")
    return rc == 0 and "FORM GATE:" in out and " 0 failed" in out and "SCALE_LABELS" not in out \
        and "scale labels: every point labelled on Q3, Q4, Q5, Q6" in out


@case("D-48 SCALE_LABELS positive: the fillable PDF prints a label under every point (Q3) and opt_out 'Not applicable' is the last choice (Q6)", "SCALE_LABELS")
def _():
    import fitz
    lines, say = quiet()
    g, files = BF.run(example("synthetic_training_survey"), os.path.join(TMP, "sl_survey"), say=say)
    doc = fitz.open(files[0])
    p, ws = widgets_of(doc, "Q3")
    ws.sort(key=lambda w: w.rect.x0)
    labels = BF.RULES["scale_point_labels"][2]
    under = len(ws) == 5 and all(label_under(p, w, t) for w, t in zip(ws, labels))
    p6, w6 = widgets_of(doc, "Q6")
    w6.sort(key=lambda w: w.rect.x0)
    last = len(w6) == 6 and w6[-1].on_state() == "Not_applicable" and label_under(p6, w6[-1], "Not applicable")
    return g.failed == 0 and under and last


@case("D-48 SCALE_LABELS positive: scale.labels overrides; a 7-point scale with 7 labels has no warning and prints each label", "SCALE_LABELS")
def _():
    import fitz
    s = survey(); it = s["sections"][0]["items"]
    seven = ["Very poor", "Poor", "Weak", "Fair", "Solid", "Strong", "Excellent"]
    for k in (1, 2, 3):
        it[k]["scale"] = {"min": 1, "max": 7, "labels": seven}
    g = BF.gate(s)
    lines, say = quiet()
    g2, files = BF.run(s, os.path.join(TMP, "sl_seven"), say=say)
    p, ws = widgets_of(fitz.open(files[0]), "Q2")
    ws.sort(key=lambda w: w.rect.x0)
    return "SCALE_LABELS" not in g.codes("WARN") and g.failed == 0 and len(ws) == 7 and all(label_under(p, w, t) for w, t in zip(ws, seven))


@case("D-48 SCALE_LABELS negative: a 7-point scale without labels warns; a wrong label count fails SURVEY_ITEM; mixed scales still fail", "SCALE_LABELS")
def _():
    s = survey(); it = s["sections"][0]["items"]
    for k in (1, 2, 3):
        it[k]["scale"] = {"min": 1, "max": 7, "low": "Strongly disagree", "high": "Strongly agree"}
    g = BF.gate(s)
    it[1]["scale"]["labels"] = ["Low", "Mid", "High"]
    g2 = BF.gate(s)
    it[1]["scale"] = {"min": 1, "max": 5, "low": "Strongly disagree", "high": "Strongly agree"}
    return has(g, "WARN", "SCALE_LABELS", "Q2: 7-point scale (1-7) has no label for every point; add scale.labels with 7 labels") \
        and has(g2, "FAIL", "SURVEY_ITEM", "Q2: scale.labels needs 7 non-empty labels, one per point (has 3)") \
        and has(BF.gate(s), "FAIL", "SURVEY_ITEM", "mixed scales (1-5, 1-7)")


@case("D-48 SCALE_LABELS opt_out: on a single item it renders as the last choice; on a text item or duplicating an option it fails", "SCALE_LABELS")
def _():
    import fitz
    s = survey(); it = s["sections"][0]["items"]
    it[0]["opt_out"] = "Prefer not to say"
    lines, say = quiet()
    g, files = BF.run(s, os.path.join(TMP, "sl_optout"), say=say)
    p, ws = widgets_of(fitz.open(files[0]), "Q1")
    ws.sort(key=lambda w: w.rect.y0)
    ok = g.failed == 0 and len(ws) == 3 and ws[-1].on_state() == "Prefer_not_to_say"
    bad = survey(); b = bad["sections"][0]["items"]
    b[6]["opt_out"] = "Not applicable"; b[0]["opt_out"] = "Morning"
    gb = BF.gate(bad)
    return ok and has(gb, "FAIL", "FORM_FIELD", "Q7: opt_out (e.g. \"Not applicable\") works only on scale and single items") \
        and has(gb, "FAIL", "SURVEY_ITEM", "Q1: duplicate options")


@case("D-49 FORM_TAGS row scenario: build_form.py on synthetic_access_request.form.json writes a PDF with /StructTreeRoot, /Lang en-US, /MarkInfo, /Title, DisplayDocTitle and /Tabs /S on every page", "FORM_TAGS")
def _():
    import fitz
    out = os.path.join(TMP, "tags_access")
    rc, txt = cli(os.path.join(EX, "synthetic_access_request.form.json"), "--out", out)
    doc = fitz.open(out + ".pdf"); cat = doc.pdf_catalog()
    tabs = all(doc.xref_get_key(p.xref, "Tabs") == ("name", "/S") for p in doc)
    tips = all((w.field_label or "").strip() for p in doc for w in p.widgets())
    return rc == 0 and "FORM_TAGS" not in txt and "PDF tags: /StructTreeRoot, /Lang en-US" in txt and doc.xref_get_key(cat, "StructTreeRoot")[0] == "xref" \
        and doc.xref_get_key(cat, "Lang") == ("string", "en-US") and doc.xref_get_key(cat, "MarkInfo/Marked")[1] == "true" \
        and doc.xref_get_key(cat, "ViewerPreferences/DisplayDocTitle")[1] == "true" and doc.metadata["title"] == "Synthetic System Access Request Form" and tabs and tips


@case("D-49 FORM_TAGS positive: Document root with P and Form children; every widget has a StructParent; one Form element per field per page", "FORM_TAGS")
def _():
    import fitz
    out = os.path.join(TMP, "tags_tree"); lines, say = quiet()
    BF.run(example("synthetic_training_survey"), out, say=say)
    doc = fitz.open(out + ".pdf"); cat = doc.pdf_catalog()
    root = int(doc.xref_get_key(cat, "StructTreeRoot")[1].split()[0])
    top = int(doc.xref_get_key(root, "K")[1].split()[0])
    kids = [int(x) for x in re.findall(r"(\d+) 0 R", doc.xref_get_key(top, "K")[1])]
    kinds = [doc.xref_get_key(k, "S")[1] for k in kids]
    per_page = sum(len({w.field_name.split(".")[0] for w in p.widgets()}) for p in doc)
    sp = all(doc.xref_get_key(w.xref, "StructParent")[0] == "int" for p in doc for w in p.widgets())
    return doc.xref_get_key(top, "S")[1] == "/Document" and kinds.count("/P") == len(doc) and kinds.count("/Form") == per_page and sp


@case("D-49 FORM_TAGS negative: a built PDF with /Lang, /MarkInfo and page 2 /Tabs removed fails and names each", "FORM_TAGS")
def _():
    import fitz
    src = os.path.join(TMP, "tags_src.pdf"); BF.build_pdf(example("synthetic_access_request"), src)
    doc = fitz.open(src); cat = doc.pdf_catalog()
    doc.xref_set_key(cat, "Lang", "null"); doc.xref_set_key(cat, "MarkInfo", "null"); doc.xref_set_key(doc[1].xref, "Tabs", "null")
    bad = os.path.join(TMP, "tags_bad.pdf"); doc.save(bad)
    g, line = BF.form_tags(bad)
    return line is None and has(g, "FAIL", "FORM_TAGS", "missing: /Lang en-US; /MarkInfo Marked true; /Tabs /S on page 2")


@case("D-49 FORM_TAGS negative: an untagged build (no tag pass, no title) fails on /StructTreeRoot and /Title", "FORM_TAGS")
def _():
    doc = BF.Builder(form()).build(); p = os.path.join(TMP, "tags_none.pdf"); doc.save(p)
    g, _ = BF.form_tags(p)
    return has(g, "FAIL", "FORM_TAGS", "missing: /StructTreeRoot; /Lang en-US; /MarkInfo Marked true; /Title; /ViewerPreferences DisplayDocTitle; /Tabs /S on page 1")


def read_csv(path):
    import csv
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.reader(fh))


@case("D-50 FORMS_EXPORT row scenario: build_form.py on synthetic_training_survey.form.json writes <Name>.msforms.csv, one row per item, same order and wording as the PDF", "FORMS_EXPORT")
def _():
    import fitz
    f = example("synthetic_training_survey"); out = os.path.join(TMP, "Training_Survey")
    rc, txt = cli(os.path.join(EX, "synthetic_training_survey.form.json"), "--out", out)
    rows = read_csv(out + ".msforms.csv")
    items = [it for _, it in BF.all_items(f)]
    pdf = " ".join(p.get_text() for p in fitz.open(out + ".pdf")).replace("\n", " ")
    pos = [pdf.find("%d. %s" % (n, it["label"])) for n, it in enumerate(items, 1)]
    r = {row[0]: row for row in rows[1:]}
    return rc == 0 and rows[0] == ["Question", "Type", "Options", "Required", "Help", "Branching"] and len(rows) == len(items) + 1 \
        and [x[0] for x in rows[1:]] == [it["label"] for it in items] and -1 not in pos and pos == sorted(pos) \
        and r[items[0]["label"]][1:4] == ["Choice", "; ".join(items[0]["options"]), "Yes"] \
        and r[items[2]["label"]][2] == "1 Strongly disagree; 2 Disagree; 3 Neutral; 4 Agree; 5 Strongly agree" \
        and r[items[5]["label"]][2].endswith("5 Strongly agree; Not applicable") and r[items[5]["label"]][4] == items[5]["help"] \
        and all(x[5] == "" for x in rows[1:]) and "create a new form and paste each row of Training_Survey.msforms.csv" in txt \
        and "Forms does not import this file automatically" in txt


@case("D-50 FORMS_EXPORT positive: forms get the file too (types, checkbox Yes, signature/date pair as two rows)", "FORMS_EXPORT")
def _():
    rows = BF.msforms_rows(example("synthetic_access_request"))
    return len(rows) == 10 and rows[3][1] == "Date" and rows[4][1] == "Choice (multiple answers)" and rows[7][2] == "Yes" \
        and [r[0] for r in rows[8:]] == ["Manager signature", "Date signed"]


@case("D-50 FORMS_EXPORT negative: a refused build and --check write no .msforms.csv", "FORMS_EXPORT")
def _():
    out = os.path.join(TMP, "msf_refused"); lines, say = quiet()
    g, files = BF.run(row_scenario_survey(), out, say=say)
    cwd = tempfile.mkdtemp(dir=TMP)
    r = subprocess.run([sys.executable, os.path.join(HERE, "build_form.py"), os.path.join(EX, "synthetic_training_survey.form.json"), "--check"], cwd=cwd, capture_output=True)
    return files is None and not os.path.exists(out + ".msforms.csv") and r.returncode == 0 and os.listdir(cwd) == []


def branch_survey():
    s = survey(); it = s["sections"][0]["items"]
    it[1] = {"id": "Q2", "type": "single", "label": "Did you use the synthetic lookup tool this week?", "required": True, "options": ["Yes", "No"]}
    it[2] = {"id": "Q3", "type": "single", "label": "How often did you use the lookup tool?", "show_if": {"Q2": "Yes"}, "options": ["Daily", "Weekly"]}
    return s


@case("D-51 FORM_BRANCH row scenario (synthetic): Q3 show_if {Q2: Yes} prints 'Answer Q3 only if you answered Yes to Q2' and fills the Branching column; no submit button", "FORM_BRANCH")
def _():
    import fitz
    out = os.path.join(TMP, "branch_row")
    rc, txt = cli(dump(branch_survey(), "branch_row.form.json"), "--out", out)
    doc = fitz.open(out + ".pdf")
    text = " ".join(p.get_text() for p in doc).replace("\n", " ")
    rows = read_csv(out + ".msforms.csv")
    buttons = [w for p in doc for w in p.widgets() if w.field_type == fitz.PDF_WIDGET_TYPE_BUTTON]
    return rc == 0 and "FORM GATE:" in txt and " 0 failed" in txt and "Answer Q3 only if you answered Yes to Q2" in text \
        and rows[3][5] == "Answer Q3 only if you answered Yes to Q2" and [r[5] for r in rows[1:] if r[5]] == [rows[3][5]] and not buttons


@case("D-51 FORM_BRANCH positive: displayed numbers (not ids) in surveys, the quoted label in forms, several answers joined with 'or'", "FORM_BRANCH")
def _():
    s = branch_survey(); it = s["sections"][0]["items"]
    it[1]["id"], it[2]["id"] = "used_tool", "tool_freq"
    it[2]["show_if"] = {"used_tool": ["Yes", "No"]}
    f = form(); f["sections"][0]["items"][1]["show_if"] = {"F3": "Replacement"}
    f["sections"][0]["items"][2], f["sections"][0]["items"][1] = f["sections"][0]["items"][1], f["sections"][0]["items"][2]
    g1, g2 = BF.gate(s), BF.gate(f)
    t1 = BF.branch_text(s, it[2]); t2 = BF.branch_text(f, f["sections"][0]["items"][2])
    return "FORM_BRANCH" not in g1.codes() + g2.codes() and t1 == "Answer Q3 only if you answered Yes or No to Q2" \
        and t2 == "Answer this question only if you answered Replacement to \"Request type\""


@case("D-51 FORM_BRANCH negative: unknown item, unknown option, a text item, a later item and a malformed show_if each fail; the build refuses", "FORM_BRANCH")
def _():
    msgs = {}
    for key, k, cond in (("item", 2, {"Q9": "Yes"}), ("option", 2, {"Q2": "Maybe"}), ("text", 7, {"Q7": "Yes"}), ("later", 2, {"Q5": "Lookups"}), ("shape", 2, "Q2=Yes")):
        s = branch_survey(); s["sections"][0]["items"][k]["show_if"] = cond
        msgs[key] = BF.gate(s)
    s = branch_survey(); s["sections"][0]["items"][2]["show_if"] = {"Q9": "Yes"}
    lines, say = quiet()
    g, files = BF.run(s, os.path.join(TMP, "branch_refused"), say=say)
    return has(msgs["item"], "FAIL", "FORM_BRANCH", "Q3: show_if names unknown item Q9") \
        and has(msgs["option"], "FAIL", "FORM_BRANCH", "Q3: show_if names option \"Maybe\" that Q2 does not have (options: Yes, No)") \
        and has(msgs["text"], "FAIL", "FORM_BRANCH", "Q8: show_if item Q7 is a longtext item with no fixed answers") \
        and has(msgs["later"], "FAIL", "FORM_BRANCH", "Q3: show_if item Q5 must come before Q3") \
        and has(msgs["shape"], "FAIL", "FORM_BRANCH", "Q3: show_if must name one item and its answer") \
        and files is None and not os.path.exists(os.path.join(TMP, "branch_refused.msforms.csv"))


# ---------------------------------------------------------------- D-53 Excel tables (synthetic)
import build_sheet as BS


def sheet_cli(*args):
    r = subprocess.run([sys.executable, os.path.join(HERE, "build_sheet.py")] + list(args), capture_output=True, text=True, cwd=TMP)
    return r.returncode, r.stdout + r.stderr


def tracker():
    return json.load(open(os.path.join(EX, "synthetic_tracker.json")))


def built(name, doc=None):
    out = os.path.join(TMP, name); l, s = quiet()
    g, files = BS.run(doc or tracker(), out, say=s)
    return out + ".xlsx", g, l


def edited(src, name, fn):
    from openpyxl import load_workbook
    wb = load_workbook(src); fn(wb); dst = os.path.join(TMP, name); wb.save(dst)
    return BS.check_xlsx(dst)


@case("D-53 XLSX_TABLE row scenario (synthetic): tracker.json with 3 statuses and 5 rows -> Name.xlsx with a named table, frozen header, Status dropdown from one list, no merged cells, Definitions tab, brand header colors", "XLSX_TABLE")
def _():
    from openpyxl import load_workbook
    rc, out = sheet_cli(os.path.join(EX, "synthetic_tracker.json"), "--out", "Synthetic_Tracker")
    rc2, out2 = sheet_cli("Synthetic_Tracker.xlsx", "--check")
    wb = load_workbook(os.path.join(TMP, "Synthetic_Tracker.xlsx")); ws = wb["Tracker"]
    dvs = ws.data_validations.dataValidation
    hdr = ws["A1"]
    C = BS.set_brand(os.path.join(EX, "synthetic_tracker.json"))["colors"]
    return (rc == 0 and rc2 == 0 and "XLSX_TABLE GATE:" in out and " 0 failed; wrote Synthetic_Tracker.xlsx (Tracker: 5 rows; Lists; Definitions)" in out
            and " 0 failed; Synthetic_Tracker.xlsx" in out2 and wb.sheetnames == ["Tracker", "Lists", "Definitions"]
            and dict(ws.tables.items()) == {"tblTracker": "A1:F6"} and ws.freeze_panes == "A2"
            and len(dvs) == 1 and dvs[0].formula1 == "=Lists!$A$2:$A$4" and str(dvs[0].sqref) == "E2:E500"
            and [c.value for c in wb["Lists"]["A"]] == ["Status", "Not Started", "In Progress", "Done"] and wb["Lists"].sheet_state == "hidden"
            and wb["Tracker"].sheet_state == wb["Definitions"].sheet_state == "visible"
            and not any(w.merged_cells.ranges for w in wb.worksheets) and "tblDefinitions" in dict(wb["Definitions"].tables.items())
            and hdr.fill.fgColor.rgb[-6:] == C["dark"] and hdr.font.color.rgb[-6:] == C["white"] and hdr.font.bold)


@case("D-53 XLSX_TABLE positive (synthetic): a RACI table shares one RACI list across every role column and a data dictionary uses a Classification list; both pass", "XLSX_TABLE")
def _():
    raci = {"meta": {"title": "Synthetic Month-End RACI", "format": "raci", "sources": ["Synthetic example"]},
            "lists": {"RACI": [{"value": v, "definition": d} for v, d in (("R", "Does the work"), ("A", "Owns the result; one per activity"), ("C", "Asked before"), ("I", "Told after"))]},
            "columns": [{"name": "Activity", "definition": "The work."}] + [{"name": r, "list": "RACI", "definition": "Role %s." % r} for r in ("Analyst", "Manager", "Director")],
            "rows": [["Close the synthetic ledger", "R", "A", "I"], ["Approve the synthetic report", "C", "R", "A"]]}
    dd = {"meta": {"title": "Synthetic Claims Dictionary", "format": "data-dictionary", "sources": ["Synthetic example"]},
          "lists": {"Classification": [{"value": "Internal", "definition": "Internal use."}, {"value": "PHI", "definition": "Protected health information; flag it."}]},
          "columns": [{"name": "Field", "definition": "Field name."}, {"name": "Description", "definition": "What it holds."}, {"name": "Classification", "list": "Classification", "definition": "Data class."}],
          "rows": [{"Field": "claim_id", "Description": "Synthetic claim key", "Classification": "Internal"}]}
    p1, g1, _ = built("Synthetic_RACI", raci); p2, g2, _ = built("Synthetic_Dictionary", dd)
    from openpyxl import load_workbook
    ws = load_workbook(p1)["RACI"]; dv = ws.data_validations.dataValidation
    return (g1.failed == 0 and g2.failed == 0 and len(dv) == 1 and str(dv[0].sqref) == "B2:B500 C2:C500 D2:D500"
            and dict(ws.tables.items()) == {"tblRACI": "A1:D3"} and "tblDataDictionary" in dict(load_workbook(p2)["Data Dictionary"].tables.items()))


@case("D-53 XLSX_TABLE positive: Tracker is a base-catalog format on the xlsx renderer; data dictionary and RACI plan pdf and xlsx (BUILD line unchanged, ALSO xlsx line)", "XLSX_TABLE")
def _():
    env = dict(os.environ, MATERIAL_OVERLAY=os.path.join(BF.SKILL, "assets", "catalog", "overlay.json"))
    def plan(req, name):
        r = subprocess.run([sys.executable, os.path.join(HERE, "plan_material.py"), req, "--out", name], capture_output=True, text=True, cwd=TMP, env=env)
        return r.stdout, json.load(open(os.path.join(TMP, name)))
    base = json.load(open(os.path.join(BF.SKILL, "references", "format-catalog.json")))["formats"]["tracker"]
    t, tp = plan("issue log", "t.plan.json"); d, dp = plan("data dictionary", "d.plan.json"); r, rp = plan("RACI", "r.plan.json")
    also = "ALSO: material-composer -> xlsx; build scripts/build_sheet.py <sheet.json> --out <Name>; gate scripts/build_sheet.py <Name>.xlsx --check"
    return (base["renderer"] == "xlsx" and "FORMAT: Tracker (data-collection)\nMATCH: alias (1.0)\nBUILD: material-composer -> xlsx" in t
            and "XLSX_TABLE (build_sheet.py <Name>.xlsx --check)" in t and tp["renderer"]["id"] == "xlsx"
            and "BUILD: document-builder -> pdf, docx\n" + also in d and also in r
            and dp["outputs"] == ["pdf", "docx", "xlsx"] and rp["outputs"] == ["pdf", "docx", "xlsx"] and "ALSO:" not in t)


@case("D-53 XLSX_TABLE negative (synthetic): merged cells in a reopened workbook fail and the CLI exits 1", "XLSX_TABLE")
def _():
    src, g, _ = built("merge_src")
    def merge(wb): wb["Tracker"].merge_cells("F3:F4"); wb["Definitions"].merge_cells("A20:B20")
    gm = edited(src, "merged.xlsx", merge)
    rc, out = sheet_cli("merged.xlsx", "--check")
    return (has(gm, "FAIL", "XLSX_TABLE", "merged cells on Tracker: F3:F4") and has(gm, "FAIL", "XLSX_TABLE", "merged cells on Definitions: A20:B20")
            and rc == 1 and "FAIL [XLSX_TABLE] merged cells on Tracker: F3:F4" in out)


@case("D-53 XLSX_TABLE negative (synthetic): free-text status fails: dropdown removed, a typed value outside the list, an inline list, a source Status column with no list or a bad value (build refuses)", "XLSX_TABLE")
def _():
    src, g, _ = built("free_src")
    def nodv(wb): wb["Tracker"].data_validations.dataValidation = []
    def typed(wb): wb["Tracker"]["E4"] = "Blocked"
    def inline(wb):
        from openpyxl.worksheet.datavalidation import DataValidation
        ws = wb["Tracker"]; ws.data_validations.dataValidation = []
        dv = DataValidation(type="list", formula1='"Not Started,In Progress,Done"'); dv.add("E2:E500"); ws.add_data_validation(dv)
    a, b, c = edited(src, "nodv.xlsx", nodv), edited(src, "typed.xlsx", typed), edited(src, "inline.xlsx", inline)
    d1 = tracker(); d1["columns"][4].pop("list")
    d2 = tracker(); d2["rows"][2]["Status"] = "Blocked"
    out = os.path.join(TMP, "refused_tracker"); l, s = quiet(); g2, f2 = BS.run(d2, out, say=s)
    g1 = BS.gate_source(d1)
    return (has(a, "FAIL", "XLSX_TABLE", "Tracker column Status is free text: no dropdown from the Lists tab")
            and has(b, "FAIL", "XLSX_TABLE", "Tracker column Status row 4: 'Blocked' is not in the Status list (Not Started, In Progress, Done)")
            and has(c, "FAIL", "XLSX_TABLE", "dropdown (\"Not Started,In Progress,Done\") is not one list on the Lists tab")
            and has(g1, "FAIL", "XLSX_TABLE", "column Status is free text")
            and f2 is None and not os.path.exists(out + ".xlsx") and any("refused" in x for x in l) and has(g2, "FAIL", "XLSX_TABLE", "Status row 4: 'Blocked'"))


@case("D-53 XLSX_TABLE negative (synthetic): no table, unfrozen header, a header color not from the project brand palette and a missing Definitions tab each fail", "XLSX_TABLE")
def _():
    from openpyxl.styles import PatternFill
    src, g, _ = built("struct_src")
    def notable(wb): del wb["Tracker"].tables["tblTracker"]
    def unfrozen(wb): wb["Tracker"].freeze_panes = None
    def color(wb): wb["Tracker"]["B1"].fill = PatternFill("solid", fgColor="FF0000")
    def nodefs(wb): del wb["Definitions"]
    r = [edited(src, n + ".xlsx", f) for n, f in (("notable", notable), ("unfrozen", unfrozen), ("color", color), ("nodefs", nodefs))]
    return (has(r[0], "FAIL", "XLSX_TABLE", "Tracker has no named Excel table") and has(r[1], "FAIL", "XLSX_TABLE", "header row of Tracker is not frozen")
            and has(r[2], "FAIL", "XLSX_TABLE", "Tracker header B1 not in the project brand colors") and has(r[3], "FAIL", "XLSX_TABLE", "no Definitions tab"))


# ---------------------------------------------------------------- project brand (synthetic projects)
LOGO_SVG = ('<svg xmlns="http://www.w3.org/2000/svg" width="240" height="80" viewBox="0 0 240 80">'
            '<rect x="0" y="0" width="80" height="80" rx="12" fill="#0B7A75"/><rect x="96" y="24" width="140" height="32" fill="#1D3B3A"/></svg>')


def branded_project(name, logo=True):
    """A synthetic web project: Tailwind primary #0B7A75, public/logo.svg, LICENSE naming Northwind Traders Ltd."""
    root = os.path.join(TMP, name)
    files = {"package.json": json.dumps({"name": "northwind-portal", "private": True}),
             "tailwind.config.js": "module.exports = {\n  theme: {\n    extend: {\n      colors: { primary: '#0B7A75' }\n    }\n  }\n}\n",
             "LICENSE": "MIT License\n\nCopyright (c) 2025 Northwind Traders Ltd.\n"}
    if logo:
        files["public/logo.svg"] = LOGO_SVG
    for rel, text in files.items():
        p = os.path.join(root, *rel.split("/")); os.makedirs(os.path.dirname(p), exist_ok=True)
        open(p, "w", encoding="utf-8").write(text)
    os.makedirs(os.path.join(root, "docs"), exist_ok=True)
    return root


class in_project:
    """Resolve brands from this project (BRAND_PROJECT) for the duration of a case."""
    def __init__(self, root):
        self.root = root

    def __enter__(self):
        self.old = os.environ.get("BRAND_PROJECT"); os.environ["BRAND_PROJECT"] = self.root

    def __exit__(self, *a):
        os.environ["BRAND_PROJECT"] = self.old


def spans(doc):
    return [s for p in doc for bl in p.get_text("dict")["blocks"] for ln in bl.get("lines", []) for s in ln["spans"]]


@case("BRAND form + sheet in a synthetic branded project (Tailwind primary #0B7A75, logo.svg, LICENSE): headings in the brand primary, "
      "header band and xlsx headers in the brand dark, logo drawn, footer names Northwind Traders Ltd.", "BRAND")
def _():
    import fitz
    from openpyxl import load_workbook
    root = branded_project("northwind")
    src = os.path.join(root, "docs", "request.form.json"); json.dump(form(), open(src, "w", encoding="utf-8"))
    ssrc = os.path.join(root, "docs", "tracker.json"); json.dump(tracker(), open(ssrc, "w", encoding="utf-8"))
    with in_project(root):
        b = BR.load(source=src)
        l, s = quiet(); g, files = BF.run(form(), os.path.join(TMP, "nw_form"), say=s, source=src)
        l2, s2 = quiet(); g2, f2 = BS.run(tracker(), os.path.join(TMP, "nw_tracker"), say=s2, source=ssrc)
    neutral = BR.Brand({}, TMP, "neutral")
    doc = fitz.open(os.path.join(TMP, "nw_form.pdf"))
    text = " ".join(p.get_text() for p in doc)
    colors = {"%06X" % sp["color"] for sp in spans(doc)}
    fills = {"%02X%02X%02X" % tuple(round(c * 255) for c in d["fill"]) for d in doc[0].get_drawings() if d.get("fill")}
    wb = load_workbook(os.path.join(TMP, "nw_form.xlsx")); wt = load_workbook(f2[0]) if f2 else None
    hdr = wt["Tracker"]["A1"] if wt else None
    return (str(b.sources.get("primary", "")).startswith("tailwind.config.js") and b.org == "Northwind Traders Ltd." and b.logo and files and f2
            and b.colors["dark"] != neutral.colors["dark"] and b.colors["primary"] in colors and b.colors["dark"] in fills
            and bool(doc[0].get_images()) == bool(b.logo_png("light")) and "Northwind Traders Ltd." in text
            and "\u00a92026 Northwind Traders Ltd." in text
            and wb["Questions"]["A1"].fill.fgColor.rgb[-6:] == b.colors["dark"] and hdr.fill.fgColor.rgb[-6:] == b.colors["dark"]
            and "Northwind Traders Ltd." in wt["Tracker"].oddFooter.left.text and wt.properties.creator == "Northwind Traders Ltd.")


@case("BRAND neutral project: no logo drawn, no copyright line, neutral dark header band (form PDF)", "BRAND")
def _():
    import fitz
    l, s = quiet(); g, files = BF.run(form(), os.path.join(TMP, "neutral_form"), say=s)
    b = BF.cur()["brand"]
    doc = fitz.open(os.path.join(TMP, "neutral_form.pdf"))
    text = " ".join(p.get_text() for p in doc)
    return files and b.logo is None and not doc[0].get_images() and "\u00a9" not in text and "Synthetic Equipment Request Form v1.0 \u00b7 Page 1 of" in text \
        and b.colors["dark"] == BR.Brand({}, TMP, "neutral").colors["dark"]


def titlebar(html_text):
    m = re.search(r'<header class="titlebar">(.*?)</header>', html_text, re.S)
    return m.group(1) if m else ""


@case("BRAND e-learning module: the title bar shows the project logo (alt text from the brand name) and the CSS uses the brand primary; "
      "a project with no logo gets no <img> in the title bar", "BRAND")
def _():
    root = branded_project("northwind_html"); bare = branded_project("northwind_nologo", logo=False)
    out = {}
    for tag, r in (("logo", root), ("bare", bare)):
        src = os.path.join(r, "docs", "m.module.json"); json.dump(module(), open(src, "w", encoding="utf-8"))
        with in_project(r):
            b = BR.load(source=src)
            l, s = quiet(); g, files = BM.run(module(), os.path.join(TMP, "nw_mod_" + tag), base_dir=os.path.dirname(src), say=s, source=src)
        out[tag] = (b, open(files[0], encoding="utf-8").read() if files else "")
    (b1, h1), (b2, h2) = out["logo"], out["bare"]
    t1, t2 = titlebar(h1), titlebar(h2)
    return (b1.logo and '<img src="data:image/png;base64,' in t1 and 'alt="Northwind Portal logo"' in t1
            and "--primary:#%s" % b1.colors["primary"] in h1 and str(b1.sources.get("primary", "")).startswith("tailwind.config.js")
            and b2.logo is None and "<img" not in t2 and "<h1>" in t2 and "Northwind Traders Ltd." in h2)


def line_step(doc, text_start):
    """Distance between the baselines of the first two lines of the paragraph that starts with text_start."""
    lines = [(round(ln["spans"][0]["origin"][1], 2), "".join(s["text"] for s in ln["spans"]))
             for bl in doc[0].get_text("dict")["blocks"] for ln in bl.get("lines", [])]
    i = next(k for k, (_, s) in enumerate(lines) if s.startswith(text_start))
    return lines[i + 1][0] - lines[i][0]


@case("BRAND fonts: a project that ships its own TTF (fonts/segoeui.ttf) gets that font embedded in the form PDF, its own line "
      "spacing (wrapped text steps by its metrics, not the bundled fallback's), and an @font-face with the file in the HTML module", "BRAND")
def _():
    import fitz, shutil, zipfile
    src_ttf = os.path.join(os.environ.get("WINDIR", "C:\\Windows"), "Fonts", "segoeui.ttf")
    if not os.path.isfile(src_ttf):
        cand = [p for p in (BR.find_font("DejaVu Sans"), BR.find_font("Liberation Sans")) if p]
        src_ttf = cand[0] if cand else None
    if not src_ttf:
        return "UNVERIFIED"
    root = branded_project("own_font", logo=False)
    os.makedirs(os.path.join(root, "fonts"), exist_ok=True)
    shutil.copy(src_ttf, os.path.join(root, "fonts", os.path.basename(src_ttf)))
    fam = BR.font_name(src_ttf)[0]
    f = form(); f["meta"]["purpose"] = ("Request a synthetic laptop dock for a new desk, a replacement for a broken dock, or a second dock "
                                       "for a shared desk; the support desk confirms the order within two working days.")
    src = os.path.join(root, "docs", "own.form.json"); json.dump(f, open(src, "w", encoding="utf-8"))
    msrc = os.path.join(root, "docs", "own.module.json"); json.dump(module(), open(msrc, "w", encoding="utf-8"))
    with in_project(root):
        b = BR.load(source=src)
        l, s = quiet(); g, files = BF.run(f, os.path.join(TMP, "own_font"), say=s, source=src)
        lm, sm = quiet(); gm, mfiles = BM.run(module(), os.path.join(TMP, "own_font_mod"), base_dir=os.path.dirname(msrc), say=sm,
                                              scorm=True, source=msrc)
    l2, s2 = quiet(); g2, f2 = BF.run(f, os.path.join(TMP, "inter_font"), say=s2)   # neutral project: bundled fallback
    neutral = BF.cur()["brand"]
    d1, d2 = fitz.open(os.path.join(TMP, "own_font.pdf")), fitz.open(os.path.join(TMP, "inter_font.pdf"))
    fonts1 = {x[3] for p in d1 for x in p.get_fonts()}
    m1, m2 = BR.font_metrics(b.measure_path("body")), BR.font_metrics(neutral.measure_path("body"))
    st1, st2 = line_step(d1, "Request a synthetic"), line_step(d2, "Request a synthetic")
    want1, want2 = BF.PT["body"] * m1["line"] * BF.LEADING, BF.PT["body"] * m2["line"] * BF.LEADING
    html_text = open(mfiles[0], encoding="utf-8").read() if mfiles else ""
    names = zipfile.ZipFile(mfiles[-1]).namelist() if mfiles else []
    rel = "fonts/" + os.path.basename(src_ttf)
    return (files and f2 and b.fonts["body"] == fam and b.font_path("body").startswith(root)
            and any(fam in x or fam.replace(" ", "") in x for x in fonts1) and any(x.startswith("PDF font: %s (embedded)" % fam) for x in l)
            and abs(m1["line"] - m2["line"]) > 0.02 and abs(st1 - want1) < 0.2 and abs(st2 - want2) < 0.2 and abs(st1 - st2) > 0.2
            and '@font-face{font-family:"%s";font-weight:400' % fam in html_text and 'url("%s")' % rel in html_text
            and 'font-family:"%s",system-ui,sans-serif' % fam in html_text
            and os.path.isfile(os.path.join(TMP, "fonts", os.path.basename(src_ttf))) and rel in names)


def only_codes():
    valid = sorted({c for _, c, _ in CASES})
    a = sys.argv[1:]
    if "--only" not in a and not any(x.startswith("--only=") for x in a):
        return None
    arg = next((x[7:] for x in a if x.startswith("--only=")), None)
    if arg is None:
        i = a.index("--only"); arg = a[i + 1] if i + 1 < len(a) else ""
    if arg in ("", "list"):
        print("\n".join(valid)); sys.exit(0)
    req = {x.strip() for x in arg.split(",") if x.strip()}
    if req - set(valid):
        sys.stderr.write("unknown --only code(s): %s. Valid codes: %s\n" % (", ".join(sorted(req - set(valid))), ", ".join(valid))); sys.exit(2)
    return req


def main():
    only = only_codes()
    ok = bad = unv = 0
    for name, code, fn in CASES:
        if only and code not in only:
            continue
        try:
            res = fn()
            err = ""
        except Exception as e:
            res, err = False, " (%s: %s)" % (type(e).__name__, e)
        if res == "UNVERIFIED":
            print("UNVERIFIED [%s] %s (node not found; not a pass)" % (code, name))
            unv += 1
            continue
        res = bool(res)
        print("%s [%s] %s%s" % ("PASS" if res else "FAIL", code, name, err))
        ok += res; bad += not res
    print("%d passed, %d failed%s" % (ok, bad, ", %d unverified" % unv if unv else ""))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())

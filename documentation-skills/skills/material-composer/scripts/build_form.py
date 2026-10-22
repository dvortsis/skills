#!/usr/bin/env python3
"""Surveys and forms (request row D-13): form.json -> fillable PDF + .xlsx question bank, behind one gate.

  python3 build_form.py form.json --check              run the gate (pre-build checks plus a trial build for FORM_FIELDS_MATCH)
  python3 build_form.py form.json --out Name [--draft] gate, then write Name.pdf, Name.xlsx and Name.msforms.csv (refuses on FAIL unless --draft)
  --brand brand.json   use this brand profile; otherwise the brand comes from the form's project (scripts/brand.py)

Codes: SURVEY_ITEM, FORM_FIELD, FORM_FIELDS_MATCH, PLACEHOLDER, SOURCES_MISSING, TITLE_CASE, FOLLOWUP_ITEMS (D-46),
SCALE_LABELS (D-48), FORM_TAGS (D-49), FORM_BRANCH (D-51). A build also writes Name.msforms.csv (D-50, FORMS_EXPORT).
Data lists live in references/survey_rules.json. Schema and limits: references/forms-and-elearning.md.
"""
import argparse, csv, datetime, json, os, re, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from text_rules import proper  # shared copy, read-only
import style_tokens as ST
import brand as BR  # shared copy, read-only

TOKENS = ST.shipped()  # sizes; colors, fonts, logo and footer come from the project brand (set_brand)
RULES = json.load(open(os.path.join(SKILL, "references", "survey_rules.json"), encoding="utf-8"))
CHOICE_TYPES = ("single", "multi", "dropdown")
_CUR = {}


def set_brand(source=None, brand=None, project=None):
    """Resolve the project brand for this build (explicit profile > profile beside the source > detected > neutral)."""
    _CUR["t"] = ST.tokens(source, brand, project)
    return _CUR["t"]


def cur():
    return _CUR.get("t") or set_brand()


def rgb(key):
    h = cur()["colors"][key]
    return tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))


# ---------------------------------------------------------------- gate
class Gate:
    def __init__(self, name):
        self.name, self.n, self.msgs = name, 0, []

    def check(self, bad, level, code, msg):
        self.n += 1
        if bad:
            self.msgs.append((level, code, msg))
        return not bad

    @property
    def failed(self):
        return sum(1 for m in self.msgs if m[0] == "FAIL")

    def codes(self, level="FAIL"):
        return [c for lv, c, _ in self.msgs if lv == level]

    def report(self, out=print):
        for lv, code, msg in self.msgs:
            out("%s [%s] %s" % (lv, code, msg))
        out("%s GATE: %d checks, %d failed" % (self.name, self.n, self.failed))


def term_rx(terms, word=True):
    parts = []
    for t in sorted(terms, key=len, reverse=True):
        e = re.escape(t).replace("\\ ", r"\s+")
        parts.append(r"\b%s\b" % e if word and re.match(r"^\w.*\w$|^\w$", t) else e)
    return re.compile("(?i)(" + "|".join(parts) + ")")


PLACEHOLDER_RX = term_rx(RULES["placeholders"])
LEADING_RX = term_rx(RULES["leading_phrases"])
PHI_RX = term_rx(RULES["phi_terms"])
IDENTITY_RX = term_rx(RULES["identity_terms"])
OK_PHRASES_RX = term_rx(RULES["compound_phrases_ok"])
JOIN_RX = re.compile(r"\s(%s)\s" % "|".join(RULES["double_barrel_joiners"]), re.I)


def norm(s):
    return (s or "").replace("\u2019", "'").replace("\u2018", "'")


def all_items(form):
    for s in form.get("sections") or []:
        for it in s.get("items") or []:
            yield s, it


def double_barrel(label):
    """Return (left object, right object) when a label joins two things with and/or, else None."""
    t = OK_PHRASES_RX.sub(" ", norm(label))
    m = JOIN_RX.search(t)
    if not m:
        return None
    right = re.split(r"[?.,;:!()]", t[m.end():])[0].split()
    left = re.split(r"[?.,;:!()]", t[:m.start()])[-1].split()
    if not right or not left:
        return None
    k = max(1, min(len(right), 4))
    return " ".join(left[-k:]), " ".join(right[:4])


def scale_of(it):
    sc = it.get("scale") or {}
    return sc.get("min", 1), sc.get("max", 5), sc


def opt_out(it):
    """D-48: the opt-out answer ("Not applicable", "Prefer not to say") shown as the last choice on scale and single items."""
    return str(it.get("opt_out") or "").strip() if it.get("type") in ("scale", "single") else ""


def numbers(form):
    """Displayed question numbers in PDF order (a signature + date pair counts as two)."""
    return {str(it.get("id")): n for n, (_, it) in enumerate(all_items(form), 1)}


def item_ref(form, iid, nums):
    """How the PDF names an item: Q<number> in surveys (numbers are printed), the quoted label in forms."""
    if (form.get("meta") or {}).get("kind") == "survey":
        return "Q%d" % nums[iid]
    return "\"%s\"" % next((it.get("label") or "") for _, it in all_items(form) if str(it.get("id")) == iid)


def branch_text(form, it, nums=None):
    """D-51: 'Answer Q3 only if you answered Yes to Q2' for an item with show_if {"Q2": "Yes"}; '' when there is none."""
    cond = it.get("show_if")
    nums = nums or numbers(form)
    if not isinstance(cond, dict) or len(cond) != 1:
        return ""
    (src, val), = cond.items()
    if str(src) not in nums or str(it.get("id")) not in nums:
        return ""
    vals = [str(v) for v in (val if isinstance(val, list) else [val])]
    me = "Q%d" % nums[str(it.get("id"))] if (form.get("meta") or {}).get("kind") == "survey" else "this question"
    return "Answer %s only if you answered %s to %s" % (me, " or ".join(vals), item_ref(form, str(src), nums))


def answer_values(it):
    """D-51: answers a show_if may name for this item, or None when the item has no fixed answers (text, date, signature)."""
    typ = it.get("type")
    if typ in CHOICE_TYPES:
        vals = [str(o) for o in it.get("options") or []]
    elif typ == "scale":
        lo, hi, _ = scale_of(it)
        vals = [str(v) for v in range(lo, hi + 1)] if isinstance(lo, int) and isinstance(hi, int) else []
        vals += scale_point_labels(it) or []
    elif typ == "checkbox":
        vals = ["Yes"]
    else:
        return None
    return vals + ([opt_out(it)] if opt_out(it) else [])


def scale_point_labels(it):
    """D-48: one label per point: scale.labels when given, else a standard set from survey_rules.json, else None."""
    lo, hi, sc = scale_of(it)
    if not (isinstance(lo, int) and isinstance(hi, int) and lo < hi):
        return None
    n = hi - lo + 1
    given = [str(x).strip() for x in sc.get("labels") or []]
    if given:
        return given if len(given) == n and all(given) else None
    low, high = str(sc.get("low") or "").strip(), str(sc.get("high") or "").strip()
    if n == 2 and low and high:
        return [low, high]
    for s in RULES["scale_point_labels"]:
        if len(s) == n and s[0].lower() == low.lower() and s[-1].lower() == high.lower():
            return list(s)
    return None


def item_texts(form):
    m = form.get("meta") or {}
    out = [("title", m.get("title")), ("purpose", m.get("purpose")), ("audience", m.get("audience")), ("owner", m.get("owner"))]
    out += [("source", s) for s in m.get("sources") or []]
    for s in form.get("sections") or []:
        out += [("section %r" % s.get("title"), s.get("title")), ("section intro", s.get("intro"))]
        for it in s.get("items") or []:
            iid = it.get("id") or "?"
            out += [(iid, it.get("label")), (iid + " help", it.get("help"))]
            out += [(iid + " option", o) for o in it.get("options") or []]
            sc = it.get("scale") or {}
            out += [(iid + " scale", sc.get("low")), (iid + " scale", sc.get("high"))]
    return [(w, str(t)) for w, t in out if t]


def gate(form):
    g = Gate("FORM")
    meta = form.get("meta") or {}
    kind = meta.get("kind")
    g.check(kind not in ("survey", "form"), "FAIL", "FORM_FIELD", "meta.kind must be survey or form (got %r)" % kind)
    title = meta.get("title") or ""
    g.check(not title.strip(), "FAIL", "FORM_FIELD", "meta.title is empty")
    if title.strip():
        g.check(proper(title) != title, "FAIL", "TITLE_CASE", "title is not Proper Case; use \"%s\"" % proper(title))
    g.check(not [s for s in meta.get("sources") or [] if str(s).strip()], "FAIL", "SOURCES_MISSING",
            "meta.sources is empty; list the sources behind the questions (policy, request, prior survey)")
    if kind == "survey":
        g.check(not meta.get("est_minutes"), "WARN", "FORM_FIELD", "survey has no meta.est_minutes; respondents see the time to complete")
    for where, text in item_texts(form):
        m = PLACEHOLDER_RX.search(text)
        g.check(bool(m), "FAIL", "PLACEHOLDER", "%s: placeholder text '%s'" % (where, m.group(0) if m else ""))

    secs = form.get("sections") or []
    g.check(not secs or not any(s.get("items") for s in secs), "FAIL", "FORM_FIELD", "no sections with items")
    seen, scales = {}, []
    for s in secs:
        st = s.get("title") or ""
        g.check(not st.strip(), "FAIL", "FORM_FIELD", "a section has no title")
        if st.strip():
            g.check(proper(st) != st, "WARN", "TITLE_CASE", "section \"%s\" is not Proper Case; use \"%s\"" % (st, proper(st)))
        items = s.get("items") or []
        for it in items:
            iid, typ, label = str(it.get("id") or ""), it.get("type"), (it.get("label") or "").strip()
            name = iid or "(no id)"
            g.check(not iid, "FAIL", "FORM_FIELD", "an item in section \"%s\" has no id" % st)
            g.check(bool(iid) and not re.match(r"^[A-Za-z][A-Za-z0-9_-]*$", iid), "FAIL", "FORM_FIELD",
                    "%s: id must start with a letter and use only letters, digits, _ or - (it becomes the PDF field name)" % name)
            if iid:
                g.check(iid in seen, "FAIL", "FORM_FIELD", "duplicate id %s (sections \"%s\" and \"%s\")" % (iid, seen.get(iid), st))
                seen.setdefault(iid, st)
            g.check(not label, "FAIL", "FORM_FIELD", "%s: empty label (every field needs a visible label)" % name)
            g.check(typ not in RULES["item_types"], "FAIL", "FORM_FIELD", "%s: unknown type %r (use %s)" % (name, typ, ", ".join(RULES["item_types"])))
            phi = PHI_RX.search(" ".join([label, it.get("help") or "", iid]))
            g.check(bool(phi) and not meta.get("phi_approved"), "FAIL", "FORM_FIELD",
                    "%s asks for PHI (\"%s\"); remove it, or set meta.phi_approved only after privacy approval" % (name, phi.group(0) if phi else ""))
            if meta.get("anonymous"):
                idm = IDENTITY_RX.search(label)
                g.check(bool(idm), "WARN", "FORM_FIELD", "%s asks for \"%s\" in an anonymous survey" % (name, idm.group(0) if idm else ""))
            if typ in RULES["double_barrel_types"] and label:
                db = double_barrel(label)
                g.check(bool(db) and not it.get("compound_ok"), "FAIL", "SURVEY_ITEM",
                        "%s asks two things (\"%s\" and \"%s\"); split it into two items, or set \"compound_ok\": true" % ((name,) + (db or ("", ""))))
            if label:
                lm = LEADING_RX.search(norm(label))
                g.check(bool(lm), "FAIL", "SURVEY_ITEM", "%s leading wording \"%s\"; ask neutrally" % (name, lm.group(0) if lm else ""))
            if it.get("opt_out") is not None:
                g.check(typ not in ("scale", "single") or not str(it.get("opt_out")).strip(), "FAIL", "FORM_FIELD",
                        "%s: opt_out (e.g. \"Not applicable\") works only on scale and single items and needs text" % name)
            if typ in CHOICE_TYPES:
                opts = [o for o in it.get("options") or [] if str(o).strip()]
                g.check(len(opts) < 2, "FAIL", "SURVEY_ITEM", "%s: %s needs at least 2 options (has %d)" % (name, typ, len(opts)))
                allo = list(map(str, opts)) + ([opt_out(it)] if opt_out(it) else [])
                g.check(len(set(allo)) != len(allo), "FAIL", "SURVEY_ITEM", "%s: duplicate options" % name)
            if typ == "scale":
                lo, hi, sc = scale_of(it)
                ok = isinstance(lo, int) and isinstance(hi, int) and lo < hi and hi - lo + 1 <= RULES["scale_max_points"]
                g.check(not ok, "FAIL", "SURVEY_ITEM", "%s: scale needs whole numbers min < max and at most %d points (got %r-%r)" % (name, RULES["scale_max_points"], lo, hi))
                given = sc.get("labels")
                labels_ok = ok and bool(given) and scale_point_labels(it) is not None
                g.check(not (labels_ok or (str(sc.get("low") or "").strip() and str(sc.get("high") or "").strip())), "FAIL", "SURVEY_ITEM",
                        "%s: scale needs both end labels (scale.low and scale.high)" % name)
                if ok:
                    scales.append("%d-%d" % (lo, hi))
                    if given:
                        g.check(not labels_ok, "FAIL", "SURVEY_ITEM", "%s: scale.labels needs %d non-empty labels, one per point (has %d)"
                                % (name, hi - lo + 1, len(given) if isinstance(given, list) else 0))
                    else:
                        # check: SCALE_LABELS WARN
                        g.check(scale_point_labels(it) is None, "WARN", "SCALE_LABELS",
                                "%s: %d-point scale (%d-%d) has no label for every point; add scale.labels with %d labels (the PDF prints only the end labels)"
                                % (name, hi - lo + 1, lo, hi, hi - lo + 1))
        sig = [it for it in items if it.get("type") == "signature"]
        has_date = any(it.get("type") == "date" for it in items)
        for it in sig:
            g.check(not has_date, "WARN", "FORM_FIELD", "%s: signature without a date item in section \"%s\"" % (it.get("id"), st))
    uniq = sorted(set(scales), key=lambda r: tuple(int(x) for x in r.split("-")))
    g.check(len(uniq) > 1, "FAIL", "SURVEY_ITEM", "mixed scales (%s)" % ", ".join(uniq))
    branches(form, g)
    if meta.get("format") == RULES["followup_format"]:
        followup(form, g)
    return g


def branches(form, g):
    """D-51 FORM_BRANCH: show_if {"Q2": "Yes"} must name one earlier item (by id) and answers that item offers."""
    items = [it for _, it in all_items(form)]
    pos = {str(it.get("id")): i for i, it in reversed(list(enumerate(items)))}
    for i, it in enumerate(items):
        cond = it.get("show_if")
        if cond is None:
            continue
        name = str(it.get("id") or "(no id)")
        # check: FORM_BRANCH FAIL
        if not g.check(not isinstance(cond, dict) or len(cond) != 1, "FAIL", "FORM_BRANCH",
                       "%s: show_if must name one item and its answer, e.g. {\"Q2\": \"Yes\"}" % name):
            continue
        (src, val), = cond.items()
        src = str(src)
        if not g.check(src not in pos, "FAIL", "FORM_BRANCH", "%s: show_if names unknown item %s" % (name, src)):
            continue
        if not g.check(pos[src] >= i, "FAIL", "FORM_BRANCH", "%s: show_if item %s must come before %s" % (name, src, name)):
            continue
        ans = answer_values(items[pos[src]])
        if not g.check(ans is None, "FAIL", "FORM_BRANCH", "%s: show_if item %s is a %s item with no fixed answers; branch on a single, multi, dropdown, scale or checkbox item"
                       % (name, src, items[pos[src]].get("type"))):
            continue
        vals = [str(v) for v in (val if isinstance(val, list) else [val])]
        bad = [v for v in vals if v not in ans] if vals else ["(none)"]
        g.check(bool(bad), "FAIL", "FORM_BRANCH", "%s: show_if names option %s that %s does not have (options: %s)"
                % (name, ", ".join("\"%s\"" % b for b in bad), src, ", ".join(ans)))


def followup(form, g):
    """D-46 FOLLOWUP_ITEMS: a follow-up evaluation asks how often the skill is used, what gets in the way and what result it brings."""
    meta = form.get("meta") or {}
    texts = [norm("%s %s" % (it.get("label") or "", it.get("help") or "")) for _, it in all_items(form)]
    miss = [t for t, kws in RULES["followup_topics"].items() if not any(term_rx(kws).search(x) for x in texts)]
    # check: FOLLOWUP_ITEMS WARN
    g.check(bool(miss), "WARN", "FOLLOWUP_ITEMS", "missing: %s (add one item for each; keywords in survey_rules.json followup_topics)" % ", ".join(miss))
    lo, hi = RULES["followup_send_after_days"]
    d = meta.get("send_after_days")
    ok = isinstance(d, (int, float)) and not isinstance(d, bool) and lo <= d <= hi
    g.check(not ok, "WARN", "FOLLOWUP_ITEMS", "meta.send_after_days is %s; send a follow-up evaluation %d-%d days after the training"
            % ("not set" if d is None else repr(d), lo, hi))
    g.check(meta.get("kind") != "survey", "WARN", "FOLLOWUP_ITEMS", "a follow-up evaluation is a survey; set meta.kind to survey")


def fields_match(pdf_path, form, g=None):
    """FORM_FIELDS_MATCH: distinct top-level widget field names in the PDF == items (radio groups and checkbox sets once)."""
    import fitz
    g = g or Gate("FORM")
    ids = [str(it.get("id")) for _, it in all_items(form)]
    names = set()
    with fitz.open(pdf_path) as doc:
        for page in doc:
            for w in page.widgets() or []:
                names.add((w.field_name or "").split(".")[0])
    missing = [i for i in ids if i not in names]
    extra = sorted(names - set(ids))
    detail = ""
    if missing:
        detail += "; missing: " + ", ".join(missing)
    if extra:
        detail += "; extra: " + ", ".join(extra)
    g.check(len(names) != len(ids) or bool(missing), "FAIL", "FORM_FIELDS_MATCH",
            "PDF has %d field names for %d items%s" % (len(names), len(ids), detail))
    return g


# ---------------------------------------------------------------- PDF
W, H, M = 612.0, 792.0, 54.0
CW, BOTTOM = W - 2 * M, H - M
PT = TOKENS["print_pt"]
LEADING = 1.08  # paragraph line step as a multiple of the font's own single line spacing (metrics "line")


class Fonts:
    """The brand body font, embedded: the project's own file, an installed copy, or the bundled fallback (brand.py
    measure_path), so the PDF always carries a real font. Widths and vertical metrics come from the same files."""
    def __init__(self):
        import fitz
        b = cur()["brand"]
        want = b.fonts.get("body")
        self.files = {"brand": b.measure_path("body"), "brandb": b.measure_path("body", bold=True),
                      "brandi": b.measure_path("body", italic=True)}
        if not all(self.files.values()):
            raise SystemExit("STOPPED: no font file to embed for %r and the bundled fallback fonts are missing "
                             "(assets/fonts); reinstall the skill" % want)
        self.m = {k: fitz.Font(fontfile=f) for k, f in self.files.items()}
        self.metrics = {"brand": BR.font_metrics(self.files["brand"]), "brandb": BR.font_metrics(self.files["brandb"]),
                        "brandi": BR.font_metrics(self.files["brandi"])}
        self.reg, self.bold, self.ital = "brand", "brandb", "brandi"
        got = BR.font_name(self.files["brand"])[0] or want
        self.family = got
        self.label = "%s (embedded)" % got if got == want else "%s (embedded; the brand font %s is not available here)" % (got, want)

    # vertical layout from the font's own tables: line step, baseline and optical center of a first line
    def lead(self, font, size, spacing=LEADING):
        return size * self.metrics[font]["line"] * spacing

    def baseline(self, font, size, spacing=LEADING):
        m = self.metrics[font]
        return (self.lead(font, size, spacing) - size * m["line"]) / 2 + size * m["ascent"]

    def center(self, font, size, spacing=LEADING):
        m = self.metrics[font]
        return (self.lead(font, size, spacing) - size * m["line"]) / 2 + size * m["first"]

    def cap(self, font, size):
        return size * self.metrics[font]["cap"]

    def setup(self, page):
        for k, f in self.files.items():
            page.insert_font(fontname=k, fontfile=f)

    def width(self, text, font, size):
        return self.m[font].text_length(text, fontsize=size)

    def wrap(self, text, font, size, width):
        lines = []
        for para in str(text).split("\n"):
            cur = ""
            for w in para.split():
                t = (cur + " " + w).strip()
                if cur and self.width(t, font, size) > width:
                    lines.append(cur); cur = w
                else:
                    cur = t
            lines.append(cur)
        return lines


def logo_png(width=480):
    """The project logo as a print-size PNG with its pixel size, or (None, None) when the brand has no readable logo."""
    b = cur()["brand"]
    raw = b.logo_png("light", width)
    if not raw:
        return None, None
    try:
        import io
        from PIL import Image
        return raw, Image.open(io.BytesIO(raw)).size
    except Exception:
        return raw, b.logo_size("light")


def on_name(text, used):
    n = re.sub(r"[^A-Za-z0-9]+", "_", str(text)).strip("_")[:40] or "Choice"
    base, k = n, 2
    while n in used or n == "Off":
        n = "%s_%d" % (base, k); k += 1
    used.add(n)
    return n


class Builder:
    def __init__(self, form, draft=False):
        import fitz
        self.fitz, self.form, self.meta, self.draft = fitz, form, form.get("meta") or {}, draft
        self.doc = fitz.open()
        self.f = Fonts()
        self.used = {}  # font name -> characters drawn with it (for the ToUnicode rewrite)
        self.groups = []  # (kind, item, [(xref, on_name)])
        self.page, self.y, self.qn = None, M, 0
        self.nums = numbers(form)
        self.primary, self.dark, self.text, self.gray = rgb("primary"), rgb("dark"), rgb("text"), rgb("muted_text")
        self.line, self.border = rgb("card_line"), rgb("context")

    # -- primitives
    def new_page(self, cont=None):
        self.page = self.doc.new_page(width=W, height=H)
        self.f.setup(self.page)
        self.y = M
        if cont:
            self.txt(M, self.y + self.f.baseline(self.f.bold, PT["small"]), "%s (continued)" % cont, self.f.bold, PT["small"], self.primary)
            self.y += self.f.lead(self.f.bold, PT["small"]) + 8

    def txt(self, x, baseline, s, font, size, color):
        self.used.setdefault(font, set()).update(str(s))
        self.page.insert_text((x, baseline), s, fontname=font, fontsize=size, color=color)

    def para(self, x, y, s, font, size, color, width, draw, spacing=LEADING):
        """Wrapped text whose first line's top is y, spaced by the font's own line spacing times `spacing`;
        returns (height, last line width)."""
        lines = self.f.wrap(s, font, size, width)
        step, base = self.f.lead(font, size, spacing), self.f.baseline(font, size, spacing)
        if draw:
            for i, ln in enumerate(lines):
                self.txt(x, y + base + i * step, ln, font, size, color)
        return len(lines) * step, self.f.width(lines[-1], font, size)

    def field_h(self, extra=0):
        """A one-line text field: the field text's line height plus padding (never below 20 pt)."""
        return max(20.0, self.lead(self.f.reg, 10, 1.0) + 8) + extra

    def lead(self, font, size, spacing=LEADING):
        return self.f.lead(font, size, spacing)

    def widget(self, kind, name, rect, tip, item, **kw):
        fz = self.fitz
        w = fz.Widget()
        w.field_type = {"text": fz.PDF_WIDGET_TYPE_TEXT, "radio": fz.PDF_WIDGET_TYPE_RADIOBUTTON,
                        "check": fz.PDF_WIDGET_TYPE_CHECKBOX, "combo": fz.PDF_WIDGET_TYPE_COMBOBOX}[kind]
        w.field_name, w.field_label, w.rect = name, tip, fz.Rect(rect)
        w.border_color, w.border_width, w.fill_color = self.border, 0.8, (1, 1, 1)
        w.text_font, w.text_fontsize, w.text_color = "Helv", kw.get("size", 10), self.text
        if kind in ("radio", "check"):
            w.field_value = False
            w.text_color = self.dark
        if kind == "combo":
            w.choice_values = kw["choices"]
            w.field_value = kw["choices"][0]
        if kw.get("multiline"):
            w.field_flags |= fz.PDF_TX_FIELD_IS_MULTILINE
        if item.get("required"):
            w.field_flags |= fz.PDF_FIELD_IS_REQUIRED
        if kw.get("date"):
            w.script_format = 'AFDate_FormatEx("mm/dd/yyyy");'
            w.script_change = 'AFDate_KeystrokeEx("mm/dd/yyyy");'
        return self.page.add_widget(w).xref

    # -- blocks (draw=False measures)
    def header(self):
        m, f = self.meta, self.f
        title_w = CW - 28
        size = 20
        while len(f.wrap(m.get("title", ""), f.bold, size, title_w)) > 2 and size > 14:
            size -= 1
        title = m.get("title", "")
        words = title.split()
        if len(f.wrap(title, f.bold, size, title_w)) == 2 and len(words) > 2:   # balance two lines (no one-word orphan)
            best = min(range(1, len(words)), key=lambda k: max(f.width(" ".join(words[:k]), f.bold, size), f.width(" ".join(words[k:]), f.bold, size)))
            if max(f.width(" ".join(words[:best]), f.bold, size), f.width(" ".join(words[best:]), f.bold, size)) <= title_w:
                title = " ".join(words[:best]) + "\n" + " ".join(words[best:])
        th, _ = self.para(0, 0, title, f.bold, size, None, title_w, False, 0.96)
        p = self.page
        # logo row on white (sheet layout): the project logo left, about 0.45 in tall (no logo: none drawn and the row
        # keeps only the format label); format label right in the brand primary
        png, lsz = logo_png()
        lh = 32.4 if png else 18.0
        if png and lsz:
            lw = min(lh * lsz[0] / max(lsz[1], 1), CW * 0.45)
            ih = lw * lsz[1] / max(lsz[0], 1)
            p.insert_image(self.fitz.Rect(M, M + (lh - ih) / 2, M + lw, M + (lh + ih) / 2), stream=png)
        label = "Survey" if m.get("kind") == "survey" else "Form"
        lsize = PT["small"] + 2
        base = M + lh / 2 + f.cap(f.bold, lsize) / 2
        x = W - M
        if self.draft:
            x -= f.width("DRAFT", f.bold, lsize)
            self.txt(x, base, "DRAFT", f.bold, lsize, rgb("alert_text"))
            label += " \u00b7 "
        self.txt(x - f.width(label, f.bold, lsize), base, label, f.bold, lsize, self.primary)
        # full-width band in the brand dark color: title, then "Version X · Date"; 4 pt secondary-color rule under it
        top = M + lh + 8
        vsize = PT["small"] + 1
        block = th + 4 + f.lead(f.reg, vsize, 1.0)  # title lines, gap, version line: all from the font's metrics
        band = max(64, block + 26)
        p.draw_rect(self.fitz.Rect(M, top, W - M, top + band), color=None, fill=self.dark)
        p.draw_rect(self.fitz.Rect(M, top + band, W - M, top + band + 4), color=None, fill=rgb("secondary"))
        tx = M + 14
        ty = top + (band - block) / 2
        self.para(tx, ty, title, f.bold, size, (1, 1, 1), title_w, True, 0.96)
        self.txt(tx, ty + th + 4 + f.baseline(f.reg, vsize, 1.0), "Version %s \u00b7 %s" % (m.get("version", ""), m.get("date", "")), f.reg, vsize, rgb("card"))
        self.y = top + band + 4 + 12

        if m.get("purpose"):
            h, _ = self.para(M, self.y, m["purpose"], f.reg, PT["body"], self.text, CW, True)
            self.y += h + 4
        bits = []
        if m.get("audience"):
            bits.append("For: %s" % m["audience"])
        if m.get("kind") == "survey" and m.get("est_minutes"):
            bits.append("About %s minutes" % m["est_minutes"])
        if m.get("kind") == "survey" and m.get("anonymous"):
            bits.append("Anonymous: no name or ID is collected")
        if m.get("kind") == "form" and m.get("owner"):
            bits.append("Owner: %s" % m["owner"])
        if bits:
            h, _ = self.para(M, self.y, " \u00b7 ".join(bits), f.reg, PT["small"], self.gray, CW, True)
            self.y += h + 2
        if any(it.get("required") for _, it in all_items(self.form)):
            bl = self.y + f.baseline(f.reg, PT["small"])
            self.txt(M, bl, "*", f.bold, PT["small"] + 2, self.primary)
            self.txt(M + 8, bl, "Required", f.reg, PT["small"], self.gray)
            self.y += f.lead(f.reg, PT["small"]) + 2
        self.y += 6
        self.page.draw_line((M, self.y), (W - M, self.y), color=self.line, width=0.8)
        self.y += 12

    def section_head(self, s, draw):
        f, y0 = self.f, self.y
        h, _ = self.para(M, y0, s.get("title", ""), f.bold, PT["h2"], self.primary, CW, draw, 1.0)
        y = y0 + h + 2
        if draw:
            self.page.draw_line((M, y), (W - M, y), color=self.primary, width=0.8)
        y += 6
        if s.get("intro"):
            hh, _ = self.para(M, y, s["intro"], f.reg, PT["body"], self.text, CW, draw)
            y += hh + 2
        return y - y0 + 6

    def label(self, x, y, it, width, draw, number=True):
        f = self.f
        y0 = y
        bt = branch_text(self.form, it, self.nums)   # D-51: printed above the question it controls
        if bt:
            hb, _ = self.para(x, y, bt, f.ital, PT["small"], self.primary, width, draw)
            y += hb + 1
        lab = it.get("label", "")
        if number and self.meta.get("kind") == "survey":
            lab = "%d. %s" % (self.nums.get(str(it.get("id")), self.qn), lab)
        star = 10 if it.get("required") else 0
        h, lw = self.para(x, y, lab, f.bold, PT["body"], self.text, width - star, draw)
        if it.get("required") and draw:
            self.txt(x + lw + 3, y + h - f.lead(f.bold, PT["body"]) + f.baseline(f.bold, PT["body"]), "*", f.bold, PT["body"] + 2, self.primary)
        y2 = y + h
        if it.get("help"):
            hh, _ = self.para(x, y2, it["help"], f.ital, PT["small"], self.gray, width, draw)
            y2 += hh
        return y2 - y0

    def choice_list(self, it, y, draw, kind):
        f, opts = self.f, [str(o) for o in it.get("options") or []] + ([opt_out(it)] if kind == "radio" and opt_out(it) else [])
        cols = 2 if len(opts) > RULES["single_column_max_options"] else 1
        colw = CW / cols
        rows = (len(opts) + cols - 1) // cols
        y0, used, xrefs, rowy = y, set(), [], y
        for r in range(rows):
            rh = 0
            for c in range(cols):
                i = c * rows + r
                if i >= len(opts):
                    continue
                x = M + 4 + c * colw
                hh, _ = self.para(x + 18, rowy, opts[i], f.reg, PT["body"], self.text, colw - 26, draw)
                if draw:
                    c = rowy + f.center(f.reg, PT["body"])  # box centered on the option's first line
                    rect = (x, c - 5.5, x + 11, c + 5.5)
                    nm = it["id"] if kind == "radio" else "%s.%d" % (it["id"], i + 1)
                    xr = self.widget(kind, nm, rect, "%s: %s" % (it.get("label", ""), opts[i]), it)
                    xrefs.append((xr, on_name(opts[i], used)))
                rh = max(rh, hh)
            rowy += rh + 0.4 * PT["body"]
        if draw:
            self.groups.append((kind, it, xrefs))
        return rowy - y0

    def scale(self, it, y, draw):
        """Numbers above the points; a label under every point (D-48) or, when none can be derived, the two end labels.
        The opt-out answer is the last choice, set apart by a small gap."""
        f = self.f
        lo, hi, sc = scale_of(it)
        n = hi - lo + 1
        labels, oo = scale_point_labels(it), opt_out(it)
        gap = 14.0 if oo else 0.0
        colw = min(80.0 if labels else 56.0, (CW - 8 - gap) / (n + (1 if oo else 0)))
        x0, total = M + 4, colw * n
        xrefs, under = [], 0.0
        num_h = f.lead(f.reg, PT["small"], 1.0)
        rtop = y + num_h + 2              # radio row under the point numbers
        ltop = rtop + 12 + 5              # point labels under the radios
        step = f.lead(f.reg, PT["small"])
        for k in range(n + (1 if oo else 0)):
            last = k == n
            cx = x0 + colw * k + colw / 2 + (gap if last else 0)
            num = "" if last else str(lo + k)
            word = oo if last else (labels[k] if labels else "")
            if draw:
                if num:
                    self.txt(cx - f.width(num, f.reg, PT["small"]) / 2, y + f.baseline(f.reg, PT["small"], 1.0), num, f.reg, PT["small"], self.text)
                tip = "%s: %s" % (it.get("label", ""), " ".join(x for x in (num, word) if x))
                xr = self.widget("radio", it["id"], (cx - 6, rtop, cx + 6, rtop + 12), tip, it)
                xrefs.append((xr, on_name(oo, set()) if last else num))
            if word:
                lines = f.wrap(word, f.reg, PT["small"], colw - 4)
                for i, ln in enumerate(lines):
                    if draw:
                        self.txt(cx - f.width(ln, f.reg, PT["small"]) / 2, ltop + f.baseline(f.reg, PT["small"]) + i * step, ln, f.reg, PT["small"], self.gray)
                under = max(under, len(lines) * step)
        if draw:
            self.groups.append(("radio", it, xrefs))
        if labels:
            return ltop - y + under
        yl = ltop
        lw = max(60.0, total / 2 - 6)
        h1, _ = self.para(x0, yl, sc.get("low", ""), f.reg, PT["small"], self.gray, lw, draw)
        hl = self.f.wrap(sc.get("high", ""), f.reg, PT["small"], lw)
        for i, ln in enumerate(hl):
            if draw:
                self.txt(x0 + total - f.width(ln, f.reg, PT["small"]), yl + f.baseline(f.reg, PT["small"]) + i * step, ln, f.reg, PT["small"], self.gray)
        return ltop - y + max(h1, len(hl) * step, under)

    def text_field(self, it, x, y, width, draw, height=None, multiline=False, date=False):
        height = height or self.field_h()
        if draw:
            self.widget("text", it["id"], (x, y, x + width, y + height), it.get("label", ""), it, multiline=multiline, date=date)
            if date:
                self.txt(x + width + 8, y + height / 2 + self.f.cap(self.f.reg, PT["small"]) / 2, "MM/DD/YYYY", self.f.reg, PT["small"], self.gray)
        return height

    def item(self, it, draw, pair=None):
        """One item (or a signature + date pair) as one unbreakable block; returns height."""
        y0, typ, f = self.y, it.get("type"), self.f
        if draw:
            self.qn += 1
        if pair is not None:
            lw = 300
            h1 = self.label(M, y0, it, lw - 16, draw)
            if draw:
                self.qn += 1
            h2 = self.label(M + lw, y0, pair, CW - lw, draw)
            yy = y0 + max(h1, h2) + 4
            fh = self.field_h(2)
            self.text_field(it, M, yy, lw - 24, draw, fh)
            self.text_field(pair, M + lw, yy, 100, draw, fh, date=True)
            return yy + fh - y0
        if typ == "checkbox":
            h = self.label(M + 18, y0, it, CW - 18, draw)
            if draw:
                c = y0 + f.center(f.bold, PT["body"])  # box centered on the label's first line
                xr = self.widget("check", it["id"], (M + 2, c - 5.5, M + 13, c + 5.5), it.get("label", ""), it)
                self.groups.append(("single_check", it, [(xr, "Yes")]))
            return max(h, 14)
        h = self.label(M, y0, it, CW, draw) + 5
        y = y0 + h
        if typ == "single":
            h += self.choice_list(it, y, draw, "radio")
        elif typ == "multi":
            h += self.choice_list(it, y, draw, "check")
        elif typ == "scale":
            h += self.scale(it, y, draw)
        elif typ == "dropdown":
            if draw:
                self.widget("combo", it["id"], (M, y, M + 260, y + self.field_h()), it.get("label", ""), it, choices=["Select one"] + [str(o) for o in it.get("options") or []])
            h += self.field_h()
        elif typ == "longtext":
            h += self.text_field(it, M, y, CW, draw, int(it.get("rows", 4)) * self.lead(f.reg, 10, 1.15) + 8, multiline=True)
        elif typ == "date":
            h += self.text_field(it, M, y, 110, draw, date=True)
        elif typ == "signature":
            h += self.text_field(it, M, y, 280, draw, self.field_h(2))
        else:
            h += self.text_field(it, M, y, CW, draw)
        return h

    def fit(self, h, cont=None):
        if self.y + h > BOTTOM:
            self.new_page(cont)

    def build(self):
        self.new_page()
        self.header()
        for s in self.form.get("sections") or []:
            items = list(s.get("items") or [])
            blocks, i = [], 0
            while i < len(items):
                it = items[i]
                if it.get("type") == "signature" and i + 1 < len(items) and items[i + 1].get("type") == "date":
                    blocks.append((it, items[i + 1])); i += 2
                else:
                    blocks.append((it, None)); i += 1
            sh = self.section_head(s, False)
            first = self.item(blocks[0][0], False, blocks[0][1]) if blocks else 0
            self.fit(sh + first)
            self.y += self.section_head(s, True)
            for it, pair in blocks:
                h = self.item(it, False, pair)
                self.fit(h, s.get("title"))
                self.y += self.item(it, True, pair) + 14
            self.y += 6
        self.group_fields()
        self.footer()
        return self.doc

    def group_fields(self):
        """Make radio sets and multi-select checkbox sets real AcroForm groups: one parent field per item id."""
        doc = self.doc
        cat = doc.pdf_catalog()
        kind_, val = doc.xref_get_key(cat, "AcroForm/Fields")
        fields = [int(x) for x in re.findall(r"(\d+) 0 R", val)]
        for kind, it, xrefs in self.groups:
            if kind == "single_check" or not xrefs:
                continue
            kids = [x for x, _ in xrefs]
            label = it.get("label", "").replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            par = doc.get_new_xref()
            if kind == "radio":
                ff = 32768 + 16384 + (2 if it.get("required") else 0)
                doc.update_object(par, "<</FT/Btn/Ff %d/T(%s)/TU(%s)/V/Off/Kids[%s]>>" % (ff, it["id"], label, " ".join("%d 0 R" % k for k in kids)))
                for x, on in xrefs:
                    obj = doc.xref_object(x, compressed=True)
                    obj = re.sub(r"/FT\s*/Btn|/Ff\s*\d+|/T\s*\([^)]*\)|/V\s*(/Off|\(Off\))", "", obj)
                    doc.update_object(x, obj)
                    for ap in ("AP/N", "AP/D"):
                        t, v = doc.xref_get_key(x, ap)
                        if t == "dict":
                            doc.xref_set_key(x, ap, v.replace("/Yes ", "/%s " % on))
                    doc.xref_set_key(x, "Parent", "%d 0 R" % par)
            else:
                doc.update_object(par, "<</FT/Btn/T(%s)/TU(%s)/Kids[%s]>>" % (it["id"], label, " ".join("%d 0 R" % k for k in kids)))
                for n, (x, on) in enumerate(xrefs, 1):
                    doc.xref_set_key(x, "T", "(%d)" % n)
                    doc.xref_set_key(x, "Parent", "%d 0 R" % par)
            fields = [x for x in fields if x not in kids] + [par]
        doc.xref_set_key(cat, "AcroForm/Fields", "[%s]" % " ".join("%d 0 R" % x for x in fields))

    def footer(self):
        m, f = self.meta, self.f
        year = (re.match(r"(\d{4})", str(m.get("date", ""))) or [None, str(datetime.date.today().year)])[1]
        base = cur()["brand"].footer(year)  # "" when the project names no owner: no copyright line
        n = len(self.doc)
        for i, page in enumerate(self.doc, 1):
            self.page = page
            rest = "%s v%s \u00b7 Page %d of %d" % (m.get("title", ""), m.get("version", ""), i, n)
            s = "%s \u00b7 %s" % (base, rest) if base else rest
            size = 8
            page.draw_line((M, H - M + 10), (W - M, H - M + 10), color=self.line, width=0.6)
            if not base or f.width(s, f.reg, size) <= CW:
                self.txt(M, H - M + 24, s, f.reg, size, self.gray)
            else:
                self.txt(M, H - M + 22, base, f.reg, size, self.gray)
                self.txt(M, H - M + 32, rest, f.reg, size, self.gray)


# Many fonts map several code points to one glyph (space and no-break space, hyphen and U+2010, digits and look-alike
# letters in other scripts). PyMuPDF's ToUnicode then keeps one of them at random, so copied or extracted text reads wrong
# ("v1.0" came out with Armenian letters in one font). Rewrite each embedded font's ToUnicode from the characters this
# build actually drew with it: glyph -> the lowest code point drawn with that glyph.
CMAP_HEAD = ("/CIDInit /ProcSet findresource begin\n12 dict begin\nbegincmap\n/CIDSystemInfo <</Registry(Adobe)/Ordering(UCS)/Supplement 0>> def\n"
             "/CMapName /Adobe-Identity-UCS def\n/CMapType 2 def\n1 begincodespacerange\n<0000> <FFFF>\nendcodespacerange\n")
CMAP_TAIL = "endcmap\nCMapName currentdict /CMap defineresource pop\nend\nend\n"


def utf16hex(ch):
    return ch.encode("utf-16-be").hex().upper()


def fix_tounicode(doc, fonts, used):
    """fonts: Fonts (name -> file); used: {font name: set of characters drawn with it}."""
    import fitz
    done = set()
    for page in doc:
        for xref, _ext, _typ, _base, name, _enc in page.get_fonts():
            if xref in done or name not in fonts.files:
                continue
            done.add(xref)
            t, v = doc.xref_get_key(xref, "ToUnicode")
            if t != "xref":
                continue
            font = fitz.Font(fontfile=fonts.files[name])
            gid = {}
            for ch in sorted(used.get(name, ())):
                g = font.has_glyph(ord(ch))
                if g:
                    gid.setdefault(g, ch)
            pairs = sorted(gid.items())
            body = ""
            for i in range(0, len(pairs), 100):
                chunk = pairs[i:i + 100]
                body += "%d beginbfchar\n%sendbfchar\n" % (len(chunk), "".join("<%04X> <%s>\n" % (g, utf16hex(c)) for g, c in chunk))
            doc.update_stream(int(v.split()[0]), (CMAP_HEAD + body + CMAP_TAIL).encode("latin-1"))


def pdf_text(s):
    """PDF text string as UTF-16BE hex, safe for any character."""
    return "<FEFF%s>" % str(s).encode("utf-16-be").hex().upper()


def tag_pdf(doc, title):
    """D-49: basic tags for screen readers (not full PDF/UA). Catalog /StructTreeRoot, /Lang, /MarkInfo, /ViewerPreferences
    DisplayDocTitle; each page's content is one marked P (MCID 0) and each field is a Form element (OBJR to its widgets,
    /Alt = label) in widget order; /Tabs /S makes the tab order follow that structure. Field tooltips (/TU) are kept."""
    cat = doc.pdf_catalog()
    root, top = doc.get_new_xref(), doc.get_new_xref()
    kids, nums, key = [], [], 0
    for page in doc:
        xs = page.get_contents()
        data = b"\n".join(doc.xref_stream(x) for x in xs)
        doc.update_stream(xs[0], b"/P <</MCID 0>> BDC\n" + data + b"\nEMC\n")
        doc.xref_set_key(page.xref, "Contents", "%d 0 R" % xs[0])
        p_el = doc.get_new_xref()
        doc.update_object(p_el, "<</Type/StructElem/S/P/P %d 0 R/Pg %d 0 R/K 0>>" % (top, page.xref))
        doc.xref_set_key(page.xref, "StructParents", str(key))
        nums.append("%d [%d 0 R]" % (key, p_el)); key += 1
        kids.append(p_el)
        fields = {}
        for w in page.widgets() or []:
            fields.setdefault(w.field_name.split(".")[0], []).append(w)
        for name, ws in fields.items():
            el = doc.get_new_xref()
            objr = " ".join("<</Type/OBJR/Obj %d 0 R/Pg %d 0 R>>" % (w.xref, page.xref) for w in ws)
            alt = ws[0].field_label or name
            doc.update_object(el, "<</Type/StructElem/S/Form/P %d 0 R/Pg %d 0 R/Alt %s/K[%s]>>" % (top, page.xref, pdf_text(alt), objr))
            for w in ws:
                doc.xref_set_key(w.xref, "StructParent", str(key))
                nums.append("%d %d 0 R" % (key, el)); key += 1
            kids.append(el)
        doc.xref_set_key(page.xref, "Tabs", "/S")
    doc.update_object(top, "<</Type/StructElem/S/Document/P %d 0 R/T %s/K[%s]>>" % (root, pdf_text(title), " ".join("%d 0 R" % k for k in kids)))
    doc.update_object(root, "<</Type/StructTreeRoot/K %d 0 R/ParentTree <</Nums[%s]>>/ParentTreeNextKey %d>>" % (top, " ".join(nums), key))
    doc.xref_set_key(cat, "StructTreeRoot", "%d 0 R" % root)
    doc.xref_set_key(cat, "Lang", "(%s)" % RULES["pdf_tags"]["lang"])
    doc.xref_set_key(cat, "MarkInfo", "<</Marked true>>")
    doc.xref_set_key(cat, "ViewerPreferences", "<</DisplayDocTitle true>>")


def form_tags(pdf_path, g=None):
    """D-49 FORM_TAGS: reopen the built PDF and confirm the tags; returns (gate, summary line)."""
    import fitz
    g = g or Gate("FORM")
    lang = RULES["pdf_tags"]["lang"]
    miss = []
    with fitz.open(pdf_path) as doc:
        cat = doc.pdf_catalog()
        key = lambda k: doc.xref_get_key(cat, k)
        if key("StructTreeRoot")[0] != "xref":
            miss.append("/StructTreeRoot")
        if key("Lang")[1] != lang:
            miss.append("/Lang %s" % lang)
        if key("MarkInfo/Marked")[1] != "true":
            miss.append("/MarkInfo Marked true")
        if not str((doc.metadata or {}).get("title") or "").strip():
            miss.append("/Title")
        if key("ViewerPreferences/DisplayDocTitle")[1] != "true":
            miss.append("/ViewerPreferences DisplayDocTitle")
        pages = [str(i) for i, p in enumerate(doc, 1) if doc.xref_get_key(p.xref, "Tabs")[1] != "/S"]
        if pages:
            miss.append("/Tabs /S on page %s" % ", ".join(pages))
        no_tu = sorted({w.field_name.split(".")[0] for p in doc for w in p.widgets() or [] if not (w.field_label or "").strip()})
        if no_tu:
            miss.append("field tooltip (/TU) on %s" % ", ".join(no_tu))
        n = len(doc)
    # check: FORM_TAGS FAIL
    g.check(bool(miss), "FAIL", "FORM_TAGS", "PDF is not tagged for screen readers; missing: %s" % "; ".join(miss))
    return g, ("PDF tags: /StructTreeRoot, /Lang %s, /MarkInfo Marked true, /Title, DisplayDocTitle, /Tabs /S on %d of %d pages, "
               "field tooltips kept (basic tags, not full PDF/UA)" % (lang, n, n) if not miss else None)


def build_pdf(form, path, draft=False):
    b = Builder(form, draft)
    doc = b.build()
    doc.set_metadata({"title": form["meta"].get("title", ""), "author": form["meta"].get("owner", ""), "subject": form["meta"].get("purpose", ""),
                      "creator": "material-composer build_form.py"})
    doc.subset_fonts()
    fix_tounicode(doc, b.f, b.used)
    tag_pdf(doc, form["meta"].get("title", ""))
    doc.save(path, garbage=3, deflate=True)
    doc.close()
    return b.f.label


# ---------------------------------------------------------------- XLSX
def build_xlsx(form, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    T = cur()
    hdr_fill = PatternFill("solid", fgColor=T["colors"]["dark"])
    hdr_font = Font(name=T["fonts"]["body"], bold=True, color=T["colors"]["white"])
    body = Font(name=T["fonts"]["body"], color=T["colors"]["text"])
    ws.append(["ID", "Section", "Question", "Type", "Required", "Options/Scale"])
    for s, it in all_items(form):
        if it.get("type") == "scale":
            lo, hi, sc = scale_of(it)
            labels = scale_point_labels(it)
            if labels:
                opt = "; ".join("%d = %s" % (lo + k, w) for k, w in enumerate(labels))
            else:
                opt = "%s-%s (%s = %s, %s = %s)" % (lo, hi, lo, sc.get("low", ""), hi, sc.get("high", ""))
        elif it.get("type") == "date":
            opt = "MM/DD/YYYY"
        else:
            opt = "; ".join(map(str, it.get("options") or []))
        if opt_out(it):
            opt += "; " + opt_out(it)
        ws.append([it.get("id"), s.get("title"), it.get("label"), it.get("type"), "Yes" if it.get("required") else "No", opt])
    about = wb.create_sheet("About")
    about.append(["Field", "Value"])
    m = form.get("meta") or {}
    for k in ("title", "kind", "format", "send_after_days", "audience", "purpose", "owner", "version", "date", "est_minutes", "anonymous", "phi_approved"):
        v = m.get(k)
        about.append([k, "Yes" if v is True else "No" if v is False else ("" if v is None else str(v))])
    about.append(["sources", "; ".join(map(str, m.get("sources") or []))])
    about.append(["items", sum(1 for _ in all_items(form))])
    about.append(["built", "%s by material-composer build_form.py" % datetime.date.today().isoformat()])
    for sh, widths in ((ws, [10, 24, 60, 12, 10, 48]), (about, [16, 80])):
        for c in sh[1]:
            c.fill, c.font = hdr_fill, hdr_font
            c.alignment = Alignment(vertical="center")
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.font, c.alignment = body, Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(widths):
            sh.column_dimensions[chr(65 + i)].width = w
        sh.freeze_panes = "A2"
    wb.save(path)


# ---------------------------------------------------------------- Microsoft Forms question file (D-50)
def msforms_rows(form):
    """One row per item in PDF order: the label as printed (Forms numbers questions itself), type, options, required, help, branching."""
    nums, rows = numbers(form), []
    for _, it in all_items(form):
        typ = it.get("type")
        opts = []
        if typ == "scale":
            lo, hi, sc = scale_of(it)
            labels = scale_point_labels(it)
            if isinstance(lo, int) and isinstance(hi, int):
                opts = ["%d %s" % (lo + k, labels[k]) if labels else str(lo + k) for k in range(hi - lo + 1)]
                if not labels and opts:
                    opts[0] += " %s" % (sc.get("low") or ""); opts[-1] += " %s" % (sc.get("high") or "")
        elif typ in CHOICE_TYPES:
            opts = [str(o) for o in it.get("options") or []]
        elif typ == "checkbox":
            opts = ["Yes"]
        if opt_out(it):
            opts.append(opt_out(it))
        rows.append([it.get("label") or "", RULES["msforms_types"].get(typ, "Text"), "; ".join(o.strip() for o in opts),
                     "Yes" if it.get("required") else "No", it.get("help") or "", branch_text(form, it, nums)])
    return rows


def build_msforms(form, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(RULES["msforms_columns"])
        w.writerows(msforms_rows(form))


# ---------------------------------------------------------------- CLI
def scale_note(form):
    """One line naming the scale items and how their points are labelled (D-48)."""
    lab, ends = [], []
    for _, it in all_items(form):
        if it.get("type") == "scale":
            (lab if scale_point_labels(it) else ends).append(str(it.get("id")))
    if not lab and not ends:
        return None
    s = "scale labels: every point labelled on %s" % (", ".join(lab) or "none")
    return s + ("; end labels only on %s" % ", ".join(ends) if ends else "")


def run(form, out=None, draft=False, check=False, say=print, source=None, brand=None):
    """source: the form.json path (its project supplies the brand); brand: an explicit brand.json."""
    set_brand(source, brand)
    g = gate(form)
    if check or out is None:
        with tempfile.TemporaryDirectory() as td:
            p = os.path.join(td, "trial.pdf")
            try:
                build_pdf(form, p, True)
                fields_match(p, form, g)
                form_tags(p, g)
            except Exception as e:
                g.check(True, "FAIL", "FORM_FIELDS_MATCH", "trial build failed: %s" % e)
        g.report(say)
        if scale_note(form):
            say(scale_note(form))
        return g, None
    if g.failed and not draft:
        g.report(say)
        say("refused: fix the FAIL lines or pass --draft")
        return g, None
    pdf, xlsx, msf = out + ".pdf", out + ".xlsx", out + ".msforms.csv"
    font = build_pdf(form, pdf, draft)
    build_xlsx(form, xlsx)
    build_msforms(form, msf)
    fields_match(pdf, form, g)
    _, tags = form_tags(pdf, g)
    g.report(say)
    say("PDF font: %s" % font)
    if tags:
        say(tags)
    if scale_note(form):
        say(scale_note(form))
    say("wrote %s, %s and %s%s" % (pdf, xlsx, msf, " (DRAFT)" if draft else ""))
    say("Microsoft Forms: create a new form and paste each row of %s (question, then its options; type, required and branching "
        "tell you which settings to pick). Forms does not import this file automatically." % os.path.basename(msf))
    return g, (pdf, xlsx, msf)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("form")
    ap.add_argument("--out")
    ap.add_argument("--draft", action="store_true")
    ap.add_argument("--check", action="store_true")
    ap.add_argument("--brand", help="brand profile (brand.json); default: the form's project brand")
    a = ap.parse_args()
    form = json.load(open(a.form, encoding="utf-8"))
    if not a.check and not a.out:
        ap.error("--out Name or --check is required")
    g, files = run(form, a.out, a.draft, a.check, source=a.form, brand=a.brand)
    print(cur()["brand"].lines()[0])
    return 1 if g.failed and not (a.draft and files) else 0


if __name__ == "__main__":
    sys.exit(main())

#!/usr/bin/env python3
"""Excel tables (request row D-53): tracker, data-dictionary and RACI tables -> one <Name>.xlsx, behind one gate.

  python3 build_sheet.py sheet.json --out Name    writes Name.xlsx; refuses (no file) on FAIL
  python3 build_sheet.py sheet.json --check       checks the source and a trial build
  python3 build_sheet.py Name.xlsx --check        reopens a built workbook and checks it
  --brand brand.json   use this brand profile; otherwise the brand comes from the source's project (scripts/brand.py)

Workbook: a first tab holding one named Excel table (header row frozen, header in the brand's dark color with white
text, the brand body font, the brand footer on printed pages), a Lists tab
(one column per dropdown list, the only source of every dropdown) and a Definitions tab (every column and list value).
Code: XLSX_TABLE (FAIL) on merged cells, free-text status (no dropdown, or a dropdown not from the Lists tab),
values outside their list, a missing table, unfrozen header, header colors not from the project brand palette, missing tabs.
Data lists live in references/material_rules.json "sheet". Schema: references/forms-and-elearning.md (Excel tables).
"""
import argparse, datetime, json, os, re, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import style_tokens as ST

RULES = json.load(open(os.path.join(SKILL, "references", "material_rules.json"), encoding="utf-8"))["sheet"]
_CUR = {}


def set_brand(source=None, brand=None, project=None):
    """Resolve the project brand for this build or check (explicit profile > beside the source > detected > neutral)."""
    _CUR["t"] = ST.tokens(source, brand, project)
    return _CUR["t"]


def cur():
    return _CUR.get("t") or set_brand()


_MEASURE = {}


def units(text, bold=False):
    """Width of text in Excel column units (widths of "0"), measured with the brand body font's own file (brand.py
    measure_path: the project's font, an installed copy or the bundled fallback); characters when it can't be read."""
    s = str(text if text is not None else "")
    b = cur()["brand"]
    path = b.measure_path("body", bold=bold)
    key = (path, bold)
    if key not in _MEASURE:
        try:
            from PIL import ImageFont
            f = ImageFont.truetype(path, 110)
            _MEASURE[key] = (f, f.getlength("0") or 1.0)
        except Exception:
            _MEASURE[key] = None
    m = _MEASURE[key]
    if not m or not s:
        return float(len(s))
    return max(m[0].getlength(line) for line in s.split("\n")) / m[1]


class Gate:
    def __init__(self, name):
        self.name, self.n, self.msgs = name, 0, []

    def add(self, level, code, msg):
        self.msgs.append((level, code, msg))

    def check(self, bad, msg):
        self.n += 1
        if bad:
            self.add("FAIL", "XLSX_TABLE", msg)  # check: XLSX_TABLE FAIL
        return bool(bad)

    @property
    def failed(self):
        return sum(1 for lv, _, _ in self.msgs if lv == "FAIL")

    def codes(self):
        return sorted({c for lv, c, _ in self.msgs if lv == "FAIL"})

    def report(self, out=print):
        for lv, c, m in self.msgs:
            out("%s [%s] %s" % (lv, c, m))


def is_status(name):
    return str(name).strip().lower() in RULES["status_headers"]


def tab_of(meta):
    return (RULES["tab_labels"].get(str(meta.get("format", "")).lower()) or meta.get("tab") or "Table")[:31]


def table_name(tab):
    return "tbl" + re.sub(r"[^A-Za-z0-9]", "", tab)


def rows_of(doc):
    cols = [c["name"] for c in doc.get("columns") or []]
    out = []
    for r in doc.get("rows") or []:
        out.append([r.get(c, "") for c in cols] if isinstance(r, dict) else list(r))
    return cols, out


def gate_source(doc, g=None):
    """Check sheet.json before building."""
    g = g or Gate("XLSX_TABLE")
    cols, rows = rows_of(doc)
    lists = doc.get("lists") or {}
    g.check(not cols, "no columns")
    g.check(not rows, "no rows")
    g.check(len(set(c.lower() for c in cols)) != len(cols), "duplicate column names: %s" % ", ".join(cols))
    for name, vals in lists.items():
        v = [str(x.get("value", "")) for x in vals]
        g.check(not v or "" in v or len(set(v)) != len(v), "list %s needs unique, non-empty values" % name)
        g.check(any(not x.get("definition") for x in vals), "list %s: every value needs a definition (Definitions tab)" % name)
    for c in doc.get("columns") or []:
        g.check(not c.get("definition"), "column %s has no definition (Definitions tab)" % c.get("name"))
        if is_status(c.get("name")):
            g.check(not c.get("list"), "column %s is free text; give it a list from \"lists\" (dropdown)" % c.get("name"))
        if c.get("list"):
            g.check(c["list"] not in lists, "column %s uses list %s, which is not in \"lists\"" % (c.get("name"), c["list"]))
    for k, r in enumerate(rows, start=2):
        if g.check(len(r) != len(cols), "row %d has %d cells for %d columns" % (k, len(r), len(cols))):
            continue
        for c, v in zip(doc.get("columns") or [], r):
            allowed = [str(x.get("value")) for x in lists.get(c.get("list"), [])]
            if c.get("list") in lists and str(v).strip():
                g.check(str(v) not in allowed, "%s row %d: '%s' is not in the %s list (%s)" % (c["name"], k, v, c["list"], ", ".join(allowed)))
    return g


def cell_value(v, col):
    if col.get("type") == "date" and isinstance(v, str) and re.match(r"^\d{4}-\d{2}-\d{2}$", v):
        return datetime.date(*map(int, v.split("-")))
    return v


def write(doc, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo
    meta = doc.get("meta") or {}
    cols, rows = rows_of(doc)
    spec = doc["columns"]
    lists = doc.get("lists") or {}
    T = cur()
    C, b = T["colors"], T["brand"]
    fill = PatternFill("solid", fgColor=C[RULES["header_fill"]])
    hfont = Font(name=T["fonts"]["body"], bold=True, color=C[RULES["header_text"]])
    body = Font(name=T["fonts"]["body"], color=C["text"])
    year = (re.match(r"(\d{4})", str(meta.get("date", ""))) or [None, None])[1]
    owner = b.footer(year)  # "" when the project names no owner: no copyright line

    def table(ws, head, data, name, widths=None):
        ws.append(head)
        for r in data:
            ws.append(r)
        for j, h in enumerate(head, start=1):
            c = ws.cell(row=1, column=j)
            c.fill, c.font, c.alignment = fill, hfont, Alignment(vertical="center", wrap_text=True)
            w = max([units(h, bold=True)] + [units(r[j - 1]) for r in data if j - 1 < len(r)])
            ws.column_dimensions[L(j)].width = round(min(max(w + 2, 10), RULES["max_width"]), 1)
        for row in ws.iter_rows(min_row=2, max_row=len(data) + 1):
            for c in row:
                c.font, c.alignment = body, Alignment(horizontal="left", vertical="top", wrap_text=True)
                if isinstance(c.value, datetime.date):
                    c.number_format = "yyyy-mm-dd"
        t = Table(displayName=name, ref="A1:%s%d" % (L(len(head)), len(data) + 1))
        t.tableStyleInfo = TableStyleInfo(name=RULES["table_style"], showRowStripes=True)
        ws.add_table(t)
        ws.freeze_panes = "A2"
        ws.print_title_rows = "1:1"
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        if owner:
            ws.oddFooter.left.text = owner.replace("&", "&&")
        ws.oddFooter.right.text = "Page &P of &N"

    wb = Workbook()
    ws = wb.active
    ws.title = tab_of(meta)
    table(ws, cols, [[cell_value(v, c) for v, c in zip(r, spec)] for r in rows], table_name(ws.title))
    wl = wb.create_sheet(RULES["lists_tab"])
    wl.sheet_state = "hidden" if RULES.get("lists_hidden") else "visible"
    where = {}
    for j, (name, vals) in enumerate(lists.items(), start=1):
        wl.cell(row=1, column=j, value=name).fill = fill
        wl.cell(row=1, column=j).font = hfont
        for k, x in enumerate(vals, start=2):
            wl.cell(row=k, column=j, value=str(x["value"])).font = body
        wl.column_dimensions[L(j)].width = round(min(max([units(name, bold=True)] + [units(x["value"]) for x in vals]) + 2, RULES["max_width"]), 1)
        where[name] = "=%s!$%s$2:$%s$%d" % (RULES["lists_tab"], L(j), L(j), len(vals) + 1)
    wl.freeze_panes = "A2"
    last = max(len(rows) + 1, RULES["validate_rows"])
    for name in lists:
        used = [L(j) for j, c in enumerate(spec, start=1) if c.get("list") == name]
        if not used:
            continue
        dv = DataValidation(type="list", formula1=where[name], allow_blank=True, showErrorMessage=True, errorStyle="stop",
                            errorTitle="Choose from the list", error="Choose a %s value; the list is on the %s tab." % (name, RULES["lists_tab"]))
        for col in used:
            dv.add("%s2:%s%d" % (col, col, last))
        ws.add_data_validation(dv)
    wd = wb.create_sheet(RULES["definitions_tab"])
    defs = [[c["name"], "Column", c.get("list", ""), c.get("definition", "")] for c in spec]
    defs += [[str(x["value"]), "List value", name, x.get("definition", "")] for name, vals in lists.items() for x in vals]
    defs += [[k.replace("_", " ").title(), "About", "", "; ".join(v) if isinstance(v, list) else str(v)]
             for k, v in meta.items() if k in RULES["about_keys"] and v]
    table(wd, RULES["definitions_header"], defs, table_name(RULES["definitions_tab"]))
    wb.properties.title = str(meta.get("title") or ws.title)
    wb.properties.creator = b.org or b.display_name or "material-composer build_sheet.py"
    wb.active = 0
    wb.save(path)


def check_xlsx(path, g=None):
    """Reopen a built workbook and gate it (XLSX_TABLE)."""
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string, get_column_letter as L
    g = g or Gate("XLSX_TABLE")
    C = cur()["colors"]
    try:
        wb = load_workbook(path)
    except Exception as e:  # noqa: BLE001 - any unreadable workbook is a gate failure
        g.check(True, "%s does not open as a workbook (%s)" % (os.path.basename(path), e))
        return g
    for ws in wb.worksheets:
        g.check(bool(ws.merged_cells.ranges), "merged cells on %s: %s (merged cells break sorting, filtering and screen readers)"
                % (ws.title, ", ".join(sorted(str(r) for r in ws.merged_cells.ranges))))
    names = wb.sheetnames
    for t in (RULES["lists_tab"], RULES["definitions_tab"]):
        g.check(t not in names, "no %s tab" % t)
    ws = wb.worksheets[0]
    if g.check(not ws.tables, "%s has no named Excel table" % ws.title):
        return g
    t = list(ws.tables.values())[0]
    m = re.match(r"^A1:([A-Z]+)(\d+)$", t.ref)
    if g.check(not m, "table %s does not start at A1 (%s)" % (t.displayName, t.ref)):
        return g
    ncol, nrow = column_index_from_string(m.group(1)), int(m.group(2))
    used = max([r for r in range(1, ws.max_row + 1) if any(ws.cell(row=r, column=c).value not in (None, "") for c in range(1, ncol + 1))] or [1])
    g.check(used > nrow, "rows %d-%d of %s are outside table %s (%s)" % (nrow + 1, used, ws.title, t.displayName, t.ref))
    g.check(ws.freeze_panes != "A2", "header row of %s is not frozen (freeze panes at A2)" % ws.title)
    head = [str(ws.cell(row=1, column=c).value or "") for c in range(1, ncol + 1)]
    palette = {C[RULES["header_fill"]].upper()}
    for w in [ws] + [wb[x] for x in (RULES["definitions_tab"],) if x in names]:
        bad = [c.coordinate for c in w[1] if c.value not in (None, "") and (str(c.fill.fgColor.rgb or "")[-6:].upper() not in palette
                                                                           or str(c.font.color.rgb if c.font.color else "")[-6:].upper() != C[RULES["header_text"]])]
        g.check(bool(bad), "%s header %s not in the project brand colors (%s fill #%s, %s text)" % (w.title, ", ".join(bad), RULES["header_fill"], C[RULES["header_fill"]], RULES["header_text"]))
    lists = {}
    if RULES["lists_tab"] in names:
        wl = wb[RULES["lists_tab"]]
        for c in range(1, wl.max_column + 1):
            h = wl.cell(row=1, column=c).value
            vals, r = [], 2
            while wl.cell(row=r, column=c).value not in (None, ""):
                vals.append(str(wl.cell(row=r, column=c).value)); r += 1
            if h:
                lists[str(h)] = (L(c), vals)
    need = {}
    if RULES["definitions_tab"] in names:
        for r in wb[RULES["definitions_tab"]].iter_rows(min_row=2, values_only=True):
            if r and len(r) > 2 and r[1] == "Column" and r[2]:
                need[str(r[0])] = str(r[2])
    by_col = {}
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list":
            continue
        f = re.match(r"^=?'?%s'?!\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)$" % re.escape(RULES["lists_tab"]), str(dv.formula1 or ""))
        src = next((n for n, (col, vals) in lists.items() if f and f.group(1) == col == f.group(3)), None)
        if src:
            v = lists[src][1]
            g.check(int(f.group(2)) != 2 or int(f.group(4)) != len(v) + 1, "dropdown %s covers %s but the %s list is %s2:%s%d"
                    % (dv.formula1, "%s%s:%s%s" % f.groups(), src, lists[src][0], lists[src][0], len(v) + 1))
        for rg in dv.sqref.ranges:
            for c in range(rg.min_col, rg.max_col + 1):
                by_col.setdefault(c, []).append((src, rg.min_row, rg.max_row, dv.formula1))
    for c, h in enumerate(head, start=1):
        must = need.get(h) or (h if is_status(h) else None)
        dvs = by_col.get(c, [])
        if must and g.check(not dvs, "%s column %s is free text: no dropdown from the %s tab" % (ws.title, h, RULES["lists_tab"])):
            continue
        for src, lo, hi, f in dvs:
            if g.check(src is None, "%s column %s dropdown (%s) is not one list on the %s tab" % (ws.title, h, f, RULES["lists_tab"])):
                continue
            g.check(lo > 2 or hi < nrow, "%s column %s dropdown covers rows %d-%d; the table runs to row %d" % (ws.title, h, lo, hi, nrow))
            allowed = lists[src][1]
            for r in range(2, nrow + 1):
                v = ws.cell(row=r, column=c).value
                if v not in (None, ""):
                    g.check(str(v) not in allowed, "%s column %s row %d: '%s' is not in the %s list (%s)" % (ws.title, h, r, v, src, ", ".join(allowed)))
    return g


def run(doc, out=None, check=False, say=print, source=None, brand=None):
    """source: the sheet.json path (its project supplies the brand); brand: an explicit brand.json."""
    set_brand(source, brand)
    g = gate_source(doc)
    files = None
    if not g.failed:
        path = out + ".xlsx" if out and not check else os.path.join(tempfile.mkdtemp(prefix="sheet_"), "trial.xlsx")
        write(doc, path)
        check_xlsx(path, g)
        if g.failed and not check:
            os.remove(path)
        elif not check:
            files = [path]
    g.report(say)
    cols, rows = rows_of(doc)
    tail = ("refused: %s.xlsx not written" % os.path.basename(out or "")) if g.failed and out and not check else (
        "wrote %s (%s: %d rows; %s; %s)" % (os.path.basename(files[0]), tab_of(doc.get("meta") or {}), len(rows), RULES["lists_tab"], RULES["definitions_tab"]) if files else "checked")
    say("XLSX_TABLE GATE: %d checks, %d failed; %s" % (g.n, g.failed, tail))
    return g, files


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("src", help="sheet.json, or a built .xlsx with --check")
    ap.add_argument("--out", help="output name without .xlsx")
    ap.add_argument("--check", action="store_true")
    ap.add_argument("--brand", help="brand profile (brand.json); default: the project brand of the source or workbook")
    a = ap.parse_args()
    if a.src.lower().endswith(".xlsx"):
        set_brand(a.src, a.brand)
        g = check_xlsx(a.src)
        g.report()
        print("XLSX_TABLE GATE: %d checks, %d failed; %s" % (g.n, g.failed, os.path.basename(a.src)))
        return 1 if g.failed else 0
    with open(a.src, encoding="utf-8") as fh:
        doc = json.load(fh)
    if not a.check and not a.out:
        a.out = re.sub(r"[^A-Za-z0-9]+", "_", str((doc.get("meta") or {}).get("title") or "Sheet")).strip("_")
    g, files = run(doc, a.out, check=a.check, source=a.src, brand=a.brand)
    print(cur()["brand"].lines()[0])
    return 1 if g.failed else 0


if __name__ == "__main__":
    sys.exit(main())

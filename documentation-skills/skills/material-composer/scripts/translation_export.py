#!/usr/bin/env python3
"""Text-only Word export for written translation (TRANSLATION_READY).

  translation_export.py SOURCE.doc.json|SOURCE.md [--out-dir DIR]
Writes <Name>_for_translation.docx: title, headings, paragraphs, lists, steps, tables, callouts and FAQ answers as
plain editable text. Images, figures and diagrams are left out (and counted) so every word can be quoted and
translated. Send it through the written translation process named in the plan; never machine translation.
Prints "TRANSLATION EXPORT: <path> (<n> text blocks, <m> images left out)"."""
import argparse, json, os, re, sys

INLINE = re.compile(r"\*\*(.+?)\*\*|`([^`]+)`|\[([^\]]+)\]\([^)]+\)")


def plain(s):
    return INLINE.sub(lambda m: m.group(1) or m.group(2) or m.group(3) or "", str(s or "")).strip()


class Out:
    def __init__(self):
        import docx
        self.d = docx.Document(); self.n = 0; self.images = 0

    def head(self, text, level):
        if plain(text):
            self.d.add_heading(plain(text), min(level, 4)); self.n += 1

    def para(self, text, style=None):
        if plain(text):
            self.d.add_paragraph(plain(text), style=style); self.n += 1

    def table(self, header, rows):
        rows = [r for r in rows if r]
        if not (header or rows):
            return
        cols = max(len(r) for r in ([header] if header else []) + rows)
        t = self.d.add_table(rows=0, cols=cols); t.style = "Table Grid"
        for r in ([header] if header else []) + rows:
            cells = t.add_row().cells
            for i, c in enumerate(r):
                cells[i].text = plain(c)
        self.n += 1


def doc_blocks(o, blocks):
    for b in blocks:
        k = b.get("type")
        if k in ("paragraph", "code"):
            o.para(b.get("text") or b.get("code"))
        elif k in ("bullets", "numbered"):
            for x in b.get("items", []):
                o.para(x, "List Bullet" if k == "bullets" else "List Number")
        elif k == "procedure":
            o.para(b.get("title"))
            for s in b.get("steps", []):
                o.para(s.get("text"), "List Number"); o.para(s.get("result"))
                o.images += 1 if s.get("image") else 0
        elif k == "table":
            o.para(b.get("caption")); o.table(b.get("columns") or [], b.get("rows", []))
        elif k in ("callout", "card"):
            o.para(b.get("title")); o.para(b.get("text"))
            for x in b.get("items", []):
                o.para(x, "List Bullet")
        elif k == "faq":
            for q in b.get("items", []):
                o.para(q.get("q")); o.para(q.get("a"))
        elif k == "step_diagram":
            o.images += 1
            for s in b.get("steps", []):
                o.para(s.get("label"), "List Number")
        elif k == "figure":
            o.images += 1; o.para(b.get("caption"))


def doc_sections(o, secs, level):
    for s in secs:
        o.head(s.get("title"), level); doc_blocks(o, s.get("blocks", [])); doc_sections(o, s.get("sections", []), level + 1)


def from_md(o, text):
    m = re.match(r"^---\n(.*?)\n---\n?", text, re.S)
    if m:
        t = re.search(r"(?m)^title:\s*(.+)$", m.group(1))
        if t:
            o.head(t.group(1), 0)
        text = text[m.end():]
    rows = []
    for line in text.splitlines() + [""]:
        if line.startswith("|"):
            if not re.match(r"^\|\s*:?-", line):
                rows.append([c.strip() for c in line.strip().strip("|").split("|")])
            continue
        if rows:
            o.table(rows[0], rows[1:]); rows = []
        if re.search(r"!\[[^\]]*\]\([^)]+\)", line):
            o.images += 1; line = re.sub(r"!\[[^\]]*\]\([^)]+\)", "", line)
        h = re.match(r"^(#{1,4})\s+(.+)$", line)
        if h:
            o.head(h.group(2), len(h.group(1)))
        elif re.match(r"^\s*[-*]\s+", line):
            o.para(re.sub(r"^\s*[-*]\s+(\[[ x]\]\s*)?", "", line), "List Bullet")
        elif re.match(r"^\s*\d+[.)]\s+", line):
            o.para(re.sub(r"^\s*\d+[.)]\s+", "", line), "List Number")
        else:
            o.para(line)


def export(src, out_dir=None):
    o = Out()
    if src.endswith(".json"):
        d = json.load(open(src, encoding="utf-8")); m = d.get("meta", {})
        o.head(m.get("title"), 0); o.para(m.get("subtitle")); doc_sections(o, d.get("sections", []), 1)
    else:
        from_md(o, open(src, encoding="utf-8").read())
    name = os.path.basename(src).split(".")[0]
    path = os.path.join(out_dir or os.path.dirname(os.path.abspath(src)), "%s_for_translation.docx" % name)
    o.d.save(path)
    print("TRANSLATION EXPORT: %s (%d text blocks, %d images left out)" % (path, o.n, o.images))
    return path


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("source"); ap.add_argument("--out-dir")
    a = ap.parse_args()
    export(a.source, a.out_dir)
    return 0


if __name__ == "__main__":
    sys.exit(main())
